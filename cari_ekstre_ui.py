"""Stok detaylı cari hesap ekstresi — örnek Excel düzenine yakın.

Ana kolonlar: Tarih | (Açıklama alanı: Kodu…Tutar) | Borç | Alacak | Bakiye
Fatura altında stok satırları ayrı kolonlarda; borç/alacak/bakiye yalnız belge satırında.
"""

from __future__ import annotations

import html
import tempfile
import webbrowser
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
import tkinter as tk

from database.rapor_service import RaporService
from ui_takvim import takvim_butonu

STOKLU_TURLER = (
    "Satış Faturası",
    "Alış Faturası",
    "Satış İadesi",
    "Alış İadesi",
)

# Örnek Excel: KODU | ÜRÜN ADI | MİKTAR | BİRİM | FİYAT | TUTAR
STOK_ALT_BASLIK = ("Kodu", "Ürün Adı", "Miktar", "Birim", "Fiyat", "Tutar")


def _para(tutar) -> str:
    if tutar is None:
        return ""
    d = Decimal(str(tutar))
    if d == 0:
        return ""
    return f"{float(d):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _para_zorunlu(tutar) -> str:
    d = Decimal(str(tutar or 0))
    return f"{float(d):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _miktar(deger) -> str:
    d = Decimal(str(deger or 0))
    if d == d.to_integral_value():
        return str(int(d))
    return (
        f"{float(d):,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")
        .rstrip("0")
        .rstrip(",")
    )


def _tarih(t) -> str:
    if not t:
        return ""
    if isinstance(t, str):
        return t
    return t.strftime("%d.%m.%Y")


def _belge_baslik(belge: dict) -> str:
    """Örnek: 'S-113 NOLU FATURA TOPLAMI' / tahsilat açıklaması."""
    tur = (belge.get("tur") or "").strip()
    no = (belge.get("belge_no") or "").strip()
    aciklama = (belge.get("aciklama") or "").strip()
    stoklu = tur in STOKLU_TURLER
    if stoklu and no:
        return f"{no} NOLU FATURA TOPLAMI"
    if aciklama:
        return aciklama.upper() if len(aciklama) < 80 else aciklama
    if tur and no:
        return f"{tur} — {no}"
    return (tur or no or "Hareket").upper()


def _iskontolu_kdvli_birim_fiyat(satir: dict) -> Decimal | None:
    miktar = Decimal(str(satir.get("miktar") or 0))
    genel = Decimal(str(satir.get("genel") or 0))
    if miktar != 0:
        return genel / miktar
    birim_fiyat = Decimal(str(satir.get("birim_fiyat") or 0))
    if birim_fiyat == 0:
        return None
    iskonto = Decimal(str(satir.get("iskonto_orani") or 0))
    kdv = Decimal(str(satir.get("kdv_orani") or 0))
    net_birim = birim_fiyat * (Decimal("1") - iskonto / Decimal("100"))
    return net_birim * (Decimal("1") + kdv / Decimal("100"))


def _stok_degerleri(satir: dict) -> tuple[str, str, str, str, str, str]:
    fiyat = _iskontolu_kdvli_birim_fiyat(satir)
    return (
        (satir.get("urun_kodu") or "").strip(),
        (satir.get("urun_adi") or "").strip(),
        _miktar(satir.get("miktar")),
        (satir.get("birim") or "").strip(),
        _para_zorunlu(fiyat) if fiyat is not None else "",
        _para_zorunlu(satir.get("genel")),
    )


class StokDetayliCariEkstreDialog(tk.Toplevel):
    """Örnek Excel: Tarih + Açıklama(Kodu…Tutar) + Borç + Alacak + Bakiye."""

    def __init__(self, parent, cari):
        super().__init__(parent)
        if cari is None:
            raise ValueError("Cari gerekli.")
        self.cari = cari
        self._ham = None
        self._zoom = 1.0
        self._tam_ekran = False
        self._onceki_geometry = None
        self._oturum_genislik = None  # kullanıcı sürükleyince oturumda tutulur
        self.title(f"Stok Detaylı Cari Ekstre — {cari.cari_kodu} {cari.unvan}")
        self.geometry("1280x700")
        self.minsize(1000, 520)
        self.resizable(True, True)
        self.transient(parent)
        self.grab_set()

        ust = ttk.Frame(self, padding=(12, 10, 12, 6))
        ust.pack(fill="x")
        ttk.Label(
            ust,
            text="STOK DETAYLI CARİ HESAP EKSTRESİ",
            font=("Segoe UI", 12, "bold"),
        ).pack(anchor="w")
        self.ozet = ttk.Label(
            ust,
            text=f"{cari.cari_kodu} — {cari.unvan}",
            font=("Segoe UI", 10),
            foreground="#1a237e",
        )
        self.ozet.pack(anchor="w", pady=(2, 0))

        filtre = ttk.Frame(self, padding=(12, 0, 12, 6))
        filtre.pack(fill="x")
        ttk.Label(filtre, text="Başlangıç:").pack(side="left")
        bas_c = ttk.Frame(filtre)
        bas_c.pack(side="left", padx=(4, 8))
        self.tarih_bas = ttk.Entry(bas_c, width=11)
        self.tarih_bas.pack(side="left")
        takvim_butonu(bas_c, self.tarih_bas, on_select=self.yenile)

        ttk.Label(filtre, text="Bitiş:").pack(side="left")
        bit_c = ttk.Frame(filtre)
        bit_c.pack(side="left", padx=(4, 8))
        self.tarih_bit = ttk.Entry(bit_c, width=11)
        self.tarih_bit.pack(side="left")
        takvim_butonu(bit_c, self.tarih_bit, on_select=self.yenile)

        ttk.Label(filtre, text="Stok ara:").pack(side="left", padx=(4, 0))
        self.stok_ara = ttk.Entry(filtre, width=16)
        self.stok_ara.pack(side="left", padx=4)
        self.stok_ara.bind("<Return>", lambda _e: self.yenile())

        ttk.Button(filtre, text="Uygula", command=self.yenile).pack(side="left", padx=(8, 4))
        ttk.Button(filtre, text="Tümünü Aç", command=self._hepsini_ac).pack(side="left", padx=2)
        ttk.Button(filtre, text="Tümünü Kapat", command=self._hepsini_kapat).pack(side="left", padx=2)
        ttk.Button(filtre, text="Excel", command=self._excel_aktar).pack(side="left", padx=(12, 0))

        zoom_bar = ttk.Frame(filtre)
        zoom_bar.pack(side="left", padx=(16, 0))
        ttk.Button(zoom_bar, text="−", width=3, command=self._zoom_azalt).pack(side="left")
        self.zoom_lbl = ttk.Label(zoom_bar, text="100%", width=5, anchor="center")
        self.zoom_lbl.pack(side="left", padx=2)
        ttk.Button(zoom_bar, text="+", width=3, command=self._zoom_arttir).pack(side="left")
        ttk.Button(
            filtre, text="Kolonları Sıfırla", command=self._kolonlari_sifirla
        ).pack(side="left", padx=(10, 0))

        kapat_btn = ttk.Button(filtre, text="✕", width=3, command=self.destroy)
        kapat_btn.pack(side="right")
        self.tam_ekran_btn = ttk.Button(
            filtre, text="□", width=3, command=self._tam_ekran_degistir
        )
        self.tam_ekran_btn.pack(side="right", padx=(0, 2))
        try:
            _win_font = ("Segoe UI", 9)
            self.tam_ekran_btn.configure(font=_win_font)
            kapat_btn.configure(font=_win_font)
        except tk.TclError:
            pass
        ttk.Button(filtre, text="Yazdır", command=self._yazdir).pack(side="right", padx=(0, 8))

        tablo_cerceve = ttk.Frame(self, padding=(12, 0, 12, 8))
        tablo_cerceve.pack(fill="both", expand=True)

        self._stil = ttk.Style(self)
        try:
            self._stil.map(
                "StokEkstre.Treeview",
                background=[("selected", "#cfe2f3")],
                foreground=[("selected", "#000000")],
            )
        except tk.TclError:
            pass

        # Örnek Excel A..M: tarih | kod..tutar | borç | alacak | bakiye
        self._kolonlar = (
            "tarih",
            "kod",
            "urun",
            "miktar",
            "birim",
            "fiyat",
            "tutar",
            "borc",
            "alacak",
            "bakiye",
        )
        self._kolon_ayar = {
            "tarih": ("Tarih", 96, "center"),
            "kod": ("Kodu", 90, "w"),
            "urun": ("Ürün Adı", 280, "w"),
            "miktar": ("Miktar", 72, "e"),
            "birim": ("Birim", 64, "w"),
            "fiyat": ("Fiyat", 96, "e"),
            "tutar": ("Tutar", 100, "e"),
            "borc": ("Borç", 110, "e"),
            "alacak": ("Alacak", 110, "e"),
            "bakiye": ("Bakiye", 110, "e"),
        }
        self.tablo = ttk.Treeview(
            tablo_cerceve,
            columns=self._kolonlar,
            show="headings",
            style="StokEkstre.Treeview",
            selectmode="browse",
        )
        self._uygula_zoom()
        # Başlık ayırıcılarını sürükleyerek kolon genişlet/daralt (ttk.Treeview)
        self._kolon_surukleniyor = False
        self.tablo.bind("<ButtonPress-1>", self._kolon_ayirici_basla, add="+")
        self.tablo.bind("<ButtonRelease-1>", self._kolon_genislik_kaydet, add="+")

        dikey = ttk.Scrollbar(tablo_cerceve, orient="vertical", command=self.tablo.yview)
        yatay = ttk.Scrollbar(tablo_cerceve, orient="horizontal", command=self.tablo.xview)
        self.tablo.configure(yscrollcommand=dikey.set, xscrollcommand=yatay.set)
        self.tablo.grid(row=0, column=0, sticky="nsew")
        dikey.grid(row=0, column=1, sticky="ns")
        yatay.grid(row=1, column=0, sticky="ew")
        tablo_cerceve.rowconfigure(0, weight=1)
        tablo_cerceve.columnconfigure(0, weight=1)

        alt = ttk.Frame(self, padding=(12, 0, 12, 12))
        alt.pack(fill="x")
        self.alt_ozet = ttk.Label(alt, text="", font=("Segoe UI", 9, "bold"), foreground="#1a237e")
        self.alt_ozet.pack(side="right")

        self.after(50, self.yenile)

    def _varsayilan_genislikler(self) -> dict:
        z = self._zoom
        return {
            k: max(36, int(round(self._kolon_ayar[k][1] * z))) for k in self._kolonlar
        }

    def _kolon_minwidth(self, kolon: str) -> int:
        """Sürükleyerek daraltmaya izin verecek makul alt sınırlar."""
        z = self._zoom
        if kolon in ("borc", "alacak", "bakiye"):
            return max(48, int(round(48 * z)))
        if kolon in ("fiyat", "tutar", "miktar"):
            return max(40, int(round(40 * z)))
        if kolon == "urun":
            return max(72, int(round(72 * z)))
        if kolon in ("tarih", "kod"):
            return max(40, int(round(40 * z)))
        return max(32, int(round(32 * z)))

    def _kolon_stretch(self, kolon: str) -> bool:
        # Açıklama/ürün alanı pencere büyüyünce yer kaplasın; Borç/Alacak/Bakiye sabit kalsın
        return kolon == "urun"

    def _kolonlari_uygula(self, genislikler: dict | None = None) -> None:
        gen = genislikler if genislikler is not None else (
            self._oturum_genislik or self._varsayilan_genislikler()
        )
        self.tablo.column("#0", width=0, minwidth=0, stretch=False)
        for kolon in self._kolonlar:
            baslik, _, hiza = self._kolon_ayar[kolon]
            w = max(self._kolon_minwidth(kolon), int(gen.get(kolon, 80)))
            self.tablo.heading(kolon, text=baslik, anchor=hiza)
            self.tablo.column(
                kolon,
                width=w,
                minwidth=self._kolon_minwidth(kolon),
                anchor=hiza,
                stretch=self._kolon_stretch(kolon),
            )

    def _kolon_ayirici_basla(self, event) -> None:
        try:
            self._kolon_surukleniyor = self.tablo.identify_region(event.x, event.y) == "separator"
        except tk.TclError:
            self._kolon_surukleniyor = False

    def _kolon_genislik_kaydet(self, event=None) -> None:
        """Başlık ayırıcısı sürüklenince genişlikleri oturumda sakla."""
        surukleme = getattr(self, "_kolon_surukleniyor", False)
        self._kolon_surukleniyor = False
        if event is not None and not surukleme:
            try:
                bolge = self.tablo.identify_region(event.x, event.y)
            except tk.TclError:
                bolge = ""
            if bolge != "separator":
                return
        self._oturum_genislik = {
            k: int(self.tablo.column(k, "width")) for k in self._kolonlar
        }

    def _kolonlari_sifirla(self) -> None:
        self._oturum_genislik = None
        self._kolonlari_uygula(self._varsayilan_genislikler())

    def _uygula_zoom(self):
        z = self._zoom
        stil = self._stil
        row_h = max(18, int(round(22 * z)))
        f_govde = max(7, int(round(9 * z)))
        f_baslik = max(7, int(round(9 * z)))
        f_belge = max(8, int(round(10 * z)))
        f_stok = max(7, int(round(9 * z)))
        stil.configure(
            "StokEkstre.Treeview",
            rowheight=row_h,
            font=("Segoe UI", f_govde),
            fieldbackground="#ffffff",
        )
        stil.configure(
            "StokEkstre.Treeview.Heading",
            font=("Segoe UI", f_baslik, "bold"),
            padding=(max(2, int(4 * z)), max(2, int(4 * z))),
        )
        # Zoom değişince varsayılan oranlara dön; kullanıcı sürüklemesi oturumda kalır
        # yalnızca zoom sabitken. Zoom’da oturum genişliklerini oranla ölçekle.
        if self._oturum_genislik is not None and hasattr(self, "_zoom_onceki"):
            oran = z / self._zoom_onceki if self._zoom_onceki else 1.0
            if abs(oran - 1.0) > 1e-6:
                self._oturum_genislik = {
                    k: max(self._kolon_minwidth(k), int(round(v * oran)))
                    for k, v in self._oturum_genislik.items()
                }
        self._zoom_onceki = z
        self._kolonlari_uygula()
        self.tablo.tag_configure(
            "belge", background="#e8eef5", font=("Segoe UI", f_belge, "bold")
        )
        self.tablo.tag_configure(
            "belge_cift", background="#f5f7fa", font=("Segoe UI", f_belge, "bold")
        )
        self.tablo.tag_configure(
            "stok_baslik",
            background="#cfd8dc",
            foreground="#263238",
            font=("Segoe UI", f_stok, "bold"),
        )
        self.tablo.tag_configure(
            "stok", background="#ffffff", foreground="#37474f", font=("Segoe UI", f_stok)
        )
        self.tablo.tag_configure(
            "stok_cift",
            background="#fafbfc",
            foreground="#37474f",
            font=("Segoe UI", f_stok),
        )
        if hasattr(self, "zoom_lbl"):
            self.zoom_lbl.configure(text=f"{int(round(z * 100))}%")

    def _zoom_arttir(self):
        if self._zoom >= 1.4:
            return
        self._zoom = round(min(1.4, self._zoom + 0.2), 1)
        self._uygula_zoom()

    def _zoom_azalt(self):
        if self._zoom <= 0.8:
            return
        self._zoom = round(max(0.8, self._zoom - 0.2), 1)
        self._uygula_zoom()

    def _ekran_boyutuna_yay(self):
        w = self.winfo_screenwidth()
        h = self.winfo_screenheight()
        self.geometry(f"{w}x{h}+0+0")

    def _tam_ekran_degistir(self):
        if not self._tam_ekran:
            self._onceki_geometry = self.geometry()
            try:
                self.wm_state("zoomed")
            except tk.TclError:
                try:
                    self.state("zoomed")
                except tk.TclError:
                    self._ekran_boyutuna_yay()
            self._tam_ekran = True
            self.tam_ekran_btn.configure(text="❐")
        else:
            try:
                self.wm_state("normal")
            except tk.TclError:
                try:
                    self.state("normal")
                except tk.TclError:
                    pass
            if self._onceki_geometry:
                self.geometry(self._onceki_geometry)
            self._tam_ekran = False
            self.tam_ekran_btn.configure(text="□")

    def _belge_degerleri(self, tarih, baslik, borc, alacak, bakiye) -> tuple:
        # Açıklama alanı (kod…tutar): belge başlığı Ürün Adı kolonunda (örnekteki birleşik alan)
        return (
            _tarih(tarih),
            "",
            baslik,
            "",
            "",
            "",
            "",
            _para(borc),
            _para(alacak),
            _para(bakiye),
        )

    def _stok_baslik_degerleri(self) -> tuple:
        return ("",) + STOK_ALT_BASLIK + ("", "", "")

    def _stok_satir_degerleri(self, satir: dict) -> tuple:
        kod, urun, miktar, birim, fiyat, tutar = _stok_degerleri(satir)
        return ("", kod, urun, miktar, birim, fiyat, tutar, "", "", "")

    def _tarih_araligi(self):
        baslangic = bitis = None
        try:
            bas = (self.tarih_bas.get() or "").strip()
            bit = (self.tarih_bit.get() or "").strip()
            if bas:
                baslangic = datetime.strptime(bas, "%d.%m.%Y").date()
            if bit:
                bitis = datetime.strptime(bit, "%d.%m.%Y").date()
        except ValueError:
            messagebox.showerror("Tarih", "Tarihleri gg.aa.yyyy formatında girin.", parent=self)
            return None, None, False
        if baslangic and bitis and baslangic > bitis:
            messagebox.showwarning("Tarih", "Başlangıç bitişten sonra olamaz.", parent=self)
            return None, None, False
        return baslangic, bitis, True

    def yenile(self):
        baslangic, bitis, ok = self._tarih_araligi()
        if not ok:
            return
        if self._ham is None:
            try:
                self._ham = RaporService.stok_detayli_ekstre(self.cari.id)
            except ValueError as hata:
                messagebox.showerror("Ekstre", str(hata), parent=self)
                return

        rapor = self._ham
        stok_filtre = (self.stok_ara.get() or "").strip().casefold()

        for item in self.tablo.get_children():
            self.tablo.delete(item)

        belgeler = list(reversed(rapor.get("belgeler") or []))
        gosterilen = 0
        borc_t = Decimal("0")
        alacak_t = Decimal("0")
        son_bakiye = Decimal("0")

        for belge in belgeler:
            tarih = belge.get("tarih")
            if baslangic and tarih and tarih < baslangic:
                continue
            if bitis and tarih and tarih > bitis:
                continue

            stok_satirlari = list(belge.get("stok_satirlari") or [])
            stoklu = belge.get("tur") in STOKLU_TURLER
            if stok_filtre:
                stok_satirlari = [
                    s
                    for s in stok_satirlari
                    if stok_filtre in (s.get("urun_kodu") or "").casefold()
                    or stok_filtre in (s.get("urun_adi") or "").casefold()
                ]
                if stoklu and not stok_satirlari:
                    continue
                if not stoklu:
                    continue

            borc = Decimal(str(belge.get("borc") or 0))
            alacak = Decimal(str(belge.get("alacak") or 0))
            bakiye = Decimal(str(belge.get("bakiye") or 0))
            tag = "belge" if gosterilen % 2 == 0 else "belge_cift"
            parent = self.tablo.insert(
                "",
                "end",
                values=self._belge_degerleri(
                    tarih, _belge_baslik(belge), borc, alacak, bakiye
                ),
                tags=(tag,),
                open=True,
            )
            gosterilen += 1
            borc_t += borc
            alacak_t += alacak
            son_bakiye = bakiye

            if stoklu and stok_satirlari:
                self.tablo.insert(
                    parent,
                    "end",
                    values=self._stok_baslik_degerleri(),
                    tags=("stok_baslik",),
                )
                for i, s in enumerate(stok_satirlari):
                    stok_tag = "stok" if i % 2 == 0 else "stok_cift"
                    self.tablo.insert(
                        parent,
                        "end",
                        values=self._stok_satir_degerleri(s),
                        tags=(stok_tag,),
                    )

        self.ozet.configure(
            text=(
                f"{rapor['cari'].cari_kodu} — {rapor['cari'].unvan}  |  "
                f"Genel bakiye: {_para(rapor.get('bakiye')) or '0,00'}"
            )
        )
        self.alt_ozet.configure(
            text=(
                f"Belge: {gosterilen}  |  "
                f"Borç: {_para(borc_t) or '0,00'}  |  "
                f"Alacak: {_para(alacak_t) or '0,00'}  |  "
                f"Son bakiye: {_para(son_bakiye) or '0,00'}"
            )
        )

    def _hepsini_ac(self):
        for item in self.tablo.get_children():
            self.tablo.item(item, open=True)

    def _hepsini_kapat(self):
        for item in self.tablo.get_children():
            self.tablo.item(item, open=False)

    def _filtrelenmis_belgeler(self):
        if self._ham is None:
            self.yenile()
        if self._ham is None:
            return None
        baslangic, bitis, ok = self._tarih_araligi()
        if not ok:
            return None
        stok_filtre = (self.stok_ara.get() or "").strip().casefold()
        belgeler = list(reversed(self._ham.get("belgeler") or []))
        sonuc = []
        for belge in belgeler:
            tarih = belge.get("tarih")
            if baslangic and tarih and tarih < baslangic:
                continue
            if bitis and tarih and tarih > bitis:
                continue
            stok_satirlari = list(belge.get("stok_satirlari") or [])
            stoklu = belge.get("tur") in STOKLU_TURLER
            if stok_filtre:
                stok_satirlari = [
                    s
                    for s in stok_satirlari
                    if stok_filtre in (s.get("urun_kodu") or "").casefold()
                    or stok_filtre in (s.get("urun_adi") or "").casefold()
                ]
                if stoklu and not stok_satirlari:
                    continue
                if not stoklu:
                    continue
            sonuc.append((belge, stok_satirlari if stoklu else []))
        return sonuc

    def _ekstre_html(self, belgeler) -> str:
        """A4 yazdır — örnek Excel: Tarih | Açıklama(6 alt kolon) | Borç | Alacak | Bakiye."""
        rapor = self._ham
        cari = rapor["cari"]
        bas = (self.tarih_bas.get() or "").strip() or "—"
        bit = (self.tarih_bit.get() or "").strip() or "—"
        bakiye = _para(rapor.get("bakiye")) or "0,00"

        bloklar = []
        for belge, stok_satirlari in belgeler:
            tarih = html.escape(_tarih(belge.get("tarih")))
            aciklama = html.escape(_belge_baslik(belge))
            borc = html.escape(_para(belge.get("borc")) or "")
            alacak = html.escape(_para(belge.get("alacak")) or "")
            bakiye_s = html.escape(_para(belge.get("bakiye")) or "")
            blok = [
                "<tr class='belge'>",
                f"<td class='c'>{tarih}</td>",
                f"<td class='aciklama' colspan='6'>{aciklama}</td>",
                f"<td class='n'>{borc}</td>",
                f"<td class='n'>{alacak}</td>",
                f"<td class='n'>{bakiye_s}</td>",
                "</tr>",
            ]
            if stok_satirlari:
                blok.append(
                    "<tr class='stok-baslik'>"
                    "<td></td>"
                    "<td>Kodu</td><td>Ürün Adı</td><td class='n'>Miktar</td>"
                    "<td>Birim</td><td class='n'>Fiyat</td><td class='n'>Tutar</td>"
                    "<td></td><td></td><td></td>"
                    "</tr>"
                )
                for i, s in enumerate(stok_satirlari):
                    kod, urun, miktar, birim, fiyat, tutar = _stok_degerleri(s)
                    cls = "stok" if i % 2 == 0 else "stok-cift"
                    blok.append(
                        f"<tr class='{cls}'>"
                        "<td></td>"
                        f"<td>{html.escape(kod)}</td>"
                        f"<td>{html.escape(urun)}</td>"
                        f"<td class='n'>{html.escape(miktar)}</td>"
                        f"<td>{html.escape(birim)}</td>"
                        f"<td class='n'>{html.escape(fiyat)}</td>"
                        f"<td class='n'>{html.escape(tutar)}</td>"
                        "<td></td><td></td><td></td>"
                        "</tr>"
                    )
            bloklar.append("\n".join(blok))

        govde = "\n".join(bloklar) or (
            "<tr><td colspan='10'>Filtrelere uygun hareket yok.</td></tr>"
        )
        baslik = "Stok Detaylı Cari Hesap Ekstresi"
        return f"""<!DOCTYPE html>
<html lang="tr">
<head>
<meta charset="utf-8"/>
<title>{html.escape(baslik)}</title>
<style>
  @page {{ size: A4 portrait; margin: 10mm; }}
  * {{ box-sizing: border-box; }}
  html, body {{
    margin: 0; padding: 0;
    width: 100%; max-width: 100%;
    overflow-x: hidden;
    font-family: 'Segoe UI', Calibri, Tahoma, sans-serif;
    color: #111;
    font-size: 9pt;
  }}
  body {{ width: 190mm; max-width: 100%; }}
  .toolbar {{ margin-bottom: 8px; }}
  .toolbar button {{ padding: 6px 14px; font-size: 13px; cursor: pointer; }}
  h1 {{ font-size: 13pt; margin: 0 0 4px; font-weight: 700; }}
  .meta {{ margin: 0 0 10px; color: #333; font-size: 8.5pt; line-height: 1.35; }}
  table.ekstre {{
    border-collapse: collapse;
    width: 100%; max-width: 100%;
    table-layout: fixed;
  }}
  table.ekstre col.c-tarih {{ width: 9%; }}
  table.ekstre col.c-kod {{ width: 8%; }}
  table.ekstre col.c-urun {{ width: 22%; }}
  table.ekstre col.c-miktar {{ width: 7%; }}
  table.ekstre col.c-birim {{ width: 6%; }}
  table.ekstre col.c-fiyat {{ width: 9%; }}
  table.ekstre col.c-tutar {{ width: 9%; }}
  table.ekstre col.c-borc {{ width: 10%; }}
  table.ekstre col.c-alacak {{ width: 10%; }}
  table.ekstre col.c-bakiye {{ width: 10%; }}
  table.ekstre > thead > tr > th {{
    border: none;
    border-bottom: 2px solid #333;
    background: #f0f0f0;
    padding: 3px 3px;
    font-size: 8pt;
    font-weight: 700;
    vertical-align: middle;
    white-space: nowrap;
  }}
  table.ekstre > tbody > tr.belge > td {{
    border: none;
    border-bottom: 1px solid #c5cdd4;
    padding: 3px 3px;
    vertical-align: middle;
    font-weight: 700;
    font-size: 9.5pt;
    background: #eef2f6;
    line-height: 1.25;
  }}
  table.ekstre > tbody > tr.belge > td.aciklama {{
    text-align: right;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
  }}
  table.ekstre > tbody > tr.stok-baslik > td {{
    border: none;
    border-bottom: 1px solid #90a4ae;
    padding: 2px 3px;
    font-weight: 700;
    font-size: 7.5pt;
    background: #cfd8dc;
    color: #263238;
  }}
  table.ekstre > tbody > tr.stok > td,
  table.ekstre > tbody > tr.stok-cift > td {{
    border: none;
    border-bottom: 1px solid #e8ecf0;
    padding: 2px 3px;
    font-weight: 400;
    font-size: 8pt;
    line-height: 1.2;
    background: #fff;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
  }}
  table.ekstre > tbody > tr.stok-cift > td {{ background: #fafbfc; }}
  th.c, td.c {{ text-align: center; }}
  th.n, td.n {{ text-align: right; }}
  th.aciklama {{ text-align: center; }}
  @media print {{
    .toolbar {{ display: none !important; }}
    html, body {{
      margin: 0 !important;
      padding: 0 !important;
      width: 100% !important;
      max-width: 100% !important;
      overflow: hidden;
    }}
  }}
</style>
</head>
<body>
  <div class="toolbar">
    <button type="button" onclick="window.print()">Yazdır</button>
  </div>
  <h1>{html.escape(baslik)}</h1>
  <p class="meta">
    <strong>{html.escape(cari.cari_kodu)}</strong> — {html.escape(cari.unvan or "")}<br/>
    Dönem: {html.escape(bas)} – {html.escape(bit)} &nbsp;·&nbsp;
    Genel bakiye: {html.escape(bakiye)}
  </p>
  <table class="ekstre">
    <colgroup>
      <col class="c-tarih"/>
      <col class="c-kod"/><col class="c-urun"/><col class="c-miktar"/>
      <col class="c-birim"/><col class="c-fiyat"/><col class="c-tutar"/>
      <col class="c-borc"/><col class="c-alacak"/><col class="c-bakiye"/>
    </colgroup>
    <thead>
      <tr>
        <th class="c">Tarih</th>
        <th class="aciklama" colspan="6">Açıklama</th>
        <th class="n">Borç</th>
        <th class="n">Alacak</th>
        <th class="n">Bakiye</th>
      </tr>
    </thead>
    <tbody>
      {govde}
    </tbody>
  </table>
</body>
</html>
"""

    def _yazdir(self):
        belgeler = self._filtrelenmis_belgeler()
        if belgeler is None:
            return
        klasor = Path(tempfile.gettempdir()) / "muhasebe_belge"
        klasor.mkdir(exist_ok=True)
        kod = "".join(
            c if c.isalnum() or c in "-_" else "_" for c in (self.cari.cari_kodu or "")
        )[:20]
        dosya = klasor / f"stok_ekstre_{kod or 'cari'}.html"
        dosya.write_text(self._ekstre_html(belgeler), encoding="utf-8")
        webbrowser.open(dosya.as_uri())
        messagebox.showinfo(
            "Yazdırma",
            "Ekstre tarayıcıda açıldı.\n"
            "Yazdır iletişim kutusu için Yazdır düğmesine basın veya Ctrl+P kullanın.\n"
            "Sayfa yönü: A4 dikey (portrait).",
            parent=self,
        )

    def _excel_aktar(self):
        """Örnek Excel ile aynı birleşik hücre düzeni (A=tarih, B:J=açıklama, K:M=tutarlar)."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            messagebox.showerror(
                "Excel",
                "openpyxl yüklü değil. Kurulum: pip install openpyxl",
                parent=self,
            )
            return
        belgeler_filt = self._filtrelenmis_belgeler()
        if belgeler_filt is None:
            return

        yol = filedialog.asksaveasfilename(
            parent=self,
            title="Ekstreyi kaydet",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"stok_ekstre_{self.cari.cari_kodu}_{date.today():%Y%m%d}.xlsx",
        )
        if not yol:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Stok Ekstre"

        # Kolon genişlikleri — örnek dosyaya yakın
        genislikler = {
            "A": 12,
            "B": 12,
            "C": 12,
            "D": 12,
            "E": 12,
            "F": 12,
            "G": 10,
            "H": 9,
            "I": 12,
            "J": 12,
            "K": 12,
            "L": 12,
            "M": 12,
        }
        for col, width in genislikler.items():
            ws.column_dimensions[col].width = width

        baslik_font = Font(name="Calibri", size=12, bold=True)
        belge_font = Font(name="Calibri", size=12, bold=True)
        stok_font = Font(name="Calibri", size=11)
        stok_baslik_font = Font(name="Calibri", size=11, bold=True)
        stok_baslik_fill = PatternFill("solid", fgColor="CFD8DC")
        sag = Alignment(horizontal="right", vertical="center")
        orta = Alignment(horizontal="center", vertical="center")

        # Üst başlık: Tarih | AÇIKLAMA (B:J) | BORÇ | ALACAK | BAKİYE
        ws["A1"] = "Tarih"
        ws["A1"].font = baslik_font
        ws["B1"] = "AÇIKLAMA"
        ws["B1"].font = Font(name="Calibri", size=11, bold=True)
        ws["B1"].alignment = orta
        ws.merge_cells("B1:J1")
        ws["K1"] = "BORÇ"
        ws["L1"] = "ALACAK"
        ws["M1"] = "BAKİYE"
        for col in ("K", "L", "M"):
            ws[f"{col}1"].font = baslik_font
            ws[f"{col}1"].alignment = sag

        row = 2
        for belge, stok_satirlari in belgeler_filt:
            tarih = belge.get("tarih")
            borc = float(belge.get("borc") or 0) or None
            alacak = float(belge.get("alacak") or 0) or None
            bakiye = float(belge.get("bakiye") or 0) or None

            ws.cell(row, 1, _tarih(tarih)).font = belge_font
            ws.cell(row, 2, _belge_baslik(belge)).font = belge_font
            ws.cell(row, 2).alignment = sag
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            if borc is not None:
                c = ws.cell(row, 11, borc)
                c.font = belge_font
                c.number_format = "#,##0.00"
                c.alignment = sag
            if alacak is not None:
                c = ws.cell(row, 12, alacak)
                c.font = belge_font
                c.number_format = "#,##0.00"
                c.alignment = sag
            if bakiye is not None:
                c = ws.cell(row, 13, bakiye)
                c.font = belge_font
                c.number_format = "#,##0.00"
                c.alignment = sag
            row += 1

            if stok_satirlari:
                # Alt başlık: Kodu | Ürün Adı (C:F) | Miktar | Birim | Fiyat | Tutar
                ws.cell(row, 2, "Kodu").font = stok_baslik_font
                ws.cell(row, 3, "Ürün Adı").font = stok_baslik_font
                ws.cell(row, 3).alignment = orta
                ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
                ws.cell(row, 7, "Miktar").font = stok_baslik_font
                ws.cell(row, 8, "Birim").font = stok_baslik_font
                ws.cell(row, 9, "Fiyat").font = stok_baslik_font
                ws.cell(row, 10, "Tutar").font = stok_baslik_font
                for col in range(2, 11):
                    ws.cell(row, col).fill = stok_baslik_fill
                row += 1

                for s in stok_satirlari:
                    kod, urun, miktar, birim, fiyat, tutar = _stok_degerleri(s)
                    ws.cell(row, 2, kod).font = stok_font
                    ws.cell(row, 3, urun).font = stok_font
                    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
                    ws.cell(row, 7, miktar).font = stok_font
                    ws.cell(row, 7).alignment = sag
                    ws.cell(row, 8, birim).font = stok_font
                    # Sayısal fiyat/tutar — string yerine float yazmayı dene
                    try:
                        fiyat_f = float(
                            str(_iskontolu_kdvli_birim_fiyat(s) or 0)
                        )
                        tutar_f = float(str(s.get("genel") or 0))
                    except (TypeError, ValueError):
                        fiyat_f = fiyat
                        tutar_f = tutar
                    c_f = ws.cell(row, 9, fiyat_f if isinstance(fiyat_f, float) else fiyat)
                    c_f.font = stok_font
                    c_f.alignment = sag
                    if isinstance(fiyat_f, float):
                        c_f.number_format = "#,##0.00"
                    c_t = ws.cell(row, 10, tutar_f if isinstance(tutar_f, float) else tutar)
                    c_t.font = stok_font
                    c_t.alignment = sag
                    if isinstance(tutar_f, float):
                        c_t.number_format = "#,##0.00"
                    row += 1

        try:
            wb.save(yol)
        except PermissionError:
            messagebox.showerror(
                "Excel",
                "Dosya kaydedilemedi. Dosya başka bir programda (ör. Excel) açık olabilir.\n"
                f"Yol: {yol}",
                parent=self,
            )
            return
        except OSError as hata:
            messagebox.showerror(
                "Excel",
                f"Dosya kaydedilemedi:\n{hata}",
                parent=self,
            )
            return

        messagebox.showinfo(
            "Excel",
            f"Stok detaylı ekstre kaydedildi.\n{Path(yol)}",
            parent=self,
        )
