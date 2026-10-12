"""Satın Alma → Raporlar ekranları (ortak filtre, tablo, toplam, Excel ve yazdırma)."""

from __future__ import annotations

import html
import os
import tempfile
import tkinter as tk
import webbrowser
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from tkinter import filedialog, messagebox, ttk
from typing import Any, Callable, Iterable

from satis_tema import ACIK_BG, BEYAZ, IKINCIL, LACIVERT, font, tk_buton, treeview_stil

# Kolon tipleri: metin, tarih, para, miktar, fiyat, yuzde, gun, gun1 (tek ondalıklı gün ortalaması),
# oran (iki ondalıklı katsayı)
_SAG = {"para", "miktar", "fiyat", "yuzde", "gun", "gun1", "oran"}
YATAY_KOLON_SINIRI = 12  # bundan geniş tablolar yatay sayfa düzeninde yazdırılır


@dataclass
class Kolon:
    anahtar: str
    baslik: str
    genislik: int = 100
    tip: str = "metin"


@dataclass
class Tablo:
    anahtar: str
    baslik: str
    kolonlar: list[Kolon]
    cift_tik: Callable[[dict], None] | None = None
    etiket: Callable[[dict], str | None] | None = None
    coklu_secim: bool = False


@dataclass
class Filtre:
    tur: str            # tarih | tedarikci | musteri | metin | combo | check | sayi | donemler
    anahtar: str
    etiket: str
    varsayilan: Any = None
    secenekler: list[tuple[str, str]] = field(default_factory=list)


def _sayi(deger, basamak=2) -> str:
    if deger is None or deger == "":
        return ""
    d = Decimal(str(deger))
    metin = f"{d:,.{basamak}f}"
    return metin.replace(",", "X").replace(".", ",").replace("X", ".")


def bicimle(deger, tip: str) -> str:
    if deger is None:
        return ""
    if isinstance(deger, str) and tip in _SAG:
        return deger
    if tip == "tarih":
        return deger.strftime("%d.%m.%Y") if isinstance(deger, (date, datetime)) else str(deger)
    if tip == "para":
        return _sayi(deger, 2)
    if tip == "fiyat":
        return _sayi(deger, 4)
    if tip == "miktar":
        d = Decimal(str(deger))
        return _sayi(d, 0 if d == d.to_integral_value() else 3)
    if tip == "yuzde":
        return f"%{_sayi(deger, 2)}"
    if tip == "gun":
        return str(int(deger))
    if tip == "gun1":
        return _sayi(deger, 1)
    if tip == "oran":
        return _sayi(deger, 2)
    return str(deger)


def _tablo_stil(tablo: ttk.Treeview) -> None:
    """Kurumsal tablo stili; 'vista' temasında başlık zemini boyanamadığından yazı lacivert olur."""
    treeview_stil(tablo)
    try:
        stil = ttk.Style(tablo)
        if stil.theme_use() != "clam":
            stil.configure("SARapor.Satis.Treeview.Heading", foreground=LACIVERT)
            tablo.configure(style="SARapor.Satis.Treeview")
    except tk.TclError:
        pass


def _tarih_oku(metin: str) -> date | None:
    metin = (metin or "").strip()
    if not metin:
        return None
    return datetime.strptime(metin, "%d.%m.%Y").date()


def _yetki(*anahtarlar) -> bool:
    try:
        from database.access import yetki_var

        return bool(yetki_var(*anahtarlar))
    except Exception:  # noqa: BLE001
        return True


def _geri(app):
    from satin_alma_ui import satin_alma_hub_goster

    return lambda: satin_alma_hub_goster(app)


class RaporEkrani:
    """Filtre şeridi + bir veya birden çok tablo + notlar + toplam + Excel/Yazdır."""

    def __init__(self, app, *, baslik: str, alt: str, filtreler: list[Filtre], tablolar: list[Tablo],
                 hesapla: Callable[[dict], dict], toplam_metni: Callable[[dict], str] | None = None,
                 ek_butonlar: list[tuple[str, Callable[["RaporEkrani"], None], str]] | None = None,
                 otomatik: bool = True, ust: Callable[[Any, str, str], tk.Misc] | None = None,
                 hafiza: str | None = None):
        self.app = app
        self.baslik = baslik
        self.tablolar_tanim = tablolar
        self.hesapla = hesapla
        self.toplam_metni = toplam_metni
        self.sonuc: dict = {}
        self.girdiler: dict[str, Any] = {}
        self._tedarikci_harita: dict[str, int] = {}
        self.donem_editoru: DonemEditoru | None = None
        self._hafiza = hafiza
        self.son_kolon: str | None = None
        self._gizli: dict[str, set[str]] = {}
        self._tur: dict[str, str] = {f.anahtar: f.tur for f in filtreler}
        self._etiket: dict[str, str] = {f.anahtar: f.etiket for f in filtreler}
        if ust is None:
            from satin_alma_ui import _liste_ust

            self.kok = _liste_ust(app, baslik, alt, _geri(app))
        else:
            self.kok = ust(app, baslik, alt)

        for f in filtreler:
            if f.tur == "donemler":
                kutu = tk.Frame(self.kok, bg=ACIK_BG)
                kutu.pack(fill="x", padx=16, pady=(8, 0))
                self.donem_editoru = DonemEditoru(kutu, app, f.varsayilan)
                self.donem_editoru.pack(fill="x")
                self.girdiler[f.anahtar] = self.donem_editoru
        filtre_cerceve = tk.Frame(self.kok, bg=ACIK_BG)
        filtre_cerceve.pack(fill="x", padx=16, pady=(8, 2))
        for i, f in enumerate([f for f in filtreler if f.tur != "donemler"]):
            self._filtre_ekle(filtre_cerceve, f, i)
        self._hafiza_yukle()

        butonlar = tk.Frame(self.kok, bg=ACIK_BG)
        butonlar.pack(fill="x", padx=16, pady=(2, 4))
        self.rapor_btn = tk_buton(butonlar, "Raporla", self.raporla, rol="ara")
        self.rapor_btn.pack(side="left", padx=(0, 6))
        tk_buton(butonlar, "Excel'e Aktar", self.excel_aktar, rol="duzenle").pack(side="left", padx=6)
        tk_buton(butonlar, "PDF'e Aktar", self.pdf_aktar, rol="duzenle").pack(side="left", padx=6)
        tk_buton(butonlar, "Yazdır / Önizle", self.yazdir, rol="yazdir").pack(side="left", padx=6)
        tk_buton(butonlar, "Sütunlar", self.kolon_secici, rol="duzenle").pack(side="left", padx=6)
        for metin, komut, rol in ek_butonlar or []:
            tk_buton(butonlar, metin, lambda k=komut: k(self), rol=rol).pack(side="left", padx=6)
        self.durum = tk.Label(butonlar, text="", bg=ACIK_BG, fg=IKINCIL, font=font(9, "normal", app))
        self.durum.pack(side="right")

        self.toplam = tk.Label(self.kok, text="", bg=ACIK_BG, fg=LACIVERT, anchor="w", justify="left",
                               font=font(10, "bold", app))
        self.notlar = tk.Label(self.kok, text="", bg=ACIK_BG, fg=IKINCIL, anchor="w", justify="left",
                               wraplength=1100, font=font(9, "normal", app))
        self.notlar.pack(side="bottom", fill="x", padx=18, pady=(0, 8))
        self.toplam.pack(side="bottom", fill="x", padx=18, pady=(2, 2))

        govde = tk.Frame(self.kok, bg=ACIK_BG)
        govde.pack(fill="both", expand=True, padx=16, pady=4)
        self.tablolar: dict[str, ttk.Treeview] = {}
        self._satirlar: dict[str, list[dict]] = {}
        self._sirala: dict[str, tuple[str | None, bool]] = {}
        if len(tablolar) > 1:
            self.defter = ttk.Notebook(govde)
            self.defter.pack(fill="both", expand=True)
            for t in tablolar:
                sayfa = ttk.Frame(self.defter)
                self.defter.add(sayfa, text=t.baslik)
                self.tablolar[t.anahtar] = self._tablo_kur(sayfa, t)
        else:
            self.defter = None
            self.tablolar[tablolar[0].anahtar] = self._tablo_kur(govde, tablolar[0])
        if otomatik:
            self.kok.after(50, self.raporla)

    # ------------------------------------------------------------------ filtreler
    def _filtre_ekle(self, parent, f: Filtre, i: int) -> None:
        kutu = tk.Frame(parent, bg=ACIK_BG)
        kutu.grid(row=i // 6, column=i % 6, sticky="w", padx=(0, 12), pady=2)
        if f.tur == "check":
            var = tk.BooleanVar(value=bool(f.varsayilan))
            ttk.Checkbutton(kutu, text=f.etiket, variable=var).pack(side="left", pady=(14, 0))
            self.girdiler[f.anahtar] = var
            return
        tk.Label(kutu, text=f.etiket, bg=ACIK_BG, fg=IKINCIL, font=font(9, "normal", self.app)).pack(anchor="w")
        if f.tur == "tarih":
            ent = ttk.Entry(kutu, width=11)
            if f.varsayilan:
                ent.insert(0, f.varsayilan.strftime("%d.%m.%Y"))
            ent.pack(side="left")
            try:
                from ui_takvim import takvim_butonu

                takvim_butonu(kutu, ent, text="▾", width=2).pack(side="left", padx=(2, 0))
            except Exception:  # noqa: BLE001
                pass
            self.girdiler[f.anahtar] = ent
        elif f.tur in ("tedarikci", "musteri"):
            degerler = ["(Tümü)"]
            try:
                from database.cari_service import CariService

                liste = CariService.aktif_tedarikciler() if f.tur == "tedarikci" else CariService.aktif_musteriler()
                for c in liste:
                    ad = f"{c.unvan} ({c.cari_kodu})"
                    self._tedarikci_harita[ad] = int(c.id)
                    degerler.append(ad)
            except Exception:  # noqa: BLE001
                pass
            cb = ttk.Combobox(kutu, values=degerler, width=32)
            cb.set("(Tümü)")
            cb.pack(side="left")
            self.girdiler[f.anahtar] = cb
        elif f.tur == "combo":
            cb = ttk.Combobox(kutu, values=[e for _v, e in f.secenekler], state="readonly", width=24)
            varsayilan = next((e for v, e in f.secenekler if v == f.varsayilan), f.secenekler[0][1])
            cb.set(varsayilan)
            cb._secenekler = f.secenekler
            cb.pack(side="left")
            self.girdiler[f.anahtar] = cb
        elif f.tur == "sayi":
            sp = ttk.Spinbox(kutu, from_=0, to=3650, width=7)
            sp.set(str(f.varsayilan or 0))
            sp.pack(side="left")
            self.girdiler[f.anahtar] = sp
        else:
            ent = ttk.Entry(kutu, width=22)
            if f.varsayilan:
                ent.insert(0, str(f.varsayilan))
            ent.pack(side="left")
            ent.bind("<Return>", lambda _e: self.raporla())
            self.girdiler[f.anahtar] = ent

    def filtre_degerleri(self) -> dict:
        sonuc: dict[str, Any] = {}
        for anahtar, w in self.girdiler.items():
            if isinstance(w, DonemEditoru):
                sonuc[anahtar] = w.deger()
            elif isinstance(w, tk.BooleanVar):
                sonuc[anahtar] = bool(w.get())
            elif isinstance(w, ttk.Combobox) and hasattr(w, "_secenekler"):
                etiket = w.get()
                sonuc[anahtar] = next((v for v, e in w._secenekler if e == etiket), w._secenekler[0][0])
            elif isinstance(w, ttk.Combobox):
                sonuc[anahtar] = self._tedarikci_harita.get(w.get())
            elif isinstance(w, ttk.Spinbox):
                try:
                    sonuc[anahtar] = int(w.get())
                except ValueError as hata:
                    raise ValueError("Gün alanına sayı girin.") from hata
            else:
                metin = w.get().strip()
                if anahtar in ("bas", "bit", "bugun") or self._tur.get(anahtar) == "tarih":
                    try:
                        sonuc[anahtar] = _tarih_oku(metin)
                    except ValueError as hata:
                        raise ValueError("Tarih GG.AA.YYYY biçiminde olmalı.") from hata
                else:
                    sonuc[anahtar] = metin
        return sonuc

    def filtre_ozeti(self) -> str:
        parcalar = []
        for anahtar, w in self.girdiler.items():
            if isinstance(w, tk.BooleanVar):
                if w.get():
                    parcalar.append(self._etiket.get(anahtar, anahtar))
                continue
            deger = w.get()
            if isinstance(w, DonemEditoru):
                parcalar.append(deger)
            elif deger and deger != "(Tümü)":
                parcalar.append(f"{self._etiket.get(anahtar, anahtar)}: {deger}")
        parcalar += [str(b) for b in self.sonuc.get("_bilgi", [])]
        return " | ".join(parcalar)

    # ------------------------------------------------------------------ filtre hafızası
    def _hafiza_yolu(self):
        from fatura_satir_kolon_prefs import ayar_dosyasi

        return ayar_dosyasi(f"rapor_filtre_{self._hafiza}")

    def _hafiza_yukle(self) -> None:
        if not self._hafiza:
            return
        try:
            import json

            yol = self._hafiza_yolu()
            kayit = json.loads(yol.read_text(encoding="utf-8")) if yol.exists() else {}
        except Exception:  # noqa: BLE001
            return
        gizli = kayit.pop("_gizli_kolonlar", None)
        if isinstance(gizli, dict):
            self._gizli = {str(t): set(map(str, v)) for t, v in gizli.items() if isinstance(v, list)}
        for anahtar, deger in kayit.items():
            w = self.girdiler.get(anahtar)
            try:
                if isinstance(w, DonemEditoru):
                    w.hafizadan(deger)
                elif isinstance(w, tk.BooleanVar):
                    w.set(bool(deger))
                elif isinstance(w, (ttk.Combobox, ttk.Spinbox)):
                    if not isinstance(w, ttk.Combobox) or deger in tuple(w.cget("values")):
                        w.set(deger)
                elif isinstance(w, ttk.Entry) and self._tur.get(anahtar) == "metin":
                    w.delete(0, "end")
                    w.insert(0, deger)
            except (tk.TclError, TypeError, ValueError):
                continue

    def _hafiza_kaydet(self) -> None:
        if not self._hafiza:
            return
        kayit: dict[str, Any] = {}
        for anahtar, w in self.girdiler.items():
            if self._tur.get(anahtar) == "tarih":
                continue
            try:
                kayit[anahtar] = w.hafiza() if isinstance(w, DonemEditoru) else w.get()
            except tk.TclError:
                continue
        if any(self._gizli.values()):
            kayit["_gizli_kolonlar"] = {t: sorted(v) for t, v in self._gizli.items() if v}
        try:
            import json

            self._hafiza_yolu().write_text(json.dumps(kayit, ensure_ascii=False, default=str), encoding="utf-8")
        except Exception:  # noqa: BLE001
            pass

    # ------------------------------------------------------------------ tablo
    def _tablo_kur(self, parent, t: Tablo) -> ttk.Treeview:
        cerceve = ttk.Frame(parent)
        cerceve.pack(fill="both", expand=True)
        tablo = ttk.Treeview(cerceve, columns=[k.anahtar for k in t.kolonlar], show="headings",
                             selectmode="extended" if t.coklu_secim else "browse")
        _tablo_stil(tablo)
        for k in t.kolonlar:
            tablo.heading(k.anahtar, text=k.baslik, command=lambda a=k.anahtar, tt=t: self._sirala_tikla(tt, a))
            tablo.column(k.anahtar, width=k.genislik, anchor="e" if k.tip in _SAG else "w", stretch=False)
        tablo.configure(displaycolumns=[k.anahtar for k in self.gorunen_kolonlar(t)])
        dikey = ttk.Scrollbar(cerceve, orient="vertical", command=tablo.yview)
        yatay = ttk.Scrollbar(cerceve, orient="horizontal", command=tablo.xview)
        tablo.configure(yscrollcommand=dikey.set, xscrollcommand=yatay.set)
        tablo.grid(row=0, column=0, sticky="nsew")
        dikey.grid(row=0, column=1, sticky="ns")
        yatay.grid(row=1, column=0, sticky="ew")
        cerceve.rowconfigure(0, weight=1)
        cerceve.columnconfigure(0, weight=1)
        tablo.tag_configure("uyari", foreground="#B42318")
        tablo.tag_configure("iyi", foreground="#067647")
        tablo.tag_configure("soluk", foreground="#667085")
        tablo.tag_configure("toplam", background="#E8ECF4", font=font(9, "bold", self.app))
        if t.cift_tik is not None:
            tablo.bind("<Double-1>", lambda e, tt=t: self._cift_tik(tt, e))
            tablo.bind("<Return>", lambda _e, tt=t: self._cift_tik(tt))
        return tablo

    def kolonlari_ayarla(self, anahtar: str, kolonlar: list[Kolon]) -> None:
        """Sonuca göre değişen kolonlar (ör. seçilen dönem sayısı kadar dönem sütunu)."""
        t = next((x for x in self.tablolar_tanim if x.anahtar == anahtar), None)
        tablo = self.tablolar.get(anahtar)
        if t is None or tablo is None:
            return
        t.kolonlar = list(kolonlar)
        tablo.delete(*tablo.get_children())
        tablo.configure(columns=[k.anahtar for k in t.kolonlar],
                        displaycolumns=[k.anahtar for k in self.gorunen_kolonlar(t)])
        for k in t.kolonlar:
            tablo.heading(k.anahtar, text=k.baslik, command=lambda a=k.anahtar, tt=t: self._sirala_tikla(tt, a))
            tablo.column(k.anahtar, width=k.genislik, anchor="e" if k.tip in _SAG else "w", stretch=False)

    def gorunen_kolonlar(self, t: Tablo) -> list[Kolon]:
        gizli = self._gizli.get(t.anahtar, set())
        gorunen = [k for k in t.kolonlar if k.anahtar not in gizli]
        return gorunen or list(t.kolonlar)

    def kolonlari_gizle(self, anahtar: str, gizli: Iterable[str]) -> None:
        t = next((x for x in self.tablolar_tanim if x.anahtar == anahtar), None)
        tablo = self.tablolar.get(anahtar)
        if t is None or tablo is None:
            return
        self._gizli[anahtar] = set(gizli)
        tablo.configure(displaycolumns=[k.anahtar for k in self.gorunen_kolonlar(t)])
        self._hafiza_kaydet()

    def _etkin_tablo(self) -> Tablo:
        if self.defter is not None:
            try:
                return self.tablolar_tanim[self.defter.index(self.defter.select())]
            except (tk.TclError, IndexError):
                pass
        return self.tablolar_tanim[0]

    def kolon_secici(self) -> tk.Toplevel:
        """Etkin tablonun sütunlarını göster/gizle. Genişlik başlık kenarından sürüklenerek değiştirilir."""
        t = self._etkin_tablo()
        pencere = tk.Toplevel(self.app)
        pencere.title(f"Sütunlar — {t.baslik}")
        pencere.configure(bg=ACIK_BG)
        pencere.transient(self.app)
        gizli = self._gizli.get(t.anahtar, set())
        cerceve = tk.Frame(pencere, bg=ACIK_BG)
        cerceve.pack(fill="both", expand=True, padx=10, pady=8)
        degiskenler: dict[str, tk.BooleanVar] = {}
        sutun_sayisi = max(1, (len(t.kolonlar) + 17) // 18)
        for i, k in enumerate(t.kolonlar):
            var = tk.BooleanVar(value=k.anahtar not in gizli)
            degiskenler[k.anahtar] = var
            ttk.Checkbutton(cerceve, text=k.baslik or k.anahtar, variable=var).grid(
                row=i % 18 if sutun_sayisi > 1 else i, column=i // 18 if sutun_sayisi > 1 else 0,
                sticky="w", padx=4)

        def uygula():
            self.kolonlari_gizle(t.anahtar, [a for a, v in degiskenler.items() if not v.get()])
            pencere.destroy()

        def hepsi(deger: bool):
            for v in degiskenler.values():
                v.set(deger)

        alt = tk.Frame(pencere, bg=ACIK_BG)
        alt.pack(fill="x", padx=10, pady=(0, 8))
        tk_buton(alt, "Tümünü Seç", lambda: hepsi(True), rol="duzenle").pack(side="left")
        tk_buton(alt, "Tümünü Kaldır", lambda: hepsi(False), rol="duzenle").pack(side="left", padx=6)
        tk_buton(alt, "Uygula", uygula, rol="ara").pack(side="right")
        pencere.uygula = uygula
        pencere.degiskenler = degiskenler
        return pencere

    def _cift_tik(self, t: Tablo, olay=None) -> None:
        """Çift tıklanan sütun ``son_kolon``da tutulur (hücre bazlı ayrıntı için); Enter'da None olur."""
        self.son_kolon = None
        if olay is not None:
            tablo = self.tablolar[t.anahtar]
            no = tablo.identify_column(olay.x)
            try:
                gorunen = tablo["displaycolumns"]
                if gorunen in ("#all", ("#all",)):
                    gorunen = tablo["columns"]
                self.son_kolon = gorunen[int(no.lstrip("#")) - 1]
            except (ValueError, IndexError, tk.TclError):
                self.son_kolon = None
        satir = self.secili(t.anahtar)
        if satir and t.cift_tik is not None:
            t.cift_tik(satir[0])

    def secili(self, anahtar: str | None = None) -> list[dict]:
        anahtar = anahtar or self.tablolar_tanim[0].anahtar
        tablo = self.tablolar[anahtar]
        satirlar = self._satirlar.get(anahtar, [])
        return [satirlar[int(i)] for i in tablo.selection() if int(i) < len(satirlar)]

    def _sirala_tikla(self, t: Tablo, kolon: str) -> None:
        onceki, azalan = self._sirala.get(t.anahtar, (None, False))
        azalan = (not azalan) if onceki == kolon else False
        self._sirala[t.anahtar] = (kolon, azalan)
        satirlar = self._satirlar.get(t.anahtar, [])
        sabit = [r for r in satirlar if r.get("_sabit")]
        satirlar = [r for r in satirlar if not r.get("_sabit")]
        tip = next((k.tip for k in t.kolonlar if k.anahtar == kolon), "metin")

        def bos_mu(v) -> bool:
            return v in (None, "") or (tip in _SAG and isinstance(v, str))

        dolu = [r for r in satirlar if not bos_mu(r.get(kolon))]
        bos = [r for r in satirlar if bos_mu(r.get(kolon))]
        try:
            dolu.sort(key=lambda r: r[kolon], reverse=azalan)
        except TypeError:
            dolu.sort(key=lambda r: str(r[kolon]), reverse=azalan)
        self._doldur(t, dolu + bos + sabit)

    def _doldur(self, t: Tablo, satirlar: list[dict]) -> None:
        tablo = self.tablolar[t.anahtar]
        self._satirlar[t.anahtar] = satirlar
        tablo.delete(*tablo.get_children())
        for i, r in enumerate(satirlar):
            etiket = t.etiket(r) if t.etiket else None
            tablo.insert("", "end", iid=str(i), tags=(etiket,) if etiket else (),
                         values=[bicimle(r.get(k.anahtar), k.tip) for k in t.kolonlar])

    # ------------------------------------------------------------------ çalıştırma
    def raporla(self, senkron: bool = False) -> None:
        try:
            filtre = self.filtre_degerleri()
        except ValueError as hata:
            messagebox.showwarning(self.baslik, str(hata), parent=self.app)
            return
        self._hafiza_kaydet()
        if senkron:
            self.goster(self.hesapla(filtre))
            return
        from stok_rapor_ui import arka_plan_rapor

        arka_plan_rapor(self.app, self.durum, _ButonSarmal(self.rapor_btn),
                        lambda: self.hesapla(filtre), self.goster, self.baslik)

    def goster(self, sonuc: dict) -> None:
        self.sonuc = sonuc or {}
        if self.sonuc.get("_baslik"):
            self.baslik = str(self.sonuc["_baslik"])
        toplam_satir = 0
        for t, kolonlar in (self.sonuc.get("_kolonlar") or {}).items():
            self.kolonlari_ayarla(t, kolonlar)
        for t in self.tablolar_tanim:
            satirlar = list(self.sonuc.get(t.anahtar) or [])
            toplam_satir += len(satirlar)
            self._doldur(t, satirlar)
            if self.defter is not None:
                self.defter.tab(self.tablolar_tanim.index(t), text=f"{t.baslik} ({len(satirlar)})")
        self.durum.configure(text=f"{toplam_satir} satır")
        self.toplam.configure(text=self.toplam_metni(self.sonuc) if self.toplam_metni else "")
        self.notlar.configure(text="\n".join(f"• {n}" for n in self.sonuc.get("notlar", [])))

    # ------------------------------------------------------------------ dışa aktarma
    def _aktarim_tablolari(self):
        for t in self.tablolar_tanim:
            yield t, self._satirlar.get(t.anahtar, [])

    def excel_aktar(self) -> None:
        if not _yetki("excel_pdf"):
            messagebox.showwarning(self.baslik, "Excel/PDF aktarma yetkiniz yok.", parent=self.app)
            return
        if not any(s for _t, s in self._aktarim_tablolari()):
            messagebox.showinfo(self.baslik, "Aktarılacak satır yok.", parent=self.app)
            return
        yol = filedialog.asksaveasfilename(
            parent=self.app, defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"{self.baslik.title().replace(' ', '_')}_{date.today():%Y%m%d}.xlsx")
        if not yol:
            return
        try:
            self.excel_yaz(yol)
        except Exception as hata:  # noqa: BLE001
            messagebox.showerror(self.baslik, f"Excel dosyası yazılamadı:\n{hata}", parent=self.app)
            return
        messagebox.showinfo(self.baslik, f"Excel dosyası kaydedildi:\n{yol}", parent=self.app)

    def excel_yaz(self, yol: str) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        wb.remove(wb.active)
        kullanilan: set[str] = set()
        for t, satirlar in self._aktarim_tablolari():
            ad = "".join("-" if c in '\\/*?:[]' else c for c in t.baslik).strip("' ")[:31] or "Sayfa"
            temel, n = ad, 2
            while ad.casefold() in kullanilan:
                ad = f"{temel[:28]} {n}"
                n += 1
            kullanilan.add(ad.casefold())
            kolonlar = self.gorunen_kolonlar(t)
            ws = wb.create_sheet(ad)
            ws.append([self.baslik])
            ws["A1"].font = Font(bold=True, size=13)
            ws.append([self.filtre_ozeti()])
            ws.append([])
            ws.append([k.baslik for k in kolonlar])
            for hucre in ws[4]:
                hucre.font = Font(bold=True)
            for r in satirlar:
                satir = []
                for k in kolonlar:
                    v = r.get(k.anahtar)
                    if isinstance(v, Decimal):
                        v = float(v)
                    elif k.tip == "metin" and v is not None and not isinstance(v, (int, float, date)):
                        v = str(v)
                    satir.append(v)
                ws.append(satir)
            for i, k in enumerate(kolonlar, 1):
                ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = max(10, k.genislik // 7)
            if len(kolonlar) > YATAY_KOLON_SINIRI:
                ws.page_setup.orientation = "landscape"
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                ws.sheet_properties.pageSetUpPr.fitToPage = True
                ws.print_title_rows = "4:4"
            ws.append([])
            if self.toplam_metni:
                ws.append([self.toplam_metni(self.sonuc)])
            for n in self.sonuc.get("notlar", []):
                ws.append([n])
        wb.save(yol)

    def pdf_aktar(self) -> None:
        if not _yetki("excel_pdf"):
            messagebox.showwarning(self.baslik, "Excel/PDF aktarma yetkiniz yok.", parent=self.app)
            return
        if not any(s for _t, s in self._aktarim_tablolari()):
            messagebox.showinfo(self.baslik, "Aktarılacak satır yok.", parent=self.app)
            return
        yol = filedialog.asksaveasfilename(
            parent=self.app, defaultextension=".pdf", filetypes=[("PDF", "*.pdf")],
            initialfile=f"{self.baslik.title().replace(' ', '_').replace('—', '-')}_{date.today():%Y%m%d}.pdf")
        if not yol:
            return
        try:
            self.pdf_yaz(yol)
        except Exception as hata:  # noqa: BLE001
            messagebox.showerror(self.baslik, f"PDF dosyası yazılamadı:\n{hata}", parent=self.app)
            return
        messagebox.showinfo(self.baslik, f"PDF dosyası kaydedildi:\n{yol}", parent=self.app)

    def pdf_yaz(self, yol: str) -> None:
        from pathlib import Path

        from invoice_print.pdf_service import html_metnini_pdfe_cevir

        html_metnini_pdfe_cevir(self.html_metni(otomatik_yazdir=False), Path(yol), belge_adi=self.baslik)

    def html_metni(self, otomatik_yazdir: bool = True) -> str:
        genis = any(len(self.gorunen_kolonlar(t)) > YATAY_KOLON_SINIRI for t in self.tablolar_tanim)
        sayfa = ("@page{size:A4 landscape;margin:8mm}body{font-size:9px}td,th{padding:2px 3px}" if genis else "")
        parca = [f"<html><head><meta charset='utf-8'><title>{html.escape(self.baslik)}</title>",
                 "<style>body{font-family:Segoe UI,Arial;font-size:11px}table{border-collapse:collapse;"
                 "margin-bottom:14px}td,th{border:1px solid #999;padding:3px 5px}th{background:#E8ECF4}"
                 f".s{{text-align:right}}tr.t td{{font-weight:bold;background:#E8ECF4}}{sayfa}</style></head>"
                 "<body onload='window.print()'>" if otomatik_yazdir else "<body>",
                 f"<h2>{html.escape(self.baslik)}</h2><p>{html.escape(self.filtre_ozeti())} — "
                 f"{datetime.now():%d.%m.%Y %H:%M}</p>"]
        for t, satirlar in self._aktarim_tablolari():
            if len(self.tablolar_tanim) > 1:
                parca.append(f"<h3>{html.escape(t.baslik)}</h3>")
            kolonlar = self.gorunen_kolonlar(t)
            parca.append("<table><tr>" + "".join(f"<th>{html.escape(k.baslik)}</th>" for k in kolonlar) + "</tr>")
            for r in satirlar:
                parca.append(("<tr class='t'>" if r.get("_sabit") else "<tr>") + "".join(
                    f"<td class='{'s' if k.tip in _SAG else ''}'>{html.escape(bicimle(r.get(k.anahtar), k.tip))}</td>"
                    for k in kolonlar) + "</tr>")
            parca.append("</table>")
        if self.toplam_metni:
            parca.append(f"<p><b>{html.escape(self.toplam_metni(self.sonuc))}</b></p>")
        parca += [f"<p>• {html.escape(n)}</p>" for n in self.sonuc.get("notlar", [])]
        parca.append("</body></html>")
        return "".join(parca)

    def yazdir(self) -> None:
        if not any(s for _t, s in self._aktarim_tablolari()):
            messagebox.showinfo(self.baslik, "Yazdırılacak satır yok.", parent=self.app)
            return
        yol = os.path.join(tempfile.gettempdir(), f"cin_rapor_{datetime.now():%Y%m%d_%H%M%S}.html")
        with open(yol, "w", encoding="utf-8") as dosya:
            dosya.write(self.html_metni())
        webbrowser.open(f"file:///{yol}")


class DonemEditoru(tk.Frame):
    """Sınırsız bağımsız dönem: hazır seçim veya özel tarih, temel dönem seçimi, ekle/kaldır.

    Değer: ``{"donemler": [(ad, bas, bit), ...], "temel": indeks}``. İlk satır ana dönemdir.
    """

    def __init__(self, parent, app, varsayilan=None):
        super().__init__(parent, bg=ACIK_BG)
        from database.satis_yonetici_rapor_service import DONEM_SECENEKLERI

        self.app = app
        self._secenekler = DONEM_SECENEKLERI
        self._temel = tk.IntVar(value=1)
        self._satirlar: list[dict[str, Any]] = []
        ust = tk.Frame(self, bg=ACIK_BG)
        ust.pack(fill="x")
        tk.Label(ust, text="Dönemler (her dönem ayrı hesaplanır; değişimler 'Temel' seçili döneme göredir)",
                 bg=ACIK_BG, fg=IKINCIL, font=font(9, "normal", app)).pack(side="left")
        tk_buton(ust, "+ Dönem Ekle", lambda: self.ekle("gecen_yil_ayni"), rol="yeni").pack(side="left", padx=8)
        self._govde = tk.Frame(self, bg=ACIK_BG)
        self._govde.pack(fill="x")
        for kod in (varsayilan or ("bu_ay", "gecen_ay")):
            self.ekle(kod)
        self._temel.set(1 if len(self._satirlar) > 1 else 0)

    def ekle(self, kod: str = "ozel", bas: date | None = None, bit: date | None = None,
             ad: str | None = None) -> None:
        satir = tk.Frame(self._govde, bg=ACIK_BG)
        satir.pack(fill="x", pady=1)
        cb = ttk.Combobox(satir, values=[e for _k, e in self._secenekler], state="readonly", width=24)
        cb.pack(side="left")
        bas_e, bit_e = ttk.Entry(satir, width=11), ttk.Entry(satir, width=11)
        tk.Label(satir, text=" ", bg=ACIK_BG).pack(side="left")
        bas_e.pack(side="left")
        tk.Label(satir, text="–", bg=ACIK_BG).pack(side="left")
        bit_e.pack(side="left")
        kayit = {"cerceve": satir, "cb": cb, "bas": bas_e, "bit": bit_e, "ad": ad}
        ttk.Radiobutton(satir, text="Temel", variable=self._temel, value=len(self._satirlar)).pack(
            side="left", padx=8)
        tk.Button(satir, text="✕", relief="flat", bg=ACIK_BG, command=lambda k=kayit: self.kaldir(k)).pack(
            side="left")
        self._satirlar.append(kayit)
        cb.bind("<<ComboboxSelected>>", lambda _e, k=kayit: (k.update(ad=None), self._hazir_uygula(k)))
        for e in (bas_e, bit_e):
            e.bind("<KeyRelease>", lambda _e, k=kayit: k.update(ad=None), add="+")
        etiket = dict(self._secenekler).get(kod, "Özel")
        cb.set(etiket)
        if bas and bit:
            self._tarih_yaz(kayit, bas, bit)
        else:
            self._hazir_uygula(kayit)

    def kaldir(self, kayit) -> None:
        if len(self._satirlar) <= 1:
            return
        i = self._satirlar.index(kayit)
        kayit["cerceve"].destroy()
        self._satirlar.pop(i)
        for j, k in enumerate(self._satirlar):
            for w in k["cerceve"].winfo_children():
                if isinstance(w, ttk.Radiobutton):
                    w.configure(value=j)
        if self._temel.get() >= len(self._satirlar):
            self._temel.set(len(self._satirlar) - 1)

    def _kod(self, kayit) -> str:
        etiket = kayit["cb"].get()
        return next((k for k, e in self._secenekler if e == etiket), "ozel")

    @staticmethod
    def _tarih_yaz(kayit, bas: date, bit: date) -> None:
        for alan, t in (("bas", bas), ("bit", bit)):
            kayit[alan].delete(0, "end")
            kayit[alan].insert(0, t.strftime("%d.%m.%Y"))

    def _hazir_uygula(self, kayit) -> None:
        from database.satis_yonetici_rapor_service import donem_araligi

        kod = self._kod(kayit)
        if kod == "ozel":
            return
        referans = None
        if kod == "gecen_yil_ayni" and self._satirlar and self._satirlar[0] is not kayit:
            try:
                ilk = self._satirlar[0]
                referans = (_tarih_oku(ilk["bas"].get()), _tarih_oku(ilk["bit"].get()))
                if None in referans:
                    referans = None
            except ValueError:
                referans = None
        self._tarih_yaz(kayit, *donem_araligi(kod, referans=referans))

    def deger(self) -> dict[str, Any]:
        donemler = []
        adlar: set[str] = set()
        for i, k in enumerate(self._satirlar):
            try:
                bas, bit = _tarih_oku(k["bas"].get()), _tarih_oku(k["bit"].get())
            except ValueError as hata:
                raise ValueError("Dönem tarihleri GG.AA.YYYY biçiminde olmalı.") from hata
            if bas is None or bit is None:
                raise ValueError(f"{i + 1}. dönemin başlangıç ve bitiş tarihini girin.")
            if bit < bas:
                raise ValueError(f"{i + 1}. dönemde bitiş başlangıçtan önce olamaz.")
            kod = self._kod(k)
            ad = k.get("ad") or (k["cb"].get() if kod != "ozel" else f"{bas:%d.%m.%y}–{bit:%d.%m.%y}")
            if ad in adlar:
                ad = f"{ad} ({i + 1})"
            adlar.add(ad)
            donemler.append((ad, bas, bit))
        return {"donemler": donemler, "temel": min(self._temel.get(), len(donemler) - 1)}

    def get(self) -> str:
        try:
            d = self.deger()
        except ValueError:
            return ""
        return "Dönemler: " + "; ".join(
            f"{ad} {bas:%d.%m.%Y}–{bit:%d.%m.%Y}" + (" (temel)" if i == d["temel"] else "")
            for i, (ad, bas, bit) in enumerate(d["donemler"]))

    def ayarla(self, donemler, temel: int = 0) -> None:
        for k in list(self._satirlar):
            k["cerceve"].destroy()
        self._satirlar.clear()
        for ad, bas, bit in donemler:
            self.ekle("ozel", bas, bit, ad)
        self._temel.set(min(int(temel), len(self._satirlar) - 1))

    def hafiza(self) -> dict[str, Any]:
        return {"kodlar": [self._kod(k) for k in self._satirlar], "temel": self._temel.get(),
                "ozel": [[k["bas"].get(), k["bit"].get()] for k in self._satirlar]}

    def hafizadan(self, kayit: dict) -> None:
        kodlar = list(kayit.get("kodlar") or [])
        if not kodlar:
            return
        ozel = kayit.get("ozel") or []
        for k in list(self._satirlar):
            k["cerceve"].destroy()
        self._satirlar.clear()
        for i, kod in enumerate(kodlar):
            if kod == "ozel" and i < len(ozel):
                try:
                    bas, bit = _tarih_oku(ozel[i][0]), _tarih_oku(ozel[i][1])
                except ValueError:
                    bas = bit = None
                self.ekle("ozel", bas, bit)
            else:
                self.ekle(kod)
        self._temel.set(min(int(kayit.get("temel") or 0), len(self._satirlar) - 1))


class _ButonSarmal:
    """``arka_plan_rapor`` ttk ``state`` bekler; tk düğmesi için uyarlayıcı."""

    def __init__(self, buton):
        self.buton = buton

    def state(self, durum):
        try:
            self.buton.configure(state="disabled" if "disabled" in durum else "normal")
        except tk.TclError:
            pass


def _yil_basi() -> date:
    return date(date.today().year, 1, 1)


def _belge_listede(app, liste_adi: str, belge_id):
    from satin_alma_ui import belgeden_listeye_git

    liste_fn = getattr(app, liste_adi, None)
    if callable(liste_fn) and belge_id:
        belgeden_listeye_git(app, liste_fn, belge_id)


def _ekran(app, baslik, kur: Callable[[], RaporEkrani]) -> RaporEkrani | None:
    if not _yetki("alis_rapor_goruntuleme", "alis_goruntuleme"):
        messagebox.showwarning(baslik, "Satın alma raporlarını görüntüleme yetkiniz yok.", parent=app)
        return None
    ekran = kur()
    if hasattr(app, "nav_sayfa_isaretle"):
        app.nav_sayfa_isaretle(lambda: _ekran(app, baslik, kur))
    return ekran


# ---------------------------------------------------------------------------
# 1) Alış fiyatı değişim raporu
# ---------------------------------------------------------------------------

def fiyat_degisim_raporu_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def hesapla(f):
        return srv.fiyat_degisim(bas=f["bas"], bit=f["bit"], kapsam=f["kapsam"], cari_id=f["cari_id"],
                                 arama=f["arama"], sadece_degisen=f["sadece_degisen"])

    def etiket(r):
        if r.get("fark") is None:
            return "soluk"
        return "uyari" if r["fark"] > 0 else ("iyi" if r["fark"] < 0 else None)

    def kur():
        return RaporEkrani(
            app, baslik="ALIŞ FİYATI DEĞİŞİM RAPORU",
            alt="Ürünlerin önceki ve son gerçekleşmiş alış fiyatı (iskontolu net, KDV hariç)",
            filtreler=[Filtre("tarih", "bas", "Son alış başlangıç", _yil_basi()),
                       Filtre("tarih", "bit", "Bitiş", date.today()),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"),
                       Filtre("combo", "kapsam", "Karşılaştırma", "tedarikci",
                              [("tedarikci", "Aynı tedarikçi içinde"), ("tumu", "Tüm tedarikçiler")]),
                       Filtre("metin", "arama", "Ürün / tedarikçi ara"),
                       Filtre("check", "sadece_degisen", "Yalnız fiyatı değişenler", False)],
            tablolar=[Tablo("satirlar", "Fiyat Değişimi", [
                Kolon("urun_kodu", "Ürün Kodu", 100), Kolon("urun_adi", "Ürün Adı", 200),
                Kolon("tedarikci", "Son Tedarikçi", 160), Kolon("onceki_tedarikci", "Önceki Tedarikçi", 150),
                Kolon("onceki_tarih", "Önceki Tarih", 85, "tarih"), Kolon("onceki_belge", "Önceki Belge", 105),
                Kolon("son_tarih", "Son Tarih", 85, "tarih"), Kolon("son_belge", "Son Belge", 105),
                Kolon("birim", "Birim", 60), Kolon("para_birimi", "PB", 45),
                Kolon("onceki_fiyat", "Önceki Net Fiyat", 110, "fiyat"), Kolon("son_fiyat", "Son Net Fiyat", 110, "fiyat"),
                Kolon("fark", "Fark", 90, "fiyat"), Kolon("yuzde", "Değişim %", 80, "yuzde"),
                Kolon("aciklama", "Açıklama", 260)], etiket=etiket)],
            hesapla=hesapla,
            toplam_metni=lambda s: (
                f"Artan: {sum(1 for r in s.get('satirlar', []) if (r.get('fark') or 0) > 0)}   "
                f"Düşen: {sum(1 for r in s.get('satirlar', []) if (r.get('fark') or 0) < 0)}   "
                f"Karşılaştırılamayan: {sum(1 for r in s.get('satirlar', []) if r.get('fark') is None)}"),
            **kw)
    return _ekran(app, "Alış Fiyatı Değişim Raporu", kur)


# ---------------------------------------------------------------------------
# 2) Tedarikçi fiyat karşılaştırması
# ---------------------------------------------------------------------------

def tedarikci_fiyat_karsilastirma_goster(app, urun_kodu: str = "", **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def hesapla(f):
        kod = (f.get("urun_kodu") or "").strip()
        if not kod:
            return {"notlar": ["Karşılaştırma için ürün kodu girin."]}
        return srv.tedarikci_karsilastirma(kod, bas=f["bas"], bit=f["bit"], teklif_dahil=f["teklif_dahil"],
                                           cevir=f["cevir"], hedef_birim=f.get("hedef_birim") or None)

    fiyat_kolonlari = [Kolon("birim", "Birim", 60), Kolon("para_birimi", "PB", 45),
                       Kolon("net_fiyat", "Net Fiyat", 100, "fiyat"),
                       Kolon("cevrilmis_fiyat", "Çevrilmiş (TL)", 105, "fiyat"),
                       Kolon("oran", "Birim Oranı", 80, "fiyat"), Kolon("kur", "Kur", 80, "fiyat"),
                       Kolon("kur_tarihi", "Kur Tarihi", 85, "tarih")]

    def kur():
        return RaporEkrani(
            app, baslik="TEDARİKÇİ FİYAT KARŞILAŞTIRMASI",
            alt="Seçilen ürünün tedarikçilerden gerçekleşmiş alış fiyatları; teklifler ayrı sekmede",
            filtreler=[Filtre("metin", "urun_kodu", "Ürün kodu", urun_kodu),
                       Filtre("tarih", "bas", "Başlangıç"), Filtre("tarih", "bit", "Bitiş"),
                       Filtre("metin", "hedef_birim", "Çevrim birimi (boş: temel)"),
                       Filtre("check", "cevir", "TL ve tek birime çevir", False),
                       Filtre("check", "teklif_dahil", "Güncel teklifleri göster", False)],
            tablolar=[
                Tablo("ozet", "Tedarikçi Son Fiyatları", [
                    Kolon("sira", "Sıra", 45, "gun"), Kolon("tedarikci", "Tedarikçi", 180),
                    Kolon("tarih", "Son Alış", 85, "tarih"), Kolon("belge_no", "Belge", 105),
                    Kolon("miktar", "Miktar", 70, "miktar"), *fiyat_kolonlari,
                    Kolon("gun", "Fiyat Yaşı (gün)", 95, "gun"), Kolon("alis_sayisi", "Alış Sayısı", 75, "gun"),
                    Kolon("cevrim_notu", "Çevrim", 220)],
                    cift_tik=lambda r: _belge_listede(app, "alis_faturalari_goster", r.get("belge_id"))),
                Tablo("gecmis", "Alış Geçmişi", [
                    Kolon("tarih", "Tarih", 85, "tarih"), Kolon("tedarikci", "Tedarikçi", 180),
                    Kolon("belge_no", "Belge", 105), Kolon("miktar", "Miktar", 70, "miktar"),
                    *fiyat_kolonlari, Kolon("gun", "Gün", 55, "gun")],
                    cift_tik=lambda r: _belge_listede(app, "alis_faturalari_goster", r.get("belge_id"))),
                Tablo("teklifler", "Güncel Teklifler (gerçekleşmemiş)", [
                    Kolon("tedarikci", "Tedarikçi", 180), Kolon("tarih", "Teklif Tarihi", 85, "tarih"),
                    Kolon("belge_no", "Teklif No", 105), Kolon("durum", "Durum", 120),
                    Kolon("miktar", "Miktar", 70, "miktar"), Kolon("birim", "Birim", 60),
                    Kolon("para_birimi", "PB", 45), Kolon("net_fiyat", "Teklif Net Fiyat", 110, "fiyat"),
                    Kolon("vade_gun", "Vade", 55, "gun"), Kolon("termin_gun", "Termin", 55, "gun"),
                    Kolon("nakliye", "Nakliye", 80, "para"), Kolon("gun", "Yaş (gün)", 70, "gun")]),
            ],
            hesapla=hesapla,
            toplam_metni=lambda s: (f"{s.get('urun_kodu', '')} {s.get('urun_adi', '')}   "
                                    + ("Sıralama: ucuzdan pahalıya" if s.get("siralanabilir") else
                                       "Sıralama yapılmadı (birim/para birimi farklı)" if s.get("ozet") else "")),
            **{"otomatik": bool(urun_kodu), **kw})
    return _ekran(app, "Tedarikçi Fiyat Karşılaştırması", kur)


# ---------------------------------------------------------------------------
# 3) Bekleyen alış siparişleri
# ---------------------------------------------------------------------------

def bekleyen_alis_siparisleri_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="BEKLEYEN ALIŞ SİPARİŞLERİ",
            alt="Teslim alınmamış sipariş kalanları, gecikme ve kalan tutar",
            filtreler=[Filtre("tarih", "bas", "Sipariş başlangıç"), Filtre("tarih", "bit", "Bitiş"),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Ürün / sipariş ara"),
                       Filtre("check", "sadece_geciken", "Yalnız gecikenler", False)],
            tablolar=[Tablo("satirlar", "Bekleyen Siparişler", [
                Kolon("tedarikci", "Tedarikçi", 170), Kolon("siparis_tarihi", "Sipariş Tarihi", 90, "tarih"),
                Kolon("siparis_no", "Sipariş No", 105), Kolon("urun_kodu", "Ürün Kodu", 95),
                Kolon("urun_adi", "Ürün Adı", 190), Kolon("birim", "Birim", 55),
                Kolon("siparis_miktari", "Sipariş", 70, "miktar"), Kolon("gelen", "Gelen", 70, "miktar"),
                Kolon("kalan", "Kalan", 70, "miktar"), Kolon("termin", "Beklenen Teslim", 95, "tarih"),
                Kolon("gecikme", "Gecikme (gün)", 90, "gun"), Kolon("kalan_tutar", "Kalan Tutar", 100, "para"),
                Kolon("aciklama", "Açıklama", 170)],
                cift_tik=lambda r: _belge_listede(app, "alis_siparisleri_goster", r.get("siparis_id")),
                etiket=lambda r: "uyari" if (r.get("gecikme") or 0) > 0 else None)],
            hesapla=lambda f: srv.bekleyen_siparisler(bas=f["bas"], bit=f["bit"], cari_id=f["cari_id"],
                                                      arama=f["arama"], sadece_geciken=f["sadece_geciken"]),
            toplam_metni=lambda s: (f"Kalan tutar (KDV hariç): {bicimle(s.get('toplam_kalan_tutar'), 'para')} TL   "
                                    f"Geciken satır: {sum(1 for r in s.get('satirlar', []) if (r.get('gecikme') or 0) > 0)}"),
            **kw)
    return _ekran(app, "Bekleyen Alış Siparişleri", kur)


# ---------------------------------------------------------------------------
# 4) Faturası gelmeyen alış irsaliyeleri
# ---------------------------------------------------------------------------

def faturasi_gelmeyen_irsaliyeler_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="FATURASI GELMEYEN ALIŞ İRSALİYELERİ",
            alt="Mal kabulü yapılmış, faturası gelmemiş veya kısmi faturalanmış irsaliyeler",
            filtreler=[Filtre("tarih", "bas", "İrsaliye başlangıç"), Filtre("tarih", "bit", "Bitiş"),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Ürün / irsaliye ara")],
            tablolar=[Tablo("satirlar", "Faturası Gelmeyenler", [
                Kolon("tedarikci", "Tedarikçi", 170), Kolon("irsaliye_tarihi", "İrsaliye Tarihi", 90, "tarih"),
                Kolon("irsaliye_no", "İrsaliye No", 105), Kolon("urun_kodu", "Ürün Kodu", 95),
                Kolon("urun_adi", "Ürün Adı", 190), Kolon("birim", "Birim", 55),
                Kolon("teslim", "Teslim Alınan", 85, "miktar"), Kolon("faturalanan", "Faturalanan", 80, "miktar"),
                Kolon("faturalanmamis", "Faturalanmamış", 90, "miktar"),
                Kolon("bekleme_gun", "Bekleme (gün)", 85, "gun"),
                Kolon("tahmini_tutar", "Tahmini Tutar", 100, "para"), Kolon("aciklama", "Açıklama", 190)],
                cift_tik=lambda r: _belge_listede(app, "alis_irsaliyeleri_goster", r.get("irsaliye_id")),
                etiket=lambda r: "soluk" if r.get("tahmini_tutar") is None else None)],
            hesapla=lambda f: srv.faturasi_gelmeyen_irsaliyeler(bas=f["bas"], bit=f["bit"], cari_id=f["cari_id"],
                                                                arama=f["arama"]),
            toplam_metni=lambda s: (f"Tahmini toplam (KDV hariç, fiyatı olan satırlar): "
                                    f"{bicimle(s.get('toplam_tahmini'), 'para')} TL   "
                                    f"Fiyatsız satır: {s.get('fiyatsiz_satir', 0)}"),
            **kw)
    return _ekran(app, "Faturası Gelmeyen Alış İrsaliyeleri", kur)


# ---------------------------------------------------------------------------
# 5) Yeniden sipariş önerisi
# ---------------------------------------------------------------------------

def _taslak_siparis(ekran: RaporEkrani) -> None:
    from database import satin_alma_rapor_service as srv

    if not _yetki("alis_siparis_duzenleme"):
        messagebox.showwarning(ekran.baslik, "Alış siparişi oluşturma yetkiniz yok.", parent=ekran.app)
        return
    secili = [r for r in ekran.secili() if (r.get("oneri") or 0) > 0]
    if not secili:
        messagebox.showinfo(ekran.baslik, "Önerilen miktarı olan satırları seçin.", parent=ekran.app)
        return
    if not messagebox.askyesno(
            ekran.baslik, f"{len(secili)} satırdan TASLAK alış siparişi hazırlansın mı?\n"
            "Sipariş onaylanmaz; tedarikçiye gönderilmeden önce kontrol edebilirsiniz.", parent=ekran.app):
        return
    try:
        sonuc = srv.taslak_siparis_olustur(secili)
    except Exception as hata:  # noqa: BLE001
        messagebox.showwarning(ekran.baslik, str(hata), parent=ekran.app)
        return
    nolar = ", ".join(str(s.get("siparis_no") or s.get("no") or s.get("id")) for s in sonuc or [])
    messagebox.showinfo(ekran.baslik, f"Taslak sipariş hazırlandı: {nolar}", parent=ekran.app)
    ekran.raporla()


def yeniden_siparis_onerisi_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="YENİDEN SİPARİŞ ÖNERİSİ",
            alt="Azalan ve tükenen ürünler; stok, bekleyen alış ve satış hızına göre öneri (otomatik sipariş yok)",
            filtreler=[Filtre("sayi", "donem_gun", "Satış dönemi (gün)", 90),
                       Filtre("sayi", "hedef_gun", "Hedef stok (gün)", 30),
                       Filtre("sayi", "tedarik_gun", "Tedarik süresi varsayımı (gün)", 0),
                       Filtre("metin", "arama", "Ürün / tedarikçi ara"),
                       Filtre("check", "musteri_siparisi_dus", "Müşteri siparişlerini düş", True)],
            tablolar=[Tablo("satirlar", "Öneriler", [
                Kolon("stok_kodu", "Ürün Kodu", 95), Kolon("stok_adi", "Ürün Adı", 190),
                Kolon("depo", "Depo", 90), Kolon("temel_birim", "Birim", 55),
                Kolon("fiziksel", "Mevcut", 70, "miktar"), Kolon("musteri_talebi", "Müşteri Sip.", 80, "miktar"),
                Kolon("kullanilabilir", "Kullanılabilir", 85, "miktar"), Kolon("minimum", "Minimum", 70, "miktar"),
                Kolon("donem_satis", "Dönem Satışı", 85, "miktar"), Kolon("gunluk_hiz", "Günlük Hız", 75, "fiyat"),
                Kolon("bekleyen_alis", "Bekleyen Alış", 85, "miktar"), Kolon("taslak_siparis", "Taslak Sip.", 75, "miktar"),
                Kolon("hedef", "Hedef Stok", 75, "miktar"), Kolon("oneri", "Öneri", 65, "miktar"),
                Kolon("alim_birimi", "Alım Birimi", 75), Kolon("son_alis_tarihi", "Son Alış", 85, "tarih"),
                Kolon("son_alis_fiyati", "Son Fiyat", 90, "fiyat"), Kolon("son_alis_pb", "PB", 45),
                Kolon("son_tedarikci", "Son Tedarikçi", 160), Kolon("aciklama", "Açıklama", 260)],
                coklu_secim=True,
                etiket=lambda r: "soluk" if r.get("veri_yetersiz") else ("uyari" if (r.get("oneri") or 0) > 0 else None))],
            hesapla=lambda f: srv.yeniden_siparis_onerisi(donem_gun=f["donem_gun"], hedef_gun=f["hedef_gun"],
                                                          musteri_siparisi_dus=f["musteri_siparisi_dus"],
                                                          arama=f["arama"], tedarik_gun=f.get("tedarik_gun") or 0),
            toplam_metni=lambda s: (f"Öneri olan ürün: {sum(1 for r in s.get('satirlar', []) if (r.get('oneri') or 0) > 0)}   "
                                    f"Veri yetersiz: {sum(1 for r in s.get('satirlar', []) if r.get('veri_yetersiz'))}"),
            ek_butonlar=[("Seçilenlerden Taslak Sipariş", _taslak_siparis, "yeni")],
            **kw)
    return _ekran(app, "Yeniden Sipariş Önerisi", kur)


# ---------------------------------------------------------------------------
# 6) Tedarikçi borç ve vade raporu
# ---------------------------------------------------------------------------

def tedarikci_borc_detay_goster(app, cari_id: int) -> tk.Toplevel:
    from database import satin_alma_rapor_service as srv
    from ui_pencere import popup_ortala

    veri = srv.tedarikci_borc_detay(int(cari_id))
    pen = tk.Toplevel(app)
    pen.title(f"Borç ve Vade Detayı — {veri['tedarikci']}")
    pen.configure(bg=BEYAZ)
    pen.geometry("1100x620")
    popup_ortala(pen, app, genislik=1100, yukseklik=620)
    tk.Label(pen, text=f"{veri['tedarikci']} — Bakiye: {bicimle(veri['bakiye'], 'para')} TL",
             bg=BEYAZ, fg=LACIVERT, font=font(12, "bold", app)).pack(anchor="w", padx=12, pady=(10, 4))
    defter = ttk.Notebook(pen)
    defter.pack(fill="both", expand=True, padx=10, pady=6)
    bolumler = (
        ("Açık Kalemler", veri["kalemler"], [("belge_no", "Belge", 110, "metin"), ("tur", "Tür", 220, "metin"),
                                             ("tarih", "Tarih", 85, "tarih"), ("vade", "Vade", 85, "tarih"),
                                             ("gun", "Gün", 55, "gun"), ("durum", "Durum", 110, "metin"),
                                             ("tutar", "Açık Tutar", 110, "para"),
                                             ("kapanis", "Kapanış", 170, "metin")]),
        ("Ödeme Sözleri (borcu azaltmaz)", veri["sozler"], [
            ("tarih", "Söz Tarihi", 85, "tarih"), ("vade", "Ödeme Tarihi", 90, "tarih"),
            ("tutar", "Tutar", 100, "para"), ("kalan", "Kalan", 100, "para"),
            ("para_birimi", "PB", 45, "metin"), ("yontem", "Yöntem", 110, "metin"), ("durum", "Durum", 110, "metin")]),
        ("Bakiye Farkı Açıklaması", veri["aciklama"], [("tur", "Hareket Türü", 320, "metin"),
                                                      ("borc", "Borç", 110, "para"), ("alacak", "Alacak", 110, "para")]),
        ("Avans / Fazla Ödeme", veri["avanslar"], [("belge_no", "Belge", 120, "metin"),
                                                  ("tarih", "Tarih", 85, "tarih"), ("tutar", "Tutar", 110, "para")]),
    )
    for baslik, satirlar, kolonlar in bolumler:
        sayfa = ttk.Frame(defter)
        defter.add(sayfa, text=f"{baslik} ({len(satirlar)})")
        tablo = ttk.Treeview(sayfa, columns=[k[0] for k in kolonlar], show="headings")
        _tablo_stil(tablo)
        for a, b, w, tip in kolonlar:
            tablo.heading(a, text=b)
            tablo.column(a, width=w, anchor="e" if tip in _SAG else "w")
        for r in satirlar:
            tablo.insert("", "end", values=[bicimle(r.get(a), tip) for a, _b, _w, tip in kolonlar])
        tablo.pack(fill="both", expand=True)
    tk_buton(pen, "Kapat", pen.destroy, rol="geri").pack(anchor="e", padx=12, pady=(0, 10))
    return pen


def tedarikci_borc_vade_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="TEDARİKÇİ BORÇ VE VADE RAPORU",
            alt="Toplam, kapanan, açık, vadesi geçmiş, bugün ve ileri vadeli borçlar; ödeme sözleri ayrı",
            filtreler=[Filtre("tarih", "bugun", "Rapor tarihi", date.today()),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Tedarikçi ara"),
                       Filtre("check", "sadece_bakiyeli", "Yalnız bakiyesi olanlar", True)],
            tablolar=[Tablo("satirlar", "Tedarikçiler", [
                Kolon("cari_kodu", "Kod", 80), Kolon("tedarikci", "Tedarikçi", 200),
                Kolon("toplam_borc", "Toplam Borç", 105, "para"), Kolon("kapanan", "Kapanan", 100, "para"),
                Kolon("acik", "Açık", 100, "para"), Kolon("vadesi_gecmis", "Vadesi Geçmiş", 105, "para"),
                Kolon("bugun", "Bugün", 90, "para"), Kolon("ileri_vade", "İleri Vade", 100, "para"),
                Kolon("avans", "Avans", 90, "para"), Kolon("bakiye", "Cari Bakiye", 105, "para"),
                Kolon("soz_tutari", "Ödeme Sözü", 100, "para"), Kolon("en_yakin_soz", "Söz Tarihi", 85, "tarih"),
                Kolon("aciklama", "Açıklama", 300)],
                cift_tik=lambda r: tedarikci_borc_detay_goster(app, r["cari_id"]),
                etiket=lambda r: "uyari" if (r.get("vadesi_gecmis") or 0) > 0 else None)],
            hesapla=lambda f: srv.tedarikci_borc_vade(cari_id=f["cari_id"], arama=f["arama"], bugun=f["bugun"],
                                                      sadece_bakiyeli=f["sadece_bakiyeli"]),
            toplam_metni=lambda s: "   ".join(
                f"{b}: {bicimle(sum((r[a] for r in s.get('satirlar', [])), Decimal('0')), 'para')}"
                for a, b in (("acik", "Açık"), ("vadesi_gecmis", "Vadesi geçmiş"), ("bugun", "Bugün"),
                             ("ileri_vade", "İleri"), ("soz_tutari", "Ödeme sözü"))) + "   (detay için çift tıklayın)",
            **kw)
    return _ekran(app, "Tedarikçi Borç ve Vade Raporu", kur)


# ---------------------------------------------------------------------------
# 7) Ürün bazında gerçek alış maliyeti
# ---------------------------------------------------------------------------

def gercek_alis_maliyeti_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    try:
        from database.access import maliyet_izinli

        if not maliyet_izinli():
            messagebox.showwarning("Gerçek Alış Maliyeti", "Maliyet görme yetkiniz yok.", parent=app)
            return None
    except ImportError:
        pass

    def kur():
        return RaporEkrani(
            app, baslik="ÜRÜN BAZINDA GERÇEK ALIŞ MALİYETİ",
            alt="İskontolu net alış + dağıtılmış nakliye ve diğer masraflar; toplam ve birim maliyet",
            filtreler=[Filtre("tarih", "bas", "Başlangıç", _yil_basi()), Filtre("tarih", "bit", "Bitiş", date.today()),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Ürün ara"),
                       Filtre("combo", "gruplama", "Gruplama", "urun",
                              [("urun", "Ürün"), ("tedarikci", "Ürün + tedarikçi"), ("satir", "Fatura satırı")])],
            tablolar=[Tablo("satirlar", "Maliyet", [
                Kolon("urun_kodu", "Ürün Kodu", 95), Kolon("urun_adi", "Ürün Adı", 190),
                Kolon("tedarikci", "Tedarikçi", 150), Kolon("belge_no", "Belge", 100), Kolon("tarih", "Tarih", 85, "tarih"),
                Kolon("temel_birim", "Birim", 55), Kolon("miktar", "Miktar", 75, "miktar"),
                Kolon("net_alis", "Net Alış", 105, "para"), Kolon("nakliye", "Nakliye", 90, "para"),
                Kolon("diger", "Diğer Masraf", 95, "para"), Kolon("toplam_maliyet", "Toplam Maliyet", 110, "para"),
                Kolon("birim_maliyet", "Birim Maliyet", 100, "fiyat"), Kolon("durum", "Durum", 280)],
                etiket=lambda r: "uyari" if r.get("eksik") else None)],
            hesapla=lambda f: srv.gercek_alis_maliyeti(bas=f["bas"], bit=f["bit"], cari_id=f["cari_id"],
                                                       arama=f["arama"], gruplama=f["gruplama"]),
            toplam_metni=lambda s: "   ".join(
                f"{b}: {bicimle(sum((r[a] for r in s.get('satirlar', [])), Decimal('0')), 'para')}"
                for a, b in (("net_alis", "Net alış"), ("nakliye", "Nakliye"), ("diger", "Diğer"),
                             ("toplam_maliyet", "Toplam"))),
            **kw)
    return _ekran(app, "Gerçek Alış Maliyeti", kur)


# ---------------------------------------------------------------------------
# 8) Alış ve iade özeti
# ---------------------------------------------------------------------------

def alis_iade_ozeti_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="ALIŞ VE İADE ÖZETİ",
            alt="Dönem alışları ve alış iadeleri; KDV hariç, KDV ve KDV dahil ayrı",
            filtreler=[Filtre("tarih", "bas", "Başlangıç", _yil_basi()), Filtre("tarih", "bit", "Bitiş", date.today()),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Ürün / tedarikçi ara"),
                       Filtre("combo", "gruplama", "Gruplama", "urun",
                              [("urun", "Ürün"), ("grup", "Ürün grubu"), ("tedarikci", "Tedarikçi")])],
            tablolar=[Tablo("satirlar", "Özet", [
                Kolon("anahtar", "Ürün / Grup / Tedarikçi", 170), Kolon("ad", "Ad", 180), Kolon("temel_birim", "Birim", 55),
                Kolon("alis_miktar", "Alış Miktar", 80, "miktar"), Kolon("alis_haric", "Alış KDV Hariç", 105, "para"),
                Kolon("alis_kdv", "Alış KDV", 90, "para"), Kolon("alis_dahil", "Alış KDV Dahil", 105, "para"),
                Kolon("iade_miktar", "İade Miktar", 80, "miktar"), Kolon("iade_haric", "İade KDV Hariç", 105, "para"),
                Kolon("iade_kdv", "İade KDV", 90, "para"), Kolon("iade_dahil", "İade KDV Dahil", 105, "para"),
                Kolon("net_miktar", "Net Miktar", 80, "miktar"), Kolon("net_haric", "Net KDV Hariç", 105, "para"),
                Kolon("net_kdv", "Net KDV", 90, "para"), Kolon("net_dahil", "Net KDV Dahil", 105, "para")])],
            hesapla=lambda f: srv.alis_iade_ozeti(bas=f["bas"], bit=f["bit"], cari_id=f["cari_id"],
                                                  arama=f["arama"], gruplama=f["gruplama"]),
            toplam_metni=lambda s: "   ".join(
                f"{b}: {bicimle((s.get('toplam') or {}).get(a), 'para')}"
                for a, b in (("alis_haric", "Alış hariç"), ("iade_haric", "İade hariç"), ("net_haric", "Net hariç"),
                             ("net_kdv", "Net KDV"), ("net_dahil", "Net dahil"))),
            **kw)
    return _ekran(app, "Alış ve İade Özeti", kur)


def _teslim_belgesi_ac(app, tur: str | None, belge_id) -> None:
    if tur == "Fatura":
        _belge_listede(app, "alis_faturalari_goster", belge_id)
    elif tur == "İrsaliye":
        _belge_listede(app, "alis_irsaliyeleri_goster", belge_id)
    else:
        _belge_listede(app, "alis_siparisleri_goster", belge_id)


# ---------------------------------------------------------------------------
# 9) Tedarikçi teslimat performansı
# ---------------------------------------------------------------------------

def tedarikci_teslimat_performansi_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="TEDARİKÇİ TESLİMAT PERFORMANSI",
            alt="Beklenen teslim tarihi ile gerçekleşen teslimatlar; kısmi teslimatlar ayrı (çift tık: belgeye git)",
            filtreler=[Filtre("tarih", "bas", "Sipariş başlangıç", _yil_basi()),
                       Filtre("tarih", "bit", "Bitiş", date.today()),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Ürün / sipariş ara")],
            tablolar=[
                Tablo("tedarikciler", "Tedarikçi Özeti", [
                    Kolon("tedarikci", "Tedarikçi", 200), Kolon("siparis_satiri", "Sipariş Satırı", 85, "gun"),
                    Kolon("degerlendirilen", "Değerlendirilen", 95, "gun"), Kolon("zamaninda", "Zamanında", 75, "gun"),
                    Kolon("gecikmeli", "Gecikmeli", 70, "gun"), Kolon("zamaninda_oran", "Zamanında %", 85, "yuzde"),
                    Kolon("miktar_oran", "Miktar Bazında Zamanında %", 150, "yuzde"),
                    Kolon("ort_gecikme", "Ort. Gecikme (gün)", 110, "gun1"),
                    Kolon("en_cok_gecikme", "En Çok Gecikme", 95, "gun"),
                    Kolon("acik_geciken", "Açık Geciken", 85, "gun"), Kolon("kismi_satir", "Kısmi Teslim Satırı", 115, "gun"),
                    Kolon("teslimat_sayisi", "Teslimat", 65, "gun"),
                    Kolon("ort_tedarik_suresi", "Ort. Tedarik Süresi (gün)", 140, "gun1"),
                    Kolon("tarihsiz", "Teslim Tarihi Yok", 105, "gun"),
                    Kolon("suresi_dolmadi", "Süresi Dolmadı", 95, "gun")],
                    etiket=lambda r: "uyari" if (r.get("gecikmeli") or 0) > 0 else None),
                Tablo("satirlar", "Sipariş Satırları", [
                    Kolon("tedarikci", "Tedarikçi", 170), Kolon("siparis_no", "Sipariş No", 105),
                    Kolon("siparis_tarihi", "Sipariş Tarihi", 90, "tarih"), Kolon("urun_kodu", "Ürün Kodu", 95),
                    Kolon("urun_adi", "Ürün Adı", 180), Kolon("birim", "Birim", 55),
                    Kolon("siparis_miktari", "Sipariş", 70, "miktar"), Kolon("teslim", "Teslim Edilen", 90, "miktar"),
                    Kolon("kalan", "Kalan", 70, "miktar"), Kolon("termin", "Beklenen Teslim", 95, "tarih"),
                    Kolon("teslim_tarihleri", "Teslim Tarihleri", 180), Kolon("tam_teslim", "Tamamlanma", 90, "tarih"),
                    Kolon("teslimat_sayisi", "Teslimat", 65, "gun"), Kolon("gecikme", "Gecikme (gün)", 90, "gun"),
                    Kolon("durum", "Durum", 190)],
                    cift_tik=lambda r: _belge_listede(app, "alis_siparisleri_goster", r.get("siparis_id")),
                    etiket=lambda r: "uyari" if (r.get("gecikme") or 0) > 0 else (
                        "soluk" if r.get("termin") is None else None)),
                Tablo("teslimatlar", "Teslimatlar (kısmi ayrı)", [
                    Kolon("tedarikci", "Tedarikçi", 170), Kolon("siparis_no", "Sipariş No", 105),
                    Kolon("urun_kodu", "Ürün Kodu", 95), Kolon("urun_adi", "Ürün Adı", 170),
                    Kolon("teslimat_sira", "Sıra", 45, "gun"), Kolon("teslimat_turu", "Belge Türü", 75),
                    Kolon("belge_no", "Belge No", 105), Kolon("teslim_tarihi", "Teslim Tarihi", 90, "tarih"),
                    Kolon("termin", "Beklenen", 85, "tarih"), Kolon("miktar", "Miktar", 70, "miktar"),
                    Kolon("kumulatif", "Kümülatif", 75, "miktar"), Kolon("birim", "Birim", 55),
                    Kolon("gecikme", "Gecikme (gün)", 90, "gun"), Kolon("durum", "Durum", 110),
                    Kolon("kismi", "Kısmi", 60)],
                    cift_tik=lambda r: _teslim_belgesi_ac(app, r.get("teslimat_turu"), r.get("belge_id")),
                    etiket=lambda r: "uyari" if (r.get("gecikme") or 0) > 0 else None)],
            hesapla=lambda f: srv.tedarikci_teslimat_performansi(bas=f["bas"], bit=f["bit"], cari_id=f["cari_id"],
                                                                 arama=f["arama"]),
            toplam_metni=lambda s: (
                f"Değerlendirilen satır: {sum(r['degerlendirilen'] for r in s.get('tedarikciler', []))}   "
                f"Zamanında: {sum(r['zamaninda'] for r in s.get('tedarikciler', []))}   "
                f"Gecikmeli: {sum(r['gecikmeli'] for r in s.get('tedarikciler', []))}   "
                f"Teslim tarihi olmayan: {sum(r['tarihsiz'] for r in s.get('tedarikciler', []))}"),
            hafiza="sa_teslimat_performansi", **kw)
    return _ekran(app, "Tedarikçi Teslimat Performansı", kur)


# ---------------------------------------------------------------------------
# 10) Satın alma talepleri takibi
# ---------------------------------------------------------------------------

def _talep_ac(app, talep_id) -> None:
    if not talep_id:
        return
    from satin_alma_talep_ui import talep_formu_ac

    talep_formu_ac(app, int(talep_id))


def satin_alma_talep_takibi_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def baglanti_ac(r):
        if r.get("seviye") == "Sipariş":
            _belge_listede(app, "alis_siparisleri_goster", r.get("belge_id"))
        else:
            _teslim_belgesi_ac(app, r.get("_tur"), r.get("belge_id"))

    def kur():
        return RaporEkrani(
            app, baslik="SATIN ALMA TALEPLERİ TAKİBİ",
            alt="Talep → sipariş → teslimat; onaylanan, siparişe dönüşen, karşılanan ve kalan (çift tık: kaynak evrak)",
            filtreler=[Filtre("tarih", "bas", "Talep başlangıç", _yil_basi()),
                       Filtre("tarih", "bit", "Bitiş", date.today()),
                       Filtre("metin", "arama", "Talep / ürün / talep eden ara"),
                       Filtre("combo", "gorunum", "Görünüm", "tumu", [
                           ("tumu", "Tümü (iptal hariç)"), ("acik", "Açık ihtiyacı olanlar"),
                           ("onay_bekleyen", "Onay bekleyenler"), ("tamamlanan", "Tamamlananlar")])],
            tablolar=[
                Tablo("satirlar", "Talep Satırları", [
                    Kolon("talep_no", "Talep No", 105), Kolon("talep_tarihi", "Talep Tarihi", 85, "tarih"),
                    Kolon("isteyen", "Talep Eden", 130), Kolon("departman", "Departman", 100),
                    Kolon("urun_kodu", "Ürün Kodu", 95), Kolon("urun_adi", "Ürün Adı", 170),
                    Kolon("birim", "Birim", 55), Kolon("talep", "Talep", 65, "miktar"),
                    Kolon("onayli", "Onaylanan", 75, "miktar"), Kolon("siparis", "Siparişe Dönüşen", 105, "miktar"),
                    Kolon("teslim", "Teslim", 65, "miktar"), Kolon("iade", "İade", 55, "miktar"),
                    Kolon("karsilanan", "Karşılanan (net)", 100, "miktar"), Kolon("iptal", "İptal", 55, "miktar"),
                    Kolon("kalan", "Kalan (siparişe aktarılmamış)", 170, "miktar"),
                    Kolon("bekleyen_teslim", "Teslim Bekleyen", 100, "miktar"),
                    Kolon("ihtiyac_tarihi", "İhtiyaç Tarihi", 90, "tarih"), Kolon("gecikme", "Gecikme (gün)", 85, "gun"),
                    Kolon("onay_gun", "Onay Süresi (gün)", 100, "gun"), Kolon("siparisler", "Siparişler", 150),
                    Kolon("durum", "Durum", 180)],
                    cift_tik=lambda r: _talep_ac(app, r.get("talep_id")),
                    etiket=lambda r: "uyari" if r.get("gecikme") else None),
                Tablo("baglantilar", "Sipariş ve Teslimat Bağlantıları", [
                    Kolon("talep_no", "Talep No", 105), Kolon("urun_kodu", "Ürün Kodu", 95),
                    Kolon("urun_adi", "Ürün Adı", 170), Kolon("seviye", "Bağlantı", 95),
                    Kolon("belge_turu", "Belge Türü", 110), Kolon("belge_no", "Belge No", 110),
                    Kolon("tarih_turu", "Tarih Türü", 70), Kolon("tarih", "Tarih", 85, "tarih"),
                    Kolon("miktar", "Miktar", 70, "miktar"), Kolon("birim", "Birim", 55),
                    Kolon("teslim", "Teslim Payı", 80, "miktar"), Kolon("iade", "İade", 55, "miktar"),
                    Kolon("durum", "Durum", 140)],
                    cift_tik=baglanti_ac)],
            hesapla=lambda f: srv.satin_alma_talep_takibi(bas=f["bas"], bit=f["bit"], arama=f["arama"],
                                                          gorunum=f["gorunum"]),
            toplam_metni=lambda s: (
                f"Talep satırı: {len(s.get('satirlar', []))}   "
                f"Geciken ihtiyaç: {sum(1 for r in s.get('satirlar', []) if r.get('gecikme'))}   "
                f"Siparişe aktarılmamış satır: {sum(1 for r in s.get('satirlar', []) if (r.get('kalan') or 0) > 0)}"),
            hafiza="sa_talep_takibi", **kw)

    if not (_yetki("alis_talep_goruntuleme", "alis_talep_duzenleme", "alis_talep_onay", "alis_rapor_goruntuleme")):
        messagebox.showwarning("Satın Alma Talepleri Takibi", "Satın alma talep takibi için yetkiniz yok.", parent=app)
        return None
    return _ekran(app, "Satın Alma Talepleri Takibi", kur)


# ---------------------------------------------------------------------------
# 11) Açık alış faturaları ve valör (gün)
# ---------------------------------------------------------------------------

def _valor_ozet_kolonlari(ilk: list[Kolon]) -> list[Kolon]:
    from database.satin_alma_rapor_service import YAS_GRUPLARI

    k = ilk + [Kolon("para_birimi", "PB", 45), Kolon("kalem_sayisi", "Kalem", 50, "gun"),
               Kolon("acik", "Toplam Açık", 110, "para"), Kolon("acik_tl", "Açık (TL karşılığı)", 120, "para"),
               Kolon("gecmis_tutar", "Vadesi Geçmiş", 105, "para"),
               Kolon("gecmis_gun", "Ağırlıklı Gecikme (gün)", 135, "gun1"),
               Kolon("bugun_tutar", "Bugün Vadeli", 95, "para"),
               Kolon("gelmemis_tutar", "Vadesi Gelmemiş", 110, "para"),
               Kolon("gelmemis_gun", "Ağırlıklı Kalan Gün", 120, "gun1"),
               Kolon("valor", "Genel Ağırlıklı Valör (gün)", 150, "gun1"),
               Kolon("vadesiz_tutar", "Vade Tanımlanmamış", 120, "para")]
    k += [Kolon(f"g_{alt}", f"Geçmiş {ad}", 90, "para") for alt, _u, ad in YAS_GRUPLARI]
    k += [Kolon(f"i_{alt}", f"Gelecek {ad}", 90, "para") for alt, _u, ad in YAS_GRUPLARI]
    return k


def acik_alis_faturalari_valor_goster(app, **kw) -> RaporEkrani | None:
    from database import satin_alma_rapor_service as srv

    def kur():
        return RaporEkrani(
            app, baslik="AÇIK ALIŞ FATURALARI VE VALÖR (GÜN)",
            alt="Rapor tarihindeki açık alış faturaları, vade/valör günleri ve tedarikçi bazında ağırlıklı valör",
            filtreler=[Filtre("tarih", "rapor_tarihi", "Rapor tarihi", date.today()),
                       Filtre("tedarikci", "cari_id", "Tedarikçi"), Filtre("metin", "arama", "Tedarikçi ara"),
                       Filtre("combo", "durum", "Vade durumu", "tumu", [
                           ("tumu", "Tümü"), ("gecmis", "Vadesi geçmiş"), ("bugun", "Bugün vadeli"),
                           ("gelmemis", "Vadesi gelmemiş"), ("vadesiz", "Vade tanımlanmamış")]),
                       Filtre("combo", "para_birimi", "Para birimi", "", [
                           ("", "(Tümü)"), ("TRY", "TRY"), ("USD", "USD"), ("EUR", "EUR"), ("GBP", "GBP")])],
            tablolar=[
                Tablo("tedarikciler", "Tedarikçi Özeti",
                      _valor_ozet_kolonlari([Kolon("cari_kodu", "Kod", 80), Kolon("tedarikci", "Tedarikçi", 190)]),
                      cift_tik=lambda r: tedarikci_borc_detay_goster(app, r["cari_id"]),
                      etiket=lambda r: "uyari" if (r.get("gecmis_tutar") or 0) > 0 else None),
                Tablo("kalemler", "Fatura Ayrıntısı", [
                    Kolon("tedarikci", "Tedarikçi", 170), Kolon("fatura_no", "Fatura No", 105),
                    Kolon("tedarikci_fatura_no", "Tedarikçi Fat. No", 110), Kolon("belge_turu", "Tür", 110),
                    Kolon("fatura_tarihi", "Fatura Tarihi", 85, "tarih"), Kolon("vade_tarihi", "Vade Tarihi", 85, "tarih"),
                    Kolon("para_birimi", "PB", 45), Kolon("kur", "Belge Kuru", 75, "fiyat"),
                    Kolon("fatura_tutari", "Fatura Tutarı (TL)", 110, "para"),
                    Kolon("kapatilan", "Kapatılan (TL)", 100, "para"), Kolon("acik", "Kalan Açık (TL)", 105, "para"),
                    Kolon("fatura_tutari_doviz", "Fatura Tutarı (döviz)", 120, "para"),
                    Kolon("acik_doviz", "Kalan Açık (döviz)", 110, "para"),
                    Kolon("fatura_yasi", "Fatura Yaşı", 75, "gun"), Kolon("vade_gunu", "Anlaşılan Vade Günü", 120, "gun"),
                    Kolon("valor", "Valör Günü", 75, "gun"), Kolon("durum", "Durum", 130),
                    Kolon("yas_grubu", "Gün Grubu", 70), Kolon("kapanis", "Kapanış", 190)],
                    cift_tik=lambda r: _belge_listede(app, "alis_faturalari_goster", r.get("fatura_id")),
                    etiket=lambda r: "uyari" if r.get("durum") == "Vadesi geçmiş" else (
                        "soluk" if r.get("durum") == "Vade tanımlanmamış" else None)),
                Tablo("toplamlar", "Para Birimi Toplamları",
                      _valor_ozet_kolonlari([Kolon("tedarikci", "", 70)])),
                Tablo("avanslar", "Avans / Fazla Ödeme", [
                    Kolon("tedarikci", "Tedarikçi", 190), Kolon("belge_no", "Belge No", 110),
                    Kolon("tarih", "Tarih", 85, "tarih"), Kolon("tutar", "Tutar (TL)", 105, "para"),
                    Kolon("aciklama", "Açıklama", 320)]),
                Tablo("negatifler", "Negatif Hareketler", [
                    Kolon("tedarikci", "Tedarikçi", 190), Kolon("belge_no", "Belge No", 110),
                    Kolon("tarih", "Tarih", 85, "tarih"), Kolon("tutar", "Tutar (TL)", 105, "para"),
                    Kolon("aciklama", "Açıklama", 320)])],
            hesapla=lambda f: srv.acik_alis_faturalari_valor(
                rapor_tarihi=f["rapor_tarihi"], cari_id=f["cari_id"], arama=f["arama"], durum=f["durum"],
                para_birimi=f["para_birimi"]),
            toplam_metni=lambda s: "   ".join(
                f"{r['para_birimi']}: açık {bicimle(r['acik'], 'para')}, geçmiş {bicimle(r['gecmis_tutar'], 'para')}, "
                f"valör {bicimle(r['valor'], 'gun1') or '—'} gün" for r in s.get("toplamlar", [])),
            hafiza="sa_acik_fatura_valor", **kw)
    return _ekran(app, "Açık Alış Faturaları ve Valör", kur)
