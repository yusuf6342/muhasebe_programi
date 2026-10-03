"""Alış iade faturası ekranları ve çıktıları (Tk kökü gizli; geçici veritabanı).

Boş açılış, ayrı liste menüsü, kaydet-tekrar aç, küçük ekranda butonlar, kaynaksız ürün uyarı
diyaloğu, barkodla ekleme, satır silme, kaydedilmemiş kapatma uyarısı, Word/PDF çok sayfalı çıktı.
Ekran görüntüleri: C:\\CinMuhasebeBuild4\\ekran\\alis_iade_*.png (klasör varsa).
"""

from __future__ import annotations

import os
import sys
import time
import tkinter as tk
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from sqlalchemy import func, select  # noqa: E402

from database.alis_iade_faturasi_service import (  # noqa: E402
    DEVIR,
    KAYNAK_YOK_MESAJI,
    KAYNAKSIZ,
    AlisIadeFaturasiService as S,
)
from database.database import get_session  # noqa: E402
from database.models.alis_iade_faturasi import AlisIadeFaturasi  # noqa: E402
from database.models.stok import StokHareketi, StokKarti, StokLotu  # noqa: E402
from tests.stok_test_ortami import StokOrtami  # noqa: E402

D = Decimal
T_ALIS = date(2026, 3, 1)


def _ekran_dizini() -> Path | None:
    p = Path(os.environ.get("MUHASEBE_EKRAN_DIZINI") or r"C:\CinMuhasebeBuild4\ekran")
    return p if p.parent.exists() else None


class _UITaban(unittest.TestCase):
    def setUp(self):
        self.o = StokOrtami(self._testMethodName[:18])
        try:
            self.root = tk.Tk()
        except tk.TclError as hata:
            self.o.kapat()
            self.skipTest(f"Tk yok: {hata}")
        self.root.withdraw()
        self.mesajlar: list[str] = []
        self._mb = patch.multiple(
            "tkinter.messagebox",
            showinfo=lambda _b, m="", **_k: self.mesajlar.append(m),
            showwarning=lambda _b, m="", **_k: self.mesajlar.append(m),
            showerror=lambda _b, m="", **_k: self.mesajlar.append(f"HATA: {m}"),
            askyesno=lambda *a, **k: True)
        self._mb.start()

    def tearDown(self):
        self._mb.stop()
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        self.o.kapat()

    # ---- yardımcılar
    def alis(self, miktar=10, fiyat=100, lot="L1"):
        from database.alis_faturasi_service import AlisFaturasiService

        f = self.o.alis(miktar, fiyat, T_ALIS, lot=lot)
        return int(f.id), int(AlisFaturasiService.getir(f.id).satirlar[0].id)

    def pencere(self, **kw):
        from alis_iade_ui import AlisIadeFaturasiPenceresi

        p = AlisIadeFaturasiPenceresi(self.root, **kw)
        p.withdraw()
        return p

    def tedarikci_sec(self, p):
        anahtar = next(k for k, c in p._ted_map.items() if int(c.id) == self.o.ted_id)
        p.tedarikci.set(anahtar)
        p._tedarikci_degisti()

    def say(self, model) -> int:
        with get_session() as s:
            return int(s.scalar(select(func.count()).select_from(model)) or 0)

    def ekran(self, ad: str, w) -> None:
        dizin = _ekran_dizini()
        if dizin is None:
            return
        try:
            from PIL import ImageGrab

            w.deiconify()
            w.lift()
            w.attributes("-topmost", True)
            for _ in range(6):
                w.update()
                time.sleep(0.08)
            dizin.mkdir(parents=True, exist_ok=True)
            x, y = w.winfo_rootx(), w.winfo_rooty()
            ImageGrab.grab(bbox=(x, y, x + w.winfo_width(), y + w.winfo_height())).save(dizin / f"alis_iade_{ad}.png")
            w.attributes("-topmost", False)
        except Exception as hata:  # noqa: BLE001 - görüntü alınamazsa test sonucu değişmez
            print(f"ekran görüntüsü alınamadı ({ad}): {hata}")


class _SahteKaynakSec:
    """KaynakSecDialog yerine: ilk adaydan istenen temel miktarı seçer."""

    miktar = D("3")

    def __init__(self, _parent, _kod, adaylar, varsayilan=None, _birim="Adet"):
        self.result = [(adaylar[0], varsayilan or self.miktar)]

    def winfo_exists(self):
        return False


class MenuVeBosAcilisTest(_UITaban):
    def test_menu_iki_ayri_kart_ve_her_tiklama_bos_taslak(self):
        import satin_alma_ui
        from app import MuhasebeApp as App

        basliklar = [b for b, _a, _k in satin_alma_ui.SATIN_ALMA_HUB_KARTLARI]
        self.assertIn("ALIŞ İADE FATURASI", basliklar)
        self.assertIn("ALIŞ İADE FATURALARI LİSTESİ", basliklar)
        sahte = MagicMock()
        komutlar = satin_alma_ui._hub_komutlar(sahte)
        self.assertIs(komutlar["iade_yeni"], sahte.alis_iade_faturasi_ac)
        self.assertIs(komutlar["iade"], sahte.alis_iade_faturalari_goster)

        hareket0, belge0 = self.say(StokHareketi), self.say(AlisIadeFaturasi)
        p1 = App.alis_iade_faturasi_ac(self.root)
        p2 = App.alis_iade_faturasi_ac(self.root)
        for p in (p1, p2):
            p.withdraw()
            self.assertIsNone(p.iade)
            self.assertEqual(p.satirlar, [])
            self.assertEqual(p.tedarikci.get(), "")
            self.assertEqual(str(p.butonlar["onayla"]["state"]), "normal")
            self.assertEqual(str(p.butonlar["iptal"]["state"]), "disabled")
        self.assertIsNot(p1, p2)
        self.assertEqual((self.say(StokHareketi), self.say(AlisIadeFaturasi)), (hareket0, belge0))
        p1.deiconify()
        p1.geometry("1280x760+10+10")
        self.ekran("bos_acilis", p1)
        p1.destroy()
        p2.destroy()


class KaydetTekrarAcVeListeTest(_UITaban):
    def test_kaydet_belge_acik_kalir_tekrar_acilir_liste_ayri(self):
        from alis_iade_ui import AlisIadeListesiPenceresi

        _fid, ks = self.alis()
        p = self.pencere()
        self.tedarikci_sec(p)
        with patch("alis_iade_ui.KaynakSecDialog", _SahteKaynakSec), patch.object(p, "wait_window"):
            p.urun_ekle("U001")
        self.assertEqual(len(p.satirlar), 1)
        self.assertEqual(p.satirlar[0]["kaynak_fatura_satiri_id"], ks)
        self.assertTrue(p._kirli)
        self.assertTrue(p.kaydet())
        self.assertTrue(p.winfo_exists())  # kaydet belgeyi kapatmaz
        self.assertFalse(p._kirli)
        iade_id = p.iade.id
        self.assertEqual(p.iade.durum, "TASLAK")
        self.assertEqual(self.o.lot_toplam(), D("10"))
        self.assertEqual(self.o.hareket_sayisi("ALIŞ İADE ÇIKIŞ"), 0)
        p.deiconify()
        p.geometry("1280x760+10+10")
        self.ekran("taslak_kaydedildi", p)
        p.destroy()

        tekrar = self.pencere(iade_id=iade_id)
        self.assertEqual(tekrar.iade.id, iade_id)
        self.assertEqual([(s["urun_kodu"], s["miktar"], s["kaynak_fatura_satiri_id"]) for s in tekrar.satirlar],
                         [("U001", D("3"), ks)])
        self.assertEqual(tekrar.toplam_lbl["genel"]["text"], "360,00 ₺")
        tekrar.destroy()

        liste = AlisIadeListesiPenceresi(self.root)
        self.assertEqual(liste.liste.tablo.get_children(), (str(iade_id),))
        self.assertIn("TASLAK", liste.liste.tablo.item(str(iade_id), "values"))
        liste.liste.tablo.selection_set(str(iade_id))
        acilan = liste.liste.ac()
        acilan.withdraw()
        self.assertEqual(acilan.iade.id, iade_id)
        acilan.destroy()
        self.ekran("liste", liste)
        liste.destroy()

    def test_onay_ve_iptal_ekrandan(self):
        _fid, ks = self.alis()
        p = self.pencere()
        self.tedarikci_sec(p)
        with patch("alis_iade_ui.KaynakSecDialog", _SahteKaynakSec), patch.object(p, "wait_window"):
            p.urun_ekle("U001")
        p.onayla()
        self.assertIn(p.iade.durum, ("AÇIK", "KAPALI"))
        self.assertEqual(self.o.lot_toplam(), D("7"))
        self.assertEqual(str(p.butonlar["onayla"]["state"]), "disabled")
        self.assertEqual(str(p.butonlar["satir_sil"]["state"]), "disabled")
        self.assertEqual(str(p.butonlar["iptal"]["state"]), "normal")
        p.deiconify()
        p.geometry("1280x760+10+10")
        self.ekran("onaylandi", p)
        with patch("alis_iade_ui.simpledialog.askstring", return_value="Yanlış iade"):
            p.iptal()
        self.assertEqual(p.iade.durum, "İPTAL")
        self.assertEqual(self.o.lot_toplam(), D("10"))
        self.assertEqual(str(p.butonlar["kaydet"]["state"]), "disabled")
        p.destroy()


class KucukEkranTest(_UITaban):
    def test_kucuk_ekranda_tum_butonlar_gorunur(self):
        _fid, ks = self.alis()
        p = self.pencere()
        self.tedarikci_sec(p)
        with patch("alis_iade_ui.KaynakSecDialog", _SahteKaynakSec), patch.object(p, "wait_window"):
            for _ in range(25):
                p.urun_ekle("U001")
        for gen, yuk in ((1024, 600), (900, 540)):
            p.deiconify()
            p.minsize(1, 1)
            p.geometry(f"{gen}x{yuk}+0+0")
            for _ in range(4):
                p.update()
            W, H = p.winfo_width(), p.winfo_height()
            for ad, b in p.butonlar.items():
                with self.subTest(ekran=f"{gen}x{yuk}", buton=ad):
                    self.assertTrue(b.winfo_ismapped())
                    x = b.winfo_rootx() - p.winfo_rootx()
                    y = b.winfo_rooty() - p.winfo_rooty()
                    self.assertLessEqual(x + b.winfo_width(), W + 1)
                    self.assertLessEqual(y + b.winfo_height(), H + 1)
            genel = p.toplam_lbl["genel"]
            self.assertTrue(genel.winfo_ismapped())
            self.assertLessEqual(genel.winfo_rooty() - p.winfo_rooty() + genel.winfo_height(), H + 1)
            if gen == 1024:
                self.ekran("kucuk_ekran_1024x600", p)
        p.destroy()


class KaynaksizUyariTest(_UITaban):
    def _devir(self, miktar=5, maliyet=80):
        from database.stok_service import StokService

        StokService.stok_girisi("U001", "ANA DEPO", "", date(2026, 1, 1), D(str(miktar)), D(str(maliyet)), "DEVIR-1")
        with get_session() as s:
            return int(s.scalar(select(StokLotu.id).where(StokLotu.lot_no == "DEVIR-1")))

    def test_uyari_diyalogu_metin_ve_uc_secenek(self):
        from alis_iade_ui import KaynaksizUyariDialog

        lot_id = self._devir()
        durum = S.urun_kaynak_durumu(self.o.ted_id, "U001", "ANA DEPO")
        d = KaynaksizUyariDialog(self.root, "U001", "Test Ürün", durum)
        self.assertEqual(d.mesaj["text"], KAYNAK_YOK_MESAJI)
        self.assertEqual([d.btn_kaynak["text"], d.btn_devam["text"], d.btn_vazgec["text"]],
                         ["Kaynak Seç", "Gerekçeyle Devam Et", "Vazgeç"])
        d.wm_transient("")
        d.geometry("600x330+40+40")
        self.ekran("kaynaksiz_uyari", d)
        d.devam()  # gerekçesiz devam edilemez
        self.assertTrue(d.winfo_exists())
        self.assertIsNone(d.result)
        d.kaynak_sec()
        self.assertEqual(d.result["kaynak_durumu"], DEVIR)
        self.assertEqual(d.result["kaynak_lot_id"], lot_id)

        d = KaynaksizUyariDialog(self.root, "U001", "Test Ürün", durum)
        d.gerekce.insert(0, "Eski sistemden devir")
        d.devam()
        self.assertEqual(d.result, {"kaynak_durumu": KAYNAKSIZ, "kaynak_gerekce": "Eski sistemden devir"})
        d = KaynaksizUyariDialog(self.root, "U001", "Test Ürün", durum)
        d.vazgec()
        self.assertIsNone(d.result)

    def test_vazgec_satir_eklemez_gerekceyle_devam_kayit_ve_onay(self):
        self._devir()
        p = self.pencere()
        self.tedarikci_sec(p)

        class Vazgec:
            def __init__(self, *_a, **_k):
                self.result = None

        class Devam:
            def __init__(self, *_a, **_k):
                self.result = {"kaynak_durumu": KAYNAKSIZ, "kaynak_gerekce": "Devirden"}

        with patch("alis_iade_ui.KaynaksizUyariDialog", Vazgec), patch.object(p, "wait_window"):
            p.urun_ekle("U001")
        self.assertEqual(p.satirlar, [])
        with patch("alis_iade_ui.KaynaksizUyariDialog", Devam), patch.object(p, "wait_window"):
            p.urun_ekle("U001")
        self.assertEqual(p.satirlar[0]["kaynak_durumu"], KAYNAKSIZ)
        p.satirlar[0]["miktar"] = D("2")
        p.onayla()
        self.assertIn(p.iade.durum, ("AÇIK", "KAPALI"))
        self.assertEqual(self.o.lot_toplam(), D("3"))
        s0 = S.getir(p.iade.id).satirlar[0]
        self.assertEqual((s0.kaynak_durumu, s0.kaynak_gerekce, s0.kaynak_onaylayan), (KAYNAKSIZ, "Devirden", "admin"))
        p.destroy()


class SatirIslemleriTest(_UITaban):
    def test_barkodla_ekleme_satir_silme_ve_kapatma_uyarisi(self):
        _fid, ks = self.alis(20, 100)
        with get_session() as s:
            s.scalar(select(StokKarti).where(StokKarti.stok_kodu == "U001")).barkod = "8690000000017"
        p = self.pencere()
        self.tedarikci_sec(p)
        with patch("alis_iade_ui.KaynakSecDialog", _SahteKaynakSec), patch.object(p, "wait_window"):
            p.kod_giris.insert(0, "8690000000017")
            p.urun_ekle()
            p.urun_ekle("U001")
        self.assertEqual(len(p.satirlar), 2)
        self.assertEqual(p.kod_giris.get(), "")
        p.tablo.selection_set("1")
        p.satir_sil()
        self.assertEqual(len(p.satirlar), 1)
        with patch("alis_iade_ui.messagebox.askyesnocancel", return_value=None):
            p.kapat()
        self.assertTrue(p.winfo_exists())  # vazgeçildi, açık kaldı
        with patch("alis_iade_ui.messagebox.askyesnocancel", return_value=False):
            p.kapat()
        self.assertFalse(p.winfo_exists())
        self.assertEqual(self.say(AlisIadeFaturasi), 0)

    def test_miktar_degisince_kaynaksiz_tercih_gecersiz(self):
        from database.stok_service import StokService

        StokService.stok_girisi("U001", "ANA DEPO", "", date(2026, 1, 1), D("5"), D("80"), "DEVIR-1")
        p = self.pencere()
        self.tedarikci_sec(p)

        class Devam:
            def __init__(self, *_a, **_k):
                self.result = {"kaynak_durumu": KAYNAKSIZ, "kaynak_gerekce": "x"}

        with patch("alis_iade_ui.KaynaksizUyariDialog", Devam), patch.object(p, "wait_window"):
            p.urun_ekle("U001")

        class Duzenle:
            def __init__(self, *_a, **_k):
                self.result = {"miktar": D("2"), "birim": "Adet", "birim_fiyat": D("80"), "iskonto_orani": D("0"),
                               "iskonto_orani_2": D("0"), "iskonto_orani_3": D("0"), "kdv_orani": D("20")}

        p.tablo.selection_set("0")
        with patch("alis_iade_ui.SatirDuzenleDialog", Duzenle), patch.object(p, "wait_window"):
            p.satir_duzenle()
        self.assertIsNone(p.satirlar[0].get("kaynak_durumu"))
        self.assertIn("SEÇİLMEDİ", p.tablo.item("0", "values")[11])
        p.destroy()


class CiktiTest(_UITaban):
    def _cok_satirli_taslak(self, n=70):
        _fid, ks = self.alis(200, 77.77)
        return S.kaydet({"iade_tarihi": date(2026, 3, 5), "cari_id": self.o.ted_id, "depo": "ANA DEPO",
                         "iade_nedeni": "Hasarlı"},
                        [{"urun_kodu": "U001", "urun_adi": f"Test Ürün satır {i}", "miktar": D("1"), "birim": "Adet",
                          "birim_fiyat": D("100"), "kdv_orani": D("20"), "kaynak_fatura_satiri_id": ks}
                         for i in range(n)])

    def test_docx_cok_sayfa_baslik_tekrari_taslak_etiketi_maliyet_yok(self):
        from docx import Document

        from invoice_print.alis_iade_cikti import cikti_uret

        iade = self._cok_satirli_taslak()
        yol = cikti_uret(iade.id, "docx")
        doc = Document(str(yol))
        tablo = max(doc.tables, key=lambda t: len(t.rows))
        self.assertEqual(len(tablo.rows), 71)
        self.assertIsNotNone(tablo.rows[0]._tr.trPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"))
        ust = doc.sections[0].header.paragraphs[0].text
        self.assertIn(iade.iade_no, ust)
        self.assertIn("TASLAK", ust)
        metin = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("8.400,00", metin)  # 70 × 100 × 1,20
        self.assertNotIn("77,77", metin)
        self.assertNotIn("aliyet", metin)

    def test_pdf_cok_sayfa_toplam_ekranla_ayni_maliyet_yok(self):
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("pymupdf yok")
        from invoice_print.alis_iade_cikti import cikti_uret

        iade = self._cok_satirli_taslak()
        try:
            yol = cikti_uret(iade.id, "pdf")
        except Exception as hata:  # noqa: BLE001
            self.skipTest(f"PDF motoru (Chrome/Edge) yok: {hata}")
        import fitz

        pdf = fitz.open(str(yol))
        self.assertGreaterEqual(pdf.page_count, 2)
        sayfalar = [pdf[i].get_text() for i in range(pdf.page_count)]
        pdf.close()
        for i, metin in enumerate(sayfalar):
            with self.subTest(sayfa=i + 1):
                self.assertIn("Birim Fiyat", metin)  # tablo başlığı her sayfada
                self.assertIn(iade.iade_no, metin)  # sayfa altı bilgisi
        tum = "\n".join(sayfalar)
        self.assertIn("TASLAK", tum)
        self.assertNotIn("77,77", tum)
        p = self.pencere(iade_id=iade.id)
        ekran_genel = p.toplam_lbl["genel"]["text"].replace(" ₺", "")
        p.destroy()
        self.assertEqual(ekran_genel, "8.400,00")
        self.assertIn(ekran_genel, tum)

    def test_onayli_ve_iptal_ciktisi_etiketleri(self):
        from invoice_print.alis_iade_cikti import cikti_modeli

        _fid, ks = self.alis()
        iade = S.kaydet_ve_onayla({"iade_tarihi": date(2026, 3, 5), "cari_id": self.o.ted_id, "depo": "ANA DEPO"},
                                  [{"urun_kodu": "U001", "urun_adi": "Test Ürün", "miktar": D("3"), "birim": "Adet",
                                    "birim_fiyat": D("100"), "kdv_orani": D("20"), "kaynak_fatura_satiri_id": ks}])
        vm = cikti_modeli(iade.id)
        self.assertIsNone(vm.durum_etiketi)
        self.assertEqual(vm.toplamlar[-1], ("GENEL TOPLAM", "360,00 ₺"))
        self.assertTrue(all("maliyet" not in k for s in vm.satirlar for k in s))
        S.iptal_et(iade.id, "test")
        self.assertEqual(cikti_modeli(iade.id).durum_etiketi, "İPTAL EDİLDİ")


if __name__ == "__main__":
    unittest.main()
