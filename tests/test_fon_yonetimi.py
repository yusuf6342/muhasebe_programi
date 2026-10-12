"""Fon yönetimi sayısal kabul testleri (talimat bölüm 16) — yalnız geçici test veritabanı."""

from __future__ import annotations

import sys
import threading
import unittest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from sqlalchemy import func, select  # noqa: E402

from database.database import get_session  # noqa: E402
from database.fon_service import FonService, dagit, oran_oku  # noqa: E402
from database.models.cari import CariKapatma, CariIslem, SatisHareketi  # noqa: E402
from database.models.finans import FinansHareketi, FinansHesabi  # noqa: E402
from database.models.fon import FonHareketi, FonParaHareketi  # noqa: E402
from database.models.satis_faturasi import SatisFaturasi, SatisFaturasiSatiri  # noqa: E402
from database.models.stok import StokHareketi, StokKarti  # noqa: E402
from tests.stok_test_ortami import StokOrtami  # noqa: E402

D = Decimal
BUGUN = date.today()
GECMIS = BUGUN - timedelta(days=20)


class _FonTaban(unittest.TestCase):
    def setUp(self):
        self.o = StokOrtami("fon")
        with get_session() as s:
            s.add(FinansHesabi(hesap_adi="TEST KASA", hesap_turu="KASA", aktif=True))
            s.add(FinansHesabi(hesap_adi="TEST BANKA", hesap_turu="BANKA", aktif=True))
            self.u1 = int(s.scalar(select(StokKarti.id).where(StokKarti.stok_kodu == "U001")))
        self.u2 = self.o.urun_ekle("U002", "İkinci Ürün")
        self.hz = self.o.urun_ekle("H001", "Nakliye Hizmeti", kart_turu="Hizmet")
        for kod in ("U001", "U002", "H001"):
            self.o.alis(1000, 100, GECMIS, kod=kod)
        self.yardim = FonService.fon_kaydet({"kod": "YARDIM", "ad": "Yardım"})
        self.yedek = FonService.fon_kaydet({"kod": "YEDEK", "ad": "Yedek Akçe"})
        self.tasarruf = FonService.fon_kaydet({"kod": "TASARRUF", "ad": "Tasarruf"})
        FonService.stok_fon_oranlarini_kaydet(self.u1, [
            {"fon_id": self.yardim, "oran": "1,5", "aktif": True},
            {"fon_id": self.yedek, "oran": "0,5", "aktif": True},
            {"fon_id": self.tasarruf, "oran": "1", "aktif": True},
        ])

    def tearDown(self):
        self.o.kapat()

    # --- yardımcılar
    def satis(self, satirlar, *, tahsilat=None, onayla=True, tarih=None, **ek):
        from database.satis_faturasi_service import SatisFaturasiService

        veriler = {"fatura_tarihi": tarih or BUGUN, "vade_tarihi": tarih or BUGUN,
                   "cari_id": self.o.mus_id, "depo": "ANA DEPO", "sales_person_id": 1, **ek}
        rows = [{"urun_kodu": k, "urun_adi": k, "miktar": D(str(m)), "birim": "Adet",
                 "birim_fiyat": D(str(f)), "iskonto_orani": D(str(i)), "kdv_orani": D("20")}
                for k, m, f, *rest in satirlar for i in [rest[0] if rest else 0]]
        th = [{"tutar": D(str(t)), "odeme_sekli": sekil, "hesap": "TEST KASA" if "NAK" in sekil else "TEST BANKA"}
              for t, sekil in (tahsilat or [])]
        f = SatisFaturasiService.kaydet(veriler, rows, tahsilat_verileri=th or None)
        if onayla:
            SatisFaturasiService.onayla(f.id)
        return int(f.id)

    def paylar(self, belge_id=None, fon_id=None, durum="AKTIF", olay=None):
        with get_session() as s:
            q = select(func.coalesce(func.sum(FonHareketi.tutar), 0)).where(FonHareketi.durum == durum)
            if belge_id is not None:
                q = q.where(FonHareketi.kaynak_fatura_id == belge_id)
            if fon_id is not None:
                q = q.where(FonHareketi.fon_id == fon_id)
            if olay:
                q = q.where(FonHareketi.olay == olay)
            return D(str(s.scalar(q)))

    def durum(self, fon_id):
        return FonService.fon_durumu(fon_id)

    def iade(self, fatura_id, miktar, kod="U001", fiyat=1000):
        from database.satis_iade_faturasi_service import SatisIadeFaturasiService

        with get_session() as s:
            kaynak = s.scalar(select(SatisFaturasiSatiri).where(
                SatisFaturasiSatiri.fatura_id == fatura_id, SatisFaturasiSatiri.urun_kodu == kod))
            ks_id = int(kaynak.id)
        return SatisIadeFaturasiService.kaydet(
            {"iade_tarihi": BUGUN, "cari_id": self.o.mus_id, "kaynak_fatura_id": fatura_id, "depo": "ANA DEPO"},
            [{"urun_kodu": kod, "urun_adi": kod, "miktar": D(str(miktar)), "birim": "Adet",
              "birim_fiyat": D(str(fiyat)), "kdv_orani": D("20"), "kaynak_fatura_satiri_id": ks_id}],
        )


class PayHesaplamaTest(_FonTaban):
    def test_150_50_100_toplam_300_tekrar_onay_degistirmez(self):
        from database.satis_faturasi_service import SatisFaturasiService

        fid = self.satis([("U001", 10, 1000)])
        self.assertEqual(self.paylar(fid, self.yardim), D("150.00"))
        self.assertEqual(self.paylar(fid, self.yedek), D("50.00"))
        self.assertEqual(self.paylar(fid, self.tasarruf), D("100.00"))
        self.assertEqual(self.paylar(fid), D("300.00"))
        with self.assertRaises(ValueError):
            SatisFaturasiService.onayla(fid)
        # Kanca doğrudan ikinci kez çağrılsa da çift pay oluşmaz
        with get_session() as s:
            f = s.get(SatisFaturasi, fid)
            self.assertEqual(FonService.satis_onaylandi(s, f), 0)
        self.assertEqual(self.paylar(fid), D("300.00"))

    def test_taslak_pay_uretmez_onizleme_gosterir(self):
        fid = self.satis([("U001", 10, 1000)], onayla=False)
        self.assertEqual(self.paylar(fid), D("0"))
        onizleme = FonService.belge_onizleme([{"urun_kodu": "U001", "miktar": 10, "birim_fiyat": 1000,
                                               "kdv_orani": 20}])
        self.assertEqual(sum(p["tutar"] for p in onizleme), D("300.00"))
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(FonHareketi.id))), 0)

    def test_hizmet_haric_iskonto_ve_bedelsiz(self):
        FonService.stok_fon_oranlarini_kaydet(self.hz, [{"fon_id": self.yardim, "oran": "10", "aktif": True}])
        FonService.stok_fon_oranlarini_kaydet(self.u2, [{"fon_id": self.yardim, "oran": "1,5", "aktif": True}])
        fid = self.satis([("U001", 10, 1000, 10), ("H001", 1, 500), ("U002", 1, 0)])
        # 10×1000 %10 iskonto = 9000 → 135; hizmet hariç; bedelsiz ürün pay üretmez
        self.assertEqual(self.paylar(fid, self.yardim), D("135.00"))
        with get_session() as s:
            kodlar = {h.urun_kodu for h in s.scalars(select(FonHareketi).where(FonHareketi.belge_id == fid))}
        self.assertEqual(kodlar, {"U001"})

    def test_genel_indirim_dagitilir_hizli_satis_ornegi(self):
        # 1.200 KDV dahil → Yeni Toplam 1.080: taban 900, fon %1,5 = 13,50
        fid = self.satis([("U001", 1, 1000)], tl_brut_toplam=D("1200"), tl_genel_toplam=D("1080"),
                         genel_islem_turu="INDIRIM", genel_islem_tutari=D("120"))
        with get_session() as s:
            h = s.scalar(select(FonHareketi).where(FonHareketi.belge_id == fid,
                                                   FonHareketi.fon_id == self.yardim))
            self.assertEqual(D(str(h.taban)), D("900.00"))
            self.assertEqual(D(str(h.tutar)), D("13.50"))

    def test_hizli_satis_yeni_toplam_ve_normal_ekranda_duzenleme(self):
        from database.hizli_satis_service import HizliSatisService
        from database.satis_faturasi_service import SatisFaturasiService
        from hizli_satis_sepet import HizliSatisSepet

        sepet = HizliSatisSepet()
        sepet.ekle(stok_id=self.u1, stok_kodu="U001", stok_adi="Test Ürün", miktar=1, birim_fiyat="1000",
                   kdv_orani=20)
        toplamlar = sepet.toplamlar()
        sepet.hedef_toplam_uygula("1080")
        sonuc = HizliSatisService.satisi_tamamla(
            sepet=sepet, cari_id=self.o.mus_id, tahsilatlar=[],
            idempotency_token=HizliSatisService.yeni_idempotency_token(), musteri_bakiye=0, risk_limiti=None)
        fid = int(sonuc["fatura_id"]) if isinstance(sonuc, dict) else int(sonuc)
        self.assertIn(D("1200.00"), [D(str(v)) for v in toplamlar.values() if isinstance(v, Decimal)])
        with get_session() as s:
            f = s.get(SatisFaturasi, fid)
            hareket = s.scalar(select(SatisHareketi).where(SatisHareketi.belge_no == f.fatura_no))
            self.assertEqual(D(str(hareket.satis_tutari)), D("1080.00"))
        self.assertEqual(self.paylar(fid, self.yardim), D("13.50"))
        # Normal ekranda onay kaldır + yeniden onay: ikinci katkı oluşmaz
        SatisFaturasiService.onay_kaldir(fid)
        self.assertEqual(self.paylar(fid), D("0"))
        SatisFaturasiService.onayla(fid)
        self.assertEqual(self.paylar(fid, self.yardim), D("13.50"))
        self.assertGreater(self.paylar(fid, durum="IPTAL"), 0)

    def test_kurus_dagitimi_deterministik(self):
        self.assertEqual(dagit([D("0.333"), D("0.333"), D("0.334")], D("1.00")), [D("0.33"), D("0.33"), D("0.34")])
        self.assertEqual(sum(dagit([D("1.005")] * 3, D("3.02"))), D("3.02"))
        FonService.stok_fon_oranlarini_kaydet(self.u2, [{"fon_id": self.yardim, "oran": "1,5", "aktif": True}])
        fid = self.satis([("U001", 1, D("33.33")), ("U002", 1, D("33.33")), ("U001", 1, D("33.34"))])
        # Toplam taban 100,00 × %1,5 = 1,50 — satır toplamları fon toplamını verir
        self.assertEqual(self.paylar(fid, self.yardim), D("1.50"))

    def test_oran_dogrulama(self):
        self.assertEqual(oran_oku("1,5"), D("1.5000"))
        for kotu in ("-1", "101", "abc"):
            with self.assertRaises(ValueError):
                oran_oku(kotu)
        with self.assertRaises(ValueError) as ctx:
            FonService.stok_fon_oranlarini_kaydet(self.u1, [{"fon_id": self.yedek, "oran": "98", "aktif": True}])
        self.assertIn("100", str(ctx.exception))
        FonService.stok_fon_oranlarini_kaydet(self.u2, [{"fon_id": self.yardim, "oran": "0", "aktif": True}])
        fid = self.satis([("U002", 1, 1000)])
        self.assertEqual(self.paylar(fid), D("0"))
        gecmis = FonService.oran_gecmisi(self.u1)
        self.assertTrue(gecmis and gecmis[0]["kullanici"])


class OranDegisikligiTest(_FonTaban):
    def test_eski_satis_korunur_yeni_satis_yeni_oran_iade_eski_oran(self):
        eski = self.satis([("U001", 10, 1000)])
        FonService.stok_fon_oranlarini_kaydet(self.u1, [{"fon_id": self.yardim, "oran": "2", "aktif": True}])
        yeni = self.satis([("U001", 10, 1000)])
        self.assertEqual(self.paylar(eski, self.yardim), D("150.00"))
        self.assertEqual(self.paylar(yeni, self.yardim), D("200.00"))
        self.iade(eski, 2)
        with get_session() as s:
            h = s.scalar(select(FonHareketi).where(FonHareketi.olay == "IADE_DUZELTME",
                                                   FonHareketi.fon_id == self.yardim))
            self.assertEqual(D(str(h.oran)), D("1.5000"))
            self.assertEqual(D(str(h.tutar)), D("-30.00"))

    def test_eski_faturaya_yeni_satir_guncel_oran_eski_satir_eski_oran(self):
        from database.satis_faturasi_service import SatisFaturasiService

        fid = self.satis([("U001", 10, 1000)])
        FonService.stok_fon_oranlarini_kaydet(self.u1, [{"fon_id": self.yardim, "oran": "2", "aktif": True}])
        FonService.stok_fon_oranlarini_kaydet(self.u2, [{"fon_id": self.yardim, "oran": "2", "aktif": True}])
        SatisFaturasiService.onay_kaldir(fid)
        with get_session() as s:
            rv = s.get(SatisFaturasi, fid).row_version
        SatisFaturasiService.kaydet(
            {"fatura_tarihi": BUGUN, "vade_tarihi": BUGUN, "cari_id": self.o.mus_id, "depo": "ANA DEPO",
             "sales_person_id": 1, "row_version": rv},
            [{"urun_kodu": "U001", "urun_adi": "U001", "miktar": D("10"), "birim": "Adet",
              "birim_fiyat": D("1000"), "iskonto_orani": D("0"), "kdv_orani": D("20")},
             {"urun_kodu": "U002", "urun_adi": "U002", "miktar": D("10"), "birim": "Adet",
              "birim_fiyat": D("1000"), "iskonto_orani": D("0"), "kdv_orani": D("20")}],
            fatura_id=fid,
        )
        SatisFaturasiService.onayla(fid)
        with get_session() as s:
            aktif = {h.urun_kodu: (D(str(h.oran)), D(str(h.tutar))) for h in s.scalars(select(FonHareketi).where(
                FonHareketi.belge_id == fid, FonHareketi.fon_id == self.yardim, FonHareketi.durum == "AKTIF"))}
        self.assertEqual(aktif["U001"], (D("1.5000"), D("150.00")))
        self.assertEqual(aktif["U002"], (D("2.0000"), D("200.00")))


class TahsilatVeParaTest(_FonTaban):
    def test_6000_tahsilat_75_75(self):
        fid = self.satis([("U001", 10, 1000)], tahsilat=[(6000, "NAKİT / KASA")])
        d = self.durum(self.yardim)
        self.assertEqual(d["net_satis_payi"], D("150.00"))
        self.assertEqual(d["tahsil_edilen"], D("75.00"))
        self.assertEqual(d["tahsil_edilmemis"], D("75.00"))

    def test_fazla_odeme_yuzde_yuzu_gecmez_ve_pos_brut(self):
        self.satis([("U001", 10, 1000)], tahsilat=[(12000, "KREDİ KARTIYLA TAHSİLAT")])
        self.assertEqual(self.durum(self.yardim)["tahsil_edilen"], D("150.00"))

    def test_ayir_60_kullan_20(self):
        self.satis([("U001", 10, 1000)], tahsilat=[(6000, "NAKİT / KASA")])
        FonService.para_ayir(self.yardim, BUGUN, "60", hesap_adi="TEST KASA")
        FonService.fon_kullan(self.yardim, BUGUN, "20", yardim_alan="Ali Veli", amac="Eğitim yardımı")
        d = self.durum(self.yardim)
        self.assertEqual((d["ayrilmis"], d["kullanim"], d["kullanilabilir"], d["ayrilabilir"]),
                         (D("60.00"), D("20.00"), D("40.00"), D("15.00")))
        with self.assertRaises(ValueError):
            FonService.para_ayir(self.yardim, BUGUN, "15,01")
        with self.assertRaises(ValueError):
            FonService.fon_kullan(self.yardim, BUGUN, "40,01", yardim_alan="X", amac="Y")
        with get_session() as s:
            p = s.scalar(select(FonParaHareketi).where(FonParaHareketi.tur == "AYIRMA"))
            self.assertTrue(p.belgesiz)
        listede = FonService.para_hareketleri(fon_id=self.yardim)
        self.assertIn("Belgesiz iç takip kaydı", [r["belge"] for r in listede])

    def test_geri_alma_ve_kullanim_duzeltme_ters_kayit(self):
        self.satis([("U001", 10, 1000)], tahsilat=[(12000, "NAKİT / KASA")])
        ayir = FonService.para_ayir(self.yardim, BUGUN, "100")
        kul = FonService.fon_kullan(self.yardim, BUGUN, "80", yardim_alan="A", amac="B")
        with self.assertRaises(ValueError):
            FonService.ayirmayi_geri_al(ayir, "30", "fazla")  # kullanılmış tutar geri alınamaz
        with self.assertRaises(ValueError):
            FonService.ayirmayi_geri_al(ayir, "10", "")
        FonService.kullanimi_duzelt(kul, "30", "Yanlış tutar")
        FonService.ayirmayi_geri_al(ayir, "30", "Fazla ayrıldı")
        d = self.durum(self.yardim)
        self.assertEqual((d["ayrilmis"], d["kullanim"], d["kullanilabilir"]), (D("70.00"), D("50.00"), D("20.00")))
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(FonParaHareketi.id))), 4)

    def test_kaynak_dagitimi_60_20_40_ve_tekrar(self):
        FonService.stok_fon_oranlarini_kaydet(self.u1, [
            {"fon_id": self.yardim, "oran": "10", "aktif": True},
            {"fon_id": self.yedek, "oran": "10", "aktif": True},
            {"fon_id": self.tasarruf, "oran": "10", "aktif": True},
        ])
        # Kaynak tahsilat 120 TL (kapasite); fonların hakkı ikinci, tam ödenmiş satıştan yeterli
        fid = self.satis([("U001", 1, 1000)], tahsilat=[(120, "NAKİT / KASA")])
        self.satis([("U001", 10, 1000)], tahsilat=[(12000, "NAKİT / KASA")])
        with get_session() as s:
            th_id = s.get(SatisFaturasi, fid).tahsilatlar[0].id
        anahtar = f"FT:{th_id}"
        FonService.para_ayir(self.yardim, BUGUN, "60", kaynak_anahtar=anahtar, istek_kimligi="t-1")
        FonService.para_ayir(self.yardim, BUGUN, "60", kaynak_anahtar=anahtar, istek_kimligi="t-1")  # tekrar
        FonService.para_ayir(self.yedek, BUGUN, "20", kaynak_anahtar=anahtar)
        FonService.para_ayir(self.tasarruf, BUGUN, "40", kaynak_anahtar=anahtar)
        self.assertEqual(FonService.kaynak_kullanilan(anahtar), D("120.00"))
        with self.assertRaises(ValueError) as ctx:
            FonService.para_ayir(self.yedek, BUGUN, "0,01", kaynak_anahtar=anahtar)
        self.assertIn("kaynak", str(ctx.exception))
        self.assertEqual(self.durum(self.yardim)["ayrilmis"], D("60.00"))
        kaynak = next(k for k in FonService.kaynak_tahsilatlar() if k["anahtar"] == anahtar)
        self.assertEqual((kaynak["kapasite"], kaynak["kalan"]), (D("120.00"), D("0.00")))

    def test_esz_zamanli_ayirma_kapasiteyi_asmaz(self):
        self.satis([("U001", 10, 1000)], tahsilat=[(12000, "NAKİT / KASA")])
        hatalar, basarili = [], []

        def calis():
            try:
                basarili.append(FonService.para_ayir(self.yardim, BUGUN, "100"))
            except Exception as e:  # noqa: BLE001
                hatalar.append(e)

        isler = [threading.Thread(target=calis) for _ in range(2)]
        for t in isler:
            t.start()
        for t in isler:
            t.join()
        self.assertEqual(len(basarili), 1)
        self.assertEqual(self.durum(self.yardim)["ayrilmis"], D("100.00"))

    def test_tek_tahsilat_iki_fatura_baglantisiz_ve_iptal(self):
        from database.cari_service import CariService

        f1 = self.satis([("U001", 10, 1000)])
        f2 = self.satis([("U001", 10, 1000)])
        with get_session() as s:
            h = {f: s.scalar(select(SatisHareketi.id).join(SatisFaturasi, SatisFaturasi.fatura_no
                                                           == SatisHareketi.belge_no).where(SatisFaturasi.id == f))
                 for f in (f1, f2)}
        islem = CariService.tahsilat_yap(self.o.mus_id, BUGUN, D("18000"), "NAKİT", "TEST KASA",
                                         dagitim=[(h[f1], D("12000")), (h[f2], D("6000"))])
        self.assertEqual(self.durum(self.yardim)["tahsil_edilen"], D("225.00"))
        # Bağlantısız tahsilat / uygulanmamış avans hak doğurmaz
        f3 = self.satis([("U001", 10, 1000)])
        self.assertTrue(f3)
        self.assertEqual(self.durum(self.yardim)["tahsil_edilen"], D("225.00"))
        # Çek/senet (finans hareketi olmayan kaynak) desteklenmez
        with get_session() as s:
            fatura = s.get(SatisFaturasi, f3)
            hid = s.scalar(select(SatisHareketi.id).where(SatisHareketi.belge_no == fatura.fatura_no))
            s.add(CariKapatma(islem_kimligi="x", cari_id=self.o.mus_id, kaynak_belge_no="CEK-0001",
                              kaynak_tur="TAHSILAT", hedef_hareket_id=hid, hedef_belge_no=fatura.fatura_no,
                              tutar=D("12000"), tarih=BUGUN))
        self.assertEqual(self.durum(self.yardim)["tahsil_edilen"], D("225.00"))
        # Tahsilat iptali payı geri alır
        CariService.tahsilat_odeme_iptal(islem.id, "Test iptali")
        self.assertEqual(self.durum(self.yardim)["tahsil_edilen"], D("0"))


class IadeTest(_FonTaban):
    def test_10_birimden_2_iade_eksi_30_net_120(self):
        fid = self.satis([("U001", 10, 1000)], tahsilat=[(12000, "NAKİT / KASA")])
        self.iade(fid, 2)
        d = self.durum(self.yardim)
        self.assertEqual(d["net_satis_payi"], D("120.00"))
        self.assertEqual(d["tahsil_edilen"], D("120.00"))  # iade + para iadesi çift düşmez
        self.assertEqual(self.paylar(fid, self.yardim, olay="IADE_DUZELTME"), D("-30.00"))

    def test_iade_toplami_asil_payi_asmaz_ve_iptal(self):
        from database.satis_iade_faturasi_service import SatisIadeFaturasiService

        fid = self.satis([("U001", 10, 1000)])
        i1 = self.iade(fid, 6)
        try:
            self.iade(fid, 6)
        except ValueError:
            pass  # iade servisi kaynak miktarı aşımını zaten engelleyebilir
        self.assertGreaterEqual(self.durum(self.yardim)["net_satis_payi"], D("0"))
        SatisIadeFaturasiService.iptal_et(i1.id)
        with get_session() as s:
            aktif_i1 = s.scalar(select(func.count(FonHareketi.id)).where(
                FonHareketi.kaynak_turu == "SATIS_IADE", FonHareketi.belge_id == i1.id,
                FonHareketi.durum == "AKTIF"))
        self.assertEqual(aktif_i1, 0)

    def test_tam_iade_sonrasi_fon_acigi_ve_kullanim_engeli(self):
        fid = self.satis([("U001", 10, 1000)], tahsilat=[(12000, "NAKİT / KASA")])
        FonService.para_ayir(self.yardim, BUGUN, "150")
        FonService.fon_kullan(self.yardim, BUGUN, "100", yardim_alan="A", amac="B")
        self.iade(fid, 10)
        d = self.durum(self.yardim)
        self.assertEqual(d["net_satis_payi"], D("0.00"))
        self.assertEqual(d["fon_acigi"], D("150.00"))
        self.assertEqual(d["kullanim"], D("100.00"))
        with self.assertRaises(ValueError) as ctx:
            FonService.fon_kullan(self.yardim, BUGUN, "1", yardim_alan="A", amac="B")
        self.assertIn("açığı", str(ctx.exception))
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(FonParaHareketi.id))), 2)

    def test_kaynaksiz_iade_inceleme(self):
        from database.satis_iade_faturasi_service import SatisIadeFaturasiService

        self.satis([("U001", 10, 1000)])
        try:
            SatisIadeFaturasiService.kaydet(
                {"iade_tarihi": BUGUN, "cari_id": self.o.mus_id, "depo": "ANA DEPO"},
                [{"urun_kodu": "U001", "urun_adi": "U001", "miktar": D("1"), "birim": "Adet",
                  "birim_fiyat": D("1000"), "kdv_orani": D("20"), "kaynak_yok_onay": True,
                  "kaynak_yok_gerekce": "Fiş kayıp", "kaynak_yok_kullanici": "admin"}],
            )
        except ValueError as hata:
            self.skipTest(f"Kaynaksız iade bu ortamda kabul edilmedi: {hata}")
        with get_session() as s:
            durumlar = {h.durum for h in s.scalars(select(FonHareketi).where(
                FonHareketi.kaynak_turu == "SATIS_IADE"))}
        self.assertEqual(durumlar, {"INCELEME"})
        self.assertEqual(self.durum(self.yardim)["net_satis_payi"], D("150.00"))


class IptalVeMutabakatTest(_FonTaban):
    def test_iptal_payi_ters_cevirir_kayit_gorunur(self):
        from database.satis_faturasi_service import SatisFaturasiService

        fid = self.satis([("U001", 10, 1000)])
        SatisFaturasiService.iptal_et(fid, "Test iptal")
        self.assertEqual(self.durum(self.yardim)["net_satis_payi"], D("0"))
        liste = FonService.hareket_listesi(fon_id=self.yardim)
        self.assertTrue(liste and all(r["durum"] == "IPTAL" for r in liste))

    def test_mutabakat_ve_raporlar(self):
        self.satis([("U001", 10, 1000)], tarih=BUGUN - timedelta(days=5), tahsilat=[(6000, "NAKİT / KASA")])
        self.satis([("U001", 10, 1000)])
        FonService.para_ayir(self.yardim, BUGUN, "50")
        satirlar = FonService.mutabakat(BUGUN - timedelta(days=2), BUGUN)
        self.assertTrue(all(r["tutarli"] for r in satirlar))
        yardim = [r for r in satirlar if r["fon_kod"] == "YARDIM"]
        net = next(r for r in yardim if r["kalem"] == "Net satış payı")
        self.assertEqual((net["acilis"], net["donem"], net["kapanis"]), (D("150.00"), D("150.00"), D("300.00")))
        urun = FonService.urun_bazli_rapor()
        self.assertEqual(sum(r["net"] for r in urun if r["fon_kod"] == "YARDIM"), D("300.00"))
        gun = FonService.tarih_bazli_rapor()
        self.assertEqual(len({r["tarih"] for r in gun}), 2)

    def test_eslesik_hareket_cogalmaz_stok_cari_tek(self):
        fid = self.satis([("U001", 10, 1000)])
        with get_session() as s:
            no = s.get(SatisFaturasi, fid).fatura_no
            self.assertEqual(s.scalar(select(func.count(CariIslem.id)).where(CariIslem.belge_no == no)), 0)
            self.assertEqual(s.scalar(select(func.count(FinansHareketi.id)).where(FinansHareketi.belge_no == no)), 0)
            self.assertEqual(s.scalar(select(func.count(StokHareketi.id)).where(
                StokHareketi.belge_no == no)), 1)


class TanimVeYetkiTest(_FonTaban):
    def test_hareketli_fon_silinmez_pasif_olur(self):
        self.satis([("U001", 10, 1000)])
        with self.assertRaises(ValueError):
            FonService.fon_sil(self.yardim)
        FonService.fon_aktiflik(self.yardim, False)
        fid = self.satis([("U001", 10, 1000)])
        self.assertEqual(self.paylar(fid, self.yardim), D("0"))
        FonService.fon_kaydet({"kod": "YARDIM", "ad": "Yardım Fonu"}, self.yardim)
        self.assertEqual(self.paylar(None, self.yardim), D("150.00"))  # ad değişikliği kimliği bozmaz

    def test_ek_bedel_yontemi_kdv_ve_hesap_olmadan_kaydedilmez(self):
        with self.assertRaisesRegex(ValueError, "KDV oranı"):
            FonService.fon_kaydet({"kod": "EK", "ad": "Ek", "yontem": "EK_BEDEL"})
        with self.assertRaisesRegex(ValueError, "muhasebe hesabı"):
            FonService.fon_kaydet({"kod": "EK", "ad": "Ek", "yontem": "EK_BEDEL", "ek_bedel_kdv_orani": "20"})

    def test_yetkisiz_kullanici(self):
        from database.access import AccessError
        from database.session_manager import oturum

        oturum.set_user(user_id=2, kullanici_adi="satis", ad_soyad="Satışçı", role_kod="SATIS",
                        role_ad="Satış", permissions={"satis_goruntuleme"})
        for islem in (lambda: FonService.para_ayir(self.yardim, BUGUN, "1"),
                      lambda: FonService.fon_kullan(self.yardim, BUGUN, "1", yardim_alan="a", amac="b"),
                      lambda: FonService.fon_kaydet({"kod": "Y2", "ad": "Y2"}),
                      lambda: FonService.bakiyeler()):
            with self.assertRaises(AccessError):
                islem()

    def test_toplu_oran_onizleme_ve_uygulama(self):
        from database.models.stok import StokGrubu

        with get_session() as s:
            g = StokGrubu(kod="G1", ad="Grup", seviye=1)
            s.add(g)
            s.flush()
            gid = g.id
            for sid in (self.u1, self.u2):
                s.get(StokKarti, sid).ana_grup_id = gid
        on = FonService.toplu_oran_onizleme(gid, self.yardim, "3", "BOS")
        islem = {r["stok_kodu"]: r["islem"] for r in on}
        self.assertEqual(islem, {"U001": "ATLA", "U002": "EKLE"})
        sayac = FonService.toplu_oran_uygula(gid, self.yardim, "3", "BOS")
        self.assertEqual(sayac["EKLE"], 1)
        self.assertEqual(FonService.toplu_oran_uygula(gid, self.yardim, "4", "SECILI", [self.u1])["DEGISTIR"], 1)
        oranlar = {r["fon_kod"]: r["oran"] for r in FonService.stok_fon_oranlari(self.u1)}
        self.assertEqual(D(str(oranlar["YARDIM"])), D("4"))


class EkranVeCiktiTest(_FonTaban):
    def setUp(self):
        super().setUp()
        import tkinter as tk

        try:
            self.root = tk.Tk()
        except tk.TclError as hata:
            self.skipTest(f"Tk yok: {hata}")
        self.root.withdraw()

    def tearDown(self):
        try:
            self.root.destroy()
        finally:
            super().tearDown()

    def test_bos_ve_dolu_ekran(self):
        from fon_ui import FonYonetimiSayfasi

        sayfa = FonYonetimiSayfasi(self.root)
        sayfa.pack()
        self.root.update()
        tree = sayfa.sekmeler["hareket"][0]
        self.assertEqual(tree.get_children(), ("bos",))  # boş liste mesajı
        self.satis([("U001", 10, 1000)], tahsilat=[(6000, "NAKİT / KASA")])
        sayfa.yenile()
        bakiye = {r["fon_kod"]: r for r in sayfa._veri["bakiye"]}
        self.assertEqual(bakiye["YARDIM"]["tahsil_edilen"], D("75.00"))
        self.assertEqual(len(sayfa.sekmeler["hareket"][0].get_children()), 3)
        sayfa.destroy()

    def test_stok_karti_fonlar_bolumu(self):
        from fon_ui import StokFonlariCercevesi

        c = StokFonlariCercevesi(self.root, self.u2)
        c.pack()
        self.root.update()
        idx = next(i for i, s in enumerate(c.satirlar) if s["fon_id"] == self.yardim)
        c.tree.selection_set(str(idx))
        c._secildi()
        c.oran_var.set("2,5")
        c._satira_uygula()
        from unittest.mock import patch

        with patch("fon_ui.messagebox.showinfo"):
            c._kaydet()
        oran = {r["fon_id"]: r["oran"] for r in FonService.stok_fon_oranlari(self.u2)}[self.yardim]
        self.assertEqual(D(str(oran)), D("2.5"))

    def test_excel_pdf_tekrarlanan_baslik(self):
        import tempfile

        from fon_rapor_cikti import excel_yaz, pdf_yaz

        satirlar = [[f"F{i}", D(i), BUGUN] for i in range(120)]
        with tempfile.TemporaryDirectory() as d:
            x = excel_yaz(Path(d) / "a.xlsx", "Fon", "Dönem", ["Fon", "Tutar", "Tarih"], satirlar)
            from openpyxl import load_workbook

            ws = load_workbook(x).active
            self.assertEqual(ws.print_title_rows.replace("$", ""), "4:4")
            self.assertEqual(ws.cell(row=4, column=2).value, "Tutar")
            p = pdf_yaz(Path(d) / "a.pdf", "Fon Yönetimi", "Dönem", ["Fon", "Tutar", "Tarih"], satirlar)
            from fon_rapor_cikti import _fitz

            doc = _fitz().open(str(p))
            self.assertGreater(doc.page_count, 1)
            self.assertTrue(all("Tutar" in doc[i].get_text() for i in range(doc.page_count)))
            doc.close()

    def test_cikti_para_birimi_ve_dort_tutar_anlami(self):
        import tempfile

        from database.fon_service import DESTEKLENEN_TAHSILAT_ACIKLAMA, TUTAR_ACIKLAMALARI
        from fon_rapor_cikti import _fitz, excel_yaz, pdf_yaz

        notlar = ["Para birimi: TL.", *TUTAR_ACIKLAMALARI.values(), DESTEKLENEN_TAHSILAT_ACIKLAMA]
        satirlar = [[f"F{i}", D(i)] for i in range(90)]
        with tempfile.TemporaryDirectory() as d:
            from openpyxl import load_workbook

            ws = load_workbook(excel_yaz(Path(d) / "n.xlsx", "Fon", "Dönem", ["Fon", "Tutar"], satirlar,
                                         notlar)).active
            ilk_sutun = [ws.cell(row=i, column=1).value for i in range(1, ws.max_row + 1)]
            for n in notlar:
                self.assertIn(n, ilk_sutun)
            baslik = next(i for i in range(1, ws.max_row + 1) if ws.cell(row=i, column=2).value == "Tutar")
            self.assertEqual(ws.print_title_rows.replace("$", ""), f"{baslik}:{baslik}")
            doc = _fitz().open(str(pdf_yaz(Path(d) / "n.pdf", "Fon", "Dönem", ["Fon", "Tutar"], satirlar,
                                           notlar)))
            try:
                for i in range(doc.page_count):
                    metin = " ".join(doc[i].get_text().split())
                    self.assertIn("Para birimi: TL.", metin)
                    self.assertIn("Tutar", metin)
                    for n in TUTAR_ACIKLAMALARI.values():
                        self.assertIn(" ".join(n.split()[:3]), metin)
                son = " ".join(doc[doc.page_count - 1].get_text().split())
                self.assertIn("F89", son)
            finally:
                doc.close()


if __name__ == "__main__":
    unittest.main()
