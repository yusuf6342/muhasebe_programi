"""Tahsilat makbuzu çıktısı: A5 / A4 ikili PDF ve Word, dosya adı, üzerine yazmama, önizleme ekranı.

Yalnız geçici test veritabanı kullanır (TahsilatMakbuzuTest kurulumu); çıktı almak kayıt değiştirmemelidir.
"""

from __future__ import annotations

import sys
import tempfile
import tkinter as tk
import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import func, select

from database.database import get_session
from database.finans_service import FinansService
from database.models.cari import CariIslem
from database.models.finans import FinansHareketi, KasaMakbuzu
from invoice_print.amount_to_words import amount_to_words
from tests import test_tahsilat_makbuzu as _servis_testi

import makbuz_cikti as mc

MM = 72 / 25.4
FIRMA = {
    "unvan": "ÖRNEK TEST MOBİLYA SANAYİ VE TİCARET LTD. ŞTİ.",
    "adres": "Atatürk Cad. No: 12",
    "ilce": "Çankaya",
    "il": "Ankara",
    "telefon": "0312 000 00 00",
    "vergi_dairesi": "Çankaya",
    "vergi_no": "0000000000",
}


def _fitz():
    return mc._fitz()


class _Ornekler:
    def _kaydet(self, satirlar, aciklama=None, tur="TAHSILAT"):
        veri = {
            "tarih": date(2026, 9, 28),
            "cari_id": self.musteri_id if tur == "TAHSILAT" else self.tedarikci_id,
            "makbuz_no_otomatik": True,
            "aciklama": aciklama,
            "satirlar": satirlar,
        }
        if tur == "TAHSILAT":
            return FinansService.kasa_tahsilat_makbuzu_kaydet(veri)
        return FinansService.kasa_odeme_makbuzu_kaydet(veri)

    def _ornekleri_olustur(self):
        self.nakit = self._kaydet(
            [{"odeme_sekli": "NAKİT / KASA", "finans_hesap_id": self.kasa_id, "tutar": "1250.50"}],
            "Eylül ayı ödemesi — elden nakit.",
        )
        self.banka = self._kaydet(
            [{"odeme_sekli": "GELEN HAVALE", "hesap": self.mevduat_adi, "tutar": "18400"}], "Havale ile tahsilat"
        )
        self.kart = self._kaydet(
            [
                {
                    "odeme_sekli": "KREDİ KARTIYLA TAHSİLAT",
                    "finans_hesap_id": self.pos_id,
                    "tutar": "7999.99",
                    "kart_tipi": "KREDI_KARTI",
                    "taksit_sayisi": 6,
                }
            ],
            "Mağaza POS — 6 taksit",
        )
        return [self.nakit.id, self.banka.id, self.kart.id]

    def _durum_ozeti(self):
        with get_session() as s:
            return (
                s.scalar(select(func.count(KasaMakbuzu.id))),
                s.scalar(select(func.count(CariIslem.id))),
                s.scalar(select(func.coalesce(func.sum(CariIslem.borc), 0))),
                s.scalar(select(func.coalesce(func.sum(CariIslem.alacak), 0))),
                s.scalar(select(func.count(FinansHareketi.id))),
                s.scalar(select(func.coalesce(func.sum(FinansHareketi.tutar), 0))),
                tuple(s.scalars(select(KasaMakbuzu.makbuz_no).order_by(KasaMakbuzu.id))),
                FinansService.makbuz_no_oner(),
            )


class MakbuzCiktiTest(_Ornekler, unittest.TestCase):
    def setUp(self):
        _servis_testi.TahsilatMakbuzuTest.setUp(self)
        self.firma = mc.firma_satirlari(FIRMA)
        self.idler = self._ornekleri_olustur()
        self.veriler = mc.makbuz_cikti_verileri(self.idler, firma=self.firma)
        self._cikti = tempfile.TemporaryDirectory()
        self.klasor = Path(self._cikti.name)

    def tearDown(self):
        self._cikti.cleanup()
        _servis_testi.TahsilatMakbuzuTest.tearDown(self)

    def _pdf(self, veriler, yerlesim):
        return _fitz().open(stream=mc.pdf_uret(veriler, yerlesim), filetype="pdf")

    def _sinir_icinde(self, page):
        r = page.rect
        for blok in page.get_text("dict")["blocks"]:
            for satir in blok.get("lines", []):
                for span in satir["spans"]:
                    if not span["text"].strip():
                        continue
                    x0, y0, x1, y1 = span["bbox"]
                    self.assertGreaterEqual(x0, r.x0 + 3 * MM, span["text"])
                    self.assertLessEqual(x1, r.x1 - 3 * MM, span["text"])
                    self.assertGreaterEqual(y0, r.y0 + 3 * MM, span["text"])
                    self.assertLessEqual(y1, r.y1 - 3 * MM, span["text"])

    # —— PDF ——
    def test_a5_pdf_boyutu_ve_zorunlu_alanlar(self):
        doc = self._pdf(self.veriler[:1], mc.A5)
        self.assertEqual(len(doc), 1)
        page = doc[0]
        self.assertAlmostEqual(page.rect.width, 148 * MM, places=1)
        self.assertAlmostEqual(page.rect.height, 210 * MM, places=1)
        metin = page.get_text()
        for beklenen in (
            FIRMA["unvan"],
            "Atatürk Cad. No: 12 Çankaya / Ankara",
            "V.D.: Çankaya",
            "TAHSİLAT MAKBUZU",
            "MKB-00001",
            "28.09.2026",
            "M001",
            "ALİ ÇELİK MOBİLYA LTD",
            "1.250,50 ₺",
            "Yalnız Bin İki Yüz Elli Türk Lirası Elli Kuruştur.",
            "NAKİT / KASA",
            "TEST KASA",
            "Eylül ayı ödemesi — elden nakit.",
            "TESLİM EDEN / TAHSİL EDEN",
            "ÖDEMEYİ YAPAN",
            "FİRMA KAŞESİ",
            "Adı Soyadı:",
            "İmza:",
        ):
            self.assertIn(beklenen, metin)
        self._sinir_icinde(page)

    def test_banka_ve_kredi_karti_hesap_bilgisi(self):
        doc = self._pdf(self.veriler, mc.A5)
        self.assertEqual(len(doc), 3)
        banka, kart = doc[1].get_text(), doc[2].get_text()
        self.assertIn("GELEN HAVALE", banka)
        self.assertIn("TEST BANK — MEVDUAT", banka)
        self.assertIn("18.400,00 ₺", banka)
        self.assertIn("KREDİ KARTIYLA TAHSİLAT", kart)
        self.assertIn("TEST BANK — POS", kart)
        self.assertIn("6 taksit", kart)
        self.assertIn("7.999,99 ₺", kart)
        for page in doc:
            self._sinir_icinde(page)

    def test_a4_ikili_sayfada_en_fazla_iki_farkli_makbuz(self):
        ek = [self._kaydet([{"odeme_sekli": "NAKİT / KASA", "finans_hesap_id": self.kasa_id, "tutar": "10"}]) for _ in range(2)]
        veriler = mc.makbuz_cikti_verileri(self.idler + [m.id for m in ek], firma=self.firma)
        doc = self._pdf(veriler, mc.A4_IKILI)
        self.assertEqual(len(doc), 3)
        gorulen = []
        for page in doc:
            self.assertAlmostEqual(page.rect.width, 297 * MM, places=1)
            self.assertAlmostEqual(page.rect.height, 210 * MM, places=1)
            nolar = sorted({t for t in page.get_text().split() if t.startswith("MKB-")})
            self.assertLessEqual(len(nolar), 2)
            gorulen += nolar
            self._sinir_icinde(page)
        self.assertEqual(gorulen, [f"MKB-0000{i}" for i in range(1, 6)])
        sol = doc[0].get_text(clip=_fitz().Rect(0, 0, 148.5 * MM, 210 * MM))
        sag = doc[0].get_text(clip=_fitz().Rect(148.5 * MM, 0, 297 * MM, 210 * MM))
        self.assertIn("MKB-00001", sol)
        self.assertNotIn("MKB-00002", sol)
        self.assertIn("MKB-00002", sag)

    def test_cok_satirli_makbuz_tek_sayfaya_sigar(self):
        satirlar = [
            {"odeme_sekli": "NAKİT / KASA", "finans_hesap_id": self.kasa_id, "tutar": str(100 + i), "aciklama": "Uzun satır açıklaması " * 3}
            for i in range(9)
        ]
        m = self._kaydet(satirlar, "Çok uzun açıklama " * 20)
        veriler = mc.makbuz_cikti_verileri([m.id], firma=self.firma)
        doc = self._pdf(veriler, mc.A5)
        self.assertEqual(len(doc), 1)
        metin = doc[0].get_text()
        self.assertIn("Diğer 3 ödeme satırı", metin)
        self.assertIn(mc.para(m.tutar), metin)
        self._sinir_icinde(doc[0])

    def test_iptal_makbuz_isaretlenir(self):
        FinansService.kasa_makbuz_iptal(self.nakit.id)
        v = mc.makbuz_cikti_verileri([self.nakit.id], firma=self.firma)
        metin = self._pdf(v, mc.A5)[0].get_text()
        self.assertIn("İPTAL EDİLMİŞTİR", metin)

    def test_odeme_makbuzu_basligi_ve_imza_alanlari(self):
        m = self._kaydet(
            [{"odeme_sekli": "NAKİT / KASA", "finans_hesap_id": self.kasa_id, "tutar": "250"}], tur="ODEME"
        )
        metin = self._pdf(mc.makbuz_cikti_verileri([m.id], firma=self.firma), mc.A5)[0].get_text()
        self.assertIn("ÖDEME MAKBUZU", metin)
        self.assertIn("ÖDENEN TUTAR", metin)
        self.assertIn("TESLİM ALAN / TAHSİL EDEN", metin)

    def test_eksik_firma_bilgisi_uydurulmaz(self):
        firma = mc.firma_satirlari({"unvan": "SADECE ÜNVAN A.Ş."})
        self.assertEqual(firma["satirlar"], [])
        self.assertIsNone(firma["logo_yolu"])
        metin = self._pdf(mc.makbuz_cikti_verileri([self.nakit.id], firma=firma), mc.A5)[0].get_text()
        self.assertIn("SADECE ÜNVAN A.Ş.", metin)
        for olmayan in ("Tel:", "V.D.:", "VKN", "MERSİS", "E-posta"):
            self.assertNotIn(olmayan, metin)

    # —— Word ——
    def test_word_a5_ve_a4_ikili_duzeni(self):
        from docx import Document
        from docx.oxml.ns import qn

        tek = Document(BytesIO(mc.docx_uret(self.veriler[:1], mc.A5)))
        s = tek.sections[0]
        self.assertAlmostEqual(s.page_width.mm, 148, delta=0.2)
        self.assertAlmostEqual(s.page_height.mm, 210, delta=0.2)
        metin = "\n".join(p.text for t in tek.tables for r in t.rows for c in r.cells for p in c.paragraphs)
        for beklenen in ("TAHSİLAT MAKBUZU", "MKB-00001", "1.250,50 ₺", "Türk Lirası Elli Kuruştur.", "FİRMA KAŞESİ"):
            self.assertIn(beklenen, metin)

        toplu = Document(BytesIO(mc.docx_uret(self.veriler, mc.A5)))
        xml = toplu.element.body.xml
        self.assertEqual(xml.count('w:type="page"'), 2)

        ikili = Document(BytesIO(mc.docx_uret(self.veriler, mc.A4_IKILI)))
        s = ikili.sections[0]
        self.assertAlmostEqual(s.page_width.mm, 297, delta=0.2)
        self.assertAlmostEqual(s.page_height.mm, 210, delta=0.2)
        cols = s._sectPr.find(qn("w:cols"))
        self.assertEqual(cols.get(qn("w:num")), "2")
        xml = ikili.element.body.xml
        self.assertEqual(xml.count('w:type="column"'), 2)
        self.assertEqual(xml.count('w:type="page"'), 0)

    # —— Dosya ——
    def test_dosya_adi_makbuz_numarasiyla(self):
        self.assertEqual(mc.dosya_adi(self.veriler[:1], mc.A5, "pdf"), "Tahsilat_Makbuzu_MKB-00001.pdf")
        self.assertEqual(
            mc.dosya_adi(self.veriler, mc.A4_IKILI, "docx"),
            "Tahsilat_Makbuzlari_MKB-00001_MKB-00003_3_adet_A4_ikili.docx",
        )

    def test_ayni_adli_dosyanin_uzerine_sormadan_yazilmaz(self):
        hedef = self.klasor / mc.dosya_adi(self.veriler[:1], mc.A5, "pdf")
        hedef.write_bytes(b"KULLANICI DOSYASI")
        with self.assertRaises(FileExistsError):
            mc.pdf_kaydet(self.veriler[:1], hedef)
        with self.assertRaises(FileExistsError):
            mc.docx_kaydet(self.veriler[:1], hedef)
        self.assertEqual(hedef.read_bytes(), b"KULLANICI DOSYASI")
        self.assertEqual(mc.bos_dosya_yolu(hedef).name, "Tahsilat_Makbuzu_MKB-00001 (2).pdf")
        mc.pdf_kaydet(self.veriler[:1], hedef, uzerine_yaz=True)
        self.assertTrue(hedef.read_bytes().startswith(b"%PDF"))

    # —— Güvenlik ——
    def test_cikti_almak_kayit_ve_bakiye_degistirmez(self):
        once = self._durum_ozeti()
        bakiye = FinansService.cari_bakiye_ozeti(self.musteri_id)
        for yer in (mc.A5, mc.A4_IKILI):
            veriler = mc.makbuz_cikti_verileri(self.idler)
            mc.pdf_kaydet(veriler, self.klasor / f"a_{yer}.pdf", yer)
            mc.docx_kaydet(veriler, self.klasor / f"a_{yer}.docx", yer)
            mc.onizleme_goruntuleri(mc.pdf_uret(veriler, yer), 300)
        self.assertEqual(self._durum_ozeti(), once)
        self.assertEqual(FinansService.cari_bakiye_ozeti(self.musteri_id), bakiye)


class TutarYaziylaTest(unittest.TestCase):
    def test_turkce_yazim_ve_ek(self):
        self.assertEqual(amount_to_words(Decimal("1250.50")), "Yalnız Bin İki Yüz Elli Türk Lirası Elli Kuruştur.")
        self.assertEqual(amount_to_words(Decimal("4000")), "Yalnız Dört Bin Türk Lirasıdır.")
        self.assertEqual(amount_to_words(Decimal("300")), "Yalnız Üç Yüz Türk Lirasıdır.")
        self.assertEqual(mc.para(Decimal("18400")), "18.400,00 ₺")
        self.assertEqual(mc.para(Decimal("7999.99")), "7.999,99 ₺")


class MakbuzCiktiEkranTest(_Ornekler, unittest.TestCase):
    def setUp(self):
        _servis_testi.TahsilatMakbuzuTest.setUp(self)
        try:
            self.root = tk.Tk()
        except tk.TclError as hata:
            _servis_testi.TahsilatMakbuzuTest.tearDown(self)
            self.skipTest(f"Tk yok: {hata}")
        self.root.withdraw()
        self.idler = self._ornekleri_olustur()
        self._cikti = tempfile.TemporaryDirectory()
        self.klasor = Path(self._cikti.name)
        import kasa_makbuz_ui
        import makbuz_cikti_ui

        self.ui = makbuz_cikti_ui
        self.mesajlar = []
        self._yamalar = [
            patch.object(mod.messagebox, ad, side_effect=lambda *a, _ad=ad, **k: self.mesajlar.append((_ad, a)))
            for mod in (kasa_makbuz_ui, makbuz_cikti_ui)
            for ad in ("showerror", "showinfo", "showwarning")
        ]
        self._yamalar.append(patch.object(makbuz_cikti_ui, "dosyayi_ac", side_effect=lambda yol: self.acilan.append(Path(yol))))
        self.acilan = []
        for y in self._yamalar:
            y.start()

    def tearDown(self):
        for y in self._yamalar:
            y.stop()
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        self._cikti.cleanup()
        _servis_testi.TahsilatMakbuzuTest.tearDown(self)

    def test_onizleme_yerlesim_sayfa_ve_ciktilar(self):
        once = self._durum_ozeti()
        d = self.ui.MakbuzCiktiDialog(self.root, self.idler)
        self.root.update()
        d._ciz()
        self.assertEqual(d.sayfa_sayisi, 3)
        self.assertEqual(d.sayfa_lbl.cget("text"), "Sayfa 1 / 3")
        self.assertIsNotNone(d._resim)
        d.sayfa_git(1)
        self.assertEqual(d.sayfa, 1)
        d.yerlesim_degistir(mc.A4_IKILI)
        self.assertEqual(d.sayfa_sayisi, 2)
        self.assertIn("A4", d.ipucu_lbl.cget("text"))

        hedef = self.klasor / "secilen.pdf"
        with patch.object(self.ui.filedialog, "asksaveasfilename", return_value=str(hedef)) as kaydet, patch.object(
            self.ui.messagebox, "askyesno", return_value=False
        ):
            self.assertEqual(d.pdf_olustur(), hedef)
        self.assertTrue(kaydet.call_args.kwargs["confirmoverwrite"])
        self.assertEqual(
            kaydet.call_args.kwargs["initialfile"], "Tahsilat_Makbuzlari_MKB-00001_MKB-00003_3_adet_A4_ikili.pdf"
        )
        self.assertEqual(len(_fitz().open(hedef)), 2)

        word = self.klasor / "secilen.docx"
        with patch.object(self.ui.filedialog, "asksaveasfilename", return_value=str(word)), patch.object(
            self.ui.messagebox, "askyesno", return_value=False
        ):
            self.assertEqual(d.word_olustur(), word)
        self.assertTrue(word.read_bytes().startswith(b"PK"))

        with patch.object(self.ui.filedialog, "asksaveasfilename", return_value=""):
            self.assertIsNone(d.pdf_olustur())

        yazdirilan = d.yazdir()
        self.assertEqual(self.acilan[-1], yazdirilan)
        self.assertTrue(yazdirilan.read_bytes().startswith(b"%PDF"))
        yazdirilan.unlink(missing_ok=True)
        d.destroy()
        self.assertEqual(self._durum_ozeti(), once)
        self.assertEqual(self.mesajlar, [])

    def test_kaydedilmemis_makbuzda_once_kaydet_uyarisi(self):
        from kasa_makbuz_ui import KasaMakbuzDialog
        import kasa_makbuz_ui

        d = KasaMakbuzDialog(self.root, "TAHSILAT")
        self.root.update_idletasks()
        self.assertEqual(d.btn_yazdir.winfo_manager(), "pack")
        once = self._durum_ozeti()
        with patch.object(kasa_makbuz_ui.messagebox, "askyesno", return_value=False) as soru, patch.object(
            self.ui, "MakbuzCiktiDialog"
        ) as pencere:
            self.assertIsNone(d.yazdir())
        self.assertIn("henüz kaydedilmedi", soru.call_args.args[1])
        pencere.assert_not_called()
        self.assertEqual(self._durum_ozeti(), once)

        d.cari_var.set("ÇELİK")
        d._cari_arama.sec(0)
        d.tutar_var.set("75")
        self.assertTrue(d.satir_ekle(), self.mesajlar)
        with patch.object(kasa_makbuz_ui.messagebox, "askyesno", return_value=True), patch.object(
            self.ui, "MakbuzCiktiDialog"
        ) as pencere:
            d.yazdir()
        self.assertEqual(d.mod, "goruntule")
        pencere.assert_called_once_with(d, [d.makbuz_id], mc.A5)
        self.assertEqual(self._durum_ozeti()[0], once[0] + 1)

        with patch.object(self.ui, "MakbuzCiktiDialog") as pencere:
            d.yazdir()
        pencere.assert_called_once()
        d.destroy()

    def test_listeden_coklu_secim_sirasiyla_ciktiya_gider(self):
        import kasa_makbuz_ui
        import makbuz_cikti_ui

        root = self.root
        root.icerik = tk.Frame(root)
        root.icerik.pack(fill="both", expand=True)
        root.menu_dugmeleri = {}

        def temizle():
            for w in root.icerik.winfo_children():
                w.destroy()

        root._icerigi_temizle = temizle

        def hemen(_widget, is_, bitti, hata=None):
            bitti(is_())

        with patch("ui_bg.arka_planda", hemen):
            kasa_makbuz_ui.kasa_makbuzlari_sayfasi(root, makbuz_turu="TAHSILAT")

        def bul(w, sinif):
            if isinstance(w, sinif):
                yield w
            for c in w.winfo_children():
                yield from bul(c, sinif)

        from tkinter import ttk
        import time

        tablo = next(bul(root.icerik, ttk.Treeview))
        son = time.time() + 10
        while len(tablo.get_children()) < 3 and time.time() < son:
            root.update()
            time.sleep(0.02)
        self.assertEqual(str(tablo.cget("selectmode")), "extended")
        cocuklar = list(tablo.get_children())
        tablo.selection_set([cocuklar[2], cocuklar[0]])
        dugme = next(b for b in bul(root.icerik, tk.Button) if b.cget("text") == "Yazdır / PDF / Word")
        with patch.object(makbuz_cikti_ui, "makbuz_ciktisi_ac") as ac:
            dugme.invoke()
        ac.assert_called_once()
        self.assertEqual(ac.call_args.args[1], [int(cocuklar[0]), int(cocuklar[2])])


if __name__ == "__main__":
    unittest.main()
