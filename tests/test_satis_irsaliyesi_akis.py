"""Satış irsaliyesi uçtan uca akış testleri (izole geçici SQLite).

Kayıt/yeniden açma, tek stok çıkışı (temel birim), sevk geri alma, silme/iptal,
siparişten kısmi aktarım kalanları, fatura bağı kilitleri ve liste filtreleri.
"""

from __future__ import annotations

import sys
import unittest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import func, select

from database.database import get_session
from database.models.satis_irsaliyesi import SatisIrsaliyesi, SatisIrsaliyesiSatiri
from database.models.satis_siparisi import SatisSiparisiSatiri
from database.models.stok import StokHareketi, StokLotu
from database.satis_irsaliyesi_service import SatisIrsaliyesiService
from tests.test_fatura_miktar_birim import FaturaStokEntegrasyonTest

TARIH = date(2026, 3, 10)


class _IrsaliyeTemel(unittest.TestCase):
    """İzole SQLite + ortak yardımcılar (test içermez)."""

    setUp = FaturaStokEntegrasyonTest.setUp
    tearDown = FaturaStokEntegrasyonTest.tearDown

    # ---- yardımcılar -------------------------------------------------
    def _veri(self, **ek):
        veri = {"irsaliye_tarihi": TARIH, "cari_id": self.cari_id, "depo": "ANA DEPO"}
        veri.update(ek)
        return veri

    @staticmethod
    def _satir(miktar, birim="Adet", **ek):
        satir = {
            "urun_kodu": "MB001",
            "urun_adi": "Birim Ürün",
            "miktar": Decimal(str(miktar)),
            "birim": birim,
            "birim_fiyat": Decimal("1"),
            "kdv_orani": Decimal("20"),
        }
        satir.update(ek)
        return satir

    def _hareket_toplam(self, tur: str, belge_no: str | None = None) -> Decimal:
        with get_session() as session:
            q = select(func.coalesce(func.sum(StokHareketi.miktar), 0)).where(
                StokHareketi.hareket_turu == tur
            )
            if belge_no:
                q = q.where(StokHareketi.belge_no == belge_no)
            return Decimal(str(session.scalar(q) or 0))

    def _lot_kalan(self) -> Decimal:
        with get_session() as session:
            return Decimal(
                str(session.scalar(select(func.coalesce(func.sum(StokLotu.kalan_miktar), 0))) or 0)
            )

    def _siparis(self, miktar=10, birim="Paket", onayla=True, cari_id=None):
        from database.satis_siparisi_service import SatisSiparisiService

        return SatisSiparisiService.kaydet(
            {
                "siparis_tarihi": date(2026, 3, 1),
                "termin_tarihi": date(2026, 3, 20),
                "cari_id": cari_id or self.cari_id,
                "onayla": onayla,
            },
            [
                {
                    "urun_kodu": "MB001",
                    "urun_adi": "Birim Ürün",
                    "miktar": Decimal(str(miktar)),
                    "birim": birim,
                    "birim_satis_fiyati": Decimal("100"),
                    "kdv_orani": Decimal("20"),
                }
            ],
            [],
        )

    def _sevk_miktari(self, siparis_satiri_id) -> Decimal:
        with get_session() as session:
            return Decimal(
                str(session.get(SatisSiparisiSatiri, siparis_satiri_id).irsaliyelenen_miktar or 0)
            )

    def _fatura_kaydet(self, irsaliye, miktar):
        from database.satis_faturasi_service import SatisFaturasiService

        satir = irsaliye.satirlar[0]
        return SatisFaturasiService.kaydet(
            {
                "fatura_tarihi": TARIH,
                "vade_tarihi": TARIH,
                "cari_id": self.cari_id,
                "depo": "ANA DEPO",
                "irsaliye_id": irsaliye.id,
                "odeme_tutari": Decimal("0"),
                "sales_person_id": 1,
            },
            [
                {
                    "urun_kodu": satir.urun_kodu,
                    "urun_adi": satir.urun_adi,
                    "miktar": Decimal(str(miktar)),
                    "birim": satir.birim,
                    "birim_fiyat": Decimal("2400"),
                    "iskonto_orani": Decimal("0"),
                    "kdv_orani": Decimal("20"),
                    "irsaliye_satiri_id": satir.id,
                    "siparis_satiri_id": satir.siparis_satiri_id,
                }
            ],
        )


class SatisIrsaliyesiAkisTest(_IrsaliyeTemel):
    # ---- kayıt / doğrulama -------------------------------------------
    def test_eksik_alan_mesajlari(self):
        with self.assertRaisesRegex(ValueError, "Müşteri"):
            SatisIrsaliyesiService.kaydet(self._veri(cari_id=None), [self._satir(1)])
        with self.assertRaisesRegex(ValueError, "satır"):
            SatisIrsaliyesiService.kaydet(self._veri(), [])
        with self.assertRaisesRegex(ValueError, "1. satır"):
            SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(0)])
        with self.assertRaisesRegex(ValueError, "gelecek"):
            SatisIrsaliyesiService.kaydet(
                self._veri(irsaliye_tarihi=date.today() + timedelta(days=3)), [self._satir(1)]
            )
        with self.assertRaisesRegex(ValueError, "tarih"):
            SatisIrsaliyesiService.kaydet(self._veri(irsaliye_tarihi=None), [self._satir(1)])
        with get_session() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(SatisIrsaliyesi)), 0)

    def test_kaydet_yeniden_ac_ayni_ve_mukerrer_yok(self):
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(aciklama="Test açıklama", sevk_adresi="Depo kapısı 3"),
            [self._satir(2, "Koli"), self._satir(5, "Paket")],
        )
        self.assertEqual(irs.durum, "TASLAK")
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), 0)
        acilan = SatisIrsaliyesiService.getir(irs.id)
        self.assertEqual(acilan.aciklama, "Test açıklama")
        self.assertEqual(acilan.sevk_adresi, "Depo kapısı 3")
        self.assertEqual(
            [(s.birim, s.miktar) for s in acilan.satirlar],
            [("Koli", Decimal("2")), ("Paket", Decimal("5"))],
        )
        ilk_idler = [s.id for s in acilan.satirlar]
        # aynı kaydı tekrar kaydet (satır id'leri ile) — mükerrer satır/belge oluşmaz
        satirlar = [
            self._satir(s.miktar, s.birim, irsaliye_satiri_id=s.id) for s in acilan.satirlar
        ]
        satirlar[1]["miktar"] = Decimal("7")
        ikinci = SatisIrsaliyesiService.kaydet(self._veri(), satirlar, irsaliye_id=irs.id)
        self.assertEqual([s.id for s in ikinci.satirlar], ilk_idler)
        self.assertEqual(ikinci.satirlar[1].miktar, Decimal("7"))
        with get_session() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(SatisIrsaliyesi)), 1)
            self.assertEqual(
                session.scalar(select(func.count()).select_from(SatisIrsaliyesiSatiri)), 2
            )

    def test_ayni_irsaliye_no_ikinci_kez_kaydedilemez(self):
        SatisIrsaliyesiService.kaydet(self._veri(irsaliye_no="IRS-X1"), [self._satir(1)])
        with self.assertRaisesRegex(ValueError, "IRS-X1"):
            SatisIrsaliyesiService.kaydet(self._veri(irsaliye_no="IRS-X1"), [self._satir(1)])

    # ---- stok --------------------------------------------------------
    def test_sevk_tek_cikis_temel_birim_ve_geri_al(self):
        irs = SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(2, "Koli")])
        SatisIrsaliyesiService.sevk_et(irs.id)
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ", irs.irsaliye_no), Decimal("4800"))
        self.assertEqual(self._lot_kalan(), Decimal("200"))
        with self.assertRaises(ValueError):
            SatisIrsaliyesiService.sevk_et(irs.id)
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), Decimal("4800"))
        # sevkli belge düzenlenemez
        with self.assertRaisesRegex(ValueError, "Sevki Geri Al"):
            SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(1, "Koli")], irsaliye_id=irs.id)
        geri = SatisIrsaliyesiService.sevk_geri_al(irs.id)
        self.assertEqual(geri.durum, "TASLAK")
        self.assertFalse(geri.stok_cikis_yapildi)
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), 0)
        self.assertEqual(self._lot_kalan(), Decimal("5000"))
        duz = SatisIrsaliyesiService.kaydet(
            self._veri(),
            [self._satir(1, "Koli", irsaliye_satiri_id=geri.satirlar[0].id)],
            irsaliye_id=irs.id,
        )
        SatisIrsaliyesiService.sevk_et(duz.id)
        self.assertEqual(self._lot_kalan(), Decimal("2600"))

    def test_yetersiz_stokta_sevk_atomik(self):
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(), [self._satir(1000, "Adet"), self._satir(3, "Koli")]
        )
        with self.assertRaisesRegex(ValueError, "Sevk yapılamadı"):
            SatisIrsaliyesiService.sevk_et(irs.id)
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), 0)
        self.assertEqual(self._lot_kalan(), Decimal("5000"))
        self.assertFalse(SatisIrsaliyesiService.getir(irs.id).stok_cikis_yapildi)

    def test_iptal_ve_sil_stogu_geri_verir(self):
        a = SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(10, "Paket")])
        SatisIrsaliyesiService.sevk_et(a.id)
        self.assertEqual(self._lot_kalan(), Decimal("4000"))
        with self.assertRaisesRegex(ValueError, "nedeni"):
            SatisIrsaliyesiService.iptal_et(a.id, "")
        SatisIrsaliyesiService.iptal_et(a.id, "Müşteri vazgeçti")
        self.assertEqual(self._lot_kalan(), Decimal("5000"))
        self.assertEqual(SatisIrsaliyesiService.getir(a.id).durum, "İPTAL")

        b = SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(5, "Paket")])
        SatisIrsaliyesiService.sevk_et(b.id)
        no = SatisIrsaliyesiService.sil(b.id, "Yanlış kayıt")
        self.assertEqual(no, b.irsaliye_no)
        self.assertIsNone(SatisIrsaliyesiService.getir(b.id))
        self.assertEqual(self._lot_kalan(), Decimal("5000"))
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), 0)
        with get_session() as session:
            self.assertEqual(
                session.scalar(
                    select(func.count())
                    .select_from(SatisIrsaliyesiSatiri)
                    .where(SatisIrsaliyesiSatiri.irsaliye_id == b.id)
                ),
                0,
            )

    # ---- siparişten aktarım ------------------------------------------
    def test_siparisten_kismi_aktarim_kalanlar(self):
        siparis = self._siparis(10, "Paket")
        ssid = siparis.satirlar[0].id
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(siparis_id=siparis.id), [self._satir(4, "Paket", siparis_satiri_id=ssid)]
        )
        self.assertEqual(self._sevk_miktari(ssid), Decimal("4"))
        # düzenlemede artırma — kendi payı kalan sayılır
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(),
            [self._satir(6, "Paket", siparis_satiri_id=ssid, irsaliye_satiri_id=irs.satirlar[0].id)],
            irsaliye_id=irs.id,
        )
        self.assertEqual(self._sevk_miktari(ssid), Decimal("6"))
        # fazla sevk engeli (hem aynı belgede hem ikinci belgede)
        with self.assertRaisesRegex(ValueError, "kalan sevk"):
            SatisIrsaliyesiService.kaydet(
                self._veri(),
                [self._satir(11, "Paket", siparis_satiri_id=ssid, irsaliye_satiri_id=irs.satirlar[0].id)],
                irsaliye_id=irs.id,
            )
        with self.assertRaisesRegex(ValueError, "kalan sevk"):
            SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(5, "Paket", siparis_satiri_id=ssid)])
        self.assertEqual(self._sevk_miktari(ssid), Decimal("6"))
        # farklı birim engeli
        with self.assertRaisesRegex(ValueError, "birim"):
            SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(1, "Adet", siparis_satiri_id=ssid)])
        # sipariş satırını irsaliyeden çıkar → kalan geri açılır
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(), [self._satir(3, "Adet")], irsaliye_id=irs.id
        )
        self.assertEqual(self._sevk_miktari(ssid), Decimal("0"))
        # kalan kadar ikinci irsaliye, sonra iptal ve silme sayaçları geri alır
        ikinci = SatisIrsaliyesiService.kaydet(
            self._veri(), [self._satir(10, "Paket", siparis_satiri_id=ssid)]
        )
        self.assertEqual(self._sevk_miktari(ssid), Decimal("10"))
        SatisIrsaliyesiService.iptal_et(ikinci.id, "test")
        self.assertEqual(self._sevk_miktari(ssid), Decimal("0"))
        ucuncu = SatisIrsaliyesiService.kaydet(
            self._veri(), [self._satir(2, "Paket", siparis_satiri_id=ssid)]
        )
        SatisIrsaliyesiService.sil(ucuncu.id)
        self.assertEqual(self._sevk_miktari(ssid), Decimal("0"))
        # iptal edilmiş belge silinirse sayaç ikinci kez düşmez
        dorduncu = SatisIrsaliyesiService.kaydet(
            self._veri(), [self._satir(4, "Paket", siparis_satiri_id=ssid)]
        )
        SatisIrsaliyesiService.iptal_et(ikinci.id, "tekrar")  # zaten iptal: etkisiz
        SatisIrsaliyesiService.sil(ikinci.id)
        self.assertEqual(self._sevk_miktari(ssid), Decimal("4"))
        SatisIrsaliyesiService.sil(dorduncu.id)

    def test_taslak_ve_baska_musteri_siparisi_aktarilamaz(self):
        taslak = self._siparis(5, "Paket", onayla=False)
        with self.assertRaisesRegex(ValueError, "TASLAK"):
            SatisIrsaliyesiService.kaydet(
                self._veri(), [self._satir(1, "Paket", siparis_satiri_id=taslak.satirlar[0].id)]
            )
        from database.models.cari import Cari

        with get_session() as session:
            diger = Cari(cari_kodu="M002", unvan="Diğer Müşteri", cari_turu="Müşteri", aktif=True)
            session.add(diger)
            session.flush()
            diger_id = diger.id
        onayli_diger = self._siparis(5, "Paket", cari_id=diger_id)
        with self.assertRaisesRegex(ValueError, "başka müşteri"):
            SatisIrsaliyesiService.kaydet(
                self._veri(), [self._satir(1, "Paket", siparis_satiri_id=onayli_diger.satirlar[0].id)]
            )
        adaylar = SatisIrsaliyesiService.siparis_aktarim_adaylari(self.cari_id)
        self.assertEqual(adaylar, [])
        onayli = self._siparis(8, "Paket")
        adaylar = SatisIrsaliyesiService.siparis_aktarim_adaylari(self.cari_id)
        self.assertEqual([a["siparis_id"] for a in adaylar], [onayli.id])
        self.assertEqual(adaylar[0]["kalan_miktar"], Decimal("8"))

    def test_aktarim_secim_kalani_form_ve_kayit_payi(self):
        from database.kalan_belge_service import siparis_irsaliye_secim_satirlari
        from database.satis_siparisi_service import SatisSiparisiService

        siparis = self._siparis(10, "Paket")
        ssid = siparis.satirlar[0].id
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(), [self._satir(4, "Paket", siparis_satiri_id=ssid)]
        )
        adaylar = SatisIrsaliyesiService.siparis_aktarim_adaylari(self.cari_id, irsaliye_id=irs.id)
        self.assertEqual(adaylar[0]["kalan_miktar"], Decimal("10"))
        guncel = SatisSiparisiService.getir(siparis.id)
        # formda hâlâ 4 duruyor → yalnız 6 aktarılabilir
        secim = siparis_irsaliye_secim_satirlari(
            guncel,
            form_satirlar=[{"siparis_satiri_id": ssid, "miktar": Decimal("4")}],
            bu_irsaliye_db_satirlar=irs.satirlar,
        )
        self.assertEqual(secim[0]["kalan_miktar"], Decimal("6"))
        # formdan silinmiş (henüz kaydedilmemiş) → 10 aktarılabilir
        secim = siparis_irsaliye_secim_satirlari(
            guncel, form_satirlar=[], bu_irsaliye_db_satirlar=irs.satirlar
        )
        self.assertEqual(secim[0]["kalan_miktar"], Decimal("10"))
        # yeni (kaydedilmemiş) irsaliye formu: kayıtlı pay yok
        secim = siparis_irsaliye_secim_satirlari(guncel, form_satirlar=[])
        self.assertEqual(secim[0]["kalan_miktar"], Decimal("6"))

    # ---- fatura bağı -------------------------------------------------
    def test_faturali_irsaliye_kilitli_ve_ikinci_stok_cikisi_yok(self):
        from database.satis_faturasi_service import SatisFaturasiService

        irs = SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(2, "Koli")])
        SatisIrsaliyesiService.sevk_et(irs.id)
        irs = SatisIrsaliyesiService.getir(irs.id)
        fatura = self._fatura_kaydet(irs, 1)
        self.assertEqual(SatisIrsaliyesiService.getir(irs.id).durum, "KISMİ FATURALANDI")
        SatisFaturasiService.onayla(fatura.id)
        self.assertEqual(self._hareket_toplam("FATURA ÇIKIŞ"), 0)
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), Decimal("4800"))
        # kalan 1 Koli'den fazlası faturalanamaz
        with self.assertRaisesRegex(ValueError, "kalan"):
            self._fatura_kaydet(irs, 2)
        for islem, beklenen in (
            (lambda: SatisIrsaliyesiService.sil(irs.id), "silinemez"),
            (lambda: SatisIrsaliyesiService.iptal_et(irs.id, "x"), "iptal edilemez"),
            (lambda: SatisIrsaliyesiService.sevk_geri_al(irs.id), "geri alınamaz"),
            (
                lambda: SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(1, "Koli")], irsaliye_id=irs.id),
                "Sevki Geri Al|değiştirilemez",
            ),
        ):
            with self.assertRaisesRegex(ValueError, beklenen) as ctx:
                islem()
            if beklenen != "Sevki Geri Al|değiştirilemez":
                self.assertIn(fatura.fatura_no, str(ctx.exception))
        bagli = SatisIrsaliyesiService.bagli_faturalar(irs.id)
        self.assertEqual([b["fatura_no"] for b in bagli], [fatura.fatura_no])
        liste = SatisIrsaliyesiService.listele(evrak_no=irs.irsaliye_no)
        self.assertEqual(liste[0]["faturalama_durumu"], "Kısmen faturalandı")
        self.assertEqual(liste[0]["fatura_nolari"], [fatura.fatura_no])
        # fatura iptal → irsaliye sevk durumuna döner, stok irsaliyede kalır
        SatisFaturasiService.iptal_et(fatura.id, "test iptal")
        geri = SatisIrsaliyesiService.getir(irs.id)
        self.assertEqual(geri.durum, "SEVK EDİLDİ")
        self.assertEqual(geri.satirlar[0].faturalanan_miktar, 0)
        self.assertEqual(self._lot_kalan(), Decimal("200"))
        # iptal faturası satıra hâlâ bağlı: silme engelli, iptal serbest
        with self.assertRaisesRegex(ValueError, "silinemez"):
            SatisIrsaliyesiService.sil(irs.id)
        SatisIrsaliyesiService.iptal_et(irs.id, "fatura iptal sonrası")
        self.assertEqual(self._lot_kalan(), Decimal("5000"))

    def test_sevksiz_faturalanan_irsaliye_ikinci_kez_stok_dusmez(self):
        from database.satis_faturasi_service import SatisFaturasiService

        irs = SatisIrsaliyesiService.kaydet(self._veri(), [self._satir(1, "Koli")])
        irs = SatisIrsaliyesiService.getir(irs.id)
        fatura = self._fatura_kaydet(irs, 1)
        SatisFaturasiService.onayla(fatura.id)
        self.assertEqual(self._hareket_toplam("FATURA ÇIKIŞ"), Decimal("2400"))
        with self.assertRaisesRegex(ValueError, fatura.fatura_no):
            SatisIrsaliyesiService.sevk_et(irs.id)
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), 0)
        self.assertEqual(self._lot_kalan(), Decimal("2600"))
        self.assertEqual(SatisIrsaliyesiService.getir(irs.id).durum, "FATURALANDI")

    # ---- liste -------------------------------------------------------
    def test_liste_filtreleri(self):
        a = SatisIrsaliyesiService.kaydet(self._veri(irsaliye_no="IRS-A1"), [self._satir(1)])
        b = SatisIrsaliyesiService.kaydet(
            self._veri(irsaliye_no="IRS-B2", irsaliye_tarihi=date(2026, 4, 5)), [self._satir(2)]
        )
        SatisIrsaliyesiService.sevk_et(b.id)
        self.assertEqual(len(SatisIrsaliyesiService.listele()), 2)
        self.assertEqual(
            [r["irsaliye"].irsaliye_no for r in SatisIrsaliyesiService.listele(evrak_no="b2")],
            ["IRS-B2"],
        )
        self.assertEqual(
            [r["irsaliye"].id for r in SatisIrsaliyesiService.listele(baslangic=date(2026, 4, 1))],
            [b.id],
        )
        self.assertEqual(
            [r["irsaliye"].id for r in SatisIrsaliyesiService.listele(bitis=date(2026, 3, 31))],
            [a.id],
        )
        self.assertEqual(
            [r["irsaliye"].id for r in SatisIrsaliyesiService.listele(durum="SEVK EDİLDİ")],
            [b.id],
        )
        self.assertEqual(len(SatisIrsaliyesiService.listele(musteri="m001")), 2)
        self.assertEqual(SatisIrsaliyesiService.listele(musteri="yok-böyle"), [])
        satir = SatisIrsaliyesiService.listele(evrak_no="IRS-A1")[0]
        self.assertEqual(satir["musteri_kodu"], "M001")
        self.assertEqual(satir["depo"], "ANA DEPO")
        self.assertEqual(satir["faturalama_durumu"], "Faturalanmadı")


class _EkranTemel(_IrsaliyeTemel):
    """Gizli Tk kökü + mesaj kutuları yamalı (test içermez)."""

    def setUp(self):
        super().setUp()
        self.root = self._kok_olustur()
        self._yamalar = [
            patch(f"tkinter.messagebox.{ad}", return_value=None)
            for ad in ("showinfo", "showwarning", "showerror")
        ] + [
            patch("tkinter.messagebox.askyesno", return_value=True),
            patch("tkinter.messagebox.askyesnocancel", return_value=False),
        ]
        self.mesajlar = [p.start() for p in self._yamalar]

    def tearDown(self):
        for p in self._yamalar:
            p.stop()
        try:
            self.root.destroy()
        except Exception:
            pass
        super().tearDown()

    def _kok_olustur(self):
        import tkinter as tk

        kok = tk.Tk()
        kok.withdraw()
        return kok

    def _dialog(self, **ek):
        from satis_irsaliyesi_ui import SatisIrsaliyesiDialog

        d = SatisIrsaliyesiDialog(self.root, **ek)
        d.withdraw()
        return d

    def _cari(self):
        from database.cari_service import CariService

        return CariService.getir(self.cari_id)

    def _urun_ekle(self, d, miktar="2", birim="Koli"):
        d.urun_secildi(("MB001", "Birim Ürün", "Adet", "", "1"))
        d.satir_girdileri["miktar"].delete(0, "end")
        d.satir_girdileri["miktar"].insert(0, miktar)
        d.satir_girdileri["birim"].set(birim)
        d.satir_ekle()


class SatisIrsaliyesiEkranTest(_EkranTemel):
    """Tk ekranı: boş açılış, kayıt/yeniden açma, kilit, değişiklik izi, çıktı."""

    def test_menu_yeni_irsaliye_her_seferinde_bos(self):
        d = self._dialog()
        d._musteri_ata(self._cari())
        self._urun_ekle(d)
        self.assertTrue(d.kaydet(sessiz=True))
        d.destroy()
        yeni = self._dialog()
        self.assertIsNone(yeni.irsaliye)
        self.assertIsNone(yeni.cari)
        self.assertEqual(yeni.satirlar, [])
        self.assertEqual(yeni._musteri_var.get(), "")
        self.assertEqual(yeni.girdiler["irsaliye_tarihi"].get(), date.today().strftime("%d.%m.%Y"))
        self.assertFalse(yeni.degisiklik_var())
        yeni.destroy()

    def test_birim_secenekleri_stok_kartindan(self):
        d = self._dialog()
        d.urun_secildi(("MB001", "Birim Ürün", "Adet", "", "1"))
        self.assertEqual(list(d.satir_girdileri["birim"].cget("values")), ["Adet", "Paket", "Koli"])
        d.destroy()

    def test_eksik_alanda_kaydetmez_ve_kayitli_gibi_davranmaz(self):
        d = self._dialog()
        self._urun_ekle(d)
        self.assertFalse(d.kaydet(sessiz=True))
        self.assertIsNone(d.irsaliye)
        self.assertIn("Müşteri", str(self.mesajlar[1].call_args))
        d._musteri_ata(self._cari())
        d.girdiler["irsaliye_tarihi"].delete(0, "end")
        d.girdiler["irsaliye_tarihi"].insert(0, "31.31.2026")
        self.assertFalse(d.kaydet(sessiz=True))
        self.assertIn("geçersiz", str(self.mesajlar[1].call_args))
        with get_session() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(SatisIrsaliyesi)), 0)
        d.destroy()

    def test_kaydet_yeniden_ac_ayni_kilit_ve_degisiklik_izi(self):
        d = self._dialog()
        d._musteri_ata(self._cari())
        d.girdiler["sevk_adresi"].delete(0, "end")
        d.girdiler["sevk_adresi"].insert(0, "Teslimat Sok. 5")
        d.girdiler["aciklama"].insert(0, "Ekran testi")
        self._urun_ekle(d, "1", "Koli")
        self._urun_ekle(d, "3", "Paket")
        self.assertTrue(d.degisiklik_var())
        self.assertTrue(d.kaydet(sessiz=True))
        self.assertFalse(d.degisiklik_var())
        self.assertTrue(d.kaydet(sessiz=True))  # tekrar kaydet: mükerrer yok
        iid = d.irsaliye.id
        d.destroy()
        with get_session() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(SatisIrsaliyesi)), 1)
            self.assertEqual(session.scalar(select(func.count()).select_from(SatisIrsaliyesiSatiri)), 2)
        acik = self._dialog(irsaliye=SatisIrsaliyesiService.getir(iid))
        self.assertEqual(acik.girdiler["sevk_adresi"].get(), "Teslimat Sok. 5")
        self.assertEqual(acik.girdiler["aciklama"].get(), "Ekran testi")
        self.assertEqual(
            [(s["birim"], s["miktar"]) for s in acik.satirlar],
            [("Koli", Decimal("1")), ("Paket", Decimal("3"))],
        )
        self.assertTrue(all(s["irsaliye_satiri_id"] for s in acik.satirlar))
        self.assertFalse(acik.degisiklik_var())
        acik.girdiler["aciklama"].insert("end", " değişti")
        self.assertTrue(acik.degisiklik_var())
        acik.sevk_et()  # değişiklik önce kaydedilir, sonra sevk
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), Decimal("2700"))
        self.assertEqual(SatisIrsaliyesiService.getir(iid).aciklama, "Ekran testi değişti")
        self.assertEqual(str(acik.btn["kaydet"].cget("state")), "disabled")
        self.assertEqual(str(acik.btn["sevk"].cget("state")), "disabled")
        self.assertEqual(str(acik.btn["sevk_geri"].cget("state")), "normal")
        self.assertEqual(str(acik.satir_girdileri["miktar"].cget("state")), "disabled")
        self.assertFalse(acik.degisiklik_var())
        acik.sevk_geri_al()
        self.assertEqual(self._hareket_toplam("İRSALİYE ÇIKIŞ"), 0)
        self.assertEqual(str(acik.btn["kaydet"].cget("state")), "normal")
        acik.destroy()

    def test_siparis_satiri_birimi_kilitli_ve_liste_siparis_no(self):
        siparis = self._siparis(10, "Paket")
        ssid = siparis.satirlar[0].id
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(siparis_id=siparis.id), [self._satir(4, "Paket", siparis_satiri_id=ssid)]
        )
        d = self._dialog(irsaliye=irs)
        self.assertEqual(d.satirlar[0]["siparis_no"], siparis.siparis_no)
        self.assertEqual(d._siparis_var.get(), siparis.siparis_no)
        d.satir_tablosu.selection_set("0")
        d.satir_duzenle()
        self.assertEqual(list(d.satir_girdileri["birim"].cget("values")), ["Paket"])
        d.satir_girdileri["miktar"].delete(0, "end")
        d.satir_girdileri["miktar"].insert(0, "6")
        d.satir_ekle()
        self.assertTrue(d.kaydet(sessiz=True))
        self.assertEqual(self._sevk_miktari(ssid), Decimal("6"))
        d.destroy()

    def test_cikti_html_toplam_miktar_ve_sayfa_kurallari(self):
        from database.irsaliye_customer_view import (
            build_customer_dispatch_from_dialog,
            render_customer_dispatch_html,
        )

        d = self._dialog()
        d._musteri_ata(self._cari())
        d.girdiler["sevk_adresi"].delete(0, "end")
        d.girdiler["sevk_adresi"].insert(0, "Depo Kapısı 7")
        for _ in range(40):
            self._urun_ekle(d, "1", "Paket")
        self._urun_ekle(d, "2", "Koli")
        html = render_customer_dispatch_html(build_customer_dispatch_from_dialog(d))
        self.assertIn("Test Müşteri", html)
        self.assertIn("Depo Kapısı 7", html)
        self.assertIn("Toplam Miktar: 40 Paket · 2 Koli", html)
        self.assertIn("<td class='c'>Koli</td><td class='r miktar'>2</td>", html)
        from database.irsaliye_customer_view import miktar_metni

        self.assertEqual(miktar_metni(Decimal("10")), "10")
        self.assertEqual(miktar_metni(Decimal("2.5000")), "2,5")
        self.assertIn("display:table-header-group", html)
        self.assertIn("page-break-inside:avoid", html)
        self.assertIn("#102A43", html)
        self.assertIn("#F4C542", html)
        d.destroy()

    def test_pdf_gercek_dosya_uretir(self):
        import re
        import tempfile

        from invoice_print.pdf_service import _chrome_edge_paths

        if not _chrome_edge_paths():
            self.skipTest("Edge/Chrome yok")
        d = self._dialog()
        d._musteri_ata(self._cari())
        for _ in range(60):
            self._urun_ekle(d, "1", "Adet")
        hedef = Path(tempfile.mkdtemp()) / "irsaliye.pdf"
        with patch("irsaliye_cikti_ui.filedialog.asksaveasfilename", return_value=str(hedef)) as sor:
            d.pdf_kaydet()
        self.assertIsNotNone(d.irsaliye, "çıktıdan önce kaydedilmeli")
        self.assertTrue(sor.call_args.kwargs["confirmoverwrite"])
        self.assertRegex(sor.call_args.kwargs["initialfile"], r"^Satis_Irsaliyesi_.+_\d{4}-\d{2}-\d{2}\.pdf$")
        self.assertTrue(hedef.is_file())
        icerik = hedef.read_bytes()
        self.assertTrue(icerik.startswith(b"%PDF"))
        self.assertGreaterEqual(len(re.findall(rb"/Type\s*/Page[^s]", icerik)), 2)
        d.destroy()


def _tum_tablolar_ozeti() -> dict[str, str]:
    """Veritabanındaki her tablonun tam içeriği (çıktı almanın kayıt değiştirmediğini kanıtlamak için)."""
    from database.database import Base

    with get_session() as session:
        return {
            t.name: repr(sorted(repr(tuple(r)) for r in session.execute(select(t)).all()))
            for t in Base.metadata.sorted_tables
            if session.bind is not None and _tablo_var(session, t.name)
        }


def _tablo_var(session, ad: str) -> bool:
    from sqlalchemy import inspect

    return ad in inspect(session.bind).get_table_names()


class SatisIrsaliyesiListeTest(_EkranTemel):
    """SATIŞ İRSALİYESİ LİSTESİ: tek tıkla kimlikle açma, filtre/sıralama, hata günlüğü."""

    def _kok_olustur(self):
        import tkinter as tk

        import app as app_mod

        class _ListeKoku(tk.Tk):
            def __init__(self):
                super().__init__()
                self.geometry("1500x700+-4000+-4000")
                self.icerik = tk.Frame(self)
                self.icerik.pack(fill="both", expand=True)

            def _icerigi_temizle(self):
                for w in self.icerik.winfo_children():
                    w.destroy()

            def nav_sayfa_isaretle(self, *_a):
                pass

        for ad in (
            "satis_irsaliyeleri_goster", "_irsaliye_tik_satiri", "_irsaliye_tik_basildi",
            "_irsaliye_tik_birakildi", "_irsaliye_sirala", "irsaliye_cikti", "_irsaliye_filtreleri",
            "irsaliye_listesini_yenile", "_secili_irsaliye_id", "irsaliye_ac", "yeni_irsaliye",
            "irsaliye_sevk", "irsaliye_faturaya_cevir", "irsaliye_iptal", "irsaliye_sil",
        ):
            setattr(_ListeKoku, ad, getattr(app_mod.MuhasebeApp, ad))
        return _ListeKoku()

    def setUp(self):
        super().setUp()
        self.l1 = SatisIrsaliyesiService.kaydet(
            self._veri(irsaliye_no="IRS-L1", irsaliye_tarihi=date(2026, 3, 1)), [self._satir(1)]
        )
        self.l2 = SatisIrsaliyesiService.kaydet(
            self._veri(irsaliye_no="IRS-L2", irsaliye_tarihi=date(2026, 3, 5)), [self._satir(2)]
        )
        self.l3 = SatisIrsaliyesiService.kaydet(
            self._veri(irsaliye_no="IRS-L10", irsaliye_tarihi=date(2026, 3, 3)), [self._satir(3)]
        )
        self.root.satis_irsaliyeleri_goster()
        self.root._irs_filtre["baslangic"].delete(0, "end")
        self.root._irs_filtre["bitis"].delete(0, "end")
        self.root.irsaliye_listesini_yenile()
        self.root.update()
        self.tablo = self.root.irsaliye_tablosu
        self.acilan: list[int] = []

    def _kaydedici(self):
        acilan = self.acilan

        class _Sahte:
            def __init__(self, _parent, irsaliye=None, **_kw):
                acilan.append(int(irsaliye.id))

            def winfo_exists(self):
                return False

        return _Sahte

    def _tikla(self, x, y, *, birak_y=None):
        self.tablo.event_generate("<ButtonPress-1>", x=x, y=y)
        self.tablo.event_generate("<ButtonRelease-1>", x=x, y=y if birak_y is None else birak_y)
        self.root.update()

    def _satira_tikla(self, iid):
        x, y, _w, h = self.tablo.bbox(iid, "musteri")
        self._tikla(x + 5, y + h // 2)

    def test_tek_tik_kimlikle_acar_baslik_bosluk_ve_kaydirma_acmaz(self):
        with patch("app.SatisIrsaliyesiDialog", self._kaydedici()):
            for kayit in (self.l2, self.l1, self.l3):
                self._satira_tikla(str(kayit.id))
                self.assertEqual(self.acilan[-1], kayit.id)
            self.assertEqual(len(self.acilan), 3)
            self._tikla(40, 8)  # başlık
            son = self.tablo.get_children()[-1]
            x, y, _w, h = self.tablo.bbox(son)
            self._tikla(40, y + h + 60)  # boş alan
            ilk = self.tablo.get_children()[0]
            x1, y1, _w, h1 = self.tablo.bbox(ilk, "musteri")
            self._tikla(x1 + 5, y1 + 2, birak_y=y + h // 2)  # bir satırda bas, başkasında bırak
        self.assertEqual(len(self.acilan), 3, "başlık/boşluk/kaydırma evrak açmamalı")

    def test_filtre_ve_siralama_sonrasi_dogru_irsaliye(self):
        with patch("app.SatisIrsaliyesiDialog", self._kaydedici()):
            self.root._irs_filtre["evrak_no"].insert(0, "L2")
            self.root.irsaliye_listesini_yenile()
            self.root.update()
            self.assertEqual(self.tablo.get_children(), (str(self.l2.id),))
            self._satira_tikla(self.tablo.get_children()[0])
            self.assertEqual(self.acilan[-1], self.l2.id)

            self.root._irs_filtre["evrak_no"].delete(0, "end")
            self.root._irsaliye_sirala("no")
            self.root._irsaliye_sirala("no")  # doğal sıralama, azalan: L10 > L2 > L1
            self.root.update()
            self.assertEqual(
                self.tablo.get_children(), (str(self.l3.id), str(self.l2.id), str(self.l1.id))
            )
            self.assertIn("▼", self.tablo.heading("no", "text"))
            self._satira_tikla(self.tablo.get_children()[0])
            self.assertEqual(self.acilan[-1], self.l3.id)
            self.root._irsaliye_sirala("tarih")
            self.root.update()
            self.assertEqual(
                self.tablo.get_children(), (str(self.l1.id), str(self.l3.id), str(self.l2.id))
            )
            self._satira_tikla(self.tablo.get_children()[1])
            self.assertEqual(self.acilan[-1], self.l3.id)

    def test_durumlara_gore_acilis_kilitleri_ve_kayit_degismez(self):
        from database.satis_faturasi_service import SatisFaturasiService
        from satis_irsaliyesi_ui import SatisIrsaliyesiDialog as Gercek

        siparis = self._siparis(10, "Adet")
        taslak = SatisIrsaliyesiService.kaydet(
            self._veri(
                irsaliye_no="IRS-TS", sevk_adresi="Kuruçeşme Mah. Depo Girişi", aciklama="Taslak açıklama",
                musteri_notu="Müşteriye not", depo="ANA DEPO",
            ),
            [
                self._satir(2, "Adet", siparis_satiri_id=siparis.satirlar[0].id, urun_adi="Çekmece Rayı Ğ"),
                self._satir(1, "Koli"),
            ],
        )
        sevkli = SatisIrsaliyesiService.kaydet(self._veri(irsaliye_no="IRS-SV"), [self._satir(1, "Koli")])
        SatisIrsaliyesiService.sevk_et(sevkli.id)
        faturali = SatisIrsaliyesiService.kaydet(self._veri(irsaliye_no="IRS-FT"), [self._satir(1, "Koli")])
        SatisIrsaliyesiService.sevk_et(faturali.id)
        SatisFaturasiService.onayla(self._fatura_kaydet(SatisIrsaliyesiService.getir(faturali.id), 1).id)
        self.root.irsaliye_listesini_yenile()
        self.root.update()

        bilgiler: dict[int, dict] = {}

        def _ac(parent, irsaliye=None, **kw):
            d = Gercek(parent, irsaliye=irsaliye, **kw)
            d.withdraw()
            bilgiler[irsaliye.id] = {
                "no": d.girdiler["irsaliye_no"].get(),
                "musteri": d._musteri_var.get(),
                "sevk_adresi": d.girdiler["sevk_adresi"].get(),
                "aciklama": d.girdiler["aciklama"].get(),
                "musteri_notu": d.musteri_notu.get("1.0", "end").strip(),
                "depo": d.depo.get(),
                "satirlar": [(s["urun_adi"], s["miktar"], s["birim"], s.get("siparis_no")) for s in d.satirlar],
                "kaydet": str(d.btn["kaydet"].cget("state")),
                "sevk_geri": str(d.btn["sevk_geri"].cget("state")),
                "iptal": str(d.btn["iptal"].cget("state")),
                "sil": str(d.btn["sil"].cget("state")),
                "degisiklik": d.degisiklik_var(),
            }
            d.destroy()
            return d

        once = _tum_tablolar_ozeti()
        with patch("app.SatisIrsaliyesiDialog", _ac):
            for kayit in (taslak, sevkli, faturali):
                self._satira_tikla(str(kayit.id))
        self.assertEqual(_tum_tablolar_ozeti(), once, "açmak kayıt/stok hareketi oluşturmamalı")

        t = bilgiler[taslak.id]
        self.assertEqual(t["no"], "IRS-TS")
        self.assertIn("M001", t["musteri"])
        self.assertEqual(t["sevk_adresi"], "Kuruçeşme Mah. Depo Girişi")
        self.assertEqual(t["aciklama"], "Taslak açıklama")
        self.assertEqual(t["musteri_notu"], "Müşteriye not")
        self.assertEqual(t["depo"], "ANA DEPO")
        self.assertEqual(
            t["satirlar"],
            [
                ("Çekmece Rayı Ğ", Decimal("2"), "Adet", siparis.siparis_no),
                ("Birim Ürün", Decimal("1"), "Koli", ""),
            ],
        )
        self.assertEqual((t["kaydet"], t["sil"], t["degisiklik"]), ("normal", "normal", False))
        s = bilgiler[sevkli.id]
        self.assertEqual((s["kaydet"], s["sevk_geri"], s["iptal"]), ("disabled", "normal", "normal"))
        f = bilgiler[faturali.id]
        self.assertEqual(
            (f["kaydet"], f["sevk_geri"], f["iptal"], f["sil"]), ("disabled", "disabled", "disabled", "disabled")
        )

    def test_acma_hatasi_anlasilir_mesaj_ve_gunluk(self):
        import tempfile

        with tempfile.TemporaryDirectory() as klasor, patch(
            "hizli_satis_log.log_dizini", return_value=Path(klasor)
        ), patch("app.SatisIrsaliyesiService.getir", side_effect=RuntimeError("bozuk kayıt")):
            self._satira_tikla(str(self.l1.id))
            gunluk = (Path(klasor) / "uygulama_hata.log").read_text(encoding="utf-8")
        hata_mesaji = self.mesajlar[2]
        self.assertEqual(hata_mesaji.call_args.args[0], "İrsaliye açılamadı")
        self.assertIn("bozuk kayıt", hata_mesaji.call_args.args[1])
        self.assertIn("uygulama_hata.log", hata_mesaji.call_args.args[1])
        self.assertIn(f"id={self.l1.id}", gunluk)
        self.assertIn("RuntimeError: bozuk kayıt", gunluk)
        self.assertFalse(self.root._irs_aciliyor)
        with patch("app.SatisIrsaliyesiDialog", self._kaydedici()):
            self._satira_tikla(str(self.l1.id))
        self.assertEqual(self.acilan, [self.l1.id], "hatadan sonra liste çalışmaya devam etmeli")

    def test_listeden_cikti_secili_kimlikle(self):
        self.tablo.selection_set(str(self.l2.id))
        with patch("irsaliye_cikti_ui.cikti_al") as cikti:
            self.root.irsaliye_cikti("word")
            self.root._irs_fiyatli.set(True)
            self.root.irsaliye_cikti("pdf")
        self.assertEqual(cikti.call_args_list[0].args[1:], (self.l2.id, "word"))
        self.assertEqual(cikti.call_args_list[0].kwargs, {"fiyatli": False})
        self.assertEqual(cikti.call_args_list[1].kwargs, {"fiyatli": True})


class SatisIrsaliyesiCiktiTest(_IrsaliyeTemel):
    """A4 PDF / Word: kayıttan aynı içerik, çok sayfa, teslim bölümleri, kayıt değişmezliği."""

    UZUN_AD = (
        "Tam açılır frenli teleskopik çekmece rayı — 450 mm, yüksek taşıma kapasiteli, "
        "galvaniz kaplamalı, sessiz kapanma mekanizmalı ÖZEL ŞIK SERİ İĞNE ÜRÜNÜ"
    )

    def _kayit(self, adet: int, *, uzun=False):
        siparis = self._siparis(100, "Adet")
        satirlar = []
        for i in range(adet):
            ad = self.UZUN_AD if uzun and i % 3 == 1 else f"Ürün {i + 1:03d} Çağ Işık"
            satirlar.append(
                self._satir(
                    Decimal("10") if i % 2 == 0 else Decimal("2.5"),
                    ("Adet", "Koli", "Paket")[i % 3] if i else "Adet",
                    urun_adi=ad,
                    aciklama="Renk: Antrasit gri" if i == 0 else "",
                    siparis_satiri_id=siparis.satirlar[0].id if i == 0 else None,
                )
            )
        irs = SatisIrsaliyesiService.kaydet(
            self._veri(
                sevk_adresi="Ölçü Sok. No: 3, Kuruçeşme Mah.", sevk_il="İstanbul", sevk_ilce="Beşiktaş",
                teslim_kisi="Gökhan Öztürk", aciklama="Özenle paketlendi; koli sayısını kontrol ediniz.",
                musteri_notu="Hafta içi teslim", ic_not="GIZLI-IC-NOT", depo_notu="GIZLI-DEPO",
            ),
            satirlar,
        )
        return irs, siparis

    def _ornek_kopyala(self, *dosyalar):
        import os
        import shutil

        hedef = os.environ.get("IRSALIYE_ORNEK_DIZIN")
        if hedef:
            Path(hedef).mkdir(parents=True, exist_ok=True)
            for d in dosyalar:
                shutil.copy2(d, Path(hedef) / d.name)

    def test_model_kayittan_ic_bilgi_yok_varsayilan_miktar_esasli(self):
        from database.irsaliye_customer_view import render_customer_dispatch_html
        from invoice_print.irsaliye_cikti import cikti_modeli, varsayilan_dosya_adi

        irs, siparis = self._kayit(3)
        vm = cikti_modeli(irs.id)
        self.assertFalse(vm.is_priced)
        self.assertEqual(vm.irsaliye_no, irs.irsaliye_no)
        self.assertEqual(vm.irsaliye_tarihi, "10.03.2026")
        self.assertEqual(vm.siparis_nolari, [siparis.siparis_no])
        self.assertEqual([s.urun_adi for s in vm.satirlar], ["Ürün 001 Çağ Işık", "Ürün 002 Çağ Işık", "Ürün 003 Çağ Işık"])
        self.assertEqual(vm.toplam_miktar_goster, "10 Adet · 2,5 Koli · 10 Paket")
        html = render_customer_dispatch_html(vm)
        for beklenen in (
            "SATIŞ İRSALİYESİ", "TESLİM EDEN", "TESLİM ALAN", "Adı Soyadı", "Tarih / Saat", "İmza",
            "Ölçü Sok. No: 3, Kuruçeşme Mah. — Beşiktaş / İstanbul", "Gökhan Öztürk", siparis.siparis_no,
            "Özenle paketlendi", "@page { size: A4 portrait",
        ):
            self.assertIn(beklenen, html)
        for yasak in ("GIZLI-IC-NOT", "GIZLI-DEPO", "Birim Fiyat", "Genel Toplam"):
            self.assertNotIn(yasak, html)
        self.assertLess(html.index("Özenle paketlendi"), html.index("TESLİM EDEN"))
        fiyatli = render_customer_dispatch_html(cikti_modeli(irs.id, fiyatli=True))
        self.assertIn("Birim Fiyat", fiyatli)
        self.assertIn("Genel Toplam", fiyatli)
        self.assertEqual(varsayilan_dosya_adi(vm, "pdf"), f"Satis_Irsaliyesi_{irs.irsaliye_no}_2026-03-10.pdf")
        self.assertEqual(varsayilan_dosya_adi(vm, "docx")[-5:], ".docx")

    def _pdf_ve_word(self, adet: int, *, uzun=False, fiyatli=False):
        import tempfile

        import pymupdf
        from docx import Document

        from invoice_print.irsaliye_cikti import cikti_modeli, docx_olustur, pdf_olustur

        irs, _ = self._kayit(adet, uzun=uzun)
        vm = cikti_modeli(irs.id, fiyatli=fiyatli)
        klasor = Path(tempfile.mkdtemp())
        ek = f"{adet}_satir{'_uzun' if uzun else ''}{'_fiyatli' if fiyatli else ''}"
        pdf = pdf_olustur(vm, klasor / f"irsaliye_{ek}.pdf")
        docx = docx_olustur(vm, klasor / f"irsaliye_{ek}.docx")
        self._ornek_kopyala(pdf, docx)
        belge = pymupdf.open(str(pdf))
        sayfalar = [s.get_text() for s in belge]
        boyut = tuple(round(x) for x in belge[0].rect[2:])
        belge.close()
        return vm, sayfalar, boyut, Document(str(docx))

    @staticmethod
    def _bosluksuz(metin: str) -> str:
        return "".join(metin.split())

    def _ortak_kontroller(self, vm, sayfalar, boyut, doc):
        from docx.oxml.ns import qn

        self.assertEqual(boyut, (595, 842))
        n = len(sayfalar)
        for i, metin in enumerate(sayfalar, start=1):
            self.assertIn(f"Sayfa {i} / {n}", metin)
            self.assertIn(vm.irsaliye_no, metin)
            if "MB001" in metin:
                self.assertIn("Stok Kodu", metin, f"{i}. sayfada tablo başlığı tekrarlanmalı")
        for yazi in ("TESLİM EDEN", "TESLİM ALAN", "Adı Soyadı"):
            self.assertIn(yazi, sayfalar[-1])
            for onceki in sayfalar[:-1]:
                self.assertNotIn(yazi, onceki, "teslim bölümü son sayfada birlikte olmalı")
        tum = self._bosluksuz("".join(sayfalar))
        konum = -1
        for s in vm.satirlar:  # her ad eksiksiz ve kayıt sırasıyla (bir öncekinden sonra) bulunmalı
            konum = tum.find(self._bosluksuz(s.urun_adi), konum + 1)
            self.assertGreaterEqual(konum, 0, f"{s.sira}. satır adı kesilmiş veya sırası bozuk")
        self.assertIn(self._bosluksuz("Çağ Işık"), tum)
        self.assertIn(self._bosluksuz("Beşiktaş / İstanbul"), tum)

        # Word: düzenlenebilir metin/tablo, aynı sıra, A4, tekrarlanan başlık, bölünmeyen imza
        sec = doc.sections[0]
        self.assertEqual((round(sec.page_width.mm), round(sec.page_height.mm)), (210, 297))
        self.assertEqual(len(doc.inline_shapes), 0, "Word görüntü değil metin olmalı")
        urun = next(t for t in doc.tables if t.rows[0].cells[1].text == "Stok Kodu")
        self.assertEqual(len(urun.rows), len(vm.satirlar) + 1)
        self.assertIsNotNone(urun.rows[0]._tr.trPr.find(qn("w:tblHeader")))
        self.assertTrue(all(r._tr.trPr.find(qn("w:cantSplit")) is not None for r in urun.rows))
        self.assertEqual(
            [r.cells[2].paragraphs[0].text for r in urun.rows[1:]], [s.urun_adi for s in vm.satirlar]
        )
        self.assertEqual([r.cells[4].text for r in urun.rows[1:]], [s.miktar_goster for s in vm.satirlar])
        self.assertEqual([r.cells[3].text for r in urun.rows[1:]], [s.birim for s in vm.satirlar])
        imza = next(t for t in doc.tables if "TESLİM EDEN" in t.rows[0].cells[0].text)
        self.assertIn("TESLİM ALAN", imza.rows[0].cells[2].text)
        self.assertIn("Adı Soyadı", imza.rows[0].cells[2].text)
        self.assertIsNotNone(imza.rows[0]._tr.trPr.find(qn("w:cantSplit")))
        alt_xml = sec.footer._element.xml
        self.assertIn("PAGE", alt_xml)
        self.assertIn("NUMPAGES", alt_xml)
        govde = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn(f"Toplam Miktar: {vm.toplam_miktar_goster}", govde)
        self.assertLess(govde.index("Özenle paketlendi"), len(govde))
        self.assertNotIn("GIZLI-IC-NOT", govde)

    def test_tek_satir_pdf_word(self):
        vm, sayfalar, boyut, doc = self._pdf_ve_word(1)
        self.assertEqual(len(sayfalar), 1)
        self._ortak_kontroller(vm, sayfalar, boyut, doc)

    def test_uzun_adli_cok_satir_pdf_word(self):
        vm, sayfalar, boyut, doc = self._pdf_ve_word(9, uzun=True)
        self._ortak_kontroller(vm, sayfalar, boyut, doc)
        self.assertIn(self._bosluksuz(self.UZUN_AD), self._bosluksuz("".join(sayfalar)))

    def test_cok_sayfali_pdf_word(self):
        vm, sayfalar, boyut, doc = self._pdf_ve_word(70, uzun=True)
        self.assertGreaterEqual(len(sayfalar), 3)
        self._ortak_kontroller(vm, sayfalar, boyut, doc)

    def test_fiyatli_ayri_secenek(self):
        vm, sayfalar, boyut, doc = self._pdf_ve_word(4, fiyatli=True)
        self._ortak_kontroller(vm, sayfalar, boyut, doc)
        self.assertIn("Birim Fiyat", "".join(sayfalar))

    def test_cikti_islemleri_kayitlari_degistirmez(self):
        import tempfile
        from unittest.mock import MagicMock

        import irsaliye_cikti_ui

        irs, _ = self._kayit(1)
        SatisIrsaliyesiService.sevk_et(irs.id)
        once = _tum_tablolar_ozeti()
        klasor = Path(tempfile.mkdtemp())
        ebeveyn = MagicMock()
        with patch("irsaliye_cikti_ui.messagebox") as mb, patch("os.startfile") as ac, patch(
            "irsaliye_cikti_ui.yazici_sec", return_value="Test Yazıcı"
        ), patch("invoice_print.yazici.pdf_yazdir", return_value=1) as yaz, patch(
            "irsaliye_cikti_ui.kayit_yeri_sor", side_effect=[str(klasor / "a.pdf"), str(klasor / "a.docx")]
        ):
            self.assertTrue(irsaliye_cikti_ui.cikti_al(ebeveyn, irs.id, "onizleme").is_file())
            self.assertIsNotNone(irsaliye_cikti_ui.cikti_al(ebeveyn, irs.id, "yazdir"))
            self.assertTrue(irsaliye_cikti_ui.cikti_al(ebeveyn, irs.id, "pdf").is_file())
            self.assertTrue(irsaliye_cikti_ui.cikti_al(ebeveyn, irs.id, "word").is_file())
        mb.showerror.assert_not_called()
        ac.assert_called_once()
        self.assertEqual(yaz.call_args.args[1], "Test Yazıcı")
        self.assertEqual(_tum_tablolar_ozeti(), once, "çıktı stok/cari/sipariş kaydını değiştirmemeli")

    def test_olmayan_kayit_anlasilir_hata(self):
        from unittest.mock import MagicMock

        import irsaliye_cikti_ui

        with patch("irsaliye_cikti_ui.messagebox") as mb:
            self.assertIsNone(irsaliye_cikti_ui.cikti_al(MagicMock(), 999999, "pdf"))
        self.assertIn("bulunamadı", mb.showerror.call_args.args[1])

    def test_microsoft_print_to_pdf_ile_yazdirma(self):
        import tempfile
        import time

        import pymupdf

        from invoice_print.irsaliye_cikti import cikti_modeli, pdf_olustur
        from invoice_print.yazici import pdf_yazdir, yazicilar

        if "Microsoft Print to PDF" not in yazicilar():
            self.skipTest("Microsoft Print to PDF yok")
        irs, _ = self._kayit(40)
        klasor = Path(tempfile.mkdtemp())
        pdf = pdf_olustur(cikti_modeli(irs.id), klasor / "kaynak.pdf")
        hedef = klasor / "basilan.pdf"
        basilan = pdf_yazdir(pdf, "Microsoft Print to PDF", belge_adi="Test", cikti_dosyasi=str(hedef))
        for _ in range(60):
            if hedef.is_file() and hedef.stat().st_size > 1000:
                break
            time.sleep(0.5)
        with pymupdf.open(str(pdf)) as a, pymupdf.open(str(hedef)) as b:
            self.assertEqual(basilan, a.page_count)
            self.assertEqual(b.page_count, a.page_count)
            self.assertEqual(tuple(round(x) for x in b[0].rect[2:]), (595, 842))


class SatisIrsaliyesiEkranCiktiTest(_EkranTemel):
    def test_kaydedilmemis_degisiklikte_vazgecilirse_cikti_yok(self):
        d = self._dialog()
        d._musteri_ata(self._cari())
        self._urun_ekle(d, "1", "Adet")
        sor = self.mesajlar[3]
        sor.return_value = False
        with patch("irsaliye_cikti_ui.cikti_al") as cikti:
            d.word_kaydet()
            self.assertIsNone(d.irsaliye, "vazgeçilen yeni belge kaydedilmemeli")
            cikti.assert_not_called()
            sor.return_value = True
            d.word_kaydet()
            iid = d.irsaliye.id
            d.girdiler["aciklama"].insert(0, "Yeni not")
            sor.return_value = False
            d.onizleme()
            self.assertEqual(cikti.call_count, 1, "eski kayıtla fark ettirmeden çıktı üretilmemeli")
            sor.return_value = True
            d._fiyatli_cikti.set(True)
            d.yazdir()
        self.assertEqual(SatisIrsaliyesiService.getir(iid).aciklama, "Yeni not")
        self.assertEqual(cikti.call_args_list[0].args[1:], (iid, "word"))
        self.assertEqual(cikti.call_args_list[0].kwargs, {"fiyatli": False})
        self.assertEqual(cikti.call_args_list[1].args[1:], (iid, "yazdir"))
        self.assertEqual(cikti.call_args_list[1].kwargs, {"fiyatli": True})
        for anahtar in ("onizleme", "yazdir", "pdf", "word"):
            self.assertIn(anahtar, d.btn)
        d.destroy()


if __name__ == "__main__":
    unittest.main()
