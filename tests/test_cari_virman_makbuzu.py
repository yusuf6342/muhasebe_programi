"""Tahsilat makbuzu — Cari Virman: bağlı iki cari hareket, bakiye, düzenleme/iptal, para birimi, MKB serisi, çıktı.

Yalnız geçici test veritabanı kullanır (TahsilatMakbuzuTest kurulumu).
"""

from __future__ import annotations

import sys
import tempfile
import tkinter as tk
import unittest
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import func, select, text

import makbuz_cikti as mc
from database.cari_virman_makbuz_service import (
    ISLEM_ODEME,
    ISLEM_TAHSILAT,
    VIRMAN_METNI,
    CariVirmanMakbuzService,
)
from database.database import get_session
from database.finans_service import FinansService
from database.models.cari import Cari, CariIslem, SatisHareketi
from database.models.finans import (
    CariVirmanMakbuzu,
    FinansHareketi,
    KasaMakbuzu,
    KrediKartiOdeme,
    PosValorKaydi,
)
from tests import test_tahsilat_makbuzu as _servis_testi

MM = 72 / 25.4
TARIH = date(2026, 9, 28)
FIRMA = {"unvan": "ÖRNEK TEST MOBİLYA SANAYİ VE TİCARET LTD. ŞTİ.", "adres": "Atatürk Cad. No: 12", "il": "Ankara"}
YASAK_METINLER = (
    "nakit alındı",
    "Nakit alındı",
    "karttan çekildi",
    "NAKİT",
    "KASA / BANKA / POS",
    "POS",
    "ÖDEME YÖNTEMİ",
    "TESLİM EDEN",
    "TAHSİLAT MAKBUZU\n",
)


def _kurulum(test):
    _servis_testi.TahsilatMakbuzuTest.setUp(test)
    # Açık fatura kalemli ikinci müşteri / tedarikçi (FIFO kapanışı ve geri açılışı için)
    with get_session() as s:
        m2 = Cari(cari_kodu="M002", unvan="DENİZ İNŞAAT A.Ş.", cari_turu="Musteri")
        t2 = Cari(cari_kodu="T002", unvan="KARTAL KERESTE", cari_turu="Tedarikçi")
        s.add_all([m2, t2])
        s.flush()
        s.add(SatisHareketi(cari_id=m2.id, satis_tarihi=date(2026, 3, 1), belge_no="SF-1",
                            satis_tutari=Decimal("800"), kalan_acik_tutar=Decimal("800")))
        s.add(SatisHareketi(cari_id=t2.id, satis_tarihi=date(2026, 3, 1), belge_no="AF-1",
                            satis_tutari=Decimal("3000"), kalan_acik_tutar=Decimal("3000")))
        s.flush()
        test.musteri2_id, test.tedarikci2_id = m2.id, t2.id


def _virman(test, tutar="400", musteri=None, tedarikci=None, **ek):
    veri = {
        "tarih": TARIH,
        "musteri_id": musteri or test.musteri_id,
        "tedarikci_id": tedarikci or test.tedarikci_id,
        "tutar": tutar,
        "aciklama": "Eylül mahsubu",
        "makbuz_no_otomatik": True,
    }
    veri.update(ek)
    return CariVirmanMakbuzService.kaydet(veri)


def _bakiye(cari_id):
    return FinansService.cari_bakiye_ozeti(cari_id)


class CariVirmanServisTest(unittest.TestCase):
    def setUp(self):
        _kurulum(self)

    def tearDown(self):
        _servis_testi.TahsilatMakbuzuTest.tearDown(self)

    def _sayilar(self):
        with get_session() as s:
            return {
                "finans": s.scalar(select(func.count(FinansHareketi.id))),
                "kasa_makbuz": s.scalar(select(func.count(KasaMakbuzu.id))),
                "pos": s.scalar(select(func.count(PosValorKaydi.id))),
                "kk": s.scalar(select(func.count(KrediKartiOdeme.id))),
            }

    def _virman_islemleri(self, k):
        with get_session() as s:
            return list(
                s.scalars(
                    select(CariIslem)
                    .where(CariIslem.belge_no.in_([k.tahsilat_belge_no, k.odeme_belge_no]))
                    .order_by(CariIslem.id)
                ).all()
            )

    def _tedarikci_farki(self, cari_id):
        """Açılış onarımının (FZO) baktığı fark: açık kalemler − defter bakiyesi; virman değiştirmemeli."""
        with get_session() as s:
            acik = s.scalar(
                select(func.coalesce(func.sum(SatisHareketi.kalan_acik_tutar), 0)).where(
                    SatisHareketi.cari_id == cari_id, ~SatisHareketi.belge_no.like("FZO-%")
                )
            )
        return Decimal(str(acik)) - Decimal(str(_bakiye(cari_id)["bakiye"]))

    def test_kayit_iki_bagli_cari_hareket_ve_bakiyeler(self):
        self.assertEqual(_bakiye(self.musteri_id)["yon"], "Borçlu")
        self.assertEqual(_bakiye(self.tedarikci_id)["yon"], "Borçlu")
        once = self._sayilar()
        k = _virman(self)
        self.assertEqual(k.makbuz_no, "MKB-00001")
        self.assertRegex(k.belge_no, r"^CVR-2026-0001$")
        self.assertEqual((k.tahsilat_belge_no, k.odeme_belge_no), (f"{k.belge_no}-T", f"{k.belge_no}-O"))
        self.assertEqual(_bakiye(self.musteri_id), {"bakiye": Decimal("600"), "tutar": Decimal("600"), "yon": "Borçlu"})
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("4600"))

        islemler = self._virman_islemleri(k)
        self.assertEqual(len(islemler), 2)
        tahsilat, odeme = islemler
        self.assertEqual((tahsilat.cari_id, tahsilat.islem_turu, tahsilat.alacak, tahsilat.borc),
                         (self.musteri_id, ISLEM_TAHSILAT, Decimal("400"), Decimal("0")))
        self.assertEqual((odeme.cari_id, odeme.islem_turu, odeme.alacak, odeme.borc),
                         (self.tedarikci_id, ISLEM_ODEME, Decimal("400"), Decimal("0")))
        self.assertEqual(tahsilat.karsi_cari_id, self.tedarikci_id)
        self.assertEqual(odeme.karsi_cari_id, self.musteri_id)
        self.assertEqual((tahsilat.para_birimi, odeme.para_birimi), ("TRY", "TRY"))
        # Kasa, banka, POS, kredi kartı hareketi yok
        self.assertEqual(self._sayilar(), once)

    def test_fifo_acik_kalemleri_kapatir_iptal_tam_geri_acar(self):
        k = _virman(self, "500", musteri=self.musteri2_id, tedarikci=self.tedarikci2_id)
        self.assertEqual((k.musteri_kapanan, k.tedarikci_kapanan), (Decimal("500"), Decimal("500")))
        with get_session() as s:
            kalan = dict(s.execute(select(SatisHareketi.belge_no, SatisHareketi.kalan_acik_tutar)
                                   .where(SatisHareketi.belge_no.in_(["SF-1", "AF-1"]))).all())
        self.assertEqual(kalan, {"SF-1": Decimal("300"), "AF-1": Decimal("2500")})
        self.assertEqual(_bakiye(self.musteri2_id)["bakiye"], Decimal("300"))
        self.assertEqual(_bakiye(self.tedarikci2_id)["bakiye"], Decimal("2500"))
        CariVirmanMakbuzService.iptal(k.id)
        with get_session() as s:
            kalan = dict(s.execute(select(SatisHareketi.belge_no, SatisHareketi.kalan_acik_tutar)
                                   .where(SatisHareketi.belge_no.in_(["SF-1", "AF-1"]))).all())
            artik = s.scalar(select(func.count(SatisHareketi.id)).where(
                SatisHareketi.belge_no.like("IPT-%") | SatisHareketi.belge_no.like("CVR-%")))
        self.assertEqual(kalan, {"SF-1": Decimal("800"), "AF-1": Decimal("3000")})
        self.assertEqual(artik, 0)
        self.assertEqual(_bakiye(self.musteri2_id)["bakiye"], Decimal("800"))
        self.assertEqual(_bakiye(self.tedarikci2_id)["bakiye"], Decimal("3000"))

    def test_duzenleme_iki_tarafi_birlikte_gunceller_cift_kayit_olusmaz(self):
        k = _virman(self)
        yeni = CariVirmanMakbuzService.guncelle(
            k.id,
            {"tarih": TARIH - timedelta(days=1), "musteri_id": self.musteri_id, "tedarikci_id": self.tedarikci_id,
             "tutar": "250", "aciklama": "Düzeltildi", "makbuz_no": k.makbuz_no},
        )
        self.assertEqual((yeni.id, yeni.belge_no, yeni.makbuz_no), (k.id, k.belge_no, k.makbuz_no))
        islemler = self._virman_islemleri(yeni)
        self.assertEqual(len(islemler), 2)
        self.assertEqual({i.alacak for i in islemler}, {Decimal("250")})
        self.assertEqual({i.tarih for i in islemler}, {TARIH - timedelta(days=1)})
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("750"))
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("4750"))
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(CariVirmanMakbuzu.id))), 1)
            self.assertEqual(
                s.scalar(select(func.count(SatisHareketi.id)).where(SatisHareketi.belge_no == k.odeme_belge_no)), 1
            )

        # Taraf değiştirme: eski tedarikçinin hareketi tamamen kalkar
        CariVirmanMakbuzService.guncelle(
            k.id,
            {"tarih": TARIH, "musteri_id": self.musteri_id, "tedarikci_id": self.tedarikci2_id, "tutar": "100",
             "makbuz_no": k.makbuz_no},
        )
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("5000"))
        self.assertEqual(_bakiye(self.tedarikci2_id)["bakiye"], Decimal("2900"))
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("900"))
        self.assertEqual(len(self._virman_islemleri(k)), 2)

    def test_iptal_iki_tarafi_geri_alir_numara_kalir(self):
        k = _virman(self)
        CariVirmanMakbuzService.iptal(k.id)
        k2 = CariVirmanMakbuzService.getir(k.id)
        self.assertEqual((k2.durum, k2.makbuz_no), ("IPTAL", "MKB-00001"))
        self.assertEqual(self._virman_islemleri(k), [])
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("1000"))
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("5000"))
        with self.assertRaises(ValueError):
            CariVirmanMakbuzService.iptal(k.id)
        with self.assertRaises(ValueError):
            CariVirmanMakbuzService.guncelle(k.id, {"tarih": TARIH, "musteri_id": self.musteri_id,
                                                    "tedarikci_id": self.tedarikci_id, "tutar": "1"})
        self.assertEqual(FinansService.makbuz_no_oner(), "MKB-00002")

    def test_tedarikci_acik_kalem_farki_degismez(self):
        for cari_id in (self.tedarikci_id, self.tedarikci2_id):
            once = self._tedarikci_farki(cari_id)
            k = _virman(self, "3600", tedarikci=cari_id)
            self.assertEqual(self._tedarikci_farki(cari_id), once)
            CariVirmanMakbuzService.guncelle(k.id, {"tarih": TARIH, "musteri_id": self.musteri_id,
                                                    "tedarikci_id": cari_id, "tutar": "7000",
                                                    "makbuz_no": k.makbuz_no})
            self.assertEqual(self._tedarikci_farki(cari_id), once)
            CariVirmanMakbuzService.iptal(k.id)
            self.assertEqual(self._tedarikci_farki(cari_id), once)

    def test_tutar_bakiyeyi_asarsa_alacakli_olur_iptalde_tam_doner(self):
        k = _virman(self, "6000")
        self.assertEqual(_bakiye(self.musteri_id), {"bakiye": Decimal("-5000"), "tutar": Decimal("5000"), "yon": "Alacaklı"})
        self.assertEqual(_bakiye(self.tedarikci_id)["yon"], "Alacaklı")
        CariVirmanMakbuzService.iptal(k.id)
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("1000"))
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("5000"))

    def test_farkli_para_birimi_engellenir(self):
        with get_session() as s:
            s.add(CariIslem(cari_id=self.tedarikci_id, tarih=date(2026, 2, 1), islem_turu="Açılış", belge_no="USD-1",
                            borc=Decimal("100"), alacak=Decimal("0"), para_birimi="USD"))
        with get_session() as s:
            once = s.scalar(select(func.count(CariIslem.id)))
        with self.assertRaisesRegex(ValueError, "USD para birimli hareket var"):
            _virman(self)
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(CariIslem.id))), once)
            self.assertEqual(s.scalar(select(func.count(CariVirmanMakbuzu.id))), 0)

    def test_dogrulamalar(self):
        with self.assertRaisesRegex(ValueError, "müşteri carisi değil"):
            _virman(self, musteri=self.tedarikci_id, tedarikci=self.tedarikci2_id)
        with self.assertRaisesRegex(ValueError, "tedarikçi carisi değil"):
            _virman(self, tedarikci=self.musteri2_id)
        with self.assertRaisesRegex(ValueError, "aynı cari"):
            _virman(self, tedarikci=self.musteri_id)
        with self.assertRaisesRegex(ValueError, "sıfırdan büyük"):
            _virman(self, "0")
        with self.assertRaisesRegex(ValueError, "gelecek"):
            _virman(self, tarih=date.today() + timedelta(days=1))
        with self.assertRaisesRegex(ValueError, "müşteriyi seçin"):
            _virman(self, musteri=None, musteri_id=None)
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(CariVirmanMakbuzu.id))), 0)

    def test_makbuz_no_kasa_makbuzlariyla_ortak_ve_tekil(self):
        kasa = _servis_testi.TahsilatMakbuzuTest._nakit(self, "50")
        self.assertEqual(kasa.makbuz_no, "MKB-00001")
        k = _virman(self)
        self.assertEqual(k.makbuz_no, "MKB-00002")
        self.assertEqual(FinansService.makbuz_no_oner(), "MKB-00003")
        self.assertTrue(FinansService.makbuz_no_kullanimda_mi("mkb-00002"))
        self.assertFalse(FinansService.makbuz_no_kullanimda_mi("MKB-00002", haric_virman_id=k.id))
        with self.assertRaisesRegex(ValueError, "zaten var"):
            _virman(self, makbuz_no_otomatik=False, makbuz_no="MKB-00001")
        with self.assertRaisesRegex(ValueError, "zaten var"):
            _servis_testi.TahsilatMakbuzuTest._nakit(self, "10", makbuz_no_otomatik=False, makbuz_no="MKB-00002")
        elle = _virman(self, makbuz_no_otomatik=False, makbuz_no="V-77")
        self.assertEqual(elle.makbuz_no, "V-77")
        self.assertEqual(FinansService.makbuz_no_oner(), "MKB-00003")

    def test_eski_veritabaninda_tablo_eklenir(self):
        with self.eng.begin() as c:
            c.execute(text("DROP TABLE cari_virman_makbuzlari"))
        CariVirmanMakbuzService._hazir_motor = None
        self.assertEqual(FinansService.makbuz_no_oner(), "MKB-00001")
        self.assertFalse(FinansService.makbuz_no_kullanimda_mi("MKB-00001"))
        self.assertEqual(CariVirmanMakbuzService.listele(), [])
        self.assertEqual(_virman(self).makbuz_no, "MKB-00001")

    def test_liste_arama_ve_belge_no(self):
        k = _virman(self)
        _virman(self, "10", musteri=self.musteri2_id, tedarikci=self.tedarikci2_id)
        self.assertEqual([x.id for x in CariVirmanMakbuzService.listele(arama="yildiz")], [k.id])
        self.assertEqual(len(CariVirmanMakbuzService.listele(arama="deniz")), 1)
        self.assertEqual(len(CariVirmanMakbuzService.listele()), 2)
        self.assertEqual(CariVirmanMakbuzService.belge_no_ile(k.odeme_belge_no).id, k.id)
        self.assertEqual(CariVirmanMakbuzService.belge_no_ile(k.belge_no).musteri.cari_kodu, "M001")


class CariVirmanCiktiTest(unittest.TestCase):
    def setUp(self):
        _kurulum(self)
        self.firma = mc.firma_satirlari(FIRMA)
        self.virman = _virman(self, "1250.50")
        self.nakit = _servis_testi.TahsilatMakbuzuTest._nakit(self, "300")
        self.vk = mc.virman_kimligi(self.virman.id)

    def tearDown(self):
        _servis_testi.TahsilatMakbuzuTest.tearDown(self)

    def _pdf(self, kimlikler, yerlesim):
        veriler = mc.makbuz_cikti_verileri(kimlikler, firma=self.firma)
        return mc._fitz().open(stream=mc.pdf_uret(veriler, yerlesim), filetype="pdf")

    def _virman_metnini_denetle(self, metin):
        k = self.virman
        for beklenen in (
            "CARİ VİRMAN MAKBUZU",
            "MKB-00001",
            k.belge_no,
            "28.09.2026",
            "M001",
            "ALİ ÇELİK MOBİLYA LTD",
            "T001",
            "YILDIZ AKSESUAR",
            "VİRMAN TUTARI",
            "1.250,50 ₺",
            "Yalnız Bin İki Yüz Elli Türk Lirası Elli Kuruştur.",
            k.tahsilat_belge_no,
            k.odeme_belge_no,
            "Tahsilat evrakı",
            "Ödeme evrakı",
            "Eylül mahsubu",
            "VİRMANI DÜZENLEYEN",
            "MÜŞTERİ YETKİLİSİ",
            "TEDARİKÇİ YETKİLİSİ",
            "FİRMA KAŞESİ",
        ):
            self.assertIn(beklenen, metin)
        self.assertIn(VIRMAN_METNI, " ".join(metin.split()))
        for yasak in YASAK_METINLER:
            self.assertNotIn(yasak, metin)

    def test_pdf_a5_virman_duzeni(self):
        doc = self._pdf([self.vk], mc.A5)
        self.assertEqual(len(doc), 1)
        page = doc[0]
        self.assertAlmostEqual(page.rect.width, 148 * MM, places=1)
        self.assertAlmostEqual(page.rect.height, 210 * MM, places=1)
        metin = page.get_text()
        self._virman_metnini_denetle(metin)
        self.assertEqual(metin.count("Adı Soyadı:"), 3)
        self.assertEqual(metin.count("İmza:"), 3)
        r = page.rect
        for blok in page.get_text("dict")["blocks"]:
            for satir in blok.get("lines", []):
                for span in satir["spans"]:
                    if span["text"].strip():
                        x0, y0, x1, y1 = span["bbox"]
                        self.assertTrue(x0 >= 3 * MM and x1 <= r.x1 - 3 * MM and y1 <= r.y1 - 3 * MM, span["text"])
        # İmza alanlarında resim / çizim yok: sayfada görüntü bulunmaz
        self.assertEqual(page.get_images(), [])

    def test_toplu_karisik_ciktida_her_makbuz_kendi_basligiyla(self):
        doc = self._pdf([self.nakit.id, self.vk], mc.A5)
        self.assertEqual(len(doc), 2)
        normal, virman = doc[0].get_text(), doc[1].get_text()
        self.assertIn("TAHSİLAT MAKBUZU", normal)
        self.assertIn("NAKİT / KASA", normal)
        self.assertIn("TESLİM EDEN / TAHSİL EDEN", normal)
        self.assertNotIn("VİRMAN", normal)
        self._virman_metnini_denetle(virman)

        ikili = self._pdf([self.vk, self.nakit.id], mc.A4_IKILI)
        self.assertEqual(len(ikili), 1)
        self.assertAlmostEqual(ikili[0].rect.width, 297 * MM, places=1)
        sol = ikili[0].get_text(clip=mc._fitz().Rect(0, 0, 148.5 * MM, 210 * MM))
        sag = ikili[0].get_text(clip=mc._fitz().Rect(148.5 * MM, 0, 297 * MM, 210 * MM))
        self._virman_metnini_denetle(sol)
        self.assertIn("TAHSİLAT MAKBUZU", sag)
        self.assertIn("MKB-00002", sag)
        self.assertNotIn("MKB-00002", sol)

    def test_word_a5_ve_a4_ikili(self):
        from docx import Document
        from docx.oxml.ns import qn

        veriler = mc.makbuz_cikti_verileri([self.vk, self.nakit.id], firma=self.firma)
        tek = Document(BytesIO(mc.docx_uret(veriler[:1], mc.A5)))
        s = tek.sections[0]
        self.assertAlmostEqual(s.page_width.mm, 148, delta=0.2)
        self.assertAlmostEqual(s.page_height.mm, 210, delta=0.2)
        metin = "\n".join(p.text for t in tek.tables for r in t.rows for c in r.cells for p in c.paragraphs)
        self._virman_metnini_denetle(metin)
        imza = tek.tables[-2]
        self.assertEqual(len(imza.columns), 4)

        ikili = Document(BytesIO(mc.docx_uret(veriler, mc.A4_IKILI)))
        s = ikili.sections[0]
        self.assertAlmostEqual(s.page_width.mm, 297, delta=0.2)
        self.assertEqual(s._sectPr.find(qn("w:cols")).get(qn("w:num")), "2")
        xml = ikili.element.body.xml
        self.assertEqual(xml.count('w:type="column"'), 1)
        tum = "\n".join(p.text for t in ikili.tables for r in t.rows for c in r.cells for p in c.paragraphs)
        self.assertIn("CARİ VİRMAN MAKBUZU", tum)
        self.assertIn("TAHSİLAT MAKBUZU", tum)
        self.assertIn("NAKİT / KASA", tum)

    def test_iptal_virman_isaretlenir(self):
        CariVirmanMakbuzService.iptal(self.virman.id)
        self.assertIn("İPTAL EDİLMİŞTİR", self._pdf([self.vk], mc.A5)[0].get_text())

    def test_dosya_adlari(self):
        v = mc.makbuz_cikti_verileri([self.vk], firma=self.firma)
        self.assertEqual(mc.dosya_adi(v, mc.A5, "pdf"), "Cari_Virman_Makbuzu_MKB-00001.pdf")
        karisik = mc.makbuz_cikti_verileri([self.vk, self.nakit.id], firma=self.firma)
        self.assertEqual(
            mc.dosya_adi(karisik, mc.A4_IKILI, "docx"), "Tahsilat_Makbuzlari_MKB-00001_MKB-00002_2_adet_A4_ikili.docx"
        )

    def test_cikti_almak_kayit_degistirmez(self):
        with get_session() as s:
            once = (s.scalar(select(func.count(CariIslem.id))), s.scalar(select(func.count(CariVirmanMakbuzu.id))))
        bakiye = (_bakiye(self.musteri_id), _bakiye(self.tedarikci_id))
        with tempfile.TemporaryDirectory() as d:
            for yer in (mc.A5, mc.A4_IKILI):
                veriler = mc.makbuz_cikti_verileri([self.vk, self.nakit.id])
                mc.pdf_kaydet(veriler, Path(d) / f"v_{yer}.pdf", yer)
                mc.docx_kaydet(veriler, Path(d) / f"v_{yer}.docx", yer)
        with get_session() as s:
            sonra = (s.scalar(select(func.count(CariIslem.id))), s.scalar(select(func.count(CariVirmanMakbuzu.id))))
        self.assertEqual(sonra, once)
        self.assertEqual((_bakiye(self.musteri_id), _bakiye(self.tedarikci_id)), bakiye)


class CariVirmanEkranTest(unittest.TestCase):
    def setUp(self):
        _kurulum(self)
        try:
            self.root = tk.Tk()
        except tk.TclError as hata:
            _servis_testi.TahsilatMakbuzuTest.tearDown(self)
            self.skipTest(f"Tk yok: {hata}")
        self.root.withdraw()
        import kasa_makbuz_ui
        import makbuz_cikti_ui

        self.ui = kasa_makbuz_ui
        self.cikti_ui = makbuz_cikti_ui
        self.mesajlar = []
        self._yamalar = [
            patch.object(mod.messagebox, ad, side_effect=lambda *a, _ad=ad, **k: self.mesajlar.append((_ad, a)))
            for mod in (kasa_makbuz_ui, makbuz_cikti_ui)
            for ad in ("showerror", "showinfo", "showwarning")
        ]
        self._yamalar.append(patch("cari_kart_ui.cari_uyari_goster", return_value=None))
        for y in self._yamalar:
            y.start()

    def tearDown(self):
        for y in self._yamalar:
            y.stop()
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        _servis_testi.TahsilatMakbuzuTest.tearDown(self)

    def _virman_dialogu(self):
        d = self.ui.KasaMakbuzDialog(self.root, "TAHSILAT")
        d.islem_turu_var.set(self.ui.ISLEM_VIRMAN)
        d._islem_turu_degisti()
        self.root.update_idletasks()
        return d

    def _carileri_sec(self, d, musteri="çelik", tedarikci="aksesuar"):
        d.v_musteri_var.set(musteri)
        d._v_musteri_arama.sec(0)
        d.v_tedarikci_var.set(tedarikci)
        d._v_tedarikci_arama.sec(0)

    def test_virman_modu_kaydet_duzenle_yazdir_iptal(self):
        d = self._virman_dialogu()
        self.assertTrue(d.virman)
        self.assertEqual(d._virman_dis.winfo_manager(), "pack")
        self.assertEqual(d._satir_dis.winfo_manager(), "")
        self.assertEqual(d.cari_etiket_lbl.winfo_manager(), "")
        self.assertIn("CARİ VİRMAN", d.baslik_lbl.cget("text"))

        # Arama yalnız ilgili türde ve adın herhangi bir yerinde
        musteri_idleri = {k[2] for k in d._v_musteri_arama.ara("a")}
        tedarikci_idleri = {k[2] for k in d._v_tedarikci_arama.ara("a")}
        self.assertEqual(musteri_idleri, {self.musteri_id, self.musteri2_id})
        self.assertEqual(tedarikci_idleri, {self.tedarikci_id, self.tedarikci2_id})
        self.assertEqual([k[2] for k in d._v_musteri_arama.ara("MOBİLYA")], [self.musteri_id])

        self._carileri_sec(d)
        self.assertEqual(d._v_musteri_arama.secili_id, self.musteri_id)
        self.assertEqual(d._v_tedarikci_arama.secili_id, self.tedarikci_id)
        m_once, t_once, _ = d.virman_bakiye_metinleri()
        self.assertEqual(m_once, "1.000,00 TL  Borçlu")
        self.assertEqual(t_once, "5.000,00 TL  Borçlu")
        d.v_tutar_var.set("400")
        self.assertEqual(d.virman_bakiye_metinleri()[2], "Müşteri: 600,00 TL  Borçlu\nTedarikçi: 4.600,00 TL  Borçlu")
        d.aciklama_var.set("Mahsup")

        self.assertTrue(d.kaydet(), self.mesajlar)
        self.assertEqual((d.mod, d.virman, d.makbuz.makbuz_no), ("goruntule", True, "MKB-00001"))
        self.assertIn("CVR-2026-0001-T", d.mesaj_lbl.cget("text"))
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("600"))
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("4600"))
        # Kayıtlı virmanda işlem öncesi bakiye virmansız bakiyedir
        self.assertEqual(d.virman_bakiye_metinleri()[:2], ("1.000,00 TL  Borçlu", "5.000,00 TL  Borçlu"))
        self.assertEqual(str(d.islem_turu_cb.cget("state")), "disabled")

        d.duzenlemeye_gec()
        self.assertEqual(d.mod, "duzenle")
        d.v_tutar_var.set("250")
        self.assertTrue(d.kaydet(), self.mesajlar)
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("750"))
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("4750"))
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(CariVirmanMakbuzu.id))), 1)
            self.assertEqual(s.scalar(select(func.count(CariIslem.id)).where(CariIslem.belge_no.like("CVR-%"))), 2)

        with patch.object(self.cikti_ui, "MakbuzCiktiDialog") as pencere:
            d.yazdir()
        pencere.assert_called_once_with(d, [f"V{d.makbuz_id}"], mc.A5)

        with patch.object(self.ui.messagebox, "askyesno", return_value=True) as soru:
            d.iptal_et()
        self.assertIn("birlikte geri alınır", soru.call_args.args[1])
        self.assertEqual(d.makbuz.durum, "IPTAL")
        self.assertIn("iptal edilmiştir", d.kilit_lbl.cget("text"))
        self.assertEqual(_bakiye(self.musteri_id)["bakiye"], Decimal("1000"))
        self.assertEqual(_bakiye(self.tedarikci_id)["bakiye"], Decimal("5000"))
        self.assertEqual(self.mesajlar, [])
        d.destroy()

    def test_bakiye_asiminda_onay_istenir(self):
        d = self._virman_dialogu()
        self._carileri_sec(d)
        d.v_tutar_var.set("6000")
        with patch.object(self.ui.messagebox, "askyesno", return_value=False) as soru:
            self.assertFalse(d.kaydet())
        self.assertIn("Alacaklı", soru.call_args.args[1])
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count(CariVirmanMakbuzu.id))), 0)
        d.destroy()

    def test_eksik_secimde_kayit_yapilmaz(self):
        d = self._virman_dialogu()
        d.v_tutar_var.set("10")
        self.assertFalse(d.kaydet())
        self.assertIn("müşteriyi seçin", self.mesajlar[-1][1][1])
        d.destroy()

    def test_tahsilata_donus_ve_odeme_makbuzunda_virman_yok(self):
        d = self._virman_dialogu()
        d.islem_turu_var.set(self.ui.ISLEM_TAHSILAT)
        d._islem_turu_degisti()
        self.root.update_idletasks()
        self.assertEqual(d._satir_dis.winfo_manager(), "pack")
        self.assertEqual(d._virman_dis.winfo_manager(), "")
        self.assertEqual(d.cari_etiket_lbl.winfo_manager(), "grid")
        d.destroy()
        o = self.ui.KasaMakbuzDialog(self.root, "ODEME")
        self.assertFalse(hasattr(o, "islem_turu_cb"))
        o.destroy()
        with self.assertRaises(ValueError):
            self.ui.KasaMakbuzDialog(self.root, "ODEME", islem_turu="VIRMAN")

    def test_liste_ve_toplu_cikti(self):
        nakit = _servis_testi.TahsilatMakbuzuTest._nakit(self, "300")
        virman = _virman(self)
        root = self.root
        root.icerik = tk.Frame(root)
        root.icerik.pack(fill="both", expand=True)
        root.menu_dugmeleri = {}
        root._icerigi_temizle = lambda: [w.destroy() for w in root.icerik.winfo_children()]

        def hemen(_widget, is_, bitti, hata=None):
            bitti(is_())

        with patch("ui_bg.arka_planda", hemen):
            self.ui.kasa_makbuzlari_sayfasi(root, makbuz_turu="TAHSILAT")

        def bul(w, sinif):
            if isinstance(w, sinif):
                yield w
            for c in w.winfo_children():
                yield from bul(c, sinif)

        from tkinter import ttk

        tablo = next(bul(root.icerik, ttk.Treeview))
        cocuklar = list(tablo.get_children())
        self.assertEqual(set(cocuklar), {str(nakit.id), f"V{virman.id}"})
        self.assertEqual(tablo.set(f"V{virman.id}", "odeme"), "CARİ VİRMAN")
        self.assertIn("YILDIZ AKSESUAR", tablo.set(f"V{virman.id}", "hesap"))
        tablo.selection_set(cocuklar)
        dugme = next(b for b in bul(root.icerik, tk.Button) if b.cget("text") == "Yazdır / PDF / Word")
        with patch.object(self.cikti_ui, "makbuz_ciktisi_ac") as ac:
            dugme.invoke()
        beklenen = [c if c.startswith("V") else int(c) for c in cocuklar]
        self.assertEqual(ac.call_args.args[1], beklenen)

        onizleme = self.cikti_ui.MakbuzCiktiDialog(root, beklenen)
        root.update()
        onizleme._ciz()
        self.assertEqual(onizleme.sayfa_sayisi, 2)
        onizleme.yerlesim_degistir(mc.A4_IKILI)
        self.assertEqual(onizleme.sayfa_sayisi, 1)
        onizleme.destroy()

        tablo.selection_set([f"V{virman.id}"])
        goruntule = next(b for b in bul(root.icerik, tk.Button) if b.cget("text") == "Görüntüle")
        goruntule.invoke()
        acilan = [w for w in root.winfo_children() if isinstance(w, self.ui.KasaMakbuzDialog)]
        self.assertEqual(len(acilan), 1)
        self.assertTrue(acilan[0].virman)
        self.assertEqual(acilan[0].makbuz_id, virman.id)
        acilan[0].destroy()
        self.assertEqual(self.mesajlar, [])


if __name__ == "__main__":
    unittest.main()
