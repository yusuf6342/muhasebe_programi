"""Satış yönetici raporları — müşteri/stok kârlılığı, çoklu dönem ve bağlı raporlar.

Yalnız geçici test veritabanı (StokOrtami) kullanır; gerçek firma verisine dokunmaz.
"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from sqlalchemy import select  # noqa: E402

from database.database import get_session  # noqa: E402
from tests.stok_test_ortami import StokOrtami  # noqa: E402

D = Decimal
OCAK = ("Ocak", date(2026, 1, 1), date(2026, 1, 31))
SUBAT = ("Şubat", date(2026, 2, 1), date(2026, 2, 28))


def _satir(miktar, fiyat, birim="Adet", i1="0", i2="0", kod="U001"):
    return {"urun_kodu": kod, "urun_adi": "Test Ürün", "miktar": D(str(miktar)), "birim": birim,
            "birim_fiyat": D(str(fiyat)), "iskonto_orani": D(i1), "iskonto_orani_2": D(i2),
            "kdv_orani": D("20")}


class _Taban(unittest.TestCase):
    def setUp(self):
        from database.models.cari import Cari
        from database.models.finans import FinansHesabi

        self._ayar = tempfile.TemporaryDirectory(prefix="sy_ayar_")
        self._env = patch.dict(os.environ, {"LOCALAPPDATA": self._ayar.name})
        self._env.start()
        self.o = StokOrtami("satis_yonetici")
        with get_session() as s:
            c = Cari(cari_kodu="M002", unvan="İkinci Müşteri", cari_turu="Müşteri", aktif=True)
            s.add(c)
            s.add(FinansHesabi(hesap_adi="ANA KASA", hesap_turu="KASA", aktif=True))
            s.flush()
            self.m2 = int(c.id)
        self.o.alis(200, 10, date(2026, 1, 2))

    def tearDown(self):
        self.o.kapat()
        self._env.stop()
        self._ayar.cleanup()

    def fatura(self, tarih, satirlar, cari_id=None, onayla=True, **ek):
        from database.satis_faturasi_service import SatisFaturasiService

        veriler = {"fatura_tarihi": tarih, "vade_tarihi": tarih, "cari_id": cari_id or self.o.mus_id,
                   "depo": "ANA DEPO", "odeme_tutari": D("0"), "sales_person_id": 1}
        veriler.update(ek)
        f = SatisFaturasiService.kaydet(veriler, satirlar)
        if onayla:
            SatisFaturasiService.onayla(f.id)
        return SatisFaturasiService.getir(f.id)

    def iade(self, tarih, satirlar, kaynak_fatura_id=None):
        from database.satis_iade_faturasi_service import SatisIadeFaturasiService

        return SatisIadeFaturasiService.kaydet(
            {"iade_tarihi": tarih, "cari_id": self.o.mus_id, "kaynak_fatura_id": kaynak_fatura_id,
             "depo": "ANA DEPO", "aciklama": None, "iade_odeme_tutari": 0}, satirlar)

    def senaryo(self):
        """Ocak: çok iskontolu Koli satışı + kaynaklı iade + taslak + eksik maliyet; Şubat: M002."""
        self.f1 = self.fatura(date(2026, 1, 10), [_satir(2, 150, "Koli", "10", "5")])
        self.iade(date(2026, 1, 20), [dict(_satir(1, "128.25", "Koli"),
                                          kaynak_fatura_satiri_id=self.f1.satirlar[0].id)],
                  kaynak_fatura_id=self.f1.id)
        self.fatura(date(2026, 1, 12), [_satir(1, 50)], onayla=False)
        self.f_eksik = self.fatura(date(2026, 1, 31), [_satir(1, 30)])
        from database.models.satis_faturasi import SatisFaturasiSatiri

        with get_session() as s:
            st = s.get(SatisFaturasiSatiri, self.f_eksik.satirlar[0].id)
            st.fifo_birim_maliyeti = D("0")
        self.f_subat = self.fatura(date(2026, 2, 1), [_satir(5, 20)], cari_id=self.m2)


class KarlilikHesapTest(_Taban):
    def test_toplamlar_musteri_urun_belge_satir_esit_ve_kurallar(self):
        from database import satis_yonetici_rapor_service as srv

        self.senaryo()
        m = srv.karlilik_raporu("musteri", [OCAK])
        u = srv.karlilik_raporu("urun", [OCAK])
        satirlar = srv.satirlari_getir(OCAK[1], OCAK[2])
        detay = srv.satir_detaylari(satirlar)
        for alan in ("net", "kar", "maliyet", "iade"):
            self.assertEqual(m["toplam"][alan], u["toplam"][alan], alan)
            self.assertEqual(sum((r[alan] for r in m["satirlar"]), D("0")), m["toplam"][alan], alan)
            self.assertEqual(sum((r[alan] for r in u["satirlar"]), D("0")), m["toplam"][alan], alan)
        self.assertEqual(sum((r["net"] for r in detay), D("0")), m["toplam"]["net"])
        # Taslak (onaysız) fatura sayılmaz; 1 satış + 1 iade + eksik maliyetli satış
        self.assertEqual(len(detay), 3)
        satis = next(r for r in detay if r["belge_no"] == self.f1.fatura_no)
        self.assertEqual(satis["net"], D("256.50"))  # 300 × 0,90 × 0,95 (art arda iskonto)
        self.assertEqual(satis["temel_miktar"], D("24"))
        self.assertEqual(satis["toplam_maliyet"], D("240.00"))
        iade = next(r for r in detay if r["belge_turu"] == "İade")
        self.assertEqual(iade["toplam_maliyet"], D("-120.00"))  # kaynak satış maliyeti ters çevrildi
        self.assertIn("kaynak satış", iade["maliyet_kaynagi"])
        self.assertEqual(iade["kaynak_belge_no"], self.f1.fatura_no)
        eksik = next(r for r in detay if r["belge_no"] == self.f_eksik.fatura_no)
        self.assertIsNone(eksik["kar"])
        self.assertEqual(eksik["maliyet_durumu"], "Eksik")
        t = m["toplam"]
        self.assertEqual(t["net"], D("256.50") - D("128.25") + D("30.00"))
        self.assertEqual(t["kar"], D("16.50") - D("8.25"))
        self.assertEqual(t["eksik_satir"], 1)
        self.assertEqual(t["marj"], (t["kar"] / (t["net"] - D("30")) * 100).quantize(D("0.01")))
        self.assertEqual(u["satirlar"][0]["net_miktar"], D("24") - D("12") + D("1"))
        self.assertEqual(u["satirlar"][0]["birim"], "Adet")
        self.assertTrue(any("Kâr eksik" in n for n in m["notlar"]))
        # Eski kâr analizi ekranı da aynı motoru kullanır
        from database.satis_kar_analiz_service import KarAnalizFiltre, SatisKarAnalizService

        k = SatisKarAnalizService.analiz(KarAnalizFiltre(baslangic=OCAK[1], bitis=OCAK[2]))
        self.assertEqual(k.ozet["net_satis"], t["net"])
        self.assertEqual(k.ozet["brut_kar"], t["kar"])
        self.assertEqual(k.ozet["kar_marji"], t["marj"])

    def test_fatura_geneli_indirim_satirlara_dagitilir(self):
        from database import satis_yonetici_rapor_service as srv
        from database.satis_faturasi_service import SatisFaturasiService

        f = self.fatura(date(2026, 1, 15), [_satir(10, 20), _satir(5, 40)],
                        tl_brut_toplam=D("480.00"), tl_genel_toplam=D("456.00"), genel_islem_turu="INDIRIM",
                        genel_islem_tutari=D("24.00"), genel_islem_orani=D("5"),
                        invoice_rounding_adjustment=D("0"))
        detay = [r for r in srv.satir_detaylari(srv.satirlari_getir(OCAK[1], OCAK[2]))
                 if r["belge_no"] == f.fatura_no]
        matrah = SatisFaturasiService.urun_net_toplam(f) - D(str(f.tl_kdv))
        self.assertEqual(sum((r["net"] for r in detay), D("0")), matrah)
        self.assertTrue(all(r["belge_duzeltme"] < 0 for r in detay))
        self.assertEqual(detay[0]["belge_duzeltme"], detay[1]["belge_duzeltme"])  # 200 / 200 eşit pay

    def test_doviz_belge_kuru_korunur(self):
        from database import satis_yonetici_rapor_service as srv

        with patch("database.doviz_service.DovizService.kur_degeri", return_value=D("48")):
            f = self.fatura(date(2026, 1, 16), [dict(_satir(1, 480), tl_esas=True)], para_birimi="USD",
                            kur=D("48"), kur_tarihi=date(2026, 1, 16), kur_sabitlendi=True, kur_kaynagi="TCMB")
        r = next(x for x in srv.satir_detaylari(srv.satirlari_getir(OCAK[1], OCAK[2]))
                 if x["belge_no"] == f.fatura_no)
        self.assertEqual((r["para_birimi"], r["kur"], r["kur_tarihi"]), ("USD", D("48"), date(2026, 1, 16)))
        self.assertEqual(r["net"], D("480.00"))

    def test_hizli_satis_bir_kez_ve_kanal_filtresi(self):
        from database import satis_yonetici_rapor_service as srv
        from database.hizli_satis_service import HizliSatisService
        from database.models.stok import StokKarti
        from hizli_satis_sepet import HizliSatisSepet

        with get_session() as s:
            sid = s.scalar(select(StokKarti.id).where(StokKarti.stok_kodu == "U001"))
        sepet = HizliSatisSepet()
        sepet.ekle(stok_id=sid, stok_kodu="U001", stok_adi="Test Ürün", miktar=3, birim_fiyat="25")
        sonuc = HizliSatisService.satisi_tamamla(
            sepet=sepet, cari_id=self.o.mus_id, tahsilatlar=[],
            idempotency_token=HizliSatisService.yeni_idempotency_token(), musteri_bakiye=0, risk_limiti=None)
        self.fatura(date.today(), [_satir(1, 40)])
        bugun = date.today()
        hepsi = srv.satir_detaylari(srv.satirlari_getir(bugun, bugun))
        hizli = [r for r in hepsi if r["belge_no"] == sonuc["fatura_no"]]
        self.assertEqual(len(hizli), 1)
        self.assertEqual(hizli[0]["kanal"], "Hızlı Satış")
        self.assertEqual(len(srv.satirlari_getir(bugun, bugun, {"kanal": "hizli"})), 1)
        self.assertEqual(len(srv.satirlari_getir(bugun, bugun, {"kanal": "fatura"})), 1)


class CokluDonemTest(_Taban):
    def test_yeni_duran_temel_donem_ve_bagimsiz_toplam(self):
        from database import satis_yonetici_rapor_service as srv

        self.senaryo()
        r = srv.karlilik_raporu("musteri", [SUBAT, OCAK], temel=1, bugun=date(2026, 3, 1))
        durum = {x["cari_kodu"]: x["durum"] for x in r["satirlar"]}
        self.assertEqual(durum["M001"], "Satışı durdu")
        self.assertEqual(durum["M002"], "Yeni satış")
        m2 = next(x for x in r["satirlar"] if x["cari_kodu"] == "M002")
        self.assertEqual(m2["satis_0"], D("100.00"))
        self.assertIsNone(m2["yuzde_satis_0"])  # temel dönem 0 → yüzde yok
        self.assertIn("Temel dönem 0", m2["aciklama_0"])
        self.assertEqual(len(srv.karlilik_raporu("musteri", [SUBAT, OCAK], temel=1, sadece="yeni")["satirlar"]), 1)
        # Çakışan dönemler bağımsız: toplamlar birleştirilmez
        r2 = srv.karlilik_raporu("urun", [OCAK, ("Ocak 2. yarı", date(2026, 1, 15), date(2026, 1, 31))])
        self.assertEqual(r2["donemler"][0]["net"], r["donemler"][1]["net"])
        self.assertEqual(r2["donemler"][1]["net"], D("-128.25") + D("30.00"))
        d0 = r["donemler"][0]
        self.assertEqual((d0["gun_sayisi"], d0["durum"]), (28, "Tamamlandı"))
        devam = srv.karlilik_raporu("urun", [SUBAT], bugun=date(2026, 2, 10))["donemler"][0]
        self.assertTrue(devam["devam_ediyor"])
        self.assertEqual(devam["gunluk_net"], D("10.00"))  # 100 / 10 geçen gün

    def test_degisim_kurallari_ve_hazir_donemler(self):
        from database import satis_yonetici_rapor_service as srv

        self.assertEqual(srv.degisim(D("150"), D("100"))["yuzde"], D("50.00"))
        self.assertIsNone(srv.degisim(D("50"), D("-20"), "kar")["yuzde"])
        self.assertIsNone(srv.degisim(D("-5"), D("20"), "kar")["yuzde"])
        self.assertEqual(srv.degisim(D("-5"), D("20"), "kar")["fark"], D("-25"))
        self.assertEqual(srv.marj_puan_farki(D("12.5"), D("10")), D("2.5"))
        b = date(2026, 10, 8)  # Perşembe
        self.assertEqual(srv.donem_araligi("bu_hafta", b), (date(2026, 10, 5), date(2026, 10, 11)))
        self.assertEqual(srv.donem_araligi("gecen_hafta", b), (date(2026, 9, 28), date(2026, 10, 4)))
        self.assertEqual(srv.donem_araligi("gecen_ay", b), (date(2026, 9, 1), date(2026, 9, 30)))
        self.assertEqual(srv.donem_araligi("bu_ay", b), (date(2026, 10, 1), date(2026, 10, 31)))
        self.assertEqual(srv.donem_araligi("gecen_yil_ayni", b, (date(2028, 2, 1), date(2028, 2, 29))),
                         (date(2027, 2, 1), date(2027, 2, 28)))
        with self.assertRaises(ValueError):
            srv.donemleri_hazirla([("X", date(2026, 2, 1), date(2026, 1, 1))])


class BagliRaporlarTest(_Taban):
    def test_dusuk_marj_iskonto_iade_ve_azalan(self):
        from database import satis_yonetici_rapor_service as srv

        self.senaryo()
        self.fatura(date(2026, 1, 25), [_satir(2, 5)], cari_id=self.m2)  # zararına
        d = srv.dusuk_marjli_satislar(OCAK[1], OCAK[2], hedef_marj=10)
        durumlar = sorted(r["durum"] for r in d["satirlar"])
        self.assertEqual(durumlar, ["Hedef altı", "Maliyeti eksik", "Zararına"])
        self.assertEqual(len(srv.dusuk_marjli_satislar(OCAK[1], OCAK[2], hedef_marj=10, tur="zararina")
                             ["satirlar"]), 1)
        ii = srv.iskonto_iade_analizi(OCAK[1], OCAK[2], boyut="musteri")
        m1 = next(r for r in ii["ozet"] if r["cari_kodu"] == "M001")
        self.assertEqual(m1["iskonto"], D("43.50") - D("0"))  # 300 − 256,50
        self.assertEqual(m1["iade"], D("128.25"))
        self.assertEqual(ii["iadeler"][0]["baglanti"], "Kaynak faturaya bağlı")
        self.assertTrue(any("iade nedeni" in n for n in ii["notlar"]))
        az = srv.azalan_musteriler((date(2026, 2, 1), date(2026, 2, 28)), (OCAK[1], OCAK[2]),
                                   inaktif_gun=20, bugun=date(2026, 3, 1))
        m1a = next(r for r in az["satirlar"] if r["cari_kodu"] == "M001")
        self.assertIn("Dönemde satış yok", m1a["durum"])
        self.assertIn("Alışverişi durdu", m1a["durum"])
        self.assertEqual(m1a["son_satis"], date(2026, 1, 31))

    def test_tahsilat_performansi_ve_ozet(self):
        from database import satis_yonetici_rapor_service as srv
        from database.cari_service import CariService

        f = self.fatura(date(2026, 1, 10), [_satir(10, 20)])
        CariService.tahsilat_yap(self.o.mus_id, date(2026, 1, 20), D("100"), "NAKİT / KASA", "ANA KASA")
        r = srv.tahsilat_performansi(OCAK[1], OCAK[2], bugun=date(2026, 3, 1))
        m1 = next(x for x in r["satirlar"] if x["cari_kodu"] == "M001")
        self.assertEqual(m1["net"], D("200.00"))
        self.assertEqual(m1["tahsilat"], D("100.00"))
        self.assertEqual(m1["acik_alacak"], D(str(f.tl_genel_toplam)) - D("100"))
        self.assertEqual(m1["vadesi_gecmis"], m1["acik_alacak"])
        self.assertEqual(m1["gerceklesen_tahsil_gun"], D("10.0"))
        oz = srv.yonetici_ozeti((OCAK[1], OCAK[2]), (date(2025, 12, 1), date(2025, 12, 31)),
                                bugun=date(2026, 3, 1))
        adlar = [g["gosterge"] for g in oz["gostergeler"]]
        self.assertIn("Brüt kâr", adlar)
        self.assertNotIn("Satış katkısı (brüt kâr − bağlı POS komisyonu)", adlar)
        acik = next(g for g in oz["gostergeler"] if g["gosterge"] == "Açık alacak")
        self.assertIn("itibarıyla", acik["tur"])
        self.assertEqual(acik["donem"], m1["acik_alacak"])

    def test_bekleyen_siparis_ve_satis_hizi(self):
        from database import satis_yonetici_rapor_service as srv
        from database.models.satis_siparisi import SatisSiparisi, SatisSiparisiSatiri

        with get_session() as s:
            for no, termin, miktar, irs in (("SS-1", date(2026, 1, 5), "150", "20"),
                                            ("SS-2", date(2026, 1, 9), "100", "0")):
                sp = SatisSiparisi(siparis_no=no, siparis_tarihi=date(2026, 1, 1), termin_tarihi=termin,
                                   cari_id=self.o.mus_id, maliyet_yontemi="FIFO", durum="AÇIK", depo="ANA DEPO")
                sp.satirlar.append(SatisSiparisiSatiri(
                    urun_kodu="U001", urun_adi="Test Ürün", miktar=D(miktar), birim="Adet",
                    birim_satis_fiyati=D("20"), fifo_birim_maliyeti=D("10"), irsaliyelenen_miktar=D(irs)))
                s.add(sp)
        r = srv.bekleyen_satis_siparisleri(bugun=date(2026, 1, 7))
        s1, s2 = r["satirlar"]
        self.assertEqual((s1["siparis_no"], s1["kalan"], s1["gecikme"]), ("SS-1", D("130"), 2))
        self.assertEqual(s1["sevk_bekleyen"], D("20"))  # irsaliye var, stok çıkışı yok
        self.assertEqual(s1["stok_durumu"], "Yeterli")
        self.assertEqual((s2["stok_durumu"], s2["karsilanabilir"]), ("Kısmi", D("70.000")))  # 200 − 130
        self.assertEqual(s1["tahmini_kar"], D("1300.00"))
        self.assertEqual(len(srv.bekleyen_satis_siparisleri(bugun=date(2026, 1, 7), sadece_geciken=True)
                             ["satirlar"]), 1)
        self.o.urun_ekle("U002", "Duran Ürün")
        self.o.alis(5, 10, date(2026, 1, 2), kod="U002")
        self.fatura(date(2026, 1, 10), [_satir(20, 20)])
        h = srv.urun_satis_hizi(OCAK[1], OCAK[2], bugun=date(2026, 2, 1))
        u1 = next(x for x in h["satirlar"] if x["urun_kodu"] == "U001")
        self.assertEqual(u1["mevcut"], D("180"))
        self.assertEqual(u1["yeterlilik_gun"], int(D("180") / (D("20") / 31)))
        yok = srv.urun_satis_hizi(OCAK[1], OCAK[2], bugun=date(2026, 2, 1), tur="satis_yok")["satirlar"]
        self.assertEqual([x["urun_kodu"] for x in yok], ["U002"])
        self.assertIsNone(yok[0]["yeterlilik_gun"])

    def test_yetkisiz_kullanici_maliyet_kar_gormez(self):
        from database import satis_yonetici_rapor_service as srv
        from database.session_manager import oturum

        self.senaryo()
        oturum.set_user(user_id=2, kullanici_adi="satis", ad_soyad="Satış", role_kod="SATIS", role_ad="Satış",
                        permissions={"satis_goruntuleme"})
        r = srv.karlilik_raporu("musteri", [OCAK])
        self.assertTrue(all(x["kar"] is None and x["marj"] is None for x in r["satirlar"]))
        self.assertIsNotNone(r["satirlar"][0]["net"])
        with self.assertRaises(Exception):
            srv.dusuk_marjli_satislar(OCAK[1], OCAK[2])
        oturum.set_user(user_id=3, kullanici_adi="yok", ad_soyad="Yok", role_kod="X", role_ad="X",
                        permissions=set())
        with self.assertRaises(Exception):
            srv.karlilik_raporu("musteri", [OCAK])


class EkranTest(_Taban):
    def test_ekranlar_acilir_raporlar_ve_disa_aktarir(self):
        import tkinter as tk

        import satis_yonetici_rapor_ui as ui
        from satis_ui import RAPOR_KARTLARI, _rapor_komutlar

        self.senaryo()
        try:
            kok = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk yok")
        kok.withdraw()
        kok.icerik = tk.Frame(kok)
        kok.icerik.pack()
        kok._icerigi_temizle = lambda: [w.destroy() for w in kok.icerik.winfo_children()]
        kok.satis_raporlari_goster = lambda: None
        try:
            komutlar = _rapor_komutlar(kok)
            for _b, _a, anahtar in RAPOR_KARTLARI:
                self.assertIn(anahtar, komutlar)
            for fn in (ui.musteri_karlilik_goster, ui.stok_karlilik_goster, ui.dusuk_marj_goster,
                       ui.tahsilat_performansi_goster, ui.azalan_musteriler_goster, ui.urun_satis_hizi_goster,
                       ui.iskonto_iade_goster, ui.bekleyen_siparisler_goster, ui.yonetici_ozeti_goster):
                ekran = fn(kok, otomatik=False)
                ekran.raporla(senkron=True)
                self.assertIn("notlar", ekran.sonuc, fn.__name__)
                ekran.html_metni()
            ekran = ui.musteri_karlilik_goster(kok, otomatik=False)
            ekran.donem_editoru.ayarla([OCAK, SUBAT], temel=0)
            ekran.raporla(senkron=True)
            kolonlar = ekran.tablolar["satirlar"]["columns"]
            self.assertIn("satis_1", kolonlar)
            self.assertIn("yuzde_satis_1", kolonlar)
            yol = os.path.join(self._ayar.name, "m.xlsx")
            ekran.excel_yaz(yol)
            from openpyxl import load_workbook

            ws = load_workbook(yol).worksheets[0]
            ust = " ".join(str(ws.cell(r, 1).value or "") for r in range(1, 4))
            self.assertIn("Ocak", ust)
            self.assertIn("FIFO", ust)
            detay = ui.karlilik_detay_goster(kok, "musteri", self.o.mus_id, "Test Müşteri", OCAK[1], OCAK[2], {})
            self.assertGreater(len(detay.ekran.tablolar["kirilim"].get_children()), 0)
            detay.destroy()
        finally:
            kok.destroy()


if __name__ == "__main__":
    unittest.main()
