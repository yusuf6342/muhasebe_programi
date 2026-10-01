"""Satın alma talebi: onay akışı, siparişe aktarım payları, teslim takibi, geçiş."""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import create_engine, event, func, inspect, select, text
from sqlalchemy.orm import sessionmaker

from database.database import Base, _aktif_engine_bagla, company_db, get_session
from database.models.cari import Cari
from database.models.donem import Donem
from database.models.firma import Firma
from database.models.stok import Depo, StokKarti
from database.session_manager import oturum

BUGUN = date.today()


def _sqlite_pragma(dbapi_connection, _connection_record):
    cur = dbapi_connection.cursor()
    cur.execute("PRAGMA foreign_keys=ON")
    cur.close()


def _modelleri_yukle():
    import database.models.cari  # noqa: F401
    import database.models.stok  # noqa: F401
    import database.models.firma  # noqa: F401
    import database.models.donem  # noqa: F401
    import database.models.alis_faturasi  # noqa: F401
    import database.models.alis_iade_faturasi  # noqa: F401
    import database.models.alis_irsaliyesi  # noqa: F401
    import database.models.alis_siparisi  # noqa: F401
    import database.models.alis_masraf  # noqa: F401
    import database.models.finans  # noqa: F401
    import database.models.satis_faturasi  # noqa: F401
    import database.models.satis_iade_faturasi  # noqa: F401
    import database.models.satis_irsaliyesi  # noqa: F401
    import database.models.satis_siparisi  # noqa: F401
    import database.models.hizli_satis  # noqa: F401
    import database.models.satin_alma_talep  # noqa: F401


class _TalepTemel(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.TemporaryDirectory()
        self.db_path = Path(self._tmpdir.name) / "test_talep.db"
        eng = create_engine(f"sqlite:///{self.db_path.as_posix()}", connect_args={"check_same_thread": False})
        event.listen(eng, "connect", _sqlite_pragma)
        _modelleri_yukle()
        Base.metadata.create_all(eng)
        Session = sessionmaker(bind=eng, autoflush=False, autocommit=False, expire_on_commit=False)
        with Session() as s:
            firma = Firma(firma_kodu="TLP", unvan="Talep Test", aktif=True)
            s.add(firma)
            s.flush()
            s.add(Donem(firma_id=firma.id, donem_adi="2026", baslangic_tarihi=date(2026, 1, 1),
                        bitis_tarihi=date(2026, 12, 31), aktif=True, kapali=False, varsayilan=True))
            s.add(Depo(ad="ANA DEPO", aktif=True))
            s.add(StokKarti(stok_kodu="U001", stok_adi="Vida", birim="Adet", aktif=True, is_deleted=False))
            s.add(StokKarti(stok_kodu="U002", stok_adi="Somun", birim="Adet", aktif=True, is_deleted=False))
            s.add(Cari(cari_kodu="T001", unvan="Tedarikçi A", cari_turu="Tedarikçi", aktif=True))
            s.commit()
        _aktif_engine_bagla(eng)
        company_db._engine = eng
        company_db._session_factory = Session
        company_db._company_id = 97
        company_db._db_path = self.db_path
        self._kullanici(1, "YONETICI")
        oturum.set_company(company_id=97, firma_kodu="TLP", firma_unvan="Talep Test", firma_uid="tlp",
                           db_path=str(self.db_path))
        with get_session() as s:
            d = s.scalar(select(Donem).limit(1))
            oturum.set_period(d.id, d.donem_adi)
            self.cari_id = s.scalar(select(Cari.id).where(Cari.cari_kodu == "T001"))
        self._muhasebe_patch = patch("database.muhasebe_entegrasyon.muhasebe_hook", lambda *_a, **_k: None)
        self._muhasebe_patch.start()
        from database.satin_alma_talep_service import SatinAlmaTalepService

        SatinAlmaTalepService._hazir_motor = None

    def _kullanici(self, uid, rol, izinler=()):
        oturum.set_user(user_id=uid, kullanici_adi=f"k{uid}", ad_soyad=f"Kullanıcı {uid}", role_kod=rol,
                        role_ad=rol, permissions=set(izinler))

    def tearDown(self):
        self._muhasebe_patch.stop()
        try:
            company_db.close()
        except Exception:
            pass
        if company_db._engine is not None:
            company_db._engine.dispose()
        company_db._engine = None
        company_db._session_factory = None
        oturum.clear()
        self._tmpdir.cleanup()

    def _talep(self, satirlar=None, **ek) -> int:
        from database.satin_alma_talep_service import SatinAlmaTalepService

        veriler = {"talep_tarihi": BUGUN, "ihtiyac_tarihi": BUGUN + timedelta(days=5), "depo": "ANA DEPO",
                   "oncelik": "NORMAL", **ek}
        return SatinAlmaTalepService.kaydet(
            veriler, satirlar or [{"urun_kodu": "U001", "urun_adi": "Vida", "birim": "Adet", "miktar": "10",
                                   "aciklama": "M6x20 galvaniz", "tahmini_birim_fiyat": "2,5"}]
        )

    def _onayli_talep(self, **ek) -> tuple[int, int]:
        from database.satin_alma_talep_service import SatinAlmaTalepService

        tid = self._talep(**ek)
        SatinAlmaTalepService.onaya_gonder(tid)
        SatinAlmaTalepService.onayla(tid)
        return tid, SatinAlmaTalepService.detay(tid)["satirlar"][0]["id"]

    def _siparis(self, satirlar, siparis_id=None, row_version=None):
        from database.alis_siparisi_service import AlisSiparisiService

        return AlisSiparisiService.kaydet(
            {"siparis_tarihi": BUGUN, "termin_tarihi": BUGUN + timedelta(days=7), "cari_id": self.cari_id,
             "row_version": row_version},
            satirlar, [], siparis_id,
        )

    def _satir(self, tid) -> dict:
        from database.satin_alma_talep_service import SatinAlmaTalepService

        return SatinAlmaTalepService.detay(tid)["satirlar"][0]


class SatinAlmaTalepServisTest(_TalepTemel):
    def test_kabul_siparis_paylari_ve_teslim(self):
        from database.alis_faturasi_service import AlisFaturasiService
        from database.alis_irsaliyesi_service import AlisIrsaliyesiService
        from database.alis_siparisi_service import AlisSiparisiService
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid, sid = self._onayli_talep()
        hazir = S.siparis_satirlari_hazirla({sid: "4"})
        self.assertEqual(self._satir(tid)["kalan"], Decimal("10"))  # kayıt olmadan tüketim yok
        sip1 = self._siparis(hazir["satirlar"])
        self.assertEqual(self._satir(tid)["kalan"], Decimal("6"))
        self.assertEqual(S.detay(tid)["gorunen_durum"], "KISMEN SİPARİŞE AKTARILDI")
        with self.assertRaises(ValueError):
            S.siparis_satirlari_hazirla({sid: "7"})
        sip2 = self._siparis(S.siparis_satirlari_hazirla({sid: "6"})["satirlar"])
        self.assertEqual(self._satir(tid)["kalan"], Decimal("0"))
        self.assertEqual(S.detay(tid)["gorunen_durum"], "TAMAMEN SİPARİŞE AKTARILDI")
        with self.assertRaises(ValueError) as ctx:
            S.siparis_satirlari_hazirla({sid: "1"})
        self.assertIn("aşamaz", str(ctx.exception))
        # Payı doğrudan fazla yazmaya çalışmak da engellenir
        with self.assertRaises(ValueError):
            self._siparis([{**hazir["satirlar"][0], "miktar": Decimal("1"),
                            "talep_paylari": [{"talep_satiri_id": sid, "miktar": Decimal("1")}]}])

        # Sipariş 2'yi 1 azalt → 1 geri açılır
        s2 = AlisSiparisiService.getir(sip2.id)
        paylar = S.siparis_paylari(sip2.id)
        satir = s2.satirlar[0]
        self._siparis([{"urun_kodu": satir.urun_kodu, "urun_adi": satir.urun_adi, "miktar": Decimal("5"),
                        "birim": satir.birim, "birim_alis_fiyati": satir.birim_alis_fiyati,
                        "talep_paylari": paylar[int(satir.id)]}], siparis_id=sip2.id, row_version=s2.row_version)
        self.assertEqual(self._satir(tid)["kalan"], Decimal("1"))
        self.assertEqual(self._satir(tid)["siparis"], Decimal("9"))

        # İrsaliye ile 3 teslim + aynı 3'ün faturası → teslim 3 (6 değil)
        s1 = AlisSiparisiService.getir(sip1.id)
        irs = AlisIrsaliyesiService.kaydet(
            {"irsaliye_tarihi": BUGUN, "cari_id": self.cari_id, "siparis_id": sip1.id},
            [{"urun_kodu": "U001", "urun_adi": "Vida", "miktar": Decimal("3"), "birim": "Adet",
              "birim_fiyat": Decimal("2"), "kdv_orani": Decimal("20"), "siparis_satiri_id": s1.satirlar[0].id}],
        )
        self.assertEqual(self._satir(tid)["teslim"], Decimal("3"))
        self.assertEqual(S.detay(tid)["gorunen_durum"], "KISMEN TESLİM ALINDI")
        irs_satir = AlisIrsaliyesiService.getir(irs.id).satirlar[0]
        AlisFaturasiService.kaydet(
            {"fatura_tarihi": BUGUN, "vade_tarihi": BUGUN, "cari_id": self.cari_id, "irsaliye_id": irs.id,
             "depo": "ANA DEPO", "odeme_tutari": Decimal("0")},
            [{"urun_kodu": "U001", "urun_adi": "Vida", "miktar": Decimal("3"), "birim": "Adet",
              "birim_fiyat": Decimal("2"), "iskonto_orani": Decimal("0"), "kdv_orani": Decimal("20"),
              "irsaliye_satiri_id": irs_satir.id}],
        )
        detay = self._satir(tid)
        self.assertEqual(detay["teslim"], Decimal("3"))
        self.assertEqual([x["siparis_no"] for x in detay["siparisler"]], [sip1.siparis_no, sip2.siparis_no])

        # Sipariş iptali payı geri açar
        AlisSiparisiService.iptal_et(sip2.id)
        self.assertEqual(self._satir(tid)["kalan"], Decimal("6"))
        gecmis = [g["islem"] for g in S.detay(tid)["gecmis"]]
        self.assertIn("SİPARİŞE AKTAR", gecmis)
        self.assertIn("SİPARİŞTEN GERİ AÇ", gecmis)
        rapor = S.urun_bazli_rapor()
        self.assertEqual((rapor[0]["talep"], rapor[0]["siparis"], rapor[0]["teslim"]),
                         (Decimal("10"), Decimal("4"), Decimal("3")))

    def test_onay_akisi_kilit_red_geri_gonder(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid = self._talep()
        with self.assertRaises(ValueError):
            S.siparis_satirlari_hazirla({self._satir(tid)["id"]: "1"})  # onaysız
        S.onaya_gonder(tid)
        with self.assertRaises(ValueError):
            self._talep_guncelle(tid)
        with self.assertRaises(ValueError):
            S.reddet(tid, "")
        S.geri_gonder(tid, "Miktarı düzeltin")
        d = S.detay(tid)
        self.assertEqual((d["durum"], d["geri_gonderme_nedeni"]), ("TASLAK", "Miktarı düzeltin"))
        self._talep_guncelle(tid)
        S.onaya_gonder(tid)
        S.reddet(tid, "Bütçe yok")
        d = S.detay(tid)
        self.assertEqual((d["durum"], d["red_nedeni"]), ("REDDEDİLDİ", "Bütçe yok"))
        islemler = [g["islem"] for g in d["gecmis"]]
        self.assertEqual(islemler, ["OLUŞTUR", "ONAYA GÖNDER", "GERİ GÖNDER", "GÜNCELLE", "ONAYA GÖNDER", "REDDET"])

    def _talep_guncelle(self, tid):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        S.kaydet({"talep_tarihi": BUGUN, "depo": "ANA DEPO"},
                 [{"urun_kodu": "U001", "urun_adi": "Vida", "birim": "Adet", "miktar": "12"}], tid)

    def test_kendi_talebini_onaylayamaz(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        izinler = ("alis_talep_duzenleme", "alis_talep_onay", "alis_talep_goruntuleme", "alis_siparis_duzenleme")
        self._kullanici(5, "FINANS", izinler)
        tid = self._talep()
        S.onaya_gonder(tid)
        with self.assertRaises(ValueError) as ctx:
            S.onayla(tid)
        self.assertIn("Kendi", str(ctx.exception))
        self._kullanici(6, "FINANS", izinler)
        S.onayla(tid)
        self.assertEqual(S.detay(tid)["onaylayan"], "Kullanıcı 6")
        self._kullanici(7, "FINANS", ("alis_talep_duzenleme", "alis_talep_goruntuleme"))
        tid2 = self._talep()
        S.onaya_gonder(tid2)
        with self.assertRaises(PermissionError):
            S.onayla(tid2)
        # Fiyat yetkisi yoksa tahmini fiyat görünmez
        self.assertIsNone(S.detay(tid)["satirlar"][0]["tahmini_birim_fiyat"])
        self.assertIsNone(S.listele()[0]["tahmini_toplam"])

    def test_manuel_satir_eslestirme_zorunlu(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        with self.assertRaises(ValueError):
            self._talep([{"urun_kodu": "", "urun_adi": "", "birim": "Adet", "miktar": "2"}])
        with self.assertRaises(ValueError):
            self._talep([{"urun_kodu": "", "urun_adi": "Özel conta", "birim": "Adet", "miktar": "0"}])
        tid = self._talep([{"urun_kodu": "", "urun_adi": "Özel conta", "birim": "Adet", "miktar": "2",
                            "aciklama": "Ø40 silikon"}])
        satir = self._satir(tid)
        self.assertTrue(satir["eslesmemis"])
        with get_session() as s:
            self.assertEqual(s.scalar(select(StokKarti.id).where(StokKarti.stok_adi == "Özel conta")), None)
        S.onaya_gonder(tid)
        S.onayla(tid)
        with self.assertRaises(ValueError) as ctx:
            S.siparis_satirlari_hazirla({satir["id"]: "2"})
        self.assertIn("eşleştir", str(ctx.exception))
        S.manuel_satiri_eslestir(satir["id"], "U002")
        hazir = S.siparis_satirlari_hazirla({satir["id"]: "2"})
        self.assertEqual(hazir["satirlar"][0]["urun_kodu"], "U002")
        self.assertEqual(hazir["satirlar"][0]["aciklama"], "Ø40 silikon")

    def test_kalan_iptal_ve_numara(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid, sid = self._onayli_talep()
        self._siparis(S.siparis_satirlari_hazirla({sid: "7"})["satirlar"])
        with self.assertRaises(ValueError):
            S.kalan_iptal(sid, "1", "")
        with self.assertRaises(ValueError):
            S.kalan_iptal(sid, "4", "fazla")
        S.kalan_iptal(sid, "3", "Artık gerekmiyor")
        self.assertEqual(self._satir(tid)["kalan"], Decimal("0"))
        self.assertEqual(S.detay(tid)["gorunen_durum"], "TAMAMEN SİPARİŞE AKTARILDI")
        with self.assertRaises(ValueError):
            S.iptal_et(tid)
        n1 = self._talep(talep_no="TAL000050")
        n2 = self._talep(talep_no="TAL000050")
        nolar = {S.detay(n1)["talep_no"], S.detay(n2)["talep_no"]}
        self.assertEqual(len(nolar), 2)

    def test_liste_filtreleri_ve_uyarilar(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        t1 = self._talep(departman="Üretim")
        t2 = self._talep([{"urun_kodu": "U002", "urun_adi": "Somun", "birim": "Adet", "miktar": "5"}],
                         oncelik="ACİL")
        S.onaya_gonder(t2)
        self.assertEqual([r["id"] for r in S.listele(gorunum="onay_bekleyen")], [t2])
        self.assertEqual([r["id"] for r in S.listele(stok="vida")], [t1])
        self.assertEqual([r["id"] for r in S.listele(oncelik="ACİL")], [t2])
        self.assertEqual([r["id"] for r in S.listele(departman_depo="üretim")], [t1])
        self.assertEqual([r["id"] for r in S.listele("ONAYA GÖNDERİLDİ")], [t2])
        toplamlar = {r["id"]: r["tahmini_toplam"] for r in S.listele(no="TAL")}
        self.assertIsNone(toplamlar[t2])
        with get_session() as s:
            s.execute(text("UPDATE satin_alma_talepleri SET ihtiyac_tarihi = :d WHERE id = :i"),
                      {"d": BUGUN - timedelta(days=2), "i": t1})
        self.assertEqual([r["id"] for r in S.listele(gorunum="geciken")], [t1])
        uyarilar = S.acik_talep_uyarilari([{"urun_kodu": "U001"}])
        self.assertEqual(len(uyarilar), 1)
        self.assertEqual(S.acik_talep_uyarilari([{"urun_kodu": "U001"}], haric_talep_id=t1), [])
        bilgi = S.stok_bilgisi("U001", "ANA DEPO")
        self.assertEqual((bilgi["stok"], bilgi["bekleyen"], bilgi["diger_depo"], bilgi["en_yakin_termin"]),
                         (Decimal("0"), Decimal("0"), Decimal("0"), None))
        self.assertIsNotNone(bilgi["zaman"])

    def test_eski_sema_gecisi(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tmp = Path(self._tmpdir.name) / "eski.db"
        eng = create_engine(f"sqlite:///{tmp.as_posix()}")
        with eng.begin() as b:
            b.execute(text("CREATE TABLE stok_kartlari (id INTEGER PRIMARY KEY, stok_kodu VARCHAR(50))"))
            b.execute(text("INSERT INTO stok_kartlari (id, stok_kodu) VALUES (7, 'U001')"))
            b.execute(text(
                "CREATE TABLE satin_alma_talepleri (id INTEGER PRIMARY KEY, talep_no VARCHAR(30) UNIQUE NOT NULL, "
                "talep_tarihi DATE NOT NULL, ihtiyac_tarihi DATE, isteyen_kullanici VARCHAR(120), depo VARCHAR(100), "
                "oncelik VARCHAR(20) NOT NULL, durum VARCHAR(30) NOT NULL, aciklama TEXT, olusturma_tarihi DATETIME NOT NULL)"
            ))
            b.execute(text(
                "CREATE TABLE satin_alma_talep_satirlari (id INTEGER PRIMARY KEY, talep_id INTEGER NOT NULL, "
                "urun_kodu VARCHAR(50) NOT NULL, urun_adi VARCHAR(200) NOT NULL, birim VARCHAR(20) NOT NULL, "
                "miktar NUMERIC(18,4) NOT NULL, aciklama VARCHAR(500), depo VARCHAR(100))"
            ))
            b.execute(text("INSERT INTO satin_alma_talepleri VALUES (1,'TAL000001','2026-01-02',NULL,'x','ANA DEPO',"
                           "'NORMAL','ONAY BEKLİYOR',NULL,'2026-01-02 10:00:00')"))
            b.execute(text("INSERT INTO satin_alma_talep_satirlari VALUES (1,1,'U001','Vida','Adet',5,NULL,NULL)"))
        S.gecis_uygula(eng)
        S.gecis_uygula(eng)  # ikinci kez etkisiz
        kolonlar = {c["name"] for c in inspect(eng).get_columns("satin_alma_talep_satirlari")}
        self.assertTrue({"iptal_miktar", "stok_id", "manuel", "tahmini_birim_fiyat"} <= kolonlar)
        self.assertTrue(inspect(eng).has_table("satin_alma_talep_siparis_baglari"))
        with eng.begin() as b:
            self.assertEqual(b.execute(text("SELECT durum FROM satin_alma_talepleri")).scalar(), "ONAYA GÖNDERİLDİ")
            self.assertEqual(b.execute(text("SELECT stok_id, sira, iptal_miktar FROM satin_alma_talep_satirlari"))
                             .one(), (7, 1, 0))
        eng.dispose()


class SatinAlmaTalepKabulTest(_TalepTemel):
    """Talimat §12 kabul senaryoları."""

    def _siparis_cari(self, satirlar, cari_id):
        from database.alis_siparisi_service import AlisSiparisiService

        return AlisSiparisiService.kaydet(
            {"siparis_tarihi": BUGUN, "termin_tarihi": BUGUN + timedelta(days=7), "cari_id": cari_id},
            satirlar, [], None,
        )

    def _teslim_ve_fatura(self, siparis, miktar):
        from database.alis_faturasi_service import AlisFaturasiService
        from database.alis_irsaliyesi_service import AlisIrsaliyesiService
        from database.alis_siparisi_service import AlisSiparisiService

        ss = AlisSiparisiService.getir(siparis.id).satirlar[0]
        irs = AlisIrsaliyesiService.kaydet(
            {"irsaliye_tarihi": BUGUN, "cari_id": self.cari_id, "siparis_id": siparis.id},
            [{"urun_kodu": ss.urun_kodu, "urun_adi": ss.urun_adi, "miktar": Decimal(miktar), "birim": ss.birim,
              "birim_fiyat": Decimal("2"), "kdv_orani": Decimal("20"), "siparis_satiri_id": ss.id}],
        )
        irs_satir = AlisIrsaliyesiService.getir(irs.id).satirlar[0]
        fat = AlisFaturasiService.kaydet(
            {"fatura_tarihi": BUGUN, "vade_tarihi": BUGUN, "cari_id": self.cari_id, "irsaliye_id": irs.id,
             "depo": "ANA DEPO", "odeme_tutari": Decimal("0")},
            [{"urun_kodu": ss.urun_kodu, "urun_adi": ss.urun_adi, "miktar": Decimal(miktar), "birim": ss.birim,
              "birim_fiyat": Decimal("2"), "iskonto_orani": Decimal("0"), "kdv_orani": Decimal("20"),
              "irsaliye_satiri_id": irs_satir.id}],
        )
        return irs, fat

    def test_iade_net_teslim_ve_yeniden_acma(self):
        from database.alis_faturasi_service import AlisFaturasiService
        from database.alis_iade_faturasi_service import AlisIadeFaturasiService
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid, sid = self._onayli_talep()
        sip = self._siparis(S.siparis_satirlari_hazirla({sid: "10"})["satirlar"])
        _irs, fat = self._teslim_ve_fatura(sip, "10")
        fat_satir = AlisFaturasiService.getir(fat.id).satirlar[0]
        iade = AlisIadeFaturasiService.kaydet(
            {"iade_tarihi": BUGUN, "cari_id": self.cari_id, "kaynak_fatura_id": fat.id, "depo": "ANA DEPO"},
            [{"urun_kodu": "U001", "urun_adi": "Vida", "miktar": Decimal("2"), "birim": "Adet",
              "birim_fiyat": Decimal("2"), "kdv_orani": Decimal("20"), "kaynak_fatura_satiri_id": fat_satir.id}],
        )
        s = self._satir(tid)
        self.assertEqual((s["teslim"], s["iade"], s["net_teslim"], s["kalan"], s["yeniden_acilabilir"]),
                         (Decimal("10"), Decimal("2"), Decimal("8"), Decimal("0"), Decimal("2")))
        # İade ihtiyacı kendiliğinden açmaz; açma kullanıcı kararıdır
        with self.assertRaises(ValueError):
            S.siparis_satirlari_hazirla({sid: "1"})
        with self.assertRaises(ValueError):
            S.iadeyi_yeniden_ac(sid, "1", "")
        with self.assertRaises(ValueError):
            S.iadeyi_yeniden_ac(sid, "3", "fazla")
        S.iadeyi_yeniden_ac(sid, "2", "Hasarlı geldi, yeniden alınacak")
        s = self._satir(tid)
        self.assertEqual((s["kalan"], s["yeniden_acilan"], s["yeniden_acilabilir"]),
                         (Decimal("2"), Decimal("2"), Decimal("0")))
        with self.assertRaises(ValueError):
            S.siparis_satirlari_hazirla({sid: "3"})
        self._siparis(S.siparis_satirlari_hazirla({sid: "2"})["satirlar"])
        s = self._satir(tid)
        self.assertEqual((s["siparis"], s["kalan"]), (Decimal("12"), Decimal("0")))
        self.assertIn("İADEYİ YENİDEN AÇ", [g["islem"] for g in S.detay(tid)["gecmis"]])
        rapor = S.urun_bazli_rapor()[0]
        self.assertEqual((rapor["iade"], rapor["net_teslim"], rapor["bekleyen_teslim"], rapor["yeniden_acilan"]),
                         (Decimal("2"), Decimal("8"), Decimal("2"), Decimal("2")))
        # İade iptal edilirse iade miktarı düşer
        AlisIadeFaturasiService.iptal_et(iade.id)
        self.assertEqual(self._satir(tid)["iade"], Decimal("0"))

    def test_iki_tedarikciye_bolme_ve_iki_talepten_tek_siparis(self):
        from database.alis_irsaliyesi_service import AlisIrsaliyesiService
        from database.alis_siparisi_service import AlisSiparisiService
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        with get_session() as s:
            c2 = Cari(cari_kodu="T002", unvan="Tedarikçi B", cari_turu="Tedarikçi", aktif=True)
            s.add(c2)
            s.flush()
            cari2 = int(c2.id)
        tid, sid = self._onayli_talep()
        self._siparis_cari(S.siparis_satirlari_hazirla({sid: "6"})["satirlar"], self.cari_id)
        self._siparis_cari(S.siparis_satirlari_hazirla({sid: "4"})["satirlar"], cari2)
        s = self._satir(tid)
        self.assertEqual((s["siparis"], s["kalan"], len(s["siparisler"])), (Decimal("10"), Decimal("0"), 2))
        self.assertEqual([t["id"] for t in S.tedarikci_talepleri(cari2)], [tid])

        u2 = [{"urun_kodu": "U002", "urun_adi": "Somun", "birim": "Adet", "miktar": "5"}]
        t1 = self._talep(u2)
        t2 = self._talep([{**u2[0], "miktar": "3"}])
        for t in (t1, t2):
            S.onaya_gonder(t)
            S.onayla(t)
        s1, s2 = self._satir(t1)["id"], self._satir(t2)["id"]
        hazir = S.siparis_satirlari_hazirla({s1: "5", s2: "3"})["satirlar"]
        tek = {**hazir[0], "miktar": Decimal("8"),
               "talep_paylari": hazir[0]["talep_paylari"] + hazir[1]["talep_paylari"]}
        sip = self._siparis([tek])
        self.assertEqual((self._satir(t1)["siparis"], self._satir(t2)["siparis"]), (Decimal("5"), Decimal("3")))
        self.assertEqual({t["talep_no"] for t in S.siparis_talepleri(sip.id)},
                         {S.detay(t1)["talep_no"], S.detay(t2)["talep_no"]})
        ss = AlisSiparisiService.getir(sip.id).satirlar[0]
        AlisIrsaliyesiService.kaydet(
            {"irsaliye_tarihi": BUGUN, "cari_id": self.cari_id, "siparis_id": sip.id},
            [{"urun_kodu": "U002", "urun_adi": "Somun", "miktar": Decimal("6"), "birim": "Adet",
              "birim_fiyat": Decimal("1"), "kdv_orani": Decimal("20"), "siparis_satiri_id": ss.id}],
        )
        # 6 teslim bir kez sayılır: ilk talebe 5, ikinciye 1
        self.assertEqual((self._satir(t1)["teslim"], self._satir(t2)["teslim"]), (Decimal("5"), Decimal("1")))
        self.assertEqual(S.detay(t1)["gorunen_durum"], "TAMAMLANDI")
        self.assertEqual(S.detay(t2)["gorunen_durum"], "KISMEN TESLİM ALINDI")

    def test_eszamanli_aktarim_ve_cift_tiklama(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid, sid = self._onayli_talep()
        kullanici_a = S.siparis_satirlari_hazirla({sid: "6"})["satirlar"]
        kullanici_b = S.siparis_satirlari_hazirla({sid: "6"})["satirlar"]
        self._siparis(kullanici_a)
        with self.assertRaises(ValueError) as ctx:
            self._siparis(kullanici_b)
        self.assertIn("aşıldı", str(ctx.exception))
        with self.assertRaises(ValueError):
            self._siparis(kullanici_a)  # aynı hazırlığın ikinci kez kaydı (çift tıklama)
        s = self._satir(tid)
        self.assertEqual((s["siparis"], s["kalan"], len(s["siparisler"])), (Decimal("6"), Decimal("4"), 1))

    def test_hata_olursa_geri_alinir(self):
        from database.models.alis_siparisi import AlisSiparisi
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid, sid = self._onayli_talep()
        hazir = S.siparis_satirlari_hazirla({sid: "4"})["satirlar"]
        hatali = {**hazir[0], "birim": "Kutu"}  # birim uyuşmazlığı bağ yazımında hata verir
        with self.assertRaises(ValueError):
            self._siparis([hatali])
        with get_session() as s:
            self.assertEqual(s.scalar(select(func.count()).select_from(AlisSiparisi)), 0)
        self.assertEqual(self._satir(tid)["kalan"], Decimal("10"))

    def test_sil_ve_ek_kontrolleri(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        taslak = self._talep()
        klasor = Path(self._tmpdir.name)
        pdf = klasor / "teklif.pdf"
        pdf.write_bytes(b"%PDF-1.4 test")
        ek_id = S.ek_ekle(taslak, str(pdf))
        self.assertTrue(pdf.exists())  # kaynak korunur
        for ad, icerik in (("zararli.exe", b"MZ"), ("bos.pdf", b"")):
            (klasor / ad).write_bytes(icerik)
            with self.assertRaises(ValueError):
                S.ek_ekle(taslak, str(klasor / ad))
        with patch("database.satin_alma_talep_service.EK_AZAMI_BAYT", 4):
            with self.assertRaises(ValueError) as ctx:
                S.ek_ekle(taslak, str(pdf))
            self.assertIn("büyük", str(ctx.exception))
        yol = S.ek_yolu(ek_id)
        yol.unlink()
        with self.assertRaises(ValueError) as ctx:
            S.ek_yolu(ek_id)
        self.assertIn("bulunamadı", str(ctx.exception))
        S.ek_sil(ek_id)
        self.assertEqual(S.detay(taslak)["ekler"], [])
        S.sil(taslak)
        self.assertEqual(S.listele(), [])

        tid, sid = self._onayli_talep()
        with self.assertRaises(ValueError):
            S.sil(tid)
        S.geri_gonder(tid, "düzeltme")
        self.assertEqual(S.detay(tid)["durum"], "TASLAK")
        S.sil(tid)  # bağı olmayan taslak silinebilir
        tid2, sid2 = self._onayli_talep()
        self._siparis(S.siparis_satirlari_hazirla({sid2: "2"})["satirlar"])
        with self.assertRaises(ValueError):
            S.sil(tid2)

    def test_gorunumler_ve_rapor_kapsami(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        t1, _s1 = self._onayli_talep()
        t2, s2 = self._onayli_talep()
        t3 = self._talep()
        self._siparis(S.siparis_satirlari_hazirla({s2: "3"})["satirlar"])
        self.assertEqual([r["id"] for r in S.listele(gorunum="aktarilmamis")], [t1])
        self.assertEqual([r["id"] for r in S.listele(gorunum="kismen")], [t2])
        self.assertIn(t3, [r["id"] for r in S.listele()])
        rapor = S.urun_bazli_rapor(depo="ANA DEPO")[0]
        self.assertEqual((rapor["talep"], rapor["onayli"], rapor["siparis"], rapor["bekleyen_teslim"],
                          rapor["kalan"]),
                         (Decimal("30"), Decimal("20"), Decimal("3"), Decimal("3"), Decimal("17")))
        self.assertEqual(S.urun_bazli_rapor(depo="YOK DEPO"), [])
        self.assertEqual(S.urun_bazli_rapor(stok="somun"), [])

    def test_stok_diger_depo_ve_bekleyen_termin(self):
        from database.alis_faturasi_service import AlisFaturasiService
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        with get_session() as s:
            s.add(Depo(ad="YAN DEPO", aktif=True))
        AlisFaturasiService.kaydet(
            {"fatura_tarihi": BUGUN, "vade_tarihi": BUGUN, "cari_id": self.cari_id, "depo": "YAN DEPO",
             "odeme_tutari": Decimal("0")},
            [{"urun_kodu": "U001", "urun_adi": "Vida", "miktar": Decimal("7"), "birim": "Adet",
              "birim_fiyat": Decimal("2"), "iskonto_orani": Decimal("0"), "kdv_orani": Decimal("20")}],
        )
        tid, sid = self._onayli_talep()
        self._siparis(S.siparis_satirlari_hazirla({sid: "5"})["satirlar"])
        bilgi = S.stok_bilgisi("U001", "ANA DEPO")
        self.assertEqual((bilgi["stok"], bilgi["diger_depo"], bilgi["toplam_stok"], bilgi["bekleyen"]),
                         (Decimal("0"), Decimal("7"), Decimal("7"), Decimal("5")))
        self.assertEqual(bilgi["en_yakin_termin"], BUGUN + timedelta(days=7))
        self.assertEqual(S.depo_stoklari("U001"), [{"depo": "YAN DEPO", "miktar": Decimal("7")}])
        satir = self._satir(tid)
        self.assertEqual((satir["diger_depo_stok"], satir["siparis_termin"]), (Decimal("7"), BUGUN + timedelta(days=7)))

    def test_talebin_stok_cari_finans_etkisi_yok(self):
        from database.models.cari import CariIslem
        from database.models.finans import FinansHareketi
        from database.models.stok import StokHareketi, StokLotu
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        def say():
            with get_session() as s:
                return tuple(s.scalar(select(func.count()).select_from(m))
                             for m in (StokHareketi, StokLotu, CariIslem, FinansHareketi))

        once = say()
        tid, sid = self._onayli_talep()
        S.kalan_iptal(sid, "1", "fazla istendi")
        S.tedarikci_talepleri(self.cari_id)
        self.assertEqual(say(), once)

    def test_bagli_belge_zinciri(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid, sid = self._onayli_talep()
        sip = self._siparis(S.siparis_satirlari_hazirla({sid: "4"})["satirlar"])
        irs, fat = self._teslim_ve_fatura(sip, "4")
        belgeler = S.bagli_belgeler(tid)
        self.assertEqual([(b["tur"], b["id"]) for b in belgeler],
                         [("Sipariş", sip.id), ("İrsaliye", irs.id), ("Fatura", fat.id)])
        self.assertEqual(self._satir(tid)["teslim"], Decimal("4"))

    def test_yetki_ve_tekrar_gerekcesi(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S

        tid = S.kaydet({"talep_tarihi": BUGUN, "depo": "ANA DEPO"},
                       [{"urun_kodu": "U001", "urun_adi": "Vida", "birim": "Adet", "miktar": "3"}],
                       tekrar_gerekcesi="Önceki talep başka proje için")
        g = [x for x in S.detay(tid)["gecmis"] if x["islem"] == "TEKRAR TALEP GEREKÇESİ"]
        self.assertEqual(g[0]["detay"], "Önceki talep başka proje için")
        S.onaya_gonder(tid)
        S.onayla(tid)
        sid = self._satir(tid)["id"]
        self._kullanici(20, "DEPO", ("stok_goruntuleme",))
        with self.assertRaises(PermissionError):
            S.listele()
        with self.assertRaises(PermissionError):
            S.detay(tid)
        self._kullanici(21, "FINANS", ("alis_talep_duzenleme", "alis_talep_goruntuleme", "alis_siparis_duzenleme"))
        with self.assertRaises(PermissionError):
            S.siparis_satirlari_hazirla({sid: "1"})
        self._kullanici(22, "FINANS", ("alis_talep_goruntuleme", "alis_talep_aktarim", "alis_siparis_duzenleme"))
        self.assertEqual(len(S.siparis_satirlari_hazirla({sid: "1"})["satirlar"]), 1)
        with self.assertRaises(PermissionError):
            S.kalan_iptal(sid, "1", "gerek yok")


class SatinAlmaTalepCiktiTest(_TalepTemel):
    def test_form_modeli_fiyat_ve_filigran(self):
        from database.satin_alma_talep_service import SatinAlmaTalepService as S
        from database.talep_form_view import build_talep_form, kolonlar, render_talep_form_html
        from invoice_print.talep_cikti import varsayilan_dosya_adi

        tid = self._talep(departman="Bakım", talep_nedeni="Hat revizyonu")
        vm = build_talep_form(tid, fiyatli=False)
        self.assertEqual(vm.filigran, "TASLAK")
        self.assertNotIn("Tahmini B. Fiyat", [k for k, *_ in kolonlar(vm)])
        html_metin = render_talep_form_html(vm)
        self.assertIn("M6x20 galvaniz", html_metin)
        self.assertIn("KONTROL EDEN", html_metin)
        self.assertNotIn("2,50", html_metin)
        self.assertTrue(varsayilan_dosya_adi(vm, "pdf").startswith(f"Satin_Alma_Talebi_{vm.talep_no}_"))

        S.onaya_gonder(tid)
        S.onayla(tid)
        vm = build_talep_form(tid, fiyatli=True)
        self.assertEqual(vm.filigran, "")
        self.assertIn("Tahmini Tutar", [k for k, *_ in kolonlar(vm)])
        self.assertEqual(vm.tahmini_toplam_goster, "25,00")
        self.assertIn("Dijital olarak onaylandı", render_talep_form_html(vm))

        # Fiyat yetkisi yoksa fiyatlı çıktı istense de fiyat modele girmez
        self._kullanici(9, "GORUNTULEME", ("alis_talep_goruntuleme",))
        self.assertFalse(build_talep_form(tid, fiyatli=True).fiyatli)

    def test_docx_fiyatsiz(self):
        from docx import Document

        from database.talep_form_view import build_talep_form
        from invoice_print.talep_cikti import docx_olustur

        tid = self._talep()
        hedef = Path(self._tmpdir.name) / "talep.docx"
        docx_olustur(build_talep_form(tid, fiyatli=False), hedef)
        doc = Document(str(hedef))
        metin = "".join(doc.element.body.itertext())
        self.assertIn("SATIN ALMA TALEBİ", metin)
        self.assertIn("M6x20 galvaniz", metin)
        self.assertNotIn("Tahmini", metin)
        self.assertIn("TASLAK", "\n".join(p.text for p in doc.sections[0].header.paragraphs))

    def test_dovizli_fiyat_tahmini_not_ve_ek_listesi(self):
        from docx import Document

        from database.satin_alma_talep_service import SatinAlmaTalepService as S
        from database.talep_form_view import build_talep_form, render_talep_form_html
        from invoice_print.talep_cikti import docx_olustur

        with self.assertRaises(ValueError) as ctx:
            self._talep([{"urun_kodu": "U001", "urun_adi": "Vida", "birim": "Adet", "miktar": "2",
                          "tahmini_birim_fiyat": "3", "para_birimi": "USD"}])
        self.assertIn("kur", str(ctx.exception))
        tid = self._talep([
            {"urun_kodu": "U001", "urun_adi": "Vida", "birim": "Adet", "miktar": "2", "tahmini_birim_fiyat": "3",
             "para_birimi": "USD", "kur": "40", "fiyat_kaynagi": "Tedarikçi teklifi"},
            {"urun_kodu": "U002", "urun_adi": "Somun", "birim": "Adet", "miktar": "4", "tahmini_birim_fiyat": "5"},
        ])
        ek = Path(self._tmpdir.name) / "teknik_resim.pdf"
        ek.write_bytes(b"%PDF-1.4")
        S.ek_ekle(tid, str(ek))
        vm = build_talep_form(tid, fiyatli=True, ek_listesi=True)
        self.assertEqual([s.tutar_goster for s in vm.satirlar], ["6,00 USD", "20,00 TL"])
        self.assertEqual(vm.tahmini_toplam_goster, "260,00")  # 6 USD × 40 + 20 TL
        self.assertIn("finansal kayıt değildir", vm.tahmini_notu)
        self.assertIn("kurla", vm.tahmini_notu)
        self.assertEqual(vm.ekler, ["teknik_resim.pdf"])
        html_metin = render_talep_form_html(vm)
        self.assertIn("teknik_resim.pdf", html_metin)
        self.assertIn("SATIN ALMA TALEBİ", html_metin)
        self.assertEqual(build_talep_form(tid, fiyatli=True).ekler, [])  # ek listesi isteğe bağlı
        hedef = Path(self._tmpdir.name) / "dovizli.docx"
        docx_olustur(vm, hedef)
        metin = "".join(Document(str(hedef)).element.body.itertext())
        self.assertIn("finansal kayıt değildir", metin)
        self.assertIn("teknik_resim.pdf", metin)


class SatinAlmaTalepEkranTest(_TalepTemel):
    def test_form_onay_aktarim_ve_siparis_karti(self):
        import tkinter as tk

        from alis_ui import AlisSiparisiDialog
        from database.alis_siparisi_service import AlisSiparisiService
        from database.satin_alma_talep_service import SatinAlmaTalepService as S
        from satin_alma_talep_ui import SatinAlmaTalepDialog, TalepAktarDialog

        try:
            root = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk yok")
        root.withdraw()
        try:
            with patch("satin_alma_talep_ui.messagebox") as mb, patch("alis_ui.messagebox") as mb2:
                mb.askyesno.return_value = True
                mb2.askyesno.return_value = True
                d = SatinAlmaTalepDialog(root)
                d.withdraw()
                satir = d._yeni_satir(("U001", "Vida", "Adet"))
                satir["miktar"] = Decimal("10")
                d.satirlar.append(satir)
                d.manuel_satir_ekle_degerle("Özel conta", "Adet", "2", "Ø40")
                d._satirlari_yenile()
                d.departman.insert(0, "Bakım")
                self.assertTrue(d._kaydet())
                self.assertEqual(d._durum, "TASLAK")
                self.assertIn("kaydedildi", d.mesaj_lbl.cget("text"))
                self.assertTrue(d.onaya_gonder())
                self.assertEqual(str(d.kaydet_btn.cget("state")), "disabled")
                self.assertTrue(d.onayla())
                self.assertEqual(d._durum, "ONAYLANDI")
                self.assertEqual(str(d.siparis_btn.cget("state")), "normal")
                self.assertEqual(str(d.teklif_btn.cget("state")), "normal")
                tid = d.talep_id
                d.destroy()

                a = TalepAktarDialog(root, talep_id=tid)
                a.withdraw()
                vida = next(s for s in a.satirlar if s["urun_kodu"] == "U001")
                conta = next(s for s in a.satirlar if s["eslesmemis"])
                self.assertNotIn(conta["talep_satiri_id"], a.secili)
                with self.assertRaises(ValueError):
                    a.miktar_ayarla(vida["talep_satiri_id"], "11")
                a.miktar_ayarla(vida["talep_satiri_id"], "4")
                a.tamam()
                self.assertEqual(a.result, {vida["talep_satiri_id"]: Decimal("4")})

                hazir = S.siparis_satirlari_hazirla(a.result)["satirlar"]
                sp = AlisSiparisiDialog(root, hazir_satirlar=hazir)
                sp.withdraw()
                sp.talep_satirlarini_ekle(S.siparis_satirlari_hazirla({vida["talep_satiri_id"]: "2"})["satirlar"])
                self.assertEqual(len(sp.satirlar), 1)
                self.assertEqual(sp.satirlar[0]["miktar"], Decimal("6"))
                self.assertEqual(len(sp.satirlar[0]["talep_paylari"]), 2)
                self.assertTrue(sp._alis_bagli(sp.satirlar[0]))
                anahtar = next(k for k, v in sp.tedarikci_map.items() if int(v.id) == int(self.cari_id))
                sp.tedarikci.set(anahtar)
                sp.kaydet()
                self.assertTrue(getattr(sp, "result", False))
                self.assertEqual(self._satir(tid)["siparis"], Decimal("6"))

                sid = self._satir(tid)["siparisler"][0]["siparis_id"]
                sp2 = AlisSiparisiDialog(root, siparis=AlisSiparisiService.getir(sid))
                sp2.withdraw()
                self.assertEqual(sum(p["miktar"] for p in sp2.satirlar[0]["talep_paylari"]), Decimal("6"))
                bt = sp2.bagli_talepler()
                bt.withdraw()
                self.assertEqual([r["id"] for r in bt.kayitlar], [tid])
                bt.destroy()
                sp2.destroy()

                d2 = SatinAlmaTalepDialog(root, talep_id=tid)
                d2.withdraw()
                bb = d2.bagli_belgeler()
                bb.withdraw()
                self.assertEqual([(b["tur"], b["id"]) for b in bb.belgeler], [("Sipariş", sid)])
                bb.destroy()
                d2.destroy()
        finally:
            root.destroy()

    def test_form_kopyala_doviz_kolonlar_ve_uyarilar(self):
        import tkinter as tk

        from database.satin_alma_talep_service import SatinAlmaTalepService as S
        from satin_alma_talep_ui import SatinAlmaTalepDialog

        try:
            root = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk yok")
        root.withdraw()
        try:
            with patch.dict("os.environ", {"LOCALAPPDATA": self._tmpdir.name}), \
                    patch("satin_alma_talep_ui.messagebox") as mb:
                mb.askyesno.return_value = True
                d = SatinAlmaTalepDialog(root)
                d.withdraw()
                self.assertEqual(d.title().split(" — ")[0], "SATIN ALMA TALEBİ")
                satir = d._yeni_satir(("U001", "Vida", "Adet"))
                satir["miktar"] = Decimal("10")
                d.satirlar.append(satir)
                d._satirlari_yenile()
                d.satir_kopyala(0)
                self.assertEqual(len(d.satirlar), 2)
                self.assertNotIn("id", d.satirlar[1])
                d.hucre_uygula(1, "fiyat", "3")
                with patch("database.doviz_service.DovizService.kur_degeri", return_value=Decimal("40")):
                    d.hucre_uygula(1, "pb", "USD")
                self.assertEqual((d.satirlar[1]["para_birimi"], d.satirlar[1]["kur"]), ("USD", Decimal("40")))
                self.assertEqual(d._kur_izin(0), "TL satırda kur girilmez; önce döviz seçin.")
                d.hucre_uygula(0, "fiyat", "2")
                metin = d.toplam_metni()
                self.assertIn("20,00 TRY", metin)
                self.assertIn("30,00 USD", metin)
                self.assertIn("1.220,00 TRY", metin)  # farklı dövizler yalnız kurla TL karşılığında toplanır
                self.assertIn("finansal kayıt değildir", d.toplam_lbl.master.winfo_children()[1].cget("text"))
                self.assertTrue(d._kaydet())
                tid = d.talep_id
                vals = d.satir_tablosu.item("1", "values")
                kolonlar = [k for k, *_ in d._KOLONLAR]
                self.assertEqual(vals[kolonlar.index("kur")], "40")
                self.assertEqual(vals[kolonlar.index("tutar")], "30,00")
                self.assertTrue(d.stok_zaman_lbl.cget("text").startswith("Stok bilgisi:"))
                d.destroy()

                # Aynı ürün için ikinci talep: gerekçe sorulur, gerekçesiz kayıt yapılmaz
                d = SatinAlmaTalepDialog(root)
                d.withdraw()
                d.satirlar.append(d._yeni_satir(("U001", "Vida", "Adet")))
                d._satirlari_yenile()
                with patch.object(SatinAlmaTalepDialog, "_tekrar_gerekcesi_sor", return_value=None):
                    self.assertFalse(d._kaydet())
                self.assertIsNone(d.talep_id)
                with patch.object(SatinAlmaTalepDialog, "_tekrar_gerekcesi_sor", return_value="Farklı hat"):
                    self.assertTrue(d._kaydet())
                self.assertIn("Farklı hat", [g["detay"] for g in S.detay(d.talep_id)["gecmis"]])
                tid2 = d.talep_id
                self.assertTrue(d.talep_sil())
                with self.assertRaises(ValueError):
                    S.detay(tid2)

                # Fiyat yetkisi yoksa fiyat kolonları görünmez
                self._kullanici(30, "FINANS", ("alis_talep_duzenleme", "alis_talep_goruntuleme"))
                d = SatinAlmaTalepDialog(root, talep_id=tid)
                d.withdraw()
                gorunen = set(d.satir_tablosu["displaycolumns"])
                self.assertFalse(gorunen & set(d._FIYAT_KOLONLARI))
                self.assertIn("yetkiniz yok", d.toplam_metni())
                d.destroy()
        finally:
            root.destroy()

    def test_liste_siralama_dogru_kaydi_acar_ve_tedarikci_sekmeleri(self):
        import tkinter as tk

        from database.satin_alma_talep_service import SatinAlmaTalepService as S
        from tedarikci_evrak_ui import TedarikciEvraklariDialog

        try:
            root = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk yok")
        root.withdraw()
        try:
            t1 = self._talep(departman="Zeta", oncelik="DÜŞÜK")
            t2 = self._talep(departman="Alfa", oncelik="ACİL")
            tid, sid = self._onayli_talep(departman="Beta")
            sip = self._siparis(S.siparis_satirlari_hazirla({sid: "3"})["satirlar"])
            kok = tk.Frame(root)
            acilan = []
            with patch("satin_alma_talep_ui._liste_ust", return_value=kok), \
                    patch("satin_alma_talep_ui.SatinAlmaTalepDialog", side_effect=lambda _a, talep_id=None:
                          acilan.append(talep_id) or tk.Toplevel(root)), \
                    patch.object(root, "wait_window", lambda _w: _w.destroy()):
                from satin_alma_talep_ui import satin_alma_talep_listesi_goster

                ui = satin_alma_talep_listesi_goster(root)
                tablo = ui["tablo"]
                ui["sirala"]("departman")
                self.assertEqual([int(i) for i in tablo.get_children()], [t2, tid, t1])
                ui["sirala"]("departman")
                self.assertEqual([int(i) for i in tablo.get_children()], [t1, tid, t2])
                ui["sirala"]("oncelik")
                ilk = int(tablo.get_children()[0])
                self.assertEqual(ilk, t2)
                tablo.selection_set(str(ilk))
                ui["secili_ac"]()
                self.assertEqual(acilan, [t2])
                ui["gorunum"].set("kismen")
                ui["yenile"]()
                self.assertEqual([int(i) for i in tablo.get_children()], [tid])

            with get_session() as s:
                cari = s.get(Cari, self.cari_id)
            with patch.object(tk.Toplevel, "grab_set", lambda _self: None):
                ted = TedarikciEvraklariDialog(root, cari, sekme="siparis")
            ted.withdraw()
            self.assertEqual(ted.tablolar["siparis"].get_children(), (str(sip.id),))
            self.assertEqual(ted.tablolar["talep"].get_children(), (str(tid),))
            ted.destroy()
        finally:
            root.destroy()


if __name__ == "__main__":
    unittest.main()
