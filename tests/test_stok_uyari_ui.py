"""Stok uyarıları ekranı: açılış, arka plan liste, ayar, taslak sipariş, Excel ve menü/hub bağlantısı."""

from __future__ import annotations

import sys
import time
import tkinter as tk
import unittest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import select  # noqa: E402

from database.database import get_session  # noqa: E402
from database.models.alis_siparisi import AlisSiparisi  # noqa: E402
from tests.stok_test_ortami import StokOrtami  # noqa: E402

T1, T2 = date.today() - timedelta(days=5), date.today() - timedelta(days=4)


class _App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.withdraw()
        self.icerik = tk.Frame(self)
        self.icerik.pack(fill="both", expand=True)

    def _icerigi_temizle(self):
        for w in self.icerik.winfo_children():
            w.destroy()


class StokUyariEkranTest(unittest.TestCase):
    def setUp(self):
        self.o = StokOrtami("ui")
        try:
            self.app = _App()
        except tk.TclError as exc:
            self.o.kapat()
            self.skipTest(f"Tk yok: {exc}")
        self.hatalar: list = []
        self.app.report_callback_exception = lambda *e: self.hatalar.append(e)
        y = patch("stok_uyari_ui.messagebox")
        self.mb = y.start()
        self.mb.askyesno.return_value = False
        self.addCleanup(y.stop)
        self.o.alis(1, 100, T1)
        self.o.satis(1, 150, T2)

    def tearDown(self):
        try:
            self.app.destroy()
        except tk.TclError:
            pass
        self.o.kapat()

    def _bekle(self, kosul, sure=15.0):
        bitis = time.monotonic() + sure
        while time.monotonic() < bitis:
            self.app.update()
            if kosul():
                return True
            time.sleep(0.03)
        return False

    def _ekran(self):
        from stok_uyari_ui import StokUyariEkrani

        ekran = StokUyariEkrani(self.app, lambda: None)
        self.assertTrue(self._bekle(lambda: bool(ekran.satirlar)), f"liste dolmadı: {self.hatalar}")
        return ekran

    def test_acilis_liste_ve_ozet(self):
        ekran = self._ekran()
        self.assertEqual(self.hatalar, [])
        s = next(iter(ekran.satirlar.values()))
        self.assertEqual(s["stok_kodu"], "U001")
        iid = str(s["id"])
        self.assertEqual(ekran.tablo.set(iid, "neden"), "Tükendi")
        self.assertEqual(ekran.tablo.set(iid, "oneri"), "Kullanıcı belirleyecek")
        self.assertIn("tukendi", ekran.tablo.item(iid, "tags"))
        self.assertIn("Etkin ihtiyaç: 1", ekran.ozet_lbl.cget("text"))
        self.assertIn("rezervasyon", ekran.not_lbl.cget("text").lower())
        self.assertTrue(self._bekle(lambda: self.mb.askyesno.called))  # ilk tarama sorusu

    def test_ayar_sonrasi_oneri_ve_taslak_siparis(self):
        from database.stok_uyari_service import StokUyariService
        from stok_uyari_ui import SiparisDialog

        ekran = self._ekran()
        s = next(iter(ekran.satirlar.values()))
        StokUyariService.ayar_kaydet(s["stok_id"], None, {"minimum": "5", "hedef": "30", "alim_birimi": "Koli"})
        ekran.yenile()
        self.assertTrue(self._bekle(lambda: next(iter(ekran.satirlar.values()))["hedef"] is not None))
        s = next(iter(ekran.satirlar.values()))
        self.assertEqual(ekran.tablo.set(str(s["id"]), "oneri"), "3 Koli")
        guncel = StokUyariService.guncel_kontrol([s["id"]])
        dlg = SiparisDialog(self.app, guncel, True)
        dlg.girdiler[0]["v"]["ted"].set(dlg.cariler[0][1] if dlg.cariler[0][0] == self.o.ted_id
                                        else next(m for i, m in dlg.cariler if i == self.o.ted_id))
        dlg.girdiler[0]["v"]["fiyat"].set("1.200,50")
        dlg._olustur()
        self.assertEqual(len(dlg.olusan), 1)
        with get_session() as ses:
            sip = ses.scalar(select(AlisSiparisi))
            self.assertEqual(sip.durum, "TASLAK")
            self.assertEqual(sip.satirlar[0].birim, "Koli")
            self.assertEqual(Decimal(str(sip.satirlar[0].miktar)), Decimal("3"))
            self.assertEqual(Decimal(str(sip.satirlar[0].birim_alis_fiyati)), Decimal("1200.50"))

    def test_excel_gorunen_kolonlari_yazar(self):
        from openpyxl import load_workbook

        ekran = self._ekran()
        hedef = Path(self.o.db_path).parent / "uyari.xlsx"
        with patch("cek_senet_ui.filedialog.asksaveasfilename", return_value=str(hedef)), \
                patch("cek_senet_ui.messagebox"):
            ekran.excel()
        ws = load_workbook(hedef).active
        satirlar = list(ws.iter_rows(values_only=True))
        self.assertIn("Ürün Adı", satirlar[3])
        self.assertEqual(satirlar[3][0], "Tedarikçi")
        self.assertIn("Dönem:", satirlar[4][0])
        self.assertEqual(len(satirlar), 6)

    def test_baslik_siralama_kimligi_korur(self):
        ekran = self._ekran()
        self.o.urun_ekle("A001", "Alfa Ürün")
        self.assertTrue(self._bekle(lambda: not ekran._yukleniyor))
        ekran.yenile()
        self.assertTrue(self._bekle(lambda: len(ekran.satirlar) == 2 and not ekran._yukleniyor))
        ekran._sirala("kod")
        self.assertEqual([ekran.satirlar[i]["stok_kodu"] for i in ekran.tablo.get_children()], ["A001", "U001"])
        ekran._sirala("kod")
        cocuklar = ekran.tablo.get_children()
        self.assertEqual([ekran.satirlar[i]["stok_kodu"] for i in cocuklar], ["U001", "A001"])
        self.assertTrue(all(ekran.tablo.set(i, "kod") == ekran.satirlar[i]["stok_kodu"] for i in cocuklar))

    def test_menu_rozeti_ve_hub_karti(self):
        import stoklar_ui
        from ana_panel_ui import SolMenuDugme, stok_uyari_rozeti_guncelle

        dugme = SolMenuDugme(self.app, etiket="Stoklar", anahtar="stoklar", simge="#", vurgulu=False,
                             komut=lambda: None)
        self.app.menu_dugmeleri = {"stoklar": dugme}
        stok_uyari_rozeti_guncelle(self.app)
        self.assertEqual(dugme.lbl_text.cget("text"), "Stoklar (1)")
        self.assertIn("uyari", [a for *_x, a in stoklar_ui.STOKLAR_HUB_KARTLARI])
        kart = next(k for k in stoklar_ui._hub_kartlari_sayili() if k[2] == "uyari")
        self.assertIn("(1)", kart[0])


if __name__ == "__main__":
    unittest.main()
