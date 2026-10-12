"""Stok yatırım verimliliği (GMROI): elle hesaplanan küçük örnekle FIFO, iade, transfer, grup ve hesaplanamayan oranlar."""

from __future__ import annotations

import sys
import unittest
from datetime import date
from decimal import Decimal as D
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import func, select  # noqa: E402

from database import stok_rapor_service as srv  # noqa: E402
from database.database import get_session  # noqa: E402
from database.models.stok import Depo, StokGrubu, StokHareketi, StokKarti, StokLotu  # noqa: E402
from tests.stok_test_ortami import StokOrtami  # noqa: E402

BAS, BIT = date(2026, 9, 1), date(2026, 9, 30)
H = srv.HESAPLANAMADI


class _Taban(unittest.TestCase):
    def setUp(self):
        self.o = StokOrtami(self._testMethodName[:20])
        with get_session() as s:
            self.ana = s.scalar(select(Depo.id).where(Depo.ad == "ANA DEPO"))

    def tearDown(self):
        self.o.kapat()

    def rapor(self, **kw):
        return srv.yatirim_verimliligi(bas=BAS, bit=BIT, bugun=date(2026, 10, 8), **kw)

    @staticmethod
    def satir(r, kod):
        return next(x for x in r["urunler"] if x["urun_kodu"] == kod)

    def hareket_sayisi(self) -> int:
        with get_session() as s:
            return int(s.scalar(select(func.count(StokHareketi.id))) or 0)

    def u001_senaryo(self):
        """Dönem başı 10×100; 11.09 10×130 alış; 21.09 12 adet × 200 satış (FIFO maliyet 10×100 + 2×130)."""
        self.o.alis(10, 100, date(2026, 8, 20), lot="L1")
        self.o.alis(10, 130, date(2026, 9, 11), lot="L2")
        self.o.satis(12, 200, date(2026, 9, 21))


class FifoElHesabiTest(_Taban):
    def test_fifo_kismi_satis_donem_basi_ve_gunluk_ortalama(self):
        self.u001_senaryo()
        once = self.hareket_sayisi()
        r = self.rapor()
        u = self.satir(r, "U001")
        # Gün sonu değerleri: 1–10 Eyl 1000, 11–20 Eyl 2300, 21–30 Eyl 8×130 = 1040 → Σ 43.400 / 30 gün
        self.assertEqual((u["bas_miktar"], u["bas_deger"]), (D("10"), D("1000.00")))
        self.assertEqual((u["son_miktar"], u["son_deger"]), (D("8"), D("1040.00")))
        self.assertEqual(u["ort_stok"], D("1446.67"))
        self.assertEqual(u["net_satis"], D("2400.00"))
        self.assertEqual(u["net_miktar"], D("12"))
        self.assertEqual(u["smm"], D("1260.00"))
        self.assertEqual(u["brut_kar"], D("1140.00"))
        self.assertEqual(u["marj"], D("47.50"))
        self.assertEqual(u["devir"], D("0.87"))          # 1260 / 1446,67
        self.assertEqual(u["kalma_gun"], D("34.4"))      # 30 / 0,8710
        self.assertEqual(u["gmroi"], D("0.79"))          # 1140 / 1446,67 — yıllıklandırılmaz
        self.assertEqual(u["son_alis"], date(2026, 9, 11))
        self.assertEqual(u["son_satis"], date(2026, 9, 21))
        self.assertEqual(u["hesap_notu"], "")
        self.assertEqual(r["gun"], 30)
        self.assertIn("FIFO", r["_baslik"])
        self.assertTrue(any("günlük kapanış" in b for b in r["_bilgi"]))
        self.assertTrue(any("yıllıklandırılmaz" in n for n in r["notlar"]))
        self.assertTrue(any("net işletme kârı değildir" in n for n in r["notlar"]))
        self.assertEqual(self.hareket_sayisi(), once)

    def test_degerleme_ve_satis_raporlariyla_tutarli(self):
        self.u001_senaryo()
        r = self.rapor()
        u = self.satir(r, "U001")
        degerleme = srv.degerleme(tarih=BIT)
        self.assertEqual(u["son_deger"], degerleme["toplam_maliyet"])
        satis = srv.satis_ozeti(BAS, BIT)["U001"]
        self.assertEqual((u["net_satis"], u["smm"]), (satis["net"], satis["maliyet"]))

    def test_hareketli_ortalama_yontemi_tutarli(self):
        from database.satis_maliyet_motoru import HAREKETLI

        self.u001_senaryo()
        u = self.satir(self.rapor(yontem=HAREKETLI), "U001")
        # Ortalama 100 → 115 (20 adet); satış 12 × 115 = 1380; kalan 8 × 115 = 920
        self.assertEqual(u["smm"], D("1380.00"))
        self.assertEqual(u["son_deger"], D("920.00"))
        self.assertEqual(u["ort_stok"], D("1406.67"))    # (10×1000 + 10×2300 + 10×920) / 30
        self.assertEqual(u["gmroi"], D("0.73"))          # 1020 / 1406,67


class IadeTransferTest(_Taban):
    def test_satis_ve_alis_iadesi(self):
        from database.alis_faturasi_service import AlisFaturasiService
        from database.alis_iade_faturasi_service import AlisIadeFaturasiService
        from database.satis_faturasi_service import SatisFaturasiService
        from database.satis_iade_faturasi_service import SatisIadeFaturasiService

        self.o.urun_ekle("U002", "Kumaş", birim="Metre")
        a = self.o.alis(5, 50, date(2026, 8, 25), kod="U002", birim="Metre")
        f = self.o.satis(4, 80, date(2026, 9, 5), kod="U002", birim="Metre")
        ks = SatisFaturasiService.getir(f.id).satirlar[0]
        SatisIadeFaturasiService.kaydet(
            {"iade_tarihi": date(2026, 9, 15), "cari_id": self.o.mus_id, "depo": "ANA DEPO", "kaynak_fatura_id": f.id,
             "iade_odeme_tutari": D("0")},
            [{"urun_kodu": "U002", "urun_adi": "Kumaş", "miktar": D("1"), "birim": "Metre", "birim_fiyat": D("80"),
              "kdv_orani": D("20"), "kaynak_fatura_satiri_id": ks.id}])
        aks = AlisFaturasiService.getir(a.id).satirlar[0]
        AlisIadeFaturasiService.kaydet_ve_onayla(
            {"iade_tarihi": date(2026, 9, 18), "cari_id": self.o.ted_id, "depo": "ANA DEPO"},
            [{"urun_kodu": "U002", "urun_adi": "Kumaş", "miktar": D("1"), "birim": "Metre", "birim_fiyat": D("50"),
              "iskonto_orani": D("0"), "kdv_orani": D("20"), "kaynak_fatura_satiri_id": aks.id}])
        u = self.satir(self.rapor(), "U002")
        # Net satış 320 − 80; SMM 200 − 50; gün sonu: 4×250 + 10×50 + 3×100 + 13×50 = 2450 / 30
        self.assertEqual(u["net_satis"], D("240.00"))
        self.assertEqual(u["smm"], D("150.00"))
        self.assertEqual(u["brut_kar"], D("90.00"))
        self.assertEqual(u["net_miktar"], D("3"))
        self.assertEqual(u["son_miktar"], D("1"))
        self.assertEqual(u["son_deger"], D("50.00"))
        self.assertEqual(u["ort_stok"], D("81.67"))
        self.assertEqual(u["gmroi"], D("1.10"))
        self.assertEqual(u["birim"], "Metre")

    def test_transfer_satis_sayilmaz_depo_filtresi_ve_satissiz_stok(self):
        from database.stok_service import StokService

        self.o.urun_ekle("U003", "Aksesuar")
        self.o.alis(6, 10, date(2026, 8, 28), kod="U003")
        StokService.depo_transfer_kaydet(
            {"cikis_depo": "ANA DEPO", "giris_depo": "DEPO 2", "fis_tarihi": date(2026, 9, 10)},
            [{"urun_kodu": "U003", "urun_adi": "Aksesuar", "miktar": D("4")}])
        firma = self.satir(self.rapor(), "U003")
        self.assertEqual((firma["net_satis"], firma["smm"], firma["brut_kar"]), (D("0.00"), D("0.00"), D("0.00")))
        self.assertEqual(firma["ort_stok"], D("60.00"))
        self.assertEqual((firma["devir"], firma["gmroi"]), (D("0.00"), D("0.00")))
        self.assertEqual((firma["kalma_gun"], firma["marj"]), (H, H))
        self.assertIn("satış maliyeti yok", firma["hesap_notu"])
        d2 = self.satir(self.rapor(depo="DEPO 2"), "U003")
        self.assertEqual(d2["ort_stok"], D("28.00"))     # 9 gün 0 + 21 gün 4×10
        ana = self.satir(self.rapor(depo="ANA DEPO"), "U003")
        self.assertEqual(ana["ort_stok"], D("32.00"))    # 9 gün 60 + 21 gün 20
        self.assertEqual(ana["son_miktar"], D("2"))


class GrupVeHesaplanamayanTest(_Taban):
    def test_grup_oranlari_toplamlardan_ve_birim_karisik(self):
        from database.satis_faturasi_service import SatisFaturasiService  # noqa: F401

        with get_session() as s:
            g = StokGrubu(kod="MOB", ad="MOBİLYA", seviye=1)
            s.add(g)
            s.flush()
            gid = g.id
        self.u001_senaryo()
        u2 = self.o.urun_ekle("U002", "Kumaş", birim="Metre", ana_grup_id=gid)
        with get_session() as s:
            s.scalar(select(StokKarti).where(StokKarti.stok_kodu == "U001")).ana_grup_id = gid
        self.o.alis(2, 50, date(2026, 8, 25), kod="U002", birim="Metre")
        self.o.satis(1, 140, date(2026, 9, 16), kod="U002", birim="Metre")
        r = self.rapor()
        u1, u2r = self.satir(r, "U001"), self.satir(r, "U002")
        self.assertEqual(u2r["gmroi"], D("1.20"))        # kâr 90 / ort (15×100 + 15×50)/30 = 75
        grup = next(x for x in r["gruplar"] if x["ana_grup"] == "MOBİLYA")
        # Oran = Σ kâr / Σ ort. stok = 1230 / 1521,67 (satır GMROI'lerinin ortalaması 0,995 olurdu)
        self.assertEqual(grup["brut_kar"], D("1230.00"))
        self.assertEqual(grup["ort_stok"], D("1521.67"))
        self.assertEqual(grup["gmroi"], D("0.81"))
        self.assertEqual(grup["marj"], D("48.43"))       # 1230 / 2540
        self.assertEqual(grup["birim"], "Karışık")
        self.assertIsNone(grup["son_miktar"])
        self.assertEqual(grup["kapsam"], "2/2 ürün")
        toplam = r["urunler"][-1]
        self.assertTrue(toplam["_sabit"])
        self.assertEqual(toplam["net_satis"], u1["net_satis"] + u2r["net_satis"])
        self.assertEqual(r["gruplar"][-1]["ana_grup"], "GENEL TOPLAM")
        self.assertEqual(r["gruplar"][-1]["brut_kar"], toplam["brut_kar"])
        sadece = self.rapor(ana_grup_id=str(gid), grup_seviyesi="alt")
        self.assertEqual({x["urun_kodu"] for x in sadece["urunler"] if not x.get("_sabit")}, {"U001", "U002"})
        siralı = self.rapor(siralama="gmroi")["urunler"]
        self.assertEqual(siralı[0]["urun_kodu"], "U002")

    def test_eksik_maliyet_ve_negatif_stok_hesaplanamadi(self):
        self.o.urun_ekle("U004", "Maliyetsiz")
        self.o.alis(3, 40, date(2026, 8, 20), kod="U004")
        u5 = self.o.urun_ekle("U005", "Negatif")
        self.o.alis(1, 10, date(2026, 8, 20), kod="U005")
        with get_session() as s:
            u4 = s.scalar(select(StokKarti.id).where(StokKarti.stok_kodu == "U004"))
            for lot in s.scalars(select(StokLotu).where(StokLotu.stok_id == u4)).all():
                lot.birim_maliyet = D("0")
            for h in s.scalars(select(StokHareketi).where(StokHareketi.stok_id == u4)).all():
                h.birim_maliyet = D("0")
            s.add(StokHareketi(tarih=date(2026, 9, 12), hareket_turu="ÇIKIŞ", belge_no="MAN-1", stok_id=u5,
                               depo_id=self.ana, miktar=D("3"), birim_maliyet=D("10")))
        r = self.rapor()
        m = self.satir(r, "U004")
        self.assertEqual((m["ort_stok"], m["devir"], m["gmroi"], m["kalma_gun"]), (H, H, H, H))
        self.assertIn("eksik maliyet", m["hesap_notu"])
        n = self.satir(r, "U005")
        self.assertEqual((n["devir"], n["gmroi"]), (H, H))
        self.assertIn("Negatif stok (19 gün)", n["hesap_notu"])
        self.assertEqual((n["son_miktar"], n["son_deger"], n["ort_stok"]), (D("-2"), H, H))
        self.assertEqual(n["bas_deger"], D("10.00"))
        self.assertEqual({x["urun_kodu"] for x in r["hesaplanamayan"]} >= {"U004", "U005"}, True)
        self.assertEqual(r["toplam"]["kapsam"], "0/2 ürün")
        self.assertEqual(r["toplam"]["gmroi"], H)

    def test_devir_girisi_donem_icinde_eksik_devir_verisi(self):
        self.o.urun_ekle("U006", "Devirli")
        with get_session() as s:
            sid = s.scalar(select(StokKarti.id).where(StokKarti.stok_kodu == "U006"))
            lot = StokLotu(stok_id=sid, depo_id=self.ana, lot_no="EVB-9", tedarikci="EvoBulut açılış",
                           giris_tarihi=date(2026, 9, 15), kalan_miktar=D("2"), birim_maliyet=D("10"))
            s.add(lot)
            s.flush()
            s.add(StokHareketi(tarih=date(2026, 9, 15), hareket_turu="GİRİŞ", belge_no="EVB-9", stok_id=sid,
                               depo_id=self.ana, lot_id=lot.id, miktar=D("2"), birim_maliyet=D("10")))
        u = self.satir(self.rapor(), "U006")
        self.assertEqual(u["gmroi"], H)
        self.assertIn("Eksik devir verisi", u["hesap_notu"])


class EkranTest(_Taban):
    def test_menu_ekran_excel_ve_cift_kayit_yok(self):
        import tkinter as tk
        import tempfile

        import stok_raporlari_ui as ui
        import stoklar_ui

        self.u001_senaryo()
        try:
            kok = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk yok")
        kok.withdraw()
        kok.icerik = tk.Frame(kok)
        kok.icerik.pack()
        kok._icerigi_temizle = lambda: [w.destroy() for w in kok.icerik.winfo_children()]
        try:
            kartlar = [a for _b, _a2, a in stoklar_ui.RAPOR_KARTLARI]
            self.assertEqual(kartlar.count("verimlilik"), 1)
            sahte = type("SahteApp", (), {"__getattr__": lambda s, a: (lambda *x, **k: None)})()
            self.assertTrue(callable(stoklar_ui._rapor_komutlar(sahte)["verimlilik"]))
            ekran = ui.yatirim_verimliligi_goster(kok, otomatik=False)
            ekran.girdiler["bas"].delete(0, "end")
            ekran.girdiler["bas"].insert(0, "01.09.2026")
            ekran.girdiler["bit"].delete(0, "end")
            ekran.girdiler["bit"].insert(0, "30.09.2026")
            ekran.raporla(senkron=True)
            self.assertIn("FIFO", ekran.baslik)
            self.assertIn("U001", [r.get("urun_kodu") for r in ekran._satirlar["urunler"]])
            html = ekran.html_metni(otomatik_yazdir=False)
            self.assertIn("GMROI", html)
            self.assertIn("Maliyet yöntemi: FIFO", html)
            self.assertNotIn("window.print", html)
            with tempfile.TemporaryDirectory() as d:
                yol = str(Path(d) / "v.xlsx")
                ekran.excel_yaz(yol)
                from openpyxl import load_workbook

                wb = load_workbook(yol)
                self.assertIn("Ürün Bazında", wb.sheetnames)
                metin = " ".join(str(c.value) for c in wb["Ürün Bazında"]["A"] if c.value)
                self.assertIn("Maliyet yöntemi", metin)
        finally:
            kok.destroy()


if __name__ == "__main__":
    unittest.main()
