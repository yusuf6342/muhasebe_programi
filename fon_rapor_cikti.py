"""Fon raporları Excel / PDF çıktısı — her sayfada tekrar eden başlık satırı."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path


def hucre_metni(deger) -> str:
    if deger is None:
        return ""
    if isinstance(deger, Decimal):
        return f"{deger:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(deger, (date, datetime)):
        return deger.strftime("%d.%m.%Y")
    if isinstance(deger, bool):
        return "Evet" if deger else "Hayır"
    return str(deger)


def excel_yaz(yol: Path | str, baslik: str, alt_baslik: str, sutunlar: list[str], satirlar: list[list],
              notlar: list[str] | tuple = ()) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    yol = Path(yol)
    wb = Workbook()
    ws = wb.active
    ws.title = (baslik or "Fon")[:31]
    ws.append([baslik])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([alt_baslik])
    for n in notlar:
        ws.append([n])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, size=9)
    ws.append([f"Oluşturma: {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append(list(sutunlar))
    baslik_satiri = ws.max_row
    for hucre in ws[baslik_satiri]:
        hucre.font = Font(bold=True)
        hucre.fill = PatternFill("solid", fgColor="DDE7F3")
    for satir in satirlar:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in satir])
    ws.freeze_panes = ws.cell(row=baslik_satiri + 1, column=1)
    ws.print_title_rows = f"{baslik_satiri}:{baslik_satiri}"
    ws.page_setup.orientation = "landscape"
    for i, ad in enumerate(sutunlar, start=1):
        en = max([len(str(ad))] + [len(hucre_metni(r[i - 1])) for r in satirlar[:500] if len(r) >= i])
        ws.column_dimensions[ws.cell(row=baslik_satiri, column=i).column_letter].width = min(45, en + 2)
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)
    return yol


def _fitz():
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    return pymupdf


def _sar(metin: str, olcu, boy: float, en: float) -> list[str]:
    satirlar, mevcut = [], ""
    for kelime in metin.split():
        aday = f"{mevcut} {kelime}".strip()
        if mevcut and olcu.text_length(aday, fontsize=boy) > en:
            satirlar.append(mevcut)
            mevcut = kelime
        else:
            mevcut = aday
    if mevcut:
        satirlar.append(mevcut)
    return satirlar


def pdf_yaz(yol: Path | str, baslik: str, alt_baslik: str, sutunlar: list[str], satirlar: list[list],
            notlar: list[str] | tuple = ()) -> Path:
    from makbuz_cikti import _font_dosyalari

    fitz = _fitz()
    yol = Path(yol)
    genislik, yukseklik = 842, 595  # A4 yatay
    kenar, satir_h, boy = 28, 13, 7.5
    normal, kalin = _font_dosyalari()
    doc = fitz.open()
    metinler = [[hucre_metni(v) for v in r] for r in satirlar]
    olcu = fitz.Font(fontfile=normal) if normal else fitz.Font("helv")
    agirlik = []
    for i, ad in enumerate(sutunlar):
        en = max([olcu.text_length(ad, fontsize=boy)] + [olcu.text_length(r[i], fontsize=boy)
                                                         for r in metinler[:300] if len(r) > i])
        agirlik.append(max(30.0, min(en + 6, 220.0)))
    oran = (genislik - 2 * kenar) / sum(agirlik)
    kolonlar = [a * oran for a in agirlik]
    not_boy, not_h = 6.5, 8.5
    not_satirlari = [s for n in notlar for s in _sar(n, olcu, not_boy, genislik - 2 * kenar)]
    ek = len(not_satirlari) * not_h
    sayfa_basi = max(1, int((yukseklik - 2 * kenar - 60 - ek) // satir_h))
    parcalar = [metinler[i:i + sayfa_basi] for i in range(0, len(metinler), sayfa_basi)] or [[]]
    for no, parca in enumerate(parcalar, start=1):
        page = doc.new_page(width=genislik, height=yukseklik)
        if normal:
            page.insert_font(fontname="fr", fontfile=normal)
            page.insert_font(fontname="fb", fontfile=kalin)
            fr, fb = "fr", "fb"
        else:
            fr, fb = "helv", "hebo"
        page.insert_text((kenar, kenar + 10), baslik, fontsize=12, fontname=fb)
        page.insert_text((kenar, kenar + 24), alt_baslik, fontsize=8, fontname=fr)
        page.insert_text((genislik - kenar - 80, kenar + 10), f"Sayfa {no}/{len(parcalar)}", fontsize=8, fontname=fr)
        for i, s in enumerate(not_satirlari):
            page.insert_text((kenar, kenar + 34 + i * not_h), s, fontsize=not_boy, fontname=fr,
                             color=(0.25, 0.25, 0.25))
        y = kenar + 44 + ek
        page.draw_rect(fitz.Rect(kenar, y - 10, genislik - kenar, y + 3), color=None, fill=(0.87, 0.91, 0.95))
        x = kenar
        for ad, en in zip(sutunlar, kolonlar):
            page.insert_textbox(fitz.Rect(x + 1, y - 9, x + en - 1, y + 4), ad, fontsize=boy, fontname=fb)
            x += en
        for r in parca:
            y += satir_h
            x = kenar
            for i, en in enumerate(kolonlar):
                metin = r[i] if i < len(r) else ""
                while metin and olcu.text_length(metin, fontsize=boy) > en - 3:
                    metin = metin[:-1]
                sag = isinstance(satirlar[0][i] if satirlar and i < len(satirlar[0]) else None, Decimal)
                page.insert_textbox(fitz.Rect(x + 1, y - 9, x + en - 2, y + 4), metin, fontsize=boy,
                                    fontname=fr, align=2 if sag else 0)
                x += en
            page.draw_line((kenar, y + 3), (genislik - kenar, y + 3), color=(0.85, 0.85, 0.85), width=0.3)
    yol.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(yol))
    doc.close()
    return yol
