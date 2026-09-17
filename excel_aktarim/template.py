"""Şablon Excel üretimi."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from excel_aktarim.types import ImportTipi, getir


class ImportTemplateService:
    @staticmethod
    def olustur(import_tipi: str | ImportTipi, hedef: str | Path) -> Path:
        tip = getir(import_tipi) if isinstance(import_tipi, str) else import_tipi
        path = Path(hedef)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = tip.sablon_sayfa[:31] or "Veriler"

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Border(
            left=Side(style="thin", color="D0D7DE"),
            right=Side(style="thin", color="D0D7DE"),
            top=Side(style="thin", color="D0D7DE"),
            bottom=Side(style="thin", color="D0D7DE"),
        )

        for col, alan in enumerate(tip.alanlar, start=1):
            cell = ws.cell(1, col, alan.baslik)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin
            ornek = alan.ornek or ("" if not alan.zorunlu else "zorunlu")
            ornek_cell = ws.cell(2, col, ornek)
            ornek_cell.border = thin
            ornek_cell.font = Font(italic=True, color="666666")
            ws.column_dimensions[get_column_letter(col)].width = max(14, len(alan.baslik) + 4)

        bilgi = wb.create_sheet("Aciklama")
        bilgi["A1"] = tip.ad
        bilgi["A1"].font = Font(bold=True, size=14)
        bilgi["A2"] = tip.aciklama or "Şablonu doldurup sihirbazda seçin."
        bilgi["A4"] = "Alan"
        bilgi["B4"] = "Zorunlu"
        bilgi["C4"] = "Açıklama"
        for i, alan in enumerate(tip.alanlar, start=5):
            bilgi.cell(i, 1, alan.baslik)
            bilgi.cell(i, 2, "Evet" if alan.zorunlu else "Hayır")
            bilgi.cell(i, 3, alan.aciklama or "")
        bilgi.column_dimensions["A"].width = 28
        bilgi.column_dimensions["B"].width = 12
        bilgi.column_dimensions["C"].width = 50

        # Veri sayfası önce gelsin
        wb.move_sheet(ws, offset=-len(wb.sheetnames) + 1)
        wb.active = ws

        wb.save(path)
        return path
