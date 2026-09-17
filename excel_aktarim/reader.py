"""Excel okuyucu — openpyxl."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from excel_aktarim.normalize import temiz


@dataclass
class ExcelOkumaSonucu:
    basliklar: list[str] = field(default_factory=list)
    satirlar: list[dict[str, Any]] = field(default_factory=list)
    sayfa: str = ""
    dosya_adi: str = ""


class ExcelReaderService:
    @staticmethod
    def oku(yol: str | Path, *, sayfa: str | None = None, max_satir: int = 50_000) -> ExcelOkumaSonucu:
        path = Path(yol)
        if not path.is_file():
            raise FileNotFoundError(f"Dosya bulunamadı: {path}")
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("Yalnızca .xlsx / .xlsm dosyaları desteklenir.")

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sayfa and sayfa in wb.sheetnames:
                ws = wb[sayfa]
            else:
                tercih = ("Veriler", "Sheet1", "Sayfa1")
                ws = next((wb[n] for n in tercih if n in wb.sheetnames), wb[wb.sheetnames[0]])
            rows = ws.iter_rows(values_only=True)
            try:
                ham_baslik = next(rows)
            except StopIteration as exc:
                raise ValueError("Excel dosyası boş.") from exc

            basliklar: list[str] = []
            for i, h in enumerate(ham_baslik or ()):
                ad = temiz(h) or f"Sutun{i + 1}"
                basliklar.append(ad)

            if not any(temiz(h) for h in (ham_baslik or ())):
                raise ValueError("Başlık satırı bulunamadı.")

            satirlar: list[dict[str, Any]] = []
            for idx, row in enumerate(rows, start=2):
                if idx - 1 > max_satir:
                    raise ValueError(f"En fazla {max_satir} veri satırı okunabilir.")
                if row is None or all(v is None or temiz(v) == "" for v in row):
                    continue
                kayit: dict[str, Any] = {}
                bos = True
                for i, baslik in enumerate(basliklar):
                    deger = row[i] if i < len(row) else None
                    if temiz(deger):
                        bos = False
                    kayit[baslik] = deger
                if not bos:
                    satirlar.append(kayit)

            return ExcelOkumaSonucu(
                basliklar=basliklar,
                satirlar=satirlar,
                sayfa=ws.title,
                dosya_adi=path.name,
            )
        finally:
            wb.close()
