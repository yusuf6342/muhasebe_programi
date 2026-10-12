"""Satış Kâr Analizi — dönemsel kârlılık, sıfır maliyet, zararına satış, karşılaştırma.

Maliyet kaynakları: satış faturası satırlarında dondurulmuş alanlar
(FIFO / son alış / ortalama / ağırlıklı ortalama); irsaliyeli satışta FIFO irsaliye çıkışından.
Tahmini maliyet üretilmez; sıfır/eksik maliyet kâra katılmaz.

Gerçekleşen satış: onaylı ve iptal edilmemiş fatura. Fatura geneli indirim/masraf ve yuvarlama
muhasebe fişiyle aynı matraha (Net − KDV) göre satırlara oransal dağıtılır. Kaynağa bağlı iade
kaynak satış satırının maliyetiyle ters çevrilir.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

from sqlalchemy import or_, select
from sqlalchemy.orm import selectinload

from database.access import kar_zorunlu, maliyet_zorunlu
from database.database import get_session
from database.models.cari import Cari
from database.models.satis_faturasi import SatisFaturasi, SatisFaturasiSatiri
from database.models.satis_iade_faturasi import SatisIadeFaturasi
from database.models.stok import StokKarti
from database.satis_faturasi_service import SatisFaturasiService
from database.satis_maliyet_motoru import GECMIS_YONTEMLER, MaliyetMotoru, satis_belge_nolari
from database.turkce_normalize import turkce_normalize

MALIYET_YONTEMLERI = (
    "FIFO",
    "SON ALIŞ FİYATI",
    "ORTALAMA ALIŞ FİYATI",
    "AĞIRLIKLI ORTALAMA ALIŞ FİYATI",
    *GECMIS_YONTEMLER,
)

# Geçmiş hareketten hesaplanan yöntemlerde satır alanı yalnız yedek bilgidir; maliyeti motor verir.
_MALIYET_ALAN = {
    "FIFO": "fifo_birim_maliyeti",
    "SON ALIŞ FİYATI": "son_alis_birim_maliyeti",
    "ORTALAMA ALIŞ FİYATI": "ortalama_birim_maliyeti",
    "AĞIRLIKLI ORTALAMA ALIŞ FİYATI": "agirlikli_ortalama_birim_maliyeti",
    **{y: "fifo_birim_maliyeti" for y in GECMIS_YONTEMLER},
}

SIFIR = Decimal("0")
YUZ = Decimal("100")

_EK_ALANLAR = frozenset({
    "satir_id", "sube_id", "kanal", "ana_grup", "tali_grup", "alt_grup", "para_birimi", "kur",
    "kur_tarihi", "kur_kaynagi", "kaynak_belge_no", "kaynak_bagli",
})


def _d(v) -> Decimal:
    if v is None:
        return SIFIR
    if isinstance(v, Decimal):
        return v
    return Decimal(str(v))


def _q2(v: Decimal) -> Decimal:
    return _d(v).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _oran(pay: Decimal, payda: Decimal) -> Decimal | None:
    if payda == 0:
        return None
    return (pay / payda * YUZ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def degisim_hesapla(donem1: Decimal, donem2: Decimal) -> dict[str, Any]:
    """Örnek C/D kuralları."""
    fark = donem1 - donem2
    if donem2 == 0 and donem1 == 0:
        yuzde: Any = Decimal("0")
        etiket = "%0"
    elif donem2 == 0 and donem1 != 0:
        yuzde = None
        etiket = "Yeni / Baz yok"
    else:
        yuzde = (fark / abs(donem2) * YUZ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        etiket = f"%{yuzde}"
    if fark > 0:
        yon = "artış"
    elif fark < 0:
        yon = "azalış"
    else:
        yon = "değişmedi"
    return {"donem1": donem1, "donem2": donem2, "fark": fark, "degisim_yuzde": yuzde, "etiket": etiket, "yon": yon}


def satir_brut_iskonto_net(miktar, fiyat, i1=0, i2=0, i3=0) -> tuple[Decimal, Decimal, Decimal]:
    return SatisFaturasiService._satir_net(miktar, fiyat, i1, i2, i3)


@dataclass
class KarAnalizFiltre:
    baslangic: date
    bitis: date
    maliyet_yontemi: str = "FIFO"
    kars_baslangic: date | None = None
    kars_bitis: date | None = None
    cari_id: int | None = None
    depo: str | None = None
    urun: str | None = None
    marka: str | None = None
    rapor_grubu: str | None = None
    musteri_grubu: str | None = None
    karlilik: str = "tumu"  # tumu|karli|dusuk_marj|basa_bas|zararina|sifir_maliyet|eksik_maliyet
    min_marj: Decimal | None = None
    max_marj: Decimal | None = None
    dusuk_marj_esigi: Decimal = Decimal("10")
    iadeleri_dahil: bool = True
    sifir_maliyeti_kar_haric: bool = False
    personel: str | None = None
    sube_id: int | None = None
    kanal: str | None = None  # None | "hizli" | "fatura"
    stok_grup_id: int | None = None  # ana, tali veya alt grup
    urun_kodu: str | None = None  # birebir ürün kodu (ayrıntı inişi)


@dataclass
class SatirAnaliz:
    belge_turu: str  # SATIS | IADE
    belge_id: int
    belge_no: str
    tarih: date
    cari_id: int | None
    cari_kodu: str
    cari_unvan: str
    musteri_grubu: str
    urun_kodu: str
    urun_adi: str
    marka: str
    rapor_grubu: str
    birim: str
    depo: str
    miktar: Decimal
    brut_satis: Decimal
    iskonto: Decimal
    net_satis: Decimal
    birim_maliyet: Decimal | None
    toplam_maliyet: Decimal | None
    brut_kar: Decimal | None
    kar_marji: Decimal | None
    maliyet_uzeri_kar: Decimal | None
    maliyet_yontemi: str
    maliyet_kaynagi: str
    maliyet_kodu: str | None  # None | ZERO_COST | NULL_COST | UNIT_UNDEFINED (birim kartta tanımsız) | ...
    durum: str  # karli | dusuk_marj | basa_bas | zararina | maliyeti_belirsiz | bedelsiz
    personel: str
    satir_id: int | None = None
    sube_id: int | None = None
    kanal: str = ""
    ana_grup: str = ""
    tali_grup: str = ""
    alt_grup: str = ""
    temel_miktar: Decimal | None = None  # stok temel biriminde (çevrim yoksa None)
    temel_birim: str = ""
    belge_duzeltme: Decimal = SIFIR  # fatura geneli indirim(−)/masraf(+)/yuvarlama payı, KDV hariç
    para_birimi: str = "TRY"
    kur: Decimal | None = None
    kur_tarihi: date | None = None
    kur_kaynagi: str = ""
    maliyet_durumu: str = ""  # Kesin | Eksik | Kaynaksız iade
    kaynak_belge_no: str = ""
    kaynak_bagli: bool | None = None


@dataclass
class KarAnalizSonuc:
    filtre: KarAnalizFiltre
    satirlar: list[SatirAnaliz] = field(default_factory=list)
    ozet: dict[str, Any] = field(default_factory=dict)
    urunler: list[dict[str, Any]] = field(default_factory=list)
    musteriler: list[dict[str, Any]] = field(default_factory=list)
    belgeler: list[dict[str, Any]] = field(default_factory=list)
    zararina: list[SatirAnaliz] = field(default_factory=list)
    sifir_maliyet: list[SatirAnaliz] = field(default_factory=list)
    veri_kalitesi: list[SatirAnaliz] = field(default_factory=list)
    gunluk: list[dict[str, Any]] = field(default_factory=list)
    bulgular: list[str] = field(default_factory=list)
    karsilastirma: dict[str, Any] | None = None
    sorgu_ms: int = 0
    yenileme: datetime | None = None


class SatisKarAnalizService:
    @staticmethod
    def onceki_es_donem(baslangic: date, bitis: date) -> tuple[date, date]:
        gun = (bitis - baslangic).days + 1
        yeni_bitis = baslangic - timedelta(days=1)
        yeni_bas = yeni_bitis - timedelta(days=gun - 1)
        return yeni_bas, yeni_bitis

    @staticmethod
    def analiz(filtre: KarAnalizFiltre) -> KarAnalizSonuc:
        maliyet_zorunlu()
        kar_zorunlu()
        if filtre.bitis < filtre.baslangic:
            raise ValueError("Bitiş tarihi başlangıçtan önce olamaz.")
        if filtre.maliyet_yontemi not in _MALIYET_ALAN:
            raise ValueError(f"Desteklenmeyen maliyet yöntemi: {filtre.maliyet_yontemi}")

        t0 = datetime.now()
        satirlar = SatisKarAnalizService._satirlari_yukle(filtre)
        sonuc = SatisKarAnalizService._derle(filtre, satirlar)
        if filtre.kars_baslangic and filtre.kars_bitis:
            if filtre.kars_bitis < filtre.kars_baslangic:
                raise ValueError("Karşılaştırma bitişi başlangıçtan önce olamaz.")
            kars_filtre = KarAnalizFiltre(
                baslangic=filtre.kars_baslangic,
                bitis=filtre.kars_bitis,
                maliyet_yontemi=filtre.maliyet_yontemi,
                cari_id=filtre.cari_id,
                depo=filtre.depo,
                urun=filtre.urun,
                marka=filtre.marka,
                rapor_grubu=filtre.rapor_grubu,
                musteri_grubu=filtre.musteri_grubu,
                karlilik="tumu",
                dusuk_marj_esigi=filtre.dusuk_marj_esigi,
                iadeleri_dahil=filtre.iadeleri_dahil,
                sifir_maliyeti_kar_haric=filtre.sifir_maliyeti_kar_haric,
                personel=filtre.personel,
                sube_id=filtre.sube_id,
                kanal=filtre.kanal,
                stok_grup_id=filtre.stok_grup_id,
                urun_kodu=filtre.urun_kodu,
            )
            kars_satir = SatisKarAnalizService._satirlari_yukle(kars_filtre)
            kars = SatisKarAnalizService._derle(kars_filtre, kars_satir)
            sonuc.karsilastirma = SatisKarAnalizService._karsilastir(sonuc, kars)
            sonuc.bulgular.extend(SatisKarAnalizService._kars_bulgular(sonuc, kars))
        sonuc.sorgu_ms = int((datetime.now() - t0).total_seconds() * 1000)
        sonuc.yenileme = datetime.now()
        return sonuc

    @staticmethod
    def belge_duzeltme_paylari(fatura) -> dict[int, Decimal]:
        """Fatura geneli indirim/masraf/yuvarlama farkının satırlara KDV hariç oransal payı.

        Hedef matrah muhasebe fişiyle aynıdır: ürün Net'i − KDV (Net satır toplamından farklıysa).
        Son satır kuruş artığını alır. Fark yoksa boş sözlük. Anahtar: ``id(satir)``.
        """
        satirlar = list(getattr(fatura, "satirlar", None) or [])
        if not satirlar:
            return {}
        netler: list[tuple[Any, Decimal]] = []
        for s in satirlar:
            _b, _i, n = satir_brut_iskonto_net(
                abs(_d(s.miktar)), _d(s.birim_fiyat),
                getattr(s, "iskonto_orani", 0) or 0,
                getattr(s, "iskonto_orani_2", 0) or 0,
                getattr(s, "iskonto_orani_3", 0) or 0,
            )
            netler.append((s, n))
        toplam = SatisFaturasiService.toplam(satirlar)
        satir_genel = _d(toplam["genel_toplam"])
        urun_net = _d(SatisFaturasiService.urun_net_toplam(fatura))
        if _q2(urun_net) == _q2(satir_genel):
            return {}
        kdv = _d(getattr(fatura, "tl_kdv", None))
        if kdv <= 0:
            kdv = _d(toplam["kdv"])
        hedef = urun_net - kdv
        satir_matrah = sum((n for _s, n in netler), SIFIR)
        fark = _q2(hedef - satir_matrah)
        if fark == 0 or satir_matrah == 0:
            return {}
        paylar: dict[int, Decimal] = {}
        dagitilan = SIFIR
        for i, (s, n) in enumerate(netler):
            if i == len(netler) - 1:
                pay = fark - dagitilan
            else:
                pay = _q2(fark * n / satir_matrah)
                dagitilan += pay
            paylar[id(s)] = pay
        return paylar

    @staticmethod
    def _grup_adlari(session) -> dict[int, str]:
        try:
            from database.models.stok import StokGrubu

            return {int(g.id): g.ad or "" for g in session.scalars(select(StokGrubu)).all()}
        except Exception:  # noqa: BLE001
            return {}

    @staticmethod
    def _stok_meta(stok: StokKarti | None, grup_ad: dict[int, str]) -> dict[str, str]:
        if stok is None:
            return {"ana_grup": "", "tali_grup": "", "alt_grup": ""}

        def ad(alan):
            gid = getattr(stok, alan, None)
            return grup_ad.get(int(gid), "") if gid else ""

        return {"ana_grup": ad("ana_grup_id"), "tali_grup": ad("tali_grup_id"), "alt_grup": ad("alt_grup_id")}

    @staticmethod
    def _stok_uygun(stok: StokKarti | None, satir, filtre: KarAnalizFiltre, urun_f, marka_f, grup_f) -> bool:
        if filtre.urun_kodu and (satir.urun_kodu or "").strip() != filtre.urun_kodu.strip():
            return False
        if urun_f and urun_f not in turkce_normalize(satir.urun_kodu or "") and urun_f not in turkce_normalize(
            satir.urun_adi or ""
        ):
            return False
        if marka_f and marka_f not in turkce_normalize((stok.marka if stok else "") or ""):
            return False
        if grup_f and grup_f not in turkce_normalize((stok.rapor_grubu if stok else "") or ""):
            return False
        if filtre.stok_grup_id:
            gid = int(filtre.stok_grup_id)
            if stok is None or gid not in {
                getattr(stok, "ana_grup_id", None), getattr(stok, "tali_grup_id", None),
                getattr(stok, "alt_grup_id", None),
            }:
                return False
        return True

    @staticmethod
    def _satirlari_yukle(filtre: KarAnalizFiltre) -> list[SatirAnaliz]:
        from database.hizli_satis_service import _hizli_satis_mi

        alan = _MALIYET_ALAN[filtre.maliyet_yontemi]
        urun_f = turkce_normalize((filtre.urun or "").strip())
        depo_f = turkce_normalize((filtre.depo or "").strip())
        marka_f = turkce_normalize((filtre.marka or "").strip())
        grup_f = turkce_normalize((filtre.rapor_grubu or "").strip())
        mgrup_f = turkce_normalize((filtre.musteri_grubu or "").strip())
        personel_f = turkce_normalize((filtre.personel or "").strip())
        kanal_f = (filtre.kanal or "").strip().casefold()

        out: list[SatirAnaliz] = []
        with get_session() as session:
            stok_map: dict[str, StokKarti] = {
                s.stok_kodu: s
                for s in session.scalars(select(StokKarti)).all()
            }
            grup_ad = SatisKarAnalizService._grup_adlari(session)
            motor = None
            if filtre.maliyet_yontemi in GECMIS_YONTEMLER:
                import calendar

                ay_sonu = filtre.bitis.replace(day=calendar.monthrange(filtre.bitis.year, filtre.bitis.month)[1])
                motor = MaliyetMotoru(session, bitis=ay_sonu)

            fq = (
                select(SatisFaturasi)
                .where(
                    SatisFaturasi.fatura_tarihi >= filtre.baslangic,
                    SatisFaturasi.fatura_tarihi <= filtre.bitis,
                    SatisFaturasi.durum != "İPTAL",
                    SatisFaturasi.onaylandi.is_(True),
                    or_(SatisFaturasi.is_deleted.is_(False), SatisFaturasi.is_deleted.is_(None)),
                )
                .options(
                    selectinload(SatisFaturasi.satirlar),
                    selectinload(SatisFaturasi.cari),
                )
            )
            if filtre.cari_id:
                fq = fq.where(SatisFaturasi.cari_id == int(filtre.cari_id))
            if filtre.sube_id:
                fq = fq.where(SatisFaturasi.sube_id == int(filtre.sube_id))
            faturalar = list(session.scalars(fq).all())

            for fatura in faturalar:
                if depo_f and depo_f not in turkce_normalize(fatura.depo or ""):
                    continue
                kanal = "Hızlı Satış" if _hizli_satis_mi(fatura.aciklama) else "Fatura"
                if kanal_f == "hizli" and kanal != "Hızlı Satış":
                    continue
                if kanal_f == "fatura" and kanal != "Fatura":
                    continue
                personel = (
                    getattr(fatura, "sales_person_full_name", None)
                    or fatura.created_by_full_name
                    or fatura.created_by_username
                    or ""
                )
                if personel_f and personel_f not in turkce_normalize(personel):
                    continue
                cari = fatura.cari
                mgrup = (cari.musteri_grubu if cari else "") or ""
                if mgrup_f and mgrup_f not in turkce_normalize(mgrup):
                    continue
                paylar = SatisKarAnalizService.belge_duzeltme_paylari(fatura)
                kur_bilgi = {
                    "para_birimi": (fatura.para_birimi or "TRY").upper(),
                    "kur": _d(fatura.kur) if fatura.kur is not None else None,
                    "kur_tarihi": fatura.kur_tarihi,
                    "kur_kaynagi": fatura.kur_kaynagi or "",
                }
                for satir in fatura.satirlar:
                    stok = stok_map.get(satir.urun_kodu)
                    if not SatisKarAnalizService._stok_uygun(stok, satir, filtre, urun_f, marka_f, grup_f):
                        continue
                    if motor is not None:
                        mb, kaynak = motor.birim_maliyet(
                            filtre.maliyet_yontemi, stok, fatura.fatura_tarihi,
                            satis_belge_nolari(session, fatura, satir))
                    elif alan == "fifo_birim_maliyeti":
                        mb = SatisFaturasiService.satir_birim_maliyeti(session, satir)
                        kaynak = (
                            "irsaliye çıkış maliyeti"
                            if getattr(satir, "irsaliye_satiri_id", None)
                            and not (_d(satir.fifo_birim_maliyeti) > 0)
                            else "satır FIFO maliyeti"
                        )
                    else:
                        mb = getattr(satir, alan, None)
                        kaynak = f"satır.{alan}"
                    out.append(
                        SatisKarAnalizService._satir_analiz(
                            belge_turu="SATIS",
                            belge_id=fatura.id,
                            belge_no=fatura.fatura_no,
                            tarih=fatura.fatura_tarihi,
                            cari=cari,
                            satir=satir,
                            stok=stok,
                            depo=fatura.depo or "",
                            alan=alan,
                            yontem=filtre.maliyet_yontemi,
                            dusuk_esik=filtre.dusuk_marj_esigi,
                            personel=personel,
                            isaret=1,
                            belge_duzeltme=paylar.get(id(satir), SIFIR),
                            maliyet=(mb, kaynak),
                            ek={
                                "satir_id": satir.id,
                                "sube_id": fatura.sube_id,
                                "kanal": kanal,
                                **SatisKarAnalizService._stok_meta(stok, grup_ad),
                                **kur_bilgi,
                            },
                        )
                    )

            if filtre.iadeleri_dahil and kanal_f != "hizli":
                iq = (
                    select(SatisIadeFaturasi)
                    .where(
                        SatisIadeFaturasi.iade_tarihi >= filtre.baslangic,
                        SatisIadeFaturasi.iade_tarihi <= filtre.bitis,
                        SatisIadeFaturasi.durum != "İPTAL",
                    )
                    .options(
                        selectinload(SatisIadeFaturasi.satirlar),
                        selectinload(SatisIadeFaturasi.cari),
                        selectinload(SatisIadeFaturasi.kaynak_fatura),
                    )
                )
                if filtre.cari_id:
                    iq = iq.where(SatisIadeFaturasi.cari_id == int(filtre.cari_id))
                if filtre.sube_id:
                    iq = iq.where(SatisIadeFaturasi.sube_id == int(filtre.sube_id))
                for iade in session.scalars(iq).all():
                    if depo_f and depo_f not in turkce_normalize(iade.depo or ""):
                        continue
                    kaynak_fatura = getattr(iade, "kaynak_fatura", None)
                    personel = ""
                    if kaynak_fatura is not None:
                        personel = (
                            getattr(kaynak_fatura, "sales_person_full_name", None)
                            or getattr(kaynak_fatura, "created_by_full_name", None)
                            or ""
                        )
                    if personel_f and personel_f not in turkce_normalize(personel):
                        continue
                    cari = iade.cari
                    mgrup = (cari.musteri_grubu if cari else "") or ""
                    if mgrup_f and mgrup_f not in turkce_normalize(mgrup):
                        continue
                    kur_bilgi = {
                        "para_birimi": (iade.para_birimi or "TRY").upper(),
                        "kur": _d(iade.kur) if iade.kur is not None else None,
                        "kur_tarihi": iade.kur_tarihi,
                        "kur_kaynagi": iade.kur_kaynagi or "",
                    }
                    for satir in iade.satirlar:
                        stok = stok_map.get(satir.urun_kodu)
                        if not SatisKarAnalizService._stok_uygun(stok, satir, filtre, urun_f, marka_f, grup_f):
                            continue
                        kaynak_satir = None
                        if getattr(satir, "kaynak_fatura_satiri_id", None):
                            kaynak_satir = session.get(SatisFaturasiSatiri, int(satir.kaynak_fatura_satiri_id))
                        kaynak_no = kaynak_fatura.fatura_no if kaynak_fatura is not None else ""
                        if kaynak_satir is not None:
                            kf = session.get(SatisFaturasi, int(kaynak_satir.fatura_id))
                            if motor is not None:
                                mb, ham = motor.birim_maliyet(
                                    filtre.maliyet_yontemi, stok,
                                    kf.fatura_tarihi if kf is not None else iade.iade_tarihi,
                                    satis_belge_nolari(session, kf, kaynak_satir))
                                kaynak = (f"kaynak satış satırı maliyeti (ters çevrildi): {ham}"
                                          if mb is not None else ham)
                            else:
                                if alan == "fifo_birim_maliyeti":
                                    mb = SatisFaturasiService.satir_birim_maliyeti(session, kaynak_satir)
                                else:
                                    mb = getattr(kaynak_satir, alan, None)
                                kaynak = "kaynak satış satırı maliyeti (ters çevrildi)"
                            if not kaynak_no:
                                kaynak_no = kf.fatura_no if kf is not None else ""
                        elif motor is not None:
                            mb, ham = motor.birim_maliyet(filtre.maliyet_yontemi, stok, iade.iade_tarihi,
                                                          (iade.iade_no,), giris=True)
                            kaynak = f"kaynaksız iade — iade tarihindeki {ham}" if mb is not None else ham
                        else:
                            mb = getattr(satir, "fifo_birim_maliyeti", None)
                            kaynak = "kaynaksız iade — iade anı FIFO maliyeti"
                        out.append(
                            SatisKarAnalizService._satir_analiz(
                                belge_turu="IADE",
                                belge_id=iade.id,
                                belge_no=iade.iade_no,
                                tarih=iade.iade_tarihi,
                                cari=cari,
                                satir=satir,
                                stok=stok,
                                depo=iade.depo or "",
                                alan="fifo_birim_maliyeti",
                                yontem=filtre.maliyet_yontemi,
                                dusuk_esik=filtre.dusuk_marj_esigi,
                                personel=personel,
                                isaret=-1,
                                maliyet=(mb, kaynak),
                                ek={
                                    "satir_id": satir.id,
                                    "sube_id": iade.sube_id,
                                    "kanal": "İade",
                                    "kaynak_belge_no": kaynak_no or "",
                                    "kaynak_bagli": kaynak_satir is not None,
                                    **SatisKarAnalizService._stok_meta(stok, grup_ad),
                                    **kur_bilgi,
                                },
                            )
                        )
        return out

    @staticmethod
    def _satir_analiz(
        *,
        belge_turu: str,
        belge_id: int,
        belge_no: str,
        tarih: date,
        cari: Cari | None,
        satir,
        stok: StokKarti | None,
        depo: str,
        alan: str,
        yontem: str,
        dusuk_esik: Decimal,
        personel: str,
        isaret: int,
        belge_duzeltme: Decimal = SIFIR,
        maliyet: tuple[Any, str] | None = None,
        ek: dict[str, Any] | None = None,
    ) -> SatirAnaliz:
        ek = dict(ek or {})
        miktar = _d(satir.miktar) * isaret
        fiyat = _d(satir.birim_fiyat)
        i1 = getattr(satir, "iskonto_orani", 0) or 0
        i2 = getattr(satir, "iskonto_orani_2", 0) or 0
        i3 = getattr(satir, "iskonto_orani_3", 0) or 0
        brut, iskonto, net = satir_brut_iskonto_net(abs(_d(satir.miktar)), fiyat, i1, i2, i3)
        duzeltme = _q2(_d(belge_duzeltme)) * isaret
        brut *= isaret
        iskonto *= isaret
        net = net * isaret + duzeltme

        if maliyet is None:
            ham_maliyet, maliyet_kaynagi = getattr(satir, alan, None), f"satır.{alan}"
        else:
            ham_maliyet, maliyet_kaynagi = maliyet
        hizmet = "hizmet" in str(getattr(stok, "kart_turu", "") or "").casefold()
        maliyet_kodu = None
        birim_maliyet: Decimal | None
        if ham_maliyet is None:
            maliyet_kodu = "NULL_COST"
            birim_maliyet = None
        else:
            birim_maliyet = _d(ham_maliyet)
            if birim_maliyet < 0:
                maliyet_kodu = "INVALID_COST"
                birim_maliyet = None
            elif birim_maliyet == 0 and not hizmet:
                # Sıfır maliyet tam kâr gösterir; eksik maliyet olarak işaretlenir, kâra katılmaz.
                maliyet_kodu = "ZERO_COST"
                birim_maliyet = None

        temel_birim = (getattr(stok, "birim", None) or "Adet") if stok is not None else ""
        carpan = getattr(satir, "birim_carpani", None)
        if carpan in (None, "") or _d(carpan) <= 0:
            if stok is not None:
                from database.stok_service import StokService

                carpan = StokService.birim_carpani_bul(
                    stok, getattr(satir, "birim", None), getattr(stok, "birim", None) or "Adet"
                )
            else:
                carpan = Decimal("1")
        if carpan is None and birim_maliyet is not None:
            maliyet_kodu = "UNIT_UNDEFINED"
            birim_maliyet = None
        temel_miktar = (miktar * _d(carpan)) if carpan is not None and stok is not None else None

        if birim_maliyet is None:
            toplam_maliyet = None
            brut_kar = None
            kar_marji = None
            maliyet_uzeri = None
            durum = "maliyeti_belirsiz"
        else:
            toplam_maliyet = _q2(abs(_d(satir.miktar)) * _d(carpan) * birim_maliyet * isaret)
            brut_kar = _q2(net - toplam_maliyet)
            kar_marji = _oran(brut_kar, net) if net != 0 else None
            maliyet_uzeri = _oran(brut_kar, toplam_maliyet) if toplam_maliyet != 0 else None
            if net == 0 and abs(_d(satir.miktar)) > 0:
                durum = "bedelsiz"
            elif brut_kar < 0:
                durum = "zararina"
            elif brut_kar == 0:
                durum = "basa_bas"
            elif kar_marji is not None and kar_marji < dusuk_esik:
                durum = "dusuk_marj"
            else:
                durum = "karli"

        if birim_maliyet is None:
            maliyet_durumu = "Eksik"
        elif ek.get("kaynak_bagli") is False:
            maliyet_durumu = "Kaynaksız iade"
        else:
            maliyet_durumu = "Kesin"

        return SatirAnaliz(
            belge_turu=belge_turu,
            belge_id=belge_id,
            belge_no=belge_no,
            tarih=tarih,
            cari_id=cari.id if cari else None,
            cari_kodu=(cari.cari_kodu if cari else "") or "",
            cari_unvan=(cari.unvan if cari else "") or "",
            musteri_grubu=(cari.musteri_grubu if cari else "") or "",
            urun_kodu=satir.urun_kodu or "",
            urun_adi=satir.urun_adi or "",
            marka=(stok.marka if stok else "") or "",
            rapor_grubu=(stok.rapor_grubu if stok else "") or "",
            birim=satir.birim or (stok.birim if stok else "") or "",
            depo=depo,
            miktar=miktar,
            brut_satis=_q2(brut),
            iskonto=_q2(iskonto),
            net_satis=_q2(net),
            birim_maliyet=birim_maliyet,
            toplam_maliyet=toplam_maliyet,
            brut_kar=brut_kar,
            kar_marji=kar_marji,
            maliyet_uzeri_kar=maliyet_uzeri,
            maliyet_yontemi=yontem,
            maliyet_kaynagi=maliyet_kaynagi,
            maliyet_kodu=maliyet_kodu,
            durum=durum,
            personel=personel,
            temel_miktar=temel_miktar,
            temel_birim=temel_birim,
            belge_duzeltme=duzeltme,
            maliyet_durumu=maliyet_durumu,
            **{k: v for k, v in ek.items() if k in _EK_ALANLAR},
        )

    @staticmethod
    def _filtrele_karlilik(satirlar: list[SatirAnaliz], filtre: KarAnalizFiltre) -> list[SatirAnaliz]:
        out = []
        for s in satirlar:
            if filtre.karlilik == "karli" and s.durum != "karli":
                continue
            if filtre.karlilik == "dusuk_marj" and s.durum != "dusuk_marj":
                continue
            if filtre.karlilik == "basa_bas" and s.durum != "basa_bas":
                continue
            if filtre.karlilik == "zararina" and s.durum != "zararina":
                continue
            if filtre.karlilik == "sifir_maliyet" and s.maliyet_kodu != "ZERO_COST":
                continue
            if filtre.karlilik == "eksik_maliyet" and s.maliyet_kodu not in {
                "NULL_COST",
                "INVALID_COST",
                "ZERO_COST",
                "UNIT_UNDEFINED",
            }:
                continue
            if filtre.min_marj is not None and (s.kar_marji is None or s.kar_marji < filtre.min_marj):
                continue
            if filtre.max_marj is not None and (s.kar_marji is None or s.kar_marji > filtre.max_marj):
                continue
            out.append(s)
        return out

    @staticmethod
    def _derle(filtre: KarAnalizFiltre, satirlar: list[SatirAnaliz]) -> KarAnalizSonuc:
        filtrelenmis = SatisKarAnalizService._filtrele_karlilik(satirlar, filtre)
        sonuc = KarAnalizSonuc(filtre=filtre, satirlar=filtrelenmis)

        net = maliyet = kar = brut = iskonto = SIFIR
        miktar = SIFIR
        iade_tutar = SIFIR
        guvenilir_kar = SIFIR
        guvenilir_net = SIFIR
        maliyetli_net = SIFIR
        zarar_toplam = SIFIR
        zarar_adet = 0
        sifir_adet = 0
        sifir_urun: set[str] = set()
        sifir_satis = SIFIR
        dusuk_adet = 0
        dusuk_tutar = SIFIR

        for s in filtrelenmis:
            brut += s.brut_satis
            iskonto += s.iskonto
            net += s.net_satis
            miktar += s.miktar
            if s.belge_turu == "IADE":
                iade_tutar += abs(s.net_satis)
            if s.toplam_maliyet is not None and s.brut_kar is not None:
                if filtre.sifir_maliyeti_kar_haric and s.maliyet_kodu:
                    pass
                else:
                    maliyet += s.toplam_maliyet
                    kar += s.brut_kar
                    maliyetli_net += s.net_satis
                    if not s.maliyet_kodu:
                        guvenilir_kar += s.brut_kar
                        guvenilir_net += s.net_satis
            if s.durum == "zararina" and s.brut_kar is not None:
                zarar_toplam += s.brut_kar
                zarar_adet += 1
                sonuc.zararina.append(s)
            if s.maliyet_kodu:
                sifir_adet += 1
                sifir_urun.add(s.urun_kodu)
                sifir_satis += s.net_satis
                sonuc.sifir_maliyet.append(s)
                sonuc.veri_kalitesi.append(s)
            if s.durum == "dusuk_marj":
                dusuk_adet += 1
                dusuk_tutar += s.net_satis

        sonuc.ozet = {
            "brut_satis": _q2(brut),
            "iskonto": _q2(iskonto),
            "iade_tutari": _q2(iade_tutar),
            "net_satis": _q2(net),
            "maliyet": _q2(maliyet),
            "brut_kar": _q2(kar),
            # Marj yalnız maliyeti bilinen satışlar üzerinden; eksik maliyetli satış kâra katılmaz.
            "kar_marji": _oran(kar, maliyetli_net),
            "maliyetli_net": _q2(maliyetli_net),
            "kar_eksik": maliyetli_net != net,
            "maliyet_uzeri_kar": _oran(kar, maliyet),
            "net_miktar": miktar,
            "belge_sayisi": len({(s.belge_turu, s.belge_id) for s in filtrelenmis}),
            "satir_sayisi": len(filtrelenmis),
            "zarar_adet": zarar_adet,
            "zarar_toplam": _q2(zarar_toplam),
            "sifir_urun": len(sifir_urun),
            "sifir_satir": sifir_adet,
            "sifir_satis": _q2(sifir_satis),
            "dusuk_adet": dusuk_adet,
            "dusuk_tutar": _q2(dusuk_tutar),
            "iade_orani": _oran(iade_tutar, brut) if brut else None,
            "guvenilir_kar": _q2(guvenilir_kar),
            "guvenilir_net": _q2(guvenilir_net),
            "maliyet_yontemi": filtre.maliyet_yontemi,
            "gun_sayisi": (filtre.bitis - filtre.baslangic).days + 1,
        }

        # Ürün / müşteri / belge agregasyon
        urun: dict[str, dict[str, Any]] = {}
        musteri: dict[int | str, dict[str, Any]] = {}
        belge: dict[tuple, dict[str, Any]] = {}
        gunluk: dict[date, dict[str, Decimal]] = defaultdict(
            lambda: {"net": SIFIR, "maliyet": SIFIR, "kar": SIFIR}
        )

        for s in filtrelenmis:
            u = urun.setdefault(
                s.urun_kodu,
                {
                    "urun_kodu": s.urun_kodu,
                    "urun_adi": s.urun_adi,
                    "marka": s.marka,
                    "rapor_grubu": s.rapor_grubu,
                    "birim": s.temel_birim or s.birim,
                    "birim_cevrimsiz": False,
                    "maliyetli_net": SIFIR,
                    "miktar": SIFIR,
                    "iade_miktar": SIFIR,
                    "brut": SIFIR,
                    "iskonto": SIFIR,
                    "iade_tutar": SIFIR,
                    "net": SIFIR,
                    "maliyet": SIFIR,
                    "kar": SIFIR,
                    "belge": set(),
                    "musteri": set(),
                    "son_satis": s.tarih,
                    "eksik": 0,
                },
            )
            if s.temel_miktar is not None:
                mik_s = s.temel_miktar
            else:
                mik_s = s.miktar
                u["birim_cevrimsiz"] = True
            u["miktar"] += mik_s if s.belge_turu == "SATIS" else SIFIR
            if s.belge_turu == "IADE":
                u["iade_miktar"] += abs(mik_s)
                u["iade_tutar"] += abs(s.net_satis)
            u["brut"] += s.brut_satis
            u["iskonto"] += s.iskonto
            u["net"] += s.net_satis
            if s.toplam_maliyet is not None and s.brut_kar is not None:
                u["maliyet"] += s.toplam_maliyet
                u["kar"] += s.brut_kar
                u["maliyetli_net"] += s.net_satis
            else:
                u["eksik"] += 1
            u["belge"].add((s.belge_turu, s.belge_id))
            if s.cari_id:
                u["musteri"].add(s.cari_id)
            if s.tarih > u["son_satis"]:
                u["son_satis"] = s.tarih

            mk = s.cari_id if s.cari_id is not None else s.cari_kodu or "?"
            m = musteri.setdefault(
                mk,
                {
                    "cari_id": s.cari_id,
                    "cari_kodu": s.cari_kodu,
                    "unvan": s.cari_unvan,
                    "net": SIFIR,
                    "maliyet": SIFIR,
                    "kar": SIFIR,
                    "belge": set(),
                    "iade": SIFIR,
                    "maliyetli_net": SIFIR,
                },
            )
            m["net"] += s.net_satis
            if s.toplam_maliyet is not None and s.brut_kar is not None:
                m["maliyet"] += s.toplam_maliyet
                m["kar"] += s.brut_kar
                m["maliyetli_net"] += s.net_satis
            m["belge"].add((s.belge_turu, s.belge_id))
            if s.belge_turu == "IADE":
                m["iade"] += abs(s.net_satis)

            bk = (s.belge_turu, s.belge_id)
            b = belge.setdefault(
                bk,
                {
                    "belge_turu": s.belge_turu,
                    "belge_id": s.belge_id,
                    "belge_no": s.belge_no,
                    "tarih": s.tarih,
                    "cari_kodu": s.cari_kodu,
                    "cari_unvan": s.cari_unvan,
                    "net": SIFIR,
                    "maliyet": SIFIR,
                    "kar": SIFIR,
                    "maliyetli_net": SIFIR,
                    "personel": s.personel,
                },
            )
            b["net"] += s.net_satis
            if s.toplam_maliyet is not None and s.brut_kar is not None:
                b["maliyet"] += s.toplam_maliyet
                b["kar"] += s.brut_kar
                b["maliyetli_net"] += s.net_satis

            g = gunluk[s.tarih]
            g["net"] += s.net_satis
            if s.toplam_maliyet is not None and s.brut_kar is not None:
                g["maliyet"] += s.toplam_maliyet
                g["kar"] += s.brut_kar

        urun_list = []
        for u in urun.values():
            net_u = u["net"]
            mik = u["miktar"] - u["iade_miktar"]
            urun_list.append(
                {
                    **{k: v for k, v in u.items() if k not in {"belge", "musteri"}},
                    "net_miktar": mik,
                    "kar_marji": _oran(u["kar"], u["maliyetli_net"]),
                    "maliyet_uzeri_kar": _oran(u["kar"], u["maliyet"]),
                    "ort_fiyat": (net_u / mik) if mik else None,
                    "ort_maliyet": (u["maliyet"] / mik) if mik else None,
                    "belge_sayisi": len(u["belge"]),
                    "musteri_sayisi": len(u["musteri"]),
                    "durum": "eksik_maliyet" if u["eksik"] else ("zararina" if u["kar"] < 0 else "karli"),
                }
            )
        urun_list.sort(key=lambda x: x["kar"], reverse=True)
        sonuc.urunler = urun_list

        musteri_list = []
        for m in musteri.values():
            musteri_list.append(
                {
                    **{k: v for k, v in m.items() if k != "belge"},
                    "kar_marji": _oran(m["kar"], m["maliyetli_net"]),
                    "belge_sayisi": len(m["belge"]),
                    "ort_sepet": (m["net"] / len(m["belge"])) if m["belge"] else SIFIR,
                }
            )
        musteri_list.sort(key=lambda x: x["kar"], reverse=True)
        sonuc.musteriler = musteri_list

        belge_list = list(belge.values())
        for b in belge_list:
            b["kar_marji"] = _oran(b["kar"], b["maliyetli_net"])
        belge_list.sort(key=lambda x: x["tarih"], reverse=True)
        sonuc.belgeler = belge_list

        sonuc.gunluk = [
            {"tarih": t, "net": _q2(v["net"]), "maliyet": _q2(v["maliyet"]), "kar": _q2(v["kar"])}
            for t, v in sorted(gunluk.items())
        ]

        sonuc.bulgular = SatisKarAnalizService._bulgular(sonuc)
        return sonuc

    @staticmethod
    def _bulgular(sonuc: KarAnalizSonuc) -> list[str]:
        o = sonuc.ozet
        bulgular = []
        if o.get("sifir_satir"):
            bulgular.append(
                f"{o['sifir_satir']} satırda / {o['sifir_urun']} üründe maliyet bulunamadığı için "
                f"{o['sifir_satis']} TL satışın kârı hesaplanamadı (kâr eksik). Marj yalnız maliyeti bilinen "
                f"{o.get('maliyetli_net')} TL satış üzerinden hesaplanır."
            )
        if o.get("zarar_adet"):
            bulgular.append(
                f"{o['zarar_adet']} zararına satış satırı; toplam zarar {o['zarar_toplam']} TL."
            )
        if sonuc.urunler:
            en = sonuc.urunler[0]
            bulgular.append(
                f"Kâra en fazla katkı: {en['urun_kodu']} — {en['urun_adi']} ({en['kar']} TL)."
            )
            zarar_urun = [u for u in sonuc.urunler if u["kar"] < 0]
            if zarar_urun:
                z = min(zarar_urun, key=lambda x: x["kar"])
                bulgular.append(
                    f"Kârı en fazla düşüren: {z['urun_kodu']} — {z['urun_adi']} ({z['kar']} TL)."
                )
        marj = o.get("kar_marji")
        if marj is not None:
            bulgular.append(f"Dönem kâr marjı %{marj}; maliyet üzerine kâr {_fmt_opt(o.get('maliyet_uzeri_kar'))}.")
        return bulgular

    @staticmethod
    def _karsilastir(a: KarAnalizSonuc, b: KarAnalizSonuc) -> dict[str, Any]:
        gun1 = a.ozet["gun_sayisi"] or 1
        gun2 = b.ozet["gun_sayisi"] or 1
        alanlar = (
            "net_satis",
            "maliyet",
            "brut_kar",
            "brut_satis",
            "iskonto",
            "iade_tutari",
            "zarar_toplam",
            "sifir_satir",
            "belge_sayisi",
            "net_miktar",
        )
        gostergeler = {}
        for alan in alanlar:
            d = degisim_hesapla(_d(a.ozet.get(alan) or 0), _d(b.ozet.get(alan) or 0))
            d["gunluk1"] = _q2(_d(a.ozet.get(alan) or 0) / gun1)
            d["gunluk2"] = _q2(_d(b.ozet.get(alan) or 0) / gun2)
            gostergeler[alan] = d
        # marj puan farkı
        m1 = a.ozet.get("kar_marji")
        m2 = b.ozet.get("kar_marji")
        gostergeler["kar_marji"] = {
            "donem1": m1,
            "donem2": m2,
            "fark": (m1 - m2) if m1 is not None and m2 is not None else None,
            "etiket": "puan",
            "yon": "artış"
            if m1 is not None and m2 is not None and m1 > m2
            else ("azalış" if m1 is not None and m2 is not None and m1 < m2 else "değişmedi"),
        }
        uyari = None
        if gun1 != gun2:
            uyari = f"Dönem gün sayıları farklı (D1={gun1}, D2={gun2}); günlük ortalamalar da gösterilir."
        return {
            "gostergeler": gostergeler,
            "gun1": gun1,
            "gun2": gun2,
            "uyari": uyari,
            "donem1": {"baslangic": a.filtre.baslangic, "bitis": a.filtre.bitis},
            "donem2": {"baslangic": b.filtre.baslangic, "bitis": b.filtre.bitis},
        }

    @staticmethod
    def _kars_bulgular(a: KarAnalizSonuc, b: KarAnalizSonuc) -> list[str]:
        d = degisim_hesapla(a.ozet["net_satis"], b.ozet["net_satis"])
        satirlar = [
            f"Net satış önceki döneme göre {d['etiket']} ({d['yon']}; fark {d['fark']} TL)."
        ]
        k = degisim_hesapla(a.ozet["brut_kar"], b.ozet["brut_kar"])
        satirlar.append(f"Brüt kâr {k['fark']} TL değişti ({k['etiket']}).")
        return satirlar

    @staticmethod
    def excel_aktar(sonuc: KarAnalizSonuc, yol: str) -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Font, numbers

        wb = Workbook()
        # Özet
        ws = wb.active
        ws.title = "Ozet"
        ws["A1"] = "Satış Kâr Analizi"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Dönem: {sonuc.filtre.baslangic:%d.%m.%Y} - {sonuc.filtre.bitis:%d.%m.%Y}"
        ws["A3"] = f"Maliyet yöntemi: {sonuc.filtre.maliyet_yontemi}"
        r = 5
        for k, v in sonuc.ozet.items():
            ws.cell(r, 1, k)
            ws.cell(r, 2, float(v) if isinstance(v, Decimal) else v)
            r += 1

        def _sheet(ad, headers, rows):
            s = wb.create_sheet(ad)
            for c, h in enumerate(headers, 1):
                s.cell(1, c, h).font = Font(bold=True)
            for i, row in enumerate(rows, 2):
                for c, val in enumerate(row, 1):
                    if isinstance(val, Decimal):
                        s.cell(i, c, float(val))
                        s.cell(i, c).number_format = "#,##0.00"
                    elif isinstance(val, date):
                        s.cell(i, c, val.strftime("%d.%m.%Y"))
                    else:
                        s.cell(i, c, val)

        _sheet(
            "Urunler",
            ["Kod", "Ad", "Marka", "Miktar", "Net", "Maliyet", "Kâr", "Marj %"],
            [
                [
                    u["urun_kodu"],
                    u["urun_adi"],
                    u["marka"],
                    u["net_miktar"],
                    u["net"],
                    u["maliyet"],
                    u["kar"],
                    u["kar_marji"],
                ]
                for u in sonuc.urunler
            ],
        )
        _sheet(
            "Musteriler",
            ["Kod", "Unvan", "Net", "Maliyet", "Kâr", "Marj %", "Belge"],
            [
                [m["cari_kodu"], m["unvan"], m["net"], m["maliyet"], m["kar"], m["kar_marji"], m["belge_sayisi"]]
                for m in sonuc.musteriler
            ],
        )
        _sheet(
            "Belgeler",
            ["Tür", "No", "Tarih", "Cari", "Net", "Maliyet", "Kâr", "Marj %"],
            [
                [
                    b["belge_turu"],
                    b["belge_no"],
                    b["tarih"],
                    b["cari_unvan"],
                    b["net"],
                    b["maliyet"],
                    b["kar"],
                    b["kar_marji"],
                ]
                for b in sonuc.belgeler
            ],
        )
        _sheet(
            "Zararina",
            ["Tarih", "Belge", "Ürün", "Net", "Maliyet", "Zarar", "Marj %"],
            [
                [s.tarih, s.belge_no, s.urun_kodu, s.net_satis, s.toplam_maliyet, s.brut_kar, s.kar_marji]
                for s in sonuc.zararina
            ],
        )
        _sheet(
            "SifirMaliyet",
            ["Tarih", "Belge", "Ürün", "Net", "Kod", "Açıklama"],
            [
                [s.tarih, s.belge_no, s.urun_kodu, s.net_satis, s.maliyet_kodu, s.maliyet_kaynagi]
                for s in sonuc.sifir_maliyet
            ],
        )
        if sonuc.karsilastirma:
            rows = []
            for ad, g in sonuc.karsilastirma["gostergeler"].items():
                rows.append([ad, g.get("donem1"), g.get("donem2"), g.get("fark"), g.get("etiket"), g.get("yon")])
            _sheet("Karsilastirma", ["Gösterge", "Dönem1", "Dönem2", "Fark", "Değişim", "Yön"], rows)

        ws_f = wb.create_sheet("Filtreler")
        ws_f["A1"] = "Filtre özeti"
        ws_f["A2"] = str(sonuc.filtre)

        wb.save(yol)
        return yol


def _fmt_opt(v) -> str:
    if v is None:
        return "N/A"
    return f"%{v}"
