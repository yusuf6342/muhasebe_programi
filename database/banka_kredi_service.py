"""Banka kredileri servisi — tek banka çıkışı + bileşen bazlı gider dağıtımı.

Taksit ödemesinde bankadan yalnızca BİR çıkış hareketi oluşur (toplam tutar).
Anapara krediler hesabındaki borcu azaltır; faiz / BSMV / KKDF / komisyon /
sigorta / dosya / diğer / gecikme bileşenleri için ek banka hareketi yazılmaz,
yalnızca gider fişi kesilir.
"""

from __future__ import annotations

import calendar
import json
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from sqlalchemy import select
from sqlalchemy.orm import selectinload

from database.database import get_session
from database.finans_service import FinansService
from database.models.finans import (
    KREDI_KISMI_DAGITIM_YONTEMLERI,
    KREDI_MAX_TAKSIT,
    KREDI_ODEME_HESAP_TURLERI,
    KREDI_TURLERI,
    BankaKarti,
    BankaKrediIslemGunlugu,
    BankaKrediOdeme,
    BankaKrediPlanVersiyon,
    BankaKredisi,
    BankaKrediTaksit,
    FinansHareketi,
    FinansHesabi,
    GiderFisi,
)
from database.session_manager import oturum

KURUS = Decimal("0.01")
SIFIR = Decimal("0.00")
YAKLASAN_GUN = 7

# Ödeme bileşenleri — anapara borç azaltır, kalanı finansman gideridir
BILESEN_ALANLARI = (
    "anapara",
    "faiz",
    "bsmv",
    "kkdf",
    "komisyon",
    "sigorta",
    "dosya_masrafi",
    "diger_masraflar",
    "gecikme_faizi",
)
GIDER_ALANLARI = tuple(a for a in BILESEN_ALANLARI if a != "anapara")

ALAN_ETIKET = {
    "anapara": "Anapara",
    "faiz": "Faiz",
    "bsmv": "BSMV",
    "kkdf": "KKDF",
    "komisyon": "Komisyon",
    "sigorta": "Sigorta",
    "dosya_masrafi": "Dosya masrafı",
    "diger_masraflar": "Diğer masraflar",
    "gecikme_faizi": "Gecikme faizi",
}

# alan -> (gider_turu, açıklama, belge eki)
GIDER_TURU_ESLEME = {
    "faiz": ("KREDI_FAIZ", "Kredi faiz gideri", "FZ"),
    "bsmv": ("KREDI_BSMV", "Kredi BSMV gideri", "BS"),
    "kkdf": ("KREDI_KKDF", "Kredi KKDF gideri", "KK"),
    "komisyon": ("KREDI_KOMISYON", "Kredi komisyon gideri", "KM"),
    "sigorta": ("KREDI_SIGORTA", "Kredi sigorta gideri", "SG"),
    "dosya_masrafi": ("KREDI_DOSYA", "Kredi dosya / işlem masrafı", "DS"),
    "diger_masraflar": ("KREDI_MASRAF", "Kredi diğer finansman gideri", "DG"),
    "gecikme_faizi": ("KREDI_GECIKME", "Kredi gecikme faizi", "GC"),
}

ODEME_TURLERI = ("TAKSIT", "KISMI", "ARA", "ERKEN_KAPAMA")
ACIK_TAKSIT_DURUMLARI = ("BEKLIYOR", "YAKLASIYOR", "VADESI_GECTI", "KISMEN_ODENDI")
KAPALI_TAKSIT_DURUMLARI = ("ODENDI", "IPTAL", "YAPILANDIRILDI", "ERKEN_KAPATILDI")

HAREKET_TAKSIT_ODEME = "KREDİ TAKSİT ÖDEME"
HAREKET_ANAPARA_ODEME = "KREDİ ANAPARA ÖDEME"
HAREKET_ODEME_IPTAL = "KREDİ ÖDEME İPTAL GİRİŞ"
HAREKET_ANAPARA_IPTAL = "KREDİ ANAPARA İPTAL GİRİŞ"

ODEME_HESAP_ALT_TURLERI = ("MEVDUAT", "KMH", "KREDI_KARTI")

# Soft ALTER — modelde tanımlı yeni nullable kolonlar
KREDI_YENI_KOLONLAR = {
    "kredi_kodu": "VARCHAR(40)",
    "banka_sube": "VARCHAR(100)",
    "kredi_hesap_no": "VARCHAR(50)",
    "kullanim_amaci": "VARCHAR(200)",
    "notlar": "TEXT",
    "faiz_orani": "NUMERIC(10, 4)",
    "faiz_turu": "VARCHAR(20) DEFAULT 'SABIT'",
    "taksit_donemi": "VARCHAR(20) DEFAULT 'AYLIK'",
    "para_birimi": "VARCHAR(10) DEFAULT 'TRY'",
    "kur": "NUMERIC(18, 6)",
    "kur_tarihi": "DATE",
    "tl_karsiligi": "NUMERIC(18, 2)",
    "bitis_tarihi": "DATE",
    "plan_versiyon": "INTEGER DEFAULT 1",
    "aktif": "BOOLEAN DEFAULT 1",
}

TAKSIT_YENI_KOLONLAR = {
    "bsmv": "NUMERIC(18, 2) DEFAULT 0",
    "kkdf": "NUMERIC(18, 2) DEFAULT 0",
    "komisyon": "NUMERIC(18, 2) DEFAULT 0",
    "sigorta": "NUMERIC(18, 2) DEFAULT 0",
    "dosya_masrafi": "NUMERIC(18, 2) DEFAULT 0",
    "diger_masraflar": "NUMERIC(18, 2) DEFAULT 0",
    "gecikme_faizi": "NUMERIC(18, 2) DEFAULT 0",
    "odenen_anapara": "NUMERIC(18, 2) DEFAULT 0",
    "odenen_faiz": "NUMERIC(18, 2) DEFAULT 0",
    "odenen_masraf": "NUMERIC(18, 2) DEFAULT 0",
    "odenen_tutar": "NUMERIC(18, 2) DEFAULT 0",
    "banka_dekont_no": "VARCHAR(60)",
    "odeme_hesap_id": "INTEGER",
    "aciklama": "VARCHAR(500)",
    "transaction_id": "VARCHAR(60)",
}


class BankaKrediService:
    # --- Şema ---

    @staticmethod
    def schema_hazirla() -> None:
        from sqlalchemy import inspect, text

        from database.database import engine

        import database.models.finans  # noqa: F401

        if engine is None:
            return
        BankaKredisi.__table__.create(engine, checkfirst=True)
        BankaKrediTaksit.__table__.create(engine, checkfirst=True)
        BankaKrediOdeme.__table__.create(engine, checkfirst=True)
        BankaKrediIslemGunlugu.__table__.create(engine, checkfirst=True)
        BankaKrediPlanVersiyon.__table__.create(engine, checkfirst=True)
        GiderFisi.__table__.create(engine, checkfirst=True)

        insp = inspect(engine)
        for tablo, kolonlar in (
            ("banka_kredileri", KREDI_YENI_KOLONLAR),
            ("banka_kredi_taksitleri", TAKSIT_YENI_KOLONLAR),
        ):
            if not insp.has_table(tablo):
                continue
            mevcut = {c["name"] for c in insp.get_columns(tablo)}
            eksikler = {a: t for a, t in kolonlar.items() if a not in mevcut}
            if not eksikler:
                continue
            with engine.begin() as connection:
                for alan, tip in eksikler.items():
                    connection.execute(
                        text(f'ALTER TABLE "{tablo}" ADD COLUMN "{alan}" {tip}')
                    )

    # --- Yardımcılar ---

    @staticmethod
    def _d(val, alan="Tutar", min=None) -> Decimal:
        ham = val
        if ham is None or (isinstance(ham, str) and not ham.strip()):
            ham = 0
        try:
            sonuc = Decimal(str(ham).strip().replace(" ", "").replace(",", "."))
        except (InvalidOperation, AttributeError) as hata:
            raise ValueError(f"{alan} geçerli bir sayı olmalıdır.") from hata
        sonuc = sonuc.quantize(KURUS, rounding=ROUND_HALF_UP)
        if min is not None and sonuc < min:
            raise ValueError(f"{alan} {min} değerinden küçük olamaz.")
        return sonuc

    @staticmethod
    def taksit_toplam(bilesenler: dict) -> Decimal:
        toplam = SIFIR
        for alan in BILESEN_ALANLARI:
            toplam += BankaKrediService._d(bilesenler.get(alan) or 0, ALAN_ETIKET[alan])
        return toplam.quantize(KURUS, rounding=ROUND_HALF_UP)

    @staticmethod
    def taksit_bilesenleri(taksit) -> dict:
        def al(alan, etiket=None):
            if isinstance(taksit, dict):
                ham = taksit.get(alan)
            else:
                ham = getattr(taksit, alan, None)
            return BankaKrediService._d(ham or 0, etiket or alan)

        veri = {alan: al(alan, ALAN_ETIKET[alan]) for alan in BILESEN_ALANLARI}
        legacy = al("masraf", "Masraf")
        yeni_masraf = sum(
            (veri[a] for a in GIDER_ALANLARI if a != "faiz"), SIFIR
        )
        if yeni_masraf == 0 and legacy > 0:
            veri["diger_masraflar"] = legacy
        veri["masraf"] = legacy
        veri["toplam"] = BankaKrediService.taksit_toplam(veri)
        return veri

    @staticmethod
    def taksit_durum_hesapla(vade, odenen, toplam, mevcut_durum=None) -> str:
        durum = (mevcut_durum or "").strip().upper()
        if durum in ("IPTAL", "YAPILANDIRILDI", "ERKEN_KAPATILDI"):
            return durum
        odenen = BankaKrediService._d(odenen or 0, "Ödenen")
        toplam = BankaKrediService._d(toplam or 0, "Toplam")
        if toplam > 0 and odenen >= toplam:
            return "ODENDI"
        if odenen > 0:
            return "KISMEN_ODENDI"
        bugun = date.today()
        if vade and vade < bugun:
            return "VADESI_GECTI"
        if vade and vade <= bugun + timedelta(days=YAKLASAN_GUN):
            return "YAKLASIYOR"
        return "BEKLIYOR"

    @staticmethod
    def _kullanici() -> str:
        return (oturum.kullanici_adi or oturum.ad_soyad or "sistem")[:80]

    @staticmethod
    def _gunluk(
        session,
        islem: str,
        *,
        kredi_id=None,
        taksit_id=None,
        odeme_id=None,
        transaction_id=None,
        detay=None,
    ) -> None:
        session.add(
            BankaKrediIslemGunlugu(
                kredi_id=int(kredi_id) if kredi_id else None,
                taksit_id=int(taksit_id) if taksit_id else None,
                odeme_id=int(odeme_id) if odeme_id else None,
                transaction_id=transaction_id,
                islem=islem[:40],
                detay=detay,
                kullanici=BankaKrediService._kullanici(),
                created_at=datetime.now(),
            )
        )

    @staticmethod
    def _izin(*kodlar: str, yazma: bool = True, mesaj: str | None = None) -> None:
        from database.access import AccessError, aktif_firma_zorunlu, yetki_var

        aktif_firma_zorunlu()
        if not yetki_var(*kodlar):
            ad = ", ".join(kodlar)
            raise AccessError(mesaj or f"Bu işlem için yetkiniz yok ({ad}).")
        if not yazma:
            return
        from database.donem_service import DonemService

        if not DonemService.kayit_izinli_mi():
            raise AccessError(
                "Seçili çalışma dönemi kapalı. Yeni kayıt veya değişiklik yapılamaz."
            )

    @staticmethod
    def _iptal_izni() -> None:
        from database.access import AccessError, aktif_firma_zorunlu, yetki_var

        aktif_firma_zorunlu()
        uygun = yetki_var("banka_kredi_odeme_iptal") or (
            yetki_var("finans_duzenleme") and yetki_var("iptal")
        )
        if not uygun:
            raise AccessError(
                "Kredi ödemesi iptal etme yetkiniz yok (banka_kredi_odeme_iptal)."
            )
        from database.donem_service import DonemService

        if not DonemService.kayit_izinli_mi():
            raise AccessError(
                "Seçili çalışma dönemi kapalı. Yeni kayıt veya değişiklik yapılamaz."
            )

    @staticmethod
    def _dekont_mukerrer(session, banka_hesap_id, dekont_no, exclude_odeme_id=None) -> bool:
        dekont = (dekont_no or "").strip()
        if not dekont or not banka_hesap_id:
            return False
        q = select(BankaKrediOdeme).where(
            BankaKrediOdeme.banka_hesap_id == int(banka_hesap_id),
            BankaKrediOdeme.banka_dekont_no == dekont,
            BankaKrediOdeme.durum != "IPTAL",
        )
        if exclude_odeme_id:
            q = q.where(BankaKrediOdeme.id != int(exclude_odeme_id))
        return session.scalar(q) is not None

    @staticmethod
    def _gider_belge_no(session, on_ek="GDF") -> str:
        yil = date.today().year
        like = f"{on_ek}-{yil}-%"
        mevcut = session.scalars(
            select(GiderFisi.belge_no).where(GiderFisi.belge_no.like(like))
        ).all()
        sira = 1
        for no in mevcut:
            parca = str(no).split("-")
            for p in parca[2:3]:
                try:
                    sira = max(sira, int(p) + 1)
                except ValueError:
                    pass
        return f"{on_ek}-{yil}-{sira:04d}"

    @staticmethod
    def _kredi_hesaplari(session, kredi, banka_hesap_id=None):
        """(kart, kredi_hesap, odeme_hesap) döner."""
        kart = session.scalar(
            select(BankaKarti)
            .where(BankaKarti.id == kredi.banka_karti_id)
            .options(selectinload(BankaKarti.alt_hesaplar))
        )
        if not kart:
            raise ValueError("Banka kartı bulunamadı.")
        kredi_hesap = next(
            (h for h in kart.alt_hesaplar if (h.alt_hesap_turu or "") == "KREDILER"),
            None,
        )
        if not kredi_hesap:
            raise ValueError("Krediler hesabı eksik — banka kartını kaydedin.")
        if banka_hesap_id:
            odeme_hesap = session.scalar(
                select(FinansHesabi).where(FinansHesabi.id == int(banka_hesap_id))
            )
            if not odeme_hesap or not odeme_hesap.aktif:
                raise ValueError("Ödeme hesabı bulunamadı veya pasif.")
            if (odeme_hesap.alt_hesap_turu or "").upper() == "KREDILER":
                raise ValueError("Krediler hesabından ödeme yapılamaz.")
            alt = (odeme_hesap.alt_hesap_turu or "").upper()
            tur = (odeme_hesap.hesap_turu or "").upper()
            if tur != "KASA" and alt and alt not in ODEME_HESAP_ALT_TURLERI:
                raise ValueError("Ödeme hesabı kasa, mevduat, KMH veya kredi kartı olmalıdır.")
        else:
            odeme_tur = (kredi.odeme_hesap_turu or "MEVDUAT").upper()
            odeme_hesap = next(
                (h for h in kart.alt_hesaplar if (h.alt_hesap_turu or "") == odeme_tur),
                None,
            )
            if not odeme_hesap:
                raise ValueError(f"{odeme_tur} hesabı eksik — banka kartını kaydedin.")
        return kart, kredi_hesap, odeme_hesap

    @staticmethod
    def _odenen_bilesenler(session, taksit_id) -> dict:
        veri = {alan: SIFIR for alan in BILESEN_ALANLARI}
        if not taksit_id:
            return veri
        odemeler = session.scalars(
            select(BankaKrediOdeme).where(
                BankaKrediOdeme.taksit_id == int(taksit_id),
                BankaKrediOdeme.durum == "AKTIF",
            )
        ).all()
        for o in odemeler:
            for alan in BILESEN_ALANLARI:
                veri[alan] += BankaKrediService._d(
                    getattr(o, alan, None) or 0, ALAN_ETIKET[alan]
                )
        return veri

    @staticmethod
    def _taksit_kalan(taksit) -> Decimal:
        if (taksit.durum or "") in KAPALI_TAKSIT_DURUMLARI:
            return SIFIR
        plan = BankaKrediService.taksit_bilesenleri(taksit)
        odenen = BankaKrediService._d(taksit.odenen_tutar or 0, "Ödenen tutar")
        kalan = plan["toplam"] - odenen
        return kalan if kalan > 0 else SIFIR

    @staticmethod
    def _dagit(kalan: dict, tutar: Decimal, yontem: str) -> dict:
        y = (yontem or "ONCE_MASRAF").strip().upper()
        if y not in {k for k, _ in KREDI_KISMI_DAGITIM_YONTEMLERI}:
            raise ValueError("Geçersiz dağıtım yöntemi.")
        if y == "MANUEL":
            raise ValueError("Manuel dağıtımda ödeme bileşenleri girilmelidir.")
        sonuc = {alan: SIFIR for alan in BILESEN_ALANLARI}
        kalan_toplam = BankaKrediService.taksit_toplam(kalan)
        if kalan_toplam <= 0:
            raise ValueError("Bu taksitte ödenecek kalan tutar yok.")
        if tutar > kalan_toplam:
            raise ValueError("Ödeme tutarı kalan taksit tutarını aşıyor.")
        if y == "ONCE_MASRAF":
            sira = (
                "gecikme_faizi",
                "faiz",
                "bsmv",
                "kkdf",
                "komisyon",
                "sigorta",
                "dosya_masrafi",
                "diger_masraflar",
                "anapara",
            )
            artan = tutar
            for alan in sira:
                if artan <= 0:
                    break
                pay = kalan[alan] if kalan[alan] < artan else artan
                if pay > 0:
                    sonuc[alan] = pay
                    artan -= pay
            return sonuc
        # ORANSAL
        dolu = [alan for alan in BILESEN_ALANLARI if kalan[alan] > 0]
        biriken = SIFIR
        for i, alan in enumerate(dolu):
            if i == len(dolu) - 1:
                pay = tutar - biriken
            else:
                pay = (tutar * kalan[alan] / kalan_toplam).quantize(
                    KURUS, rounding=ROUND_HALF_UP
                )
            if pay < 0:
                pay = SIFIR
            sonuc[alan] = pay
            biriken += pay
        return sonuc

    @staticmethod
    def _muhasebe_fis_olustur(session, odeme) -> int | None:
        """Hesap eşleştirmeleri tamsa muhasebe fişi üretir; eksikse sessizce atlar."""
        from database.models.genel_muhasebe import (
            HesapPlani,
            MuhasebeFisi,
            MuhasebeFisiSatiri,
            MuhasebeHesapEsleme,
        )
        from database.muhasebe_service import MuhasebeService

        firma_id = MuhasebeService.yerel_firma_id(session)

        def hesap(anahtar):
            e = session.scalar(
                select(MuhasebeHesapEsleme).where(
                    MuhasebeHesapEsleme.firma_id == firma_id,
                    MuhasebeHesapEsleme.anahtar == anahtar,
                    MuhasebeHesapEsleme.aktif.is_(True),
                )
            )
            if e is None or not e.hesap_id:
                return None
            h = session.get(HesapPlani, e.hesap_id)
            if h is None or not h.aktif:
                return None
            return h

        banka = hesap("banka")
        krediler = (
            hesap("kredi_kisa_vadeli")
            or hesap("kredi_uzun_vadeli")
            or hesap("krediler")
            or hesap("banka_kredileri")
        )
        genel_gider = hesap("giderler") or hesap("kredi_diger_finansman")
        if banka is None or krediler is None:
            return None

        anapara = BankaKrediService._d(odeme.anapara or 0, "Anapara")
        toplam = BankaKrediService._d(odeme.toplam or 0, "Toplam")
        bilesen_satirlar = (
            ("faiz", "kredi_faiz_gideri", "Faiz gideri"),
            ("bsmv", "kredi_bsmv_gideri", "BSMV"),
            ("kkdf", "kredi_kkdf_gideri", "KKDF"),
            ("komisyon", "kredi_komisyon_gideri", "Banka komisyonu"),
            ("sigorta", "kredi_sigorta_gideri", "Sigorta"),
            ("dosya_masrafi", "kredi_dosya_masrafi", "Dosya/işlem masrafı"),
            ("diger_masraflar", "kredi_diger_finansman", "Diğer finansman"),
            ("gecikme_faizi", "kredi_gecikme_faizi", "Gecikme faizi"),
        )
        satirlar = []
        if anapara > 0:
            satirlar.append((krediler, anapara, SIFIR, "Kredi anapara ödemesi"))
        for alan, anahtar, etiket in bilesen_satirlar:
            tutar = BankaKrediService._d(getattr(odeme, alan, 0) or 0, etiket)
            if tutar <= 0:
                continue
            h = hesap(anahtar) or genel_gider
            if h is None:
                return None
            satirlar.append((h, tutar, SIFIR, etiket))
        satirlar.append((banka, SIFIR, toplam, "Banka çıkışı"))
        borc_toplam = sum((b for _h, b, _a, _ in satirlar), SIFIR)
        alacak_toplam = sum((a for _h, _b, a, _ in satirlar), SIFIR)
        if borc_toplam != alacak_toplam or borc_toplam != toplam:
            return None

        mali_yil = odeme.odeme_tarihi.year
        from database.muhasebe_service import MuhasebeFisService

        fis = MuhasebeFisi(
            firma_id=firma_id,
            donem_id=oturum.period_id,
            mali_yil=mali_yil,
            fis_no=MuhasebeFisService._sonraki_fis_no(session, firma_id, mali_yil),
            fis_tarihi=odeme.odeme_tarihi,
            fis_turu="Mahsup",
            aciklama=f"Kredi ödemesi {odeme.odeme_belge_no}",
            belge_no=odeme.odeme_belge_no,
            durum="Kesinleşmiş",
            toplam_borc=toplam,
            toplam_alacak=toplam,
            kaynak_turu="banka_kredi_odeme",
            kaynak_id=int(odeme.id),
            olusturan_kullanici_id=oturum.user_id,
        )
        session.add(fis)
        session.flush()
        for sira, (h, borc, alacak, acik) in enumerate(satirlar, start=1):
            session.add(
                MuhasebeFisiSatiri(
                    fis_id=fis.id,
                    firma_id=firma_id,
                    sira_no=sira,
                    hesap_id=h.id,
                    hesap_kodu=h.hesap_kodu,
                    hesap_adi=h.hesap_adi,
                    aciklama=acik,
                    borc=borc,
                    alacak=alacak,
                    belge_tarihi=odeme.odeme_tarihi,
                    belge_no=odeme.odeme_belge_no,
                )
            )
        session.flush()
        return int(fis.id)

    # --- Listeleme / özet ---

    @staticmethod
    def ozet_kartlar() -> dict:
        bugun = date.today()
        ay_basi = bugun.replace(day=1)
        ay_sonu = date(bugun.year, bugun.month, calendar.monthrange(bugun.year, bugun.month)[1])
        toplam_borc = SIFIR
        bu_ay_odenecek = SIFIR
        gecikmis = SIFIR
        bu_ay_anapara = SIFIR
        bu_ay_finansman = SIFIR
        gelecek_vade = None
        with get_session() as session:
            krediler = session.scalars(
                select(BankaKredisi)
                .where(BankaKredisi.durum == "KULLANDIRILDI")
                .options(selectinload(BankaKredisi.taksitler))
            ).all()
            for k in krediler:
                toplam_borc += BankaKrediService._kalan_anapara(session, k)
                for t in k.taksitler or []:
                    kalan = BankaKrediService._taksit_kalan(t)
                    if kalan <= 0:
                        continue
                    if t.vade_tarihi < bugun:
                        gecikmis += kalan
                    if ay_basi <= t.vade_tarihi <= ay_sonu:
                        bu_ay_odenecek += kalan
                    if t.vade_tarihi >= bugun and (
                        gelecek_vade is None or t.vade_tarihi < gelecek_vade
                    ):
                        gelecek_vade = t.vade_tarihi
            odemeler = session.scalars(
                select(BankaKrediOdeme).where(
                    BankaKrediOdeme.durum == "AKTIF",
                    BankaKrediOdeme.odeme_tarihi >= ay_basi,
                    BankaKrediOdeme.odeme_tarihi <= ay_sonu,
                )
            ).all()
            for o in odemeler:
                bu_ay_anapara += BankaKrediService._d(o.anapara or 0, "Anapara")
                for alan in GIDER_ALANLARI:
                    bu_ay_finansman += BankaKrediService._d(
                        getattr(o, alan, None) or 0, ALAN_ETIKET[alan]
                    )
        return {
            "toplam_borc": toplam_borc,
            "bu_ay_odenecek": bu_ay_odenecek,
            "gecikmis": gecikmis,
            "bu_ay_odenen_anapara": bu_ay_anapara,
            "bu_ay_finansman_gideri": bu_ay_finansman,
            "gelecek_taksit_tarihi": gelecek_vade,
        }

    @staticmethod
    def _kalan_anapara(session, kredi) -> Decimal:
        ana = BankaKrediService._d(kredi.ana_para or 0, "Ana para")
        odenen = SIFIR
        for t in kredi.taksitler or []:
            tutar = BankaKrediService._d(t.odenen_anapara or 0, "Ödenen anapara")
            if tutar == 0 and (t.durum or "") in ("ODENDI", "ERKEN_KAPATILDI"):
                tutar = BankaKrediService._d(t.anapara or 0, "Anapara")
            odenen += tutar
        ara = session.scalars(
            select(BankaKrediOdeme).where(
                BankaKrediOdeme.kredi_id == kredi.id,
                BankaKrediOdeme.durum == "AKTIF",
                BankaKrediOdeme.taksit_id.is_(None),
            )
        ).all()
        for o in ara:
            odenen += BankaKrediService._d(o.anapara or 0, "Anapara")
        kalan = ana - odenen
        return kalan if kalan > 0 else SIFIR

    @staticmethod
    def kredi_listesi(banka_karti_id=None, durum=None, aktif_only=True) -> list[dict]:
        with get_session() as session:
            q = (
                select(BankaKredisi)
                .options(selectinload(BankaKredisi.taksitler))
                .order_by(BankaKredisi.kullandirim_tarihi.desc(), BankaKredisi.id.desc())
            )
            if banka_karti_id:
                q = q.where(BankaKredisi.banka_karti_id == int(banka_karti_id))
            if durum:
                q = q.where(BankaKredisi.durum == str(durum).strip().upper())
            elif aktif_only:
                q = q.where(BankaKredisi.durum == "KULLANDIRILDI")
            kartlar = {
                k.id: k
                for k in session.scalars(select(BankaKarti)).all()
            }
            sonuc = []
            for k in session.scalars(q).all():
                kart = kartlar.get(k.banka_karti_id)
                acik = [
                    t for t in (k.taksitler or []) if (t.durum or "") in ACIK_TAKSIT_DURUMLARI
                ]
                kalan_taksit = sum(
                    (BankaKrediService._taksit_kalan(t) for t in acik), SIFIR
                )
                acik_sirali = sorted(
                    acik, key=lambda t: (t.vade_tarihi or date.max, t.taksit_no or 0)
                )
                aylik_taksit = SIFIR
                gelecek_vade = None
                if acik_sirali:
                    sonraki = acik_sirali[0]
                    kalan = BankaKrediService._taksit_kalan(sonraki)
                    aylik_taksit = (
                        kalan
                        if kalan > 0
                        else BankaKrediService.taksit_bilesenleri(sonraki)["toplam"]
                    )
                    gelecek_vade = sonraki.vade_tarihi
                sonuc.append({
                    "id": k.id,
                    "belge_no": k.belge_no,
                    "kredi_kodu": k.kredi_kodu or "",
                    "banka_karti_id": k.banka_karti_id,
                    "banka_adi": kart.banka_adi if kart else "",
                    "banka_sube": k.banka_sube or (kart.sube if kart else "") or "",
                    "kredi_adi": k.kredi_adi,
                    "kredi_turu": k.kredi_turu,
                    "ana_para": BankaKrediService._d(k.ana_para or 0, "Ana para"),
                    "toplam_faiz": BankaKrediService._d(k.toplam_faiz or 0, "Toplam faiz"),
                    "toplam_masraf": BankaKrediService._d(k.toplam_masraf or 0, "Toplam masraf"),
                    "faiz_orani": BankaKrediService._d(k.faiz_orani or 0, "Faiz oranı"),
                    "taksit_sayisi": k.taksit_sayisi,
                    "kalan_taksit_sayisi": len(acik),
                    "aylik_taksit_tutari": aylik_taksit,
                    "gelecek_taksit_tarihi": gelecek_vade,
                    "kullandirim_tarihi": k.kullandirim_tarihi,
                    "ilk_taksit_tarihi": k.ilk_taksit_tarihi,
                    "bitis_tarihi": k.bitis_tarihi,
                    "para_birimi": k.para_birimi or "TRY",
                    "odeme_hesap_turu": k.odeme_hesap_turu or "MEVDUAT",
                    "sozlesme_no": k.sozlesme_no or "",
                    "durum": k.durum,
                    "aktif": bool(k.aktif) if k.aktif is not None else True,
                    "kalan_anapara": BankaKrediService._kalan_anapara(session, k),
                    "kalan_borc": kalan_taksit,
                })
            return sonuc

    @staticmethod
    def kredi_detay(kredi_id) -> dict | None:
        with get_session() as session:
            k = session.scalar(
                select(BankaKredisi)
                .where(BankaKredisi.id == int(kredi_id))
                .options(selectinload(BankaKredisi.taksitler))
            )
            if not k:
                return None
            kart = session.get(BankaKarti, k.banka_karti_id)
            taksitler = []
            odenen_toplam = SIFIR
            plan_toplam = SIFIR
            for t in k.taksitler or []:
                plan = BankaKrediService.taksit_bilesenleri(t)
                odenen = BankaKrediService._d(t.odenen_tutar or 0, "Ödenen tutar")
                if odenen == 0 and (t.durum or "") == "ODENDI":
                    odenen = plan["toplam"]
                plan_toplam += plan["toplam"]
                odenen_toplam += odenen
                satir = {
                    "id": t.id,
                    "taksit_id": t.id,
                    "taksit_no": t.taksit_no,
                    "vade_tarihi": t.vade_tarihi,
                    "durum": t.durum,
                    "odenen_tutar": odenen,
                    "kalan": BankaKrediService._taksit_kalan(t),
                    "odeme_tarihi": t.odeme_tarihi,
                    "odeme_belge_no": t.odeme_belge_no,
                    "gider_belge_no": t.gider_belge_no,
                    "banka_dekont_no": t.banka_dekont_no,
                    "aciklama": t.aciklama or "",
                }
                satir.update(plan)
                taksitler.append(satir)
            ara_odeme_anapara = SIFIR
            for o in session.scalars(
                select(BankaKrediOdeme).where(
                    BankaKrediOdeme.kredi_id == k.id,
                    BankaKrediOdeme.durum == "AKTIF",
                    BankaKrediOdeme.taksit_id.is_(None),
                )
            ).all():
                ara_odeme_anapara += BankaKrediService._d(o.anapara or 0, "Anapara")
            return {
                "id": k.id,
                "belge_no": k.belge_no,
                "kredi_kodu": k.kredi_kodu or "",
                "banka_karti_id": k.banka_karti_id,
                "banka_adi": kart.banka_adi if kart else "",
                "banka_sube": k.banka_sube or (kart.sube if kart else "") or "",
                "kredi_hesap_no": k.kredi_hesap_no or "",
                "kredi_adi": k.kredi_adi,
                "kredi_turu": k.kredi_turu,
                "ana_para": BankaKrediService._d(k.ana_para or 0, "Ana para"),
                "toplam_faiz": BankaKrediService._d(k.toplam_faiz or 0, "Toplam faiz"),
                "toplam_masraf": BankaKrediService._d(k.toplam_masraf or 0, "Toplam masraf"),
                "faiz_orani": BankaKrediService._d(k.faiz_orani or 0, "Faiz oranı"),
                "faiz_turu": k.faiz_turu or "SABIT",
                "taksit_donemi": k.taksit_donemi or "AYLIK",
                "taksit_sayisi": k.taksit_sayisi,
                "kullandirim_tarihi": k.kullandirim_tarihi,
                "ilk_taksit_tarihi": k.ilk_taksit_tarihi,
                "bitis_tarihi": k.bitis_tarihi,
                "para_birimi": k.para_birimi or "TRY",
                "kur": BankaKrediService._d(k.kur or 0, "Kur"),
                "tl_karsiligi": BankaKrediService._d(k.tl_karsiligi or 0, "TL karşılığı"),
                "odeme_hesap_turu": k.odeme_hesap_turu or "MEVDUAT",
                "sozlesme_no": k.sozlesme_no or "",
                "kullanim_amaci": k.kullanim_amaci or "",
                "aciklama": k.aciklama or "",
                "notlar": k.notlar or "",
                "plan_versiyon": k.plan_versiyon or 1,
                "durum": k.durum,
                "aktif": bool(k.aktif) if k.aktif is not None else True,
                "bakiye_ozeti": {
                    "ana_para": BankaKrediService._d(k.ana_para or 0, "Ana para"),
                    "kalan_anapara": BankaKrediService._kalan_anapara(session, k),
                    "ara_odeme_anapara": ara_odeme_anapara,
                    "plan_toplam": plan_toplam,
                    "odenen_toplam": odenen_toplam,
                    "kalan_toplam": plan_toplam - odenen_toplam
                    if plan_toplam > odenen_toplam
                    else SIFIR,
                },
                "taksitler": taksitler,
            }

    @staticmethod
    def _taksit_satirlari(session, taksitler) -> list[dict]:
        kredi_map = {}
        sonuc = []
        for t in taksitler:
            kredi = kredi_map.get(t.kredi_id)
            if kredi is None:
                kredi = session.get(BankaKredisi, t.kredi_id)
                kredi_map[t.kredi_id] = kredi
            kart = session.get(BankaKarti, kredi.banka_karti_id) if kredi else None
            plan = BankaKrediService.taksit_bilesenleri(t)
            satir = {
                "taksit_id": t.id,
                "kredi_id": t.kredi_id,
                "belge_no": kredi.belge_no if kredi else "",
                "kredi_adi": kredi.kredi_adi if kredi else "",
                "banka_adi": kart.banka_adi if kart else "",
                "taksit_no": t.taksit_no,
                "taksit_sayisi": kredi.taksit_sayisi if kredi else 0,
                "vade_tarihi": t.vade_tarihi,
                "durum": t.durum,
                "odenen_tutar": BankaKrediService._d(t.odenen_tutar or 0, "Ödenen tutar"),
                "kalan": BankaKrediService._taksit_kalan(t),
                "odeme_tarihi": t.odeme_tarihi,
                "odeme_belge_no": t.odeme_belge_no or "",
                "gider_belge_no": t.gider_belge_no or "",
                "banka_dekont_no": t.banka_dekont_no or "",
                "odeme_hesap_turu": (kredi.odeme_hesap_turu if kredi else "MEVDUAT") or "MEVDUAT",
            }
            satir.update(plan)
            sonuc.append(satir)
        return sonuc

    @staticmethod
    def bekleyen_taksitler(banka_karti_id=None, kredi_id=None, gun=None, limit=500) -> list[dict]:
        with get_session() as session:
            q = (
                select(BankaKrediTaksit)
                .join(BankaKredisi, BankaKrediTaksit.kredi_id == BankaKredisi.id)
                .where(
                    BankaKredisi.durum == "KULLANDIRILDI",
                    BankaKrediTaksit.durum.in_(ACIK_TAKSIT_DURUMLARI),
                )
                .order_by(BankaKrediTaksit.vade_tarihi, BankaKrediTaksit.id)
                .limit(limit)
            )
            if banka_karti_id:
                q = q.where(BankaKredisi.banka_karti_id == int(banka_karti_id))
            if kredi_id:
                q = q.where(BankaKrediTaksit.kredi_id == int(kredi_id))
            if gun:
                q = q.where(
                    BankaKrediTaksit.vade_tarihi <= date.today() + timedelta(days=int(gun))
                )
            return BankaKrediService._taksit_satirlari(session, session.scalars(q).all())

    @staticmethod
    def odenen_taksitler(banka_karti_id=None, kredi_id=None, limit=500) -> list[dict]:
        with get_session() as session:
            q = (
                select(BankaKrediTaksit)
                .join(BankaKredisi, BankaKrediTaksit.kredi_id == BankaKredisi.id)
                .where(BankaKrediTaksit.durum.in_(("ODENDI", "KISMEN_ODENDI", "ERKEN_KAPATILDI")))
                .order_by(BankaKrediTaksit.odeme_tarihi.desc(), BankaKrediTaksit.id.desc())
                .limit(limit)
            )
            if banka_karti_id:
                q = q.where(BankaKredisi.banka_karti_id == int(banka_karti_id))
            if kredi_id:
                q = q.where(BankaKrediTaksit.kredi_id == int(kredi_id))
            return BankaKrediService._taksit_satirlari(session, session.scalars(q).all())

    @staticmethod
    def odeme_listesi(kredi_id=None, taksit_id=None, limit=500) -> list[dict]:
        with get_session() as session:
            q = (
                select(BankaKrediOdeme)
                .order_by(BankaKrediOdeme.odeme_tarihi.desc(), BankaKrediOdeme.id.desc())
                .limit(limit)
            )
            if kredi_id:
                q = q.where(BankaKrediOdeme.kredi_id == int(kredi_id))
            if taksit_id:
                q = q.where(BankaKrediOdeme.taksit_id == int(taksit_id))
            return [
                BankaKrediService._odeme_sozluk(session, o)
                for o in session.scalars(q).all()
            ]

    @staticmethod
    def _odeme_sozluk(session, odeme) -> dict:
        hesap = session.get(FinansHesabi, odeme.banka_hesap_id)
        veri = {
            "id": odeme.id,
            "odeme_id": odeme.id,
            "transaction_id": odeme.transaction_id,
            "kredi_id": odeme.kredi_id,
            "taksit_id": odeme.taksit_id,
            "odeme_turu": odeme.odeme_turu,
            "odeme_tarihi": odeme.odeme_tarihi,
            "odeme_belge_no": odeme.odeme_belge_no,
            "gider_belge_no": odeme.gider_belge_no or "",
            "banka_hesap_id": odeme.banka_hesap_id,
            "banka_hesap_adi": hesap.hesap_adi if hesap else "",
            "banka_dekont_no": odeme.banka_dekont_no or "",
            "muhasebe_fis_id": odeme.muhasebe_fis_id,
            "aciklama": odeme.aciklama or "",
            "durum": odeme.durum,
            "iptal_nedeni": odeme.iptal_nedeni or "",
            "created_by": odeme.created_by or "",
            "created_at": odeme.created_at,
        }
        for alan in BILESEN_ALANLARI:
            veri[alan] = BankaKrediService._d(
                getattr(odeme, alan, None) or 0, ALAN_ETIKET[alan]
            )
        veri["toplam"] = BankaKrediService._d(odeme.toplam or 0, "Toplam")
        veri["masraf"] = sum(
            (veri[a] for a in GIDER_ALANLARI if a != "faiz"), SIFIR
        )
        return veri

    @staticmethod
    def bagli_belgeler(kredi_id=None, taksit_id=None, odeme_id=None) -> list[dict]:
        with get_session() as session:
            q = select(BankaKrediOdeme)
            if odeme_id:
                q = q.where(BankaKrediOdeme.id == int(odeme_id))
            elif taksit_id:
                q = q.where(BankaKrediOdeme.taksit_id == int(taksit_id))
            elif kredi_id:
                q = q.where(BankaKrediOdeme.kredi_id == int(kredi_id))
            else:
                raise ValueError("Kredi, taksit veya ödeme seçilmelidir.")
            odemeler = list(
                session.scalars(q.order_by(BankaKrediOdeme.id)).all()
            )
            belgeler = []
            belge_nolar = [o.odeme_belge_no for o in odemeler]
            for o in odemeler:
                belgeler.append({
                    "belge_turu": "KREDİ ÖDEME",
                    "belge_no": o.odeme_belge_no,
                    "tarih": o.odeme_tarihi,
                    "tutar": BankaKrediService._d(o.toplam or 0, "Toplam"),
                    "durum": o.durum,
                    "aciklama": o.aciklama or "",
                    "kayit_id": o.id,
                })
            if belge_nolar:
                for h in session.scalars(
                    select(FinansHareketi)
                    .where(FinansHareketi.belge_no.in_(belge_nolar))
                    .order_by(FinansHareketi.id)
                ).all():
                    hesap = session.get(FinansHesabi, h.hesap_id)
                    belgeler.append({
                        "belge_turu": h.hareket_turu,
                        "belge_no": h.belge_no,
                        "tarih": h.tarih,
                        "tutar": BankaKrediService._d(h.tutar or 0, "Tutar"),
                        "durum": hesap.hesap_adi if hesap else "",
                        "aciklama": h.aciklama or "",
                        "kayit_id": h.id,
                    })
                for f in session.scalars(
                    select(GiderFisi)
                    .where(GiderFisi.bagli_belge_no.in_(belge_nolar))
                    .order_by(GiderFisi.id)
                ).all():
                    belgeler.append({
                        "belge_turu": f"GİDER FİŞİ ({f.gider_turu})",
                        "belge_no": f.belge_no,
                        "tarih": f.tarih,
                        "tutar": BankaKrediService._d(f.tutar or 0, "Tutar"),
                        "durum": f.durum,
                        "aciklama": f.aciklama or "",
                        "kayit_id": f.id,
                    })
            return belgeler

    # --- Kredi oluşturma / güncelleme ---

    @staticmethod
    def kredi_kullandir(veriler: dict) -> dict:
        """FinansService.banka_kredi_kullandir'ı kullanır, ek alanları sonra yazar."""
        BankaKrediService._izin("banka_kredi_duzenleme", "finans_duzenleme")
        BankaKrediService.schema_hazirla()
        plan = list(veriler.get("taksit_plani") or [])
        toplam_faiz = veriler.get("toplam_faiz") or 0
        toplam_masraf = veriler.get("toplam_masraf") or 0
        eski_plan = None
        if plan:
            eski_plan = []
            toplam_faiz = SIFIR
            toplam_masraf = SIFIR
            for satir in plan:
                faiz = BankaKrediService._d(satir.get("faiz") or 0, "Faiz", SIFIR)
                masraf = sum(
                    (
                        BankaKrediService._d(satir.get(a) or 0, ALAN_ETIKET[a], SIFIR)
                        for a in GIDER_ALANLARI
                        if a != "faiz"
                    ),
                    SIFIR,
                )
                if masraf == 0:
                    masraf = BankaKrediService._d(satir.get("masraf") or 0, "Masraf", SIFIR)
                toplam_faiz += faiz
                toplam_masraf += masraf
                eski_plan.append({
                    "taksit_no": int(satir.get("taksit_no") or 0),
                    "vade_tarihi": satir.get("vade_tarihi") or satir.get("vade"),
                    "anapara": satir.get("anapara") or 0,
                    "faiz": faiz,
                    "masraf": masraf,
                })
        temel = FinansService.banka_kredi_kullandir(
            banka_karti_id=veriler.get("banka_karti_id"),
            kredi_adi=veriler.get("kredi_adi"),
            ana_para=veriler.get("ana_para"),
            kullandirim_tarihi=veriler.get("kullandirim_tarihi") or date.today(),
            ilk_taksit_tarihi=veriler.get("ilk_taksit_tarihi"),
            taksit_sayisi=veriler.get("taksit_sayisi") or 1,
            toplam_faiz=toplam_faiz,
            toplam_masraf=toplam_masraf,
            kredi_turu=veriler.get("kredi_turu") or "ISLETME",
            odeme_hesap_turu=veriler.get("odeme_hesap_turu") or "MEVDUAT",
            sozlesme_no=veriler.get("sozlesme_no"),
            aciklama=veriler.get("aciklama"),
            taksit_plani=eski_plan,
        )
        kredi_id = temel["id"]
        with get_session() as session:
            kredi = session.get(BankaKredisi, int(kredi_id))
            BankaKrediService._ek_alanlari_yaz(kredi, veriler)
            if plan:
                taksitler = {
                    t.taksit_no: t
                    for t in session.scalars(
                        select(BankaKrediTaksit).where(
                            BankaKrediTaksit.kredi_id == kredi.id
                        )
                    ).all()
                }
                for satir in plan:
                    t = taksitler.get(int(satir.get("taksit_no") or 0))
                    if t is None:
                        continue
                    BankaKrediService._taksit_bilesen_yaz(t, satir)
            BankaKrediService._gunluk(
                session,
                "KULLANDIRIM",
                kredi_id=kredi.id,
                detay=f"{kredi.belge_no} / {kredi.ana_para}",
            )
            session.flush()
        return BankaKrediService.kredi_detay(kredi_id)

    @staticmethod
    def _ek_alanlari_yaz(kredi, veriler: dict) -> None:
        if kredi is None:
            return
        metin_alanlar = (
            "kredi_kodu",
            "banka_sube",
            "kredi_hesap_no",
            "kullanim_amaci",
            "notlar",
            "sozlesme_no",
        )
        for alan in metin_alanlar:
            if alan in veriler:
                deger = (veriler.get(alan) or "").strip() or None
                setattr(kredi, alan, deger)
        if veriler.get("faiz_orani") is not None:
            kredi.faiz_orani = BankaKrediService._d(
                veriler.get("faiz_orani") or 0, "Faiz oranı", SIFIR
            )
        if veriler.get("faiz_turu"):
            kredi.faiz_turu = str(veriler["faiz_turu"]).strip().upper()
        if veriler.get("taksit_donemi"):
            kredi.taksit_donemi = str(veriler["taksit_donemi"]).strip().upper()
        if veriler.get("para_birimi"):
            kredi.para_birimi = str(veriler["para_birimi"]).strip().upper()
        if veriler.get("kur") is not None:
            kredi.kur = BankaKrediService._d(veriler.get("kur") or 0, "Kur", SIFIR)
        if veriler.get("kur_tarihi"):
            kredi.kur_tarihi = veriler["kur_tarihi"]
        if veriler.get("tl_karsiligi") is not None:
            kredi.tl_karsiligi = BankaKrediService._d(
                veriler.get("tl_karsiligi") or 0, "TL karşılığı", SIFIR
            )
        if veriler.get("bitis_tarihi"):
            kredi.bitis_tarihi = veriler["bitis_tarihi"]
        elif kredi.bitis_tarihi is None and kredi.taksitler:
            kredi.bitis_tarihi = max(t.vade_tarihi for t in kredi.taksitler)
        if kredi.plan_versiyon is None:
            kredi.plan_versiyon = 1
        if kredi.aktif is None:
            kredi.aktif = True

    @staticmethod
    def _taksit_bilesen_yaz(taksit, satir: dict) -> None:
        for alan in BILESEN_ALANLARI:
            if alan in satir:
                setattr(
                    taksit,
                    alan,
                    BankaKrediService._d(satir.get(alan) or 0, ALAN_ETIKET[alan], SIFIR),
                )
        if any(a in satir for a in GIDER_ALANLARI if a != "faiz"):
            taksit.masraf = sum(
                (
                    BankaKrediService._d(getattr(taksit, a, None) or 0, ALAN_ETIKET[a])
                    for a in GIDER_ALANLARI
                    if a != "faiz"
                ),
                SIFIR,
            )
        elif "masraf" in satir:
            taksit.masraf = BankaKrediService._d(
                satir.get("masraf") or 0, "Masraf", SIFIR
            )
        if satir.get("aciklama"):
            taksit.aciklama = str(satir["aciklama"]).strip()[:500] or None

    @staticmethod
    def kredi_guncelle(kredi_id, veriler: dict) -> dict:
        BankaKrediService._izin("banka_kredi_duzenleme", "finans_duzenleme")
        with get_session() as session:
            kredi = session.scalar(
                select(BankaKredisi)
                .where(BankaKredisi.id == int(kredi_id))
                .options(selectinload(BankaKredisi.taksitler))
            )
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            if veriler.get("kredi_adi") is not None:
                ad = (veriler.get("kredi_adi") or "").strip()
                if not ad:
                    raise ValueError("Kredi adı zorunludur.")
                kredi.kredi_adi = ad
            if veriler.get("kredi_turu"):
                tur = str(veriler["kredi_turu"]).strip().upper()
                if tur not in {k for k, _ in KREDI_TURLERI}:
                    raise ValueError("Geçersiz kredi türü.")
                kredi.kredi_turu = tur
            if veriler.get("odeme_hesap_turu"):
                oh = str(veriler["odeme_hesap_turu"]).strip().upper()
                if oh not in {k for k, _ in KREDI_ODEME_HESAP_TURLERI}:
                    raise ValueError(
                        "Ödeme / kullandırım hesabı mevduat veya KMH olmalıdır."
                    )
                kredi.odeme_hesap_turu = oh
            if "aciklama" in veriler:
                kredi.aciklama = (veriler.get("aciklama") or "").strip() or None
            BankaKrediService._ek_alanlari_yaz(kredi, veriler)
            BankaKrediService._gunluk(
                session, "GUNCELLEME", kredi_id=kredi.id, detay=kredi.belge_no
            )
            session.flush()
        return BankaKrediService.kredi_detay(kredi_id)

    @staticmethod
    def kredi_pasife_al(kredi_id, neden=None) -> dict:
        BankaKrediService._izin("banka_kredi_duzenleme", "finans_duzenleme")
        with get_session() as session:
            kredi = session.get(BankaKredisi, int(kredi_id))
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            kredi.aktif = False
            if (kredi.durum or "") == "KULLANDIRILDI":
                kredi.durum = "PASIF"
            BankaKrediService._gunluk(
                session, "PASIFE_ALMA", kredi_id=kredi.id, detay=(neden or "")[:500]
            )
            session.flush()
        return {"kredi_id": int(kredi_id), "durum": "PASIF"}

    @staticmethod
    def kredi_kapat(kredi_id, neden=None) -> dict:
        BankaKrediService._izin("banka_kredi_duzenleme", "finans_duzenleme")
        with get_session() as session:
            kredi = session.scalar(
                select(BankaKredisi)
                .where(BankaKredisi.id == int(kredi_id))
                .options(selectinload(BankaKredisi.taksitler))
            )
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            acik = [
                t for t in (kredi.taksitler or []) if (t.durum or "") in ACIK_TAKSIT_DURUMLARI
            ]
            if acik:
                raise ValueError(
                    f"Ödenmemiş {len(acik)} taksit var. Kredi kapatılamaz."
                )
            kredi.durum = "KAPALI"
            BankaKrediService._gunluk(
                session, "KAPATMA", kredi_id=kredi.id, detay=(neden or "")[:500]
            )
            session.flush()
        return {"kredi_id": int(kredi_id), "durum": "KAPALI"}

    @staticmethod
    def odeme_plani_kaydet(kredi_id, plan: list[dict], aciklama=None) -> dict:
        """Yeni plan sürümü oluşturur; ödenmiş taksitler korunur."""
        BankaKrediService._izin("banka_kredi_duzenleme", "finans_duzenleme")
        if not plan:
            raise ValueError("Ödeme planı boş olamaz.")
        if len(plan) > KREDI_MAX_TAKSIT:
            raise ValueError(f"Taksit sayısı 1-{KREDI_MAX_TAKSIT} arasında olmalıdır.")
        satirlar = []
        for ham in plan:
            satir = {
                "taksit_no": int(ham.get("taksit_no") or 0),
                "vade_tarihi": ham.get("vade_tarihi") or ham.get("vade"),
            }
            if satir["taksit_no"] < 1:
                raise ValueError("Taksit numarası 1'den küçük olamaz.")
            if not satir["vade_tarihi"]:
                raise ValueError("Taksit vade tarihi zorunludur.")
            for alan in BILESEN_ALANLARI:
                satir[alan] = BankaKrediService._d(
                    ham.get(alan) or 0, ALAN_ETIKET[alan], SIFIR
                )
            satir["aciklama"] = (ham.get("aciklama") or "").strip() or None
            satirlar.append(satir)
        satirlar.sort(key=lambda s: s["taksit_no"])
        if len({s["taksit_no"] for s in satirlar}) != len(satirlar):
            raise ValueError("Taksit numaraları tekrar edemez.")

        with get_session() as session:
            kredi = session.scalar(
                select(BankaKredisi)
                .where(BankaKredisi.id == int(kredi_id))
                .options(selectinload(BankaKredisi.taksitler))
            )
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            ana = BankaKrediService._d(kredi.ana_para or 0, "Ana para")
            plan_ana = sum((s["anapara"] for s in satirlar), SIFIR)
            if plan_ana != ana:
                raise ValueError("Taksit anapara toplamı kredi ana parasından farklı.")

            mevcut = {t.taksit_no: t for t in (kredi.taksitler or [])}
            plan_nolar = {s["taksit_no"] for s in satirlar}
            for no, t in mevcut.items():
                if (t.durum or "") in KAPALI_TAKSIT_DURUMLARI and no not in plan_nolar:
                    raise ValueError(
                        f"{no}. taksit ödenmiş; plandan çıkarılamaz."
                    )
            for no, t in list(mevcut.items()):
                if no not in plan_nolar:
                    session.delete(t)
            for satir in satirlar:
                t = mevcut.get(satir["taksit_no"])
                if t is not None and (t.durum or "") in KAPALI_TAKSIT_DURUMLARI:
                    continue
                if t is None:
                    t = BankaKrediTaksit(
                        kredi_id=kredi.id,
                        taksit_no=satir["taksit_no"],
                        vade_tarihi=satir["vade_tarihi"],
                        durum="BEKLIYOR",
                    )
                    session.add(t)
                t.vade_tarihi = satir["vade_tarihi"]
                BankaKrediService._taksit_bilesen_yaz(t, satir)
                odenen = BankaKrediService._d(t.odenen_tutar or 0, "Ödenen tutar")
                t.durum = BankaKrediService.taksit_durum_hesapla(
                    t.vade_tarihi,
                    odenen,
                    BankaKrediService.taksit_bilesenleri(t)["toplam"],
                    None if odenen == 0 else t.durum,
                )

            kredi.taksit_sayisi = len(satirlar)
            kredi.toplam_faiz = sum((s["faiz"] for s in satirlar), SIFIR)
            kredi.toplam_masraf = sum(
                (s[a] for s in satirlar for a in GIDER_ALANLARI if a != "faiz"), SIFIR
            )
            kredi.ilk_taksit_tarihi = satirlar[0]["vade_tarihi"]
            kredi.bitis_tarihi = max(s["vade_tarihi"] for s in satirlar)

            son = session.scalars(
                select(BankaKrediPlanVersiyon)
                .where(BankaKrediPlanVersiyon.kredi_id == kredi.id)
                .order_by(BankaKrediPlanVersiyon.versiyon.desc())
            ).first()
            yeni_versiyon = (son.versiyon + 1) if son else 1
            for eski in session.scalars(
                select(BankaKrediPlanVersiyon).where(
                    BankaKrediPlanVersiyon.kredi_id == kredi.id,
                    BankaKrediPlanVersiyon.aktif.is_(True),
                )
            ).all():
                eski.aktif = False
            session.add(
                BankaKrediPlanVersiyon(
                    kredi_id=kredi.id,
                    versiyon=yeni_versiyon,
                    plan_json=json.dumps(
                        [
                            {
                                **{
                                    a: str(s[a])
                                    for a in BILESEN_ALANLARI
                                },
                                "taksit_no": s["taksit_no"],
                                "vade_tarihi": s["vade_tarihi"].isoformat()
                                if hasattr(s["vade_tarihi"], "isoformat")
                                else str(s["vade_tarihi"]),
                            }
                            for s in satirlar
                        ],
                        ensure_ascii=False,
                    ),
                    aktif=True,
                    aciklama=(aciklama or "").strip() or None,
                    created_at=datetime.now(),
                )
            )
            kredi.plan_versiyon = yeni_versiyon
            BankaKrediService._gunluk(
                session,
                "PLAN_KAYDET",
                kredi_id=kredi.id,
                detay=f"versiyon={yeni_versiyon} / satır={len(satirlar)}",
            )
            session.flush()
        return {
            "kredi_id": int(kredi_id),
            "versiyon": yeni_versiyon,
            "taksit_sayisi": len(satirlar),
        }

    @staticmethod
    def plan_versiyonlari(kredi_id) -> list[dict]:
        with get_session() as session:
            rows = session.scalars(
                select(BankaKrediPlanVersiyon)
                .where(BankaKrediPlanVersiyon.kredi_id == int(kredi_id))
                .order_by(BankaKrediPlanVersiyon.versiyon.desc())
            ).all()
            return [
                {
                    "id": v.id,
                    "versiyon": v.versiyon,
                    "aktif": bool(v.aktif),
                    "aciklama": v.aciklama or "",
                    "created_at": v.created_at,
                    "plan": json.loads(v.plan_json or "[]"),
                }
                for v in rows
            ]

    # --- Excel / CSV plan aktarımı ---

    @staticmethod
    def _tarih_coz(deger):
        if deger is None or (isinstance(deger, str) and not deger.strip()):
            raise ValueError("Taksit vade tarihi zorunludur.")
        if isinstance(deger, datetime):
            return deger.date()
        if isinstance(deger, date):
            return deger
        metin = str(deger).strip().split(" ")[0]
        for kalip in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(metin, kalip).date()
            except ValueError:
                continue
        raise ValueError(f"Geçersiz vade tarihi: {deger}")

    @staticmethod
    def odeme_plani_excel_aktar(path) -> list[dict]:
        """Excel / CSV ödeme planını okur (openpyxl yoksa csv ile)."""
        basliklar = {
            "taksit_no": "taksit_no",
            "taksit": "taksit_no",
            "no": "taksit_no",
            "vade": "vade_tarihi",
            "vade_tarihi": "vade_tarihi",
            "tarih": "vade_tarihi",
            "anapara": "anapara",
            "ana_para": "anapara",
            "ana para": "anapara",
            "faiz": "faiz",
            "bsmv": "bsmv",
            "kkdf": "kkdf",
            "komisyon": "komisyon",
            "sigorta": "sigorta",
            "dosya": "dosya_masrafi",
            "dosya_masrafi": "dosya_masrafi",
            "diger": "diger_masraflar",
            "diger_masraflar": "diger_masraflar",
            "gecikme": "gecikme_faizi",
            "gecikme_faizi": "gecikme_faizi",
            "aciklama": "aciklama",
        }

        def normalize(metin):
            ham = str(metin or "").strip().lower()
            for eski, yeni in (
                ("ı", "i"), ("İ", "i"), ("ş", "s"), ("ğ", "g"),
                ("ü", "u"), ("ö", "o"), ("ç", "c"),
            ):
                ham = ham.replace(eski, yeni)
            return ham.replace(".", "").replace("  ", " ").strip()

        satirlar = []
        yol = str(path)
        if yol.lower().endswith((".xlsx", ".xlsm")):
            try:
                from openpyxl import load_workbook
            except ImportError as hata:
                raise ValueError(
                    "Excel okumak için openpyxl kurulu değil. CSV dosyası kullanın."
                ) from hata
            wb = load_workbook(yol, data_only=True)
            ws = wb.active
            ham_satirlar = [list(r) for r in ws.iter_rows(values_only=True)]
            wb.close()
        else:
            import csv

            with open(yol, "r", encoding="utf-8-sig", newline="") as dosya:
                ornek = dosya.read(4096)
                dosya.seek(0)
                try:
                    ayirici = csv.Sniffer().sniff(ornek, delimiters=";,\t").delimiter
                except csv.Error:
                    ayirici = ";"
                ham_satirlar = [list(r) for r in csv.reader(dosya, delimiter=ayirici)]

        if not ham_satirlar:
            raise ValueError("Ödeme planı dosyası boş.")
        baslik_satiri = ham_satirlar[0]
        eslesme = {}
        for i, hucre in enumerate(baslik_satiri):
            alan = basliklar.get(normalize(hucre))
            if alan:
                eslesme[i] = alan
        if "vade_tarihi" not in eslesme.values() or "anapara" not in eslesme.values():
            raise ValueError(
                "Başlık satırında en az 'vade' ve 'anapara' kolonları bulunmalıdır."
            )

        sira = 0
        for ham in ham_satirlar[1:]:
            if not any(str(h or "").strip() for h in ham):
                continue
            sira += 1
            satir = {alan: SIFIR for alan in BILESEN_ALANLARI}
            satir["taksit_no"] = sira
            satir["vade_tarihi"] = None
            satir["aciklama"] = None
            for i, alan in eslesme.items():
                deger = ham[i] if i < len(ham) else None
                if alan == "taksit_no":
                    try:
                        satir["taksit_no"] = int(str(deger).strip() or sira)
                    except (ValueError, TypeError):
                        satir["taksit_no"] = sira
                elif alan == "vade_tarihi":
                    satir["vade_tarihi"] = BankaKrediService._tarih_coz(deger)
                elif alan == "aciklama":
                    satir["aciklama"] = (str(deger).strip() or None) if deger else None
                else:
                    satir[alan] = BankaKrediService._d(
                        deger or 0, ALAN_ETIKET[alan], SIFIR
                    )
            if satir["vade_tarihi"] is None:
                raise ValueError(f"{satir['taksit_no']}. satırda vade tarihi yok.")
            satir["toplam"] = BankaKrediService.taksit_toplam(satir)
            satirlar.append(satir)
        if not satirlar:
            raise ValueError("Ödeme planı dosyasında satır bulunamadı.")
        satirlar.sort(key=lambda s: s["taksit_no"])
        return satirlar

    # --- ÖDEME: tek banka çıkışı, çok bileşenli dağıtım ---

    @staticmethod
    def taksit_ode(
        taksit_id=None,
        odeme_tarihi=None,
        banka_hesap_id=None,
        banka_dekont_no=None,
        aciklama=None,
        bilesenler=None,
        dagitim=None,
        toplam_odeme=None,
        odeme_turu="TAKSIT",
        transaction_id=None,
        kredi_id=None,
    ) -> dict:
        BankaKrediService._izin("banka_kredi_odeme", "finans_duzenleme")
        tur = (odeme_turu or "TAKSIT").strip().upper()
        if tur not in ODEME_TURLERI:
            raise ValueError("Geçersiz ödeme türü.")
        odeme_tarihi = odeme_tarihi or date.today()
        if odeme_tarihi > date.today():
            raise ValueError("Ödeme tarihi gelecek olamaz.")
        if not taksit_id and not kredi_id:
            raise ValueError("Taksit veya kredi seçilmelidir.")
        dekont = (banka_dekont_no or "").strip() or None
        tx = (transaction_id or "").strip() or None

        with get_session() as session:
            # 4) Idempotency — aynı transaction_id ile ikinci kez yazılmaz
            if tx:
                mevcut = session.scalar(
                    select(BankaKrediOdeme).where(BankaKrediOdeme.transaction_id == tx)
                )
                if mevcut is not None:
                    return BankaKrediService._odeme_sonuc(session, mevcut)

            taksit = None
            if taksit_id:
                taksit = session.scalar(
                    select(BankaKrediTaksit)
                    .where(BankaKrediTaksit.id == int(taksit_id))
                    .options(selectinload(BankaKrediTaksit.kredi))
                )
                if not taksit:
                    raise ValueError("Taksit bulunamadı.")
                kredi = taksit.kredi
            else:
                kredi = session.scalar(
                    select(BankaKredisi)
                    .where(BankaKredisi.id == int(kredi_id))
                    .options(selectinload(BankaKredisi.taksitler))
                )
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            if (kredi.durum or "") != "KULLANDIRILDI":
                raise ValueError("Kredi aktif değil.")

            # 2) Kilit — ödenmiş taksite tekrar tam ödeme yapılamaz
            if taksit is not None:
                mevcut_durum = (taksit.durum or "").upper()
                if mevcut_durum == "ODENDI":
                    raise ValueError("Bu taksit zaten ödenmiş.")
                if mevcut_durum in ("IPTAL", "YAPILANDIRILDI", "ERKEN_KAPATILDI"):
                    raise ValueError("Bu taksit ödemeye kapalı.")

            # 3) Dekont no banka hesabı bazında tek olmalı
            kart, kredi_hesap, odeme_hesap = BankaKrediService._kredi_hesaplari(
                session, kredi, banka_hesap_id
            )
            if dekont and BankaKrediService._dekont_mukerrer(
                session, odeme_hesap.id, dekont
            ):
                raise ValueError(
                    f"Bu banka hesabında {dekont} dekont numarası zaten kullanılmış."
                )

            plan = (
                BankaKrediService.taksit_bilesenleri(taksit)
                if taksit is not None
                else {alan: SIFIR for alan in BILESEN_ALANLARI}
            )
            odenmis = BankaKrediService._odenen_bilesenler(
                session, taksit.id if taksit is not None else None
            )
            kalan = {}
            for alan in BILESEN_ALANLARI:
                fark = plan.get(alan, SIFIR) - odenmis[alan]
                kalan[alan] = fark if fark > 0 else SIFIR

            # 5) Bileşen / toplam doğrulaması
            if bilesenler:
                hedef = {
                    alan: BankaKrediService._d(
                        bilesenler.get(alan) or 0, ALAN_ETIKET[alan], SIFIR
                    )
                    for alan in BILESEN_ALANLARI
                }
            elif toplam_odeme is not None:
                tutar = BankaKrediService._d(toplam_odeme, "Ödeme tutarı", KURUS)
                hedef = BankaKrediService._dagit(kalan, tutar, dagitim or "ONCE_MASRAF")
            elif taksit is not None:
                hedef = dict(kalan)
            else:
                raise ValueError("Ödeme bileşenleri veya toplam tutar girilmelidir.")

            toplam = BankaKrediService.taksit_toplam(hedef)
            if toplam <= 0:
                raise ValueError("Ödenecek tutar sıfır olamaz.")
            if bilesenler and toplam_odeme is not None:
                beklenen = BankaKrediService._d(toplam_odeme, "Ödeme tutarı", SIFIR)
                if beklenen != toplam:
                    raise ValueError("Bileşen toplamı ödeme tutarından farklı.")
            if taksit is not None and tur in ("TAKSIT", "KISMI"):
                for alan in BILESEN_ALANLARI:
                    if alan == "gecikme_faizi":
                        continue
                    if hedef[alan] > kalan[alan]:
                        raise ValueError(
                            f"{ALAN_ETIKET[alan]} ödemesi kalan tutarı aşıyor."
                        )

            anapara = hedef["anapara"]
            gider_toplam = toplam - anapara

            # 6) Bakiye kontrolü yalnızca TOPLAM için
            odeme_hesap = session.scalar(
                select(FinansHesabi)
                .where(FinansHesabi.id == odeme_hesap.id)
                .options(
                    selectinload(FinansHesabi.hareketler),
                    selectinload(FinansHesabi.banka_karti),
                )
            )
            FinansService.cikis_kontrol(
                odeme_hesap,
                toplam,
                kart=odeme_hesap.banka_karti or kart,
                hesap_etiket="Ödeme hesabı",
            )

            odeme_belge = FinansService._finans_belge_no(session, "KTO")
            acik = (aciklama or "").strip() or (
                f"{kredi.kredi_adi} — {taksit.taksit_no}/{kredi.taksit_sayisi}. taksit"
                if taksit is not None
                else f"{kredi.kredi_adi} — {tur.replace('_', ' ').title()} ödemesi"
            )
            if dekont:
                acik = f"{acik} | Dekont: {dekont}"

            # 7) Ödeme kaydı
            odeme = BankaKrediOdeme(
                transaction_id=tx or uuid.uuid4().hex,
                kredi_id=kredi.id,
                taksit_id=taksit.id if taksit is not None else None,
                odeme_turu=tur,
                odeme_tarihi=odeme_tarihi,
                odeme_belge_no=odeme_belge,
                banka_hesap_id=odeme_hesap.id,
                banka_dekont_no=dekont,
                toplam=toplam,
                para_birimi=(kredi.para_birimi or "TRY"),
                kur=kredi.kur,
                kur_tarihi=kredi.kur_tarihi,
                aciklama=acik[:500],
                durum="AKTIF",
                created_by=BankaKrediService._kullanici(),
                created_at=datetime.now(),
            )
            for alan in BILESEN_ALANLARI:
                setattr(odeme, alan, hedef[alan])
            session.add(odeme)
            session.flush()

            # 8) Ödeme hesabından TEK çıkış — toplam tutar
            banka_hareket = FinansHareketi(
                hesap_id=odeme_hesap.id,
                tarih=odeme_tarihi,
                hareket_turu=HAREKET_TAKSIT_ODEME,
                belge_no=odeme_belge,
                tutar=toplam,
                aciklama=acik[:500],
            )
            session.add(banka_hareket)
            session.flush()
            odeme.banka_hareket_id = banka_hareket.id

            # 9) Anapara krediler hesabındaki borcu azaltır
            if anapara > 0:
                session.add(
                    FinansHareketi(
                        hesap_id=kredi_hesap.id,
                        tarih=odeme_tarihi,
                        hareket_turu=HAREKET_ANAPARA_ODEME,
                        belge_no=odeme_belge,
                        tutar=anapara,
                        aciklama=f"Anapara | {acik}"[:500],
                    )
                )

            # 10) Gider bileşenleri — ek banka hareketi YOK, sadece gider fişi
            gider_belge = None
            if gider_toplam > 0:
                gider_belge = BankaKrediService._gider_belge_no(session)
                for alan in GIDER_ALANLARI:
                    tutar_bilesen = hedef[alan]
                    if tutar_bilesen <= 0:
                        continue
                    gider_turu, gider_acik, ek = GIDER_TURU_ESLEME[alan]
                    session.add(
                        GiderFisi(
                            belge_no=f"{gider_belge}-{ek}",
                            tarih=odeme_tarihi,
                            gider_turu=gider_turu,
                            tutar=tutar_bilesen,
                            finans_hesap_id=odeme_hesap.id,
                            bagli_belge_no=odeme_belge,
                            aciklama=f"{gider_acik} | {acik}"[:500],
                            durum="AÇIK",
                        )
                    )
                odeme.gider_belge_no = gider_belge

            # 11) Muhasebe fişi — eşleştirme eksikse ödeme durmaz
            try:
                fis_id = BankaKrediService._muhasebe_fis_olustur(session, odeme)
                if fis_id:
                    odeme.muhasebe_fis_id = fis_id
            except Exception as hata:  # pragma: no cover - entegrasyon opsiyonel
                print(f"[Kredi ödeme] Muhasebe fişi oluşturulamadı: {hata}")

            # 12) Taksit güncelle
            if taksit is not None:
                yeni_odenen = {
                    alan: odenmis[alan] + hedef[alan] for alan in BILESEN_ALANLARI
                }
                taksit.odenen_anapara = yeni_odenen["anapara"]
                taksit.odenen_faiz = yeni_odenen["faiz"]
                taksit.odenen_masraf = sum(
                    (yeni_odenen[a] for a in GIDER_ALANLARI if a != "faiz"), SIFIR
                )
                taksit.odenen_tutar = BankaKrediService.taksit_toplam(yeni_odenen)
                taksit.odeme_tarihi = odeme_tarihi
                taksit.odeme_belge_no = odeme_belge
                if gider_belge:
                    taksit.gider_belge_no = gider_belge
                taksit.transaction_id = odeme.transaction_id
                taksit.banka_dekont_no = dekont or taksit.banka_dekont_no
                taksit.odeme_hesap_id = odeme_hesap.id
                taksit.durum = BankaKrediService.taksit_durum_hesapla(
                    taksit.vade_tarihi,
                    taksit.odenen_tutar,
                    plan["toplam"],
                    None,
                )

            # Erken kapama — kalan taksitler kapatılır
            if tur == "ERKEN_KAPAMA":
                for t in session.scalars(
                    select(BankaKrediTaksit).where(
                        BankaKrediTaksit.kredi_id == kredi.id,
                        BankaKrediTaksit.durum.in_(ACIK_TAKSIT_DURUMLARI),
                    )
                ).all():
                    t.durum = "ERKEN_KAPATILDI"
                    t.odeme_tarihi = odeme_tarihi
                    t.odeme_belge_no = odeme_belge
                    if gider_belge:
                        t.gider_belge_no = gider_belge

            session.flush()

            # 13) Kalan taksit yoksa kredi kapanır
            kalan_acik = session.scalars(
                select(BankaKrediTaksit).where(
                    BankaKrediTaksit.kredi_id == kredi.id,
                    BankaKrediTaksit.durum.in_(ACIK_TAKSIT_DURUMLARI),
                )
            ).first()
            if kalan_acik is None:
                kredi.durum = "KAPALI"

            # 14) Denetim kaydı
            BankaKrediService._gunluk(
                session,
                f"ODEME_{tur}",
                kredi_id=kredi.id,
                taksit_id=taksit.id if taksit is not None else None,
                odeme_id=odeme.id,
                transaction_id=odeme.transaction_id,
                detay=f"{odeme_belge} / toplam={toplam} / anapara={anapara}",
            )
            session.flush()
            return BankaKrediService._odeme_sonuc(session, odeme)

    @staticmethod
    def _odeme_sonuc(session, odeme) -> dict:
        kredi = session.get(BankaKredisi, odeme.kredi_id)
        taksit = (
            session.get(BankaKrediTaksit, odeme.taksit_id) if odeme.taksit_id else None
        )
        veri = BankaKrediService._odeme_sozluk(session, odeme)
        veri.update({
            "kredi_adi": kredi.kredi_adi if kredi else "",
            "taksit_no": taksit.taksit_no if taksit else None,
            "taksit_durum": taksit.durum if taksit else None,
            "kredi_kapandi": bool(kredi and kredi.durum == "KAPALI"),
            "kalan_anapara": BankaKrediService._kalan_anapara(session, kredi)
            if kredi
            else SIFIR,
        })
        return veri

    @staticmethod
    def kisimi_ode(
        taksit_id,
        tutar,
        odeme_tarihi=None,
        banka_hesap_id=None,
        banka_dekont_no=None,
        aciklama=None,
        dagitim="ONCE_MASRAF",
        bilesenler=None,
        transaction_id=None,
    ) -> dict:
        """Kısmi taksit ödemesi — MANUEL / ONCE_MASRAF / ORANSAL dağıtım."""
        yontem = (dagitim or "ONCE_MASRAF").strip().upper()
        if yontem == "MANUEL" and not bilesenler:
            raise ValueError("Manuel dağıtımda ödeme bileşenleri girilmelidir.")
        return BankaKrediService.taksit_ode(
            taksit_id=taksit_id,
            odeme_tarihi=odeme_tarihi,
            banka_hesap_id=banka_hesap_id,
            banka_dekont_no=banka_dekont_no,
            aciklama=aciklama,
            bilesenler=bilesenler,
            dagitim=yontem,
            toplam_odeme=tutar,
            odeme_turu="KISMI",
            transaction_id=transaction_id,
        )

    @staticmethod
    def ara_ode(
        kredi_id,
        anapara,
        odeme_tarihi=None,
        banka_hesap_id=None,
        banka_dekont_no=None,
        aciklama=None,
        faiz=0,
        masraflar=None,
        transaction_id=None,
    ) -> dict:
        """Ara (plan dışı) anapara ödemesi — kalan anaparayı azaltır."""
        bilesen = {alan: SIFIR for alan in BILESEN_ALANLARI}
        bilesen["anapara"] = BankaKrediService._d(anapara or 0, "Anapara", SIFIR)
        bilesen["faiz"] = BankaKrediService._d(faiz or 0, "Faiz", SIFIR)
        for alan, deger in (masraflar or {}).items():
            if alan not in BILESEN_ALANLARI:
                raise ValueError(f"Geçersiz masraf bileşeni: {alan}")
            bilesen[alan] = BankaKrediService._d(deger or 0, ALAN_ETIKET[alan], SIFIR)
        if BankaKrediService.taksit_toplam(bilesen) <= 0:
            raise ValueError("Ödenecek tutar sıfır olamaz.")
        return BankaKrediService.taksit_ode(
            kredi_id=kredi_id,
            odeme_tarihi=odeme_tarihi,
            banka_hesap_id=banka_hesap_id,
            banka_dekont_no=banka_dekont_no,
            aciklama=aciklama or "Ara ödeme (anapara)",
            bilesenler=bilesen,
            odeme_turu="ARA",
            transaction_id=transaction_id,
        )

    @staticmethod
    def erken_kapama_ozeti(kredi_id) -> dict:
        with get_session() as session:
            kredi = session.scalar(
                select(BankaKredisi)
                .where(BankaKredisi.id == int(kredi_id))
                .options(selectinload(BankaKredisi.taksitler))
            )
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            kalan_anapara = SIFIR
            kalan_faiz = SIFIR
            kalan_masraf = SIFIR
            adet = 0
            for t in kredi.taksitler or []:
                if (t.durum or "") not in ACIK_TAKSIT_DURUMLARI:
                    continue
                adet += 1
                plan = BankaKrediService.taksit_bilesenleri(t)
                odenmis = BankaKrediService._odenen_bilesenler(session, t.id)
                for alan in BILESEN_ALANLARI:
                    fark = plan[alan] - odenmis[alan]
                    if fark <= 0:
                        continue
                    if alan == "anapara":
                        kalan_anapara += fark
                    elif alan == "faiz":
                        kalan_faiz += fark
                    else:
                        kalan_masraf += fark
            return {
                "kredi_id": int(kredi_id),
                "kredi_adi": kredi.kredi_adi,
                "kalan_taksit_sayisi": adet,
                "kalan_anapara": kalan_anapara,
                "kalan_faiz": kalan_faiz,
                "kalan_masraf": kalan_masraf,
            }

    @staticmethod
    def erken_kapat(
        kredi_id,
        odeme_tarihi=None,
        banka_hesap_id=None,
        banka_dekont_no=None,
        aciklama=None,
        faiz=None,
        masraflar=None,
        transaction_id=None,
    ) -> dict:
        """Kalan anapara + verilen faiz/masraf ile krediyi tek çıkışta kapatır."""
        ozet = BankaKrediService.erken_kapama_ozeti(kredi_id)
        if ozet["kalan_anapara"] <= 0 and ozet["kalan_taksit_sayisi"] == 0:
            raise ValueError("Bu kredide kapatılacak taksit yok.")
        bilesen = {alan: SIFIR for alan in BILESEN_ALANLARI}
        bilesen["anapara"] = ozet["kalan_anapara"]
        bilesen["faiz"] = (
            BankaKrediService._d(faiz, "Faiz", SIFIR) if faiz is not None else SIFIR
        )
        for alan, deger in (masraflar or {}).items():
            if alan not in BILESEN_ALANLARI:
                raise ValueError(f"Geçersiz masraf bileşeni: {alan}")
            bilesen[alan] = BankaKrediService._d(deger or 0, ALAN_ETIKET[alan], SIFIR)
        if BankaKrediService.taksit_toplam(bilesen) <= 0:
            raise ValueError("Ödenecek tutar sıfır olamaz.")
        return BankaKrediService.taksit_ode(
            kredi_id=kredi_id,
            odeme_tarihi=odeme_tarihi,
            banka_hesap_id=banka_hesap_id,
            banka_dekont_no=banka_dekont_no,
            aciklama=aciklama or "Kredi erken kapama",
            bilesenler=bilesen,
            odeme_turu="ERKEN_KAPAMA",
            transaction_id=transaction_id,
        )

    # --- Ödeme iptali ---

    @staticmethod
    def odeme_iptal(odeme_id=None, belge_no=None, neden=None) -> dict:
        BankaKrediService._iptal_izni()
        gerekce = (neden or "").strip()
        if not gerekce:
            raise ValueError("İptal nedeni zorunludur.")
        if not odeme_id and not belge_no:
            raise ValueError("İptal edilecek ödeme seçilmelidir.")

        with get_session() as session:
            q = select(BankaKrediOdeme)
            if odeme_id:
                q = q.where(BankaKrediOdeme.id == int(odeme_id))
            else:
                q = q.where(BankaKrediOdeme.odeme_belge_no == str(belge_no).strip())
            odeme = session.scalar(q)
            if not odeme:
                raise ValueError("Ödeme kaydı bulunamadı.")
            if (odeme.durum or "") == "IPTAL":
                raise ValueError("Bu ödeme zaten iptal edilmiş.")

            kredi = session.scalar(
                select(BankaKredisi)
                .where(BankaKredisi.id == odeme.kredi_id)
                .options(selectinload(BankaKredisi.taksitler))
            )
            if not kredi:
                raise ValueError("Kredi bulunamadı.")
            _kart, kredi_hesap, _odeme_hesap = BankaKrediService._kredi_hesaplari(
                session, kredi
            )

            toplam = BankaKrediService._d(odeme.toplam or 0, "Toplam")
            anapara = BankaKrediService._d(odeme.anapara or 0, "Anapara")
            iptal_belge = FinansService._finans_belge_no(session, "KTI")
            iptal_acik = f"İPTAL {odeme.odeme_belge_no} | {gerekce}"[:500]
            iptal_tarihi = date.today()

            # Banka çıkışını ters kayıtla geri al
            session.add(
                FinansHareketi(
                    hesap_id=odeme.banka_hesap_id,
                    tarih=iptal_tarihi,
                    hareket_turu=HAREKET_ODEME_IPTAL,
                    belge_no=iptal_belge,
                    tutar=toplam,
                    aciklama=iptal_acik,
                )
            )
            # Anaparayı krediler hesabına geri yükle
            if anapara > 0:
                session.add(
                    FinansHareketi(
                        hesap_id=kredi_hesap.id,
                        tarih=iptal_tarihi,
                        hareket_turu=HAREKET_ANAPARA_IPTAL,
                        belge_no=iptal_belge,
                        tutar=anapara,
                        aciklama=iptal_acik,
                    )
                )
            # Gider fişleri iptal
            for fis in session.scalars(
                select(GiderFisi).where(
                    GiderFisi.bagli_belge_no == odeme.odeme_belge_no
                )
            ).all():
                fis.durum = "IPTAL"
                fis.aciklama = f"{(fis.aciklama or '')} | İPTAL: {gerekce}"[:500]
            # Muhasebe fişi iptal (varsa)
            if odeme.muhasebe_fis_id:
                try:
                    from database.muhasebe_service import MuhasebeFisService

                    MuhasebeFisService.iptal(
                        int(odeme.muhasebe_fis_id),
                        f"Kredi ödemesi iptal: {odeme.odeme_belge_no}",
                        otomatik=True,
                    )
                except Exception as hata:  # pragma: no cover
                    print(f"[Kredi ödeme] Muhasebe fişi iptal edilemedi: {hata}")

            # Ters ödeme kaydı
            ters = BankaKrediOdeme(
                transaction_id=uuid.uuid4().hex,
                kredi_id=odeme.kredi_id,
                taksit_id=odeme.taksit_id,
                odeme_turu="IPTAL",
                odeme_tarihi=iptal_tarihi,
                odeme_belge_no=iptal_belge,
                banka_hesap_id=odeme.banka_hesap_id,
                banka_dekont_no=None,
                toplam=toplam,
                para_birimi=odeme.para_birimi or "TRY",
                aciklama=iptal_acik,
                durum="IPTAL",
                iptal_nedeni=gerekce[:500],
                iptal_odeme_id=odeme.id,
                created_by=BankaKrediService._kullanici(),
                created_at=datetime.now(),
            )
            for alan in BILESEN_ALANLARI:
                setattr(ters, alan, BankaKrediService._d(
                    getattr(odeme, alan, None) or 0, ALAN_ETIKET[alan]
                ))
            session.add(ters)
            session.flush()

            odeme.durum = "IPTAL"
            odeme.iptal_nedeni = gerekce[:500]
            odeme.iptal_odeme_id = ters.id

            # Taksit durumunu kalan aktif ödemelere göre geri yükle
            if odeme.taksit_id:
                taksit = session.get(BankaKrediTaksit, odeme.taksit_id)
                if taksit is not None:
                    odenmis = BankaKrediService._odenen_bilesenler(session, taksit.id)
                    plan = BankaKrediService.taksit_bilesenleri(taksit)
                    taksit.odenen_anapara = odenmis["anapara"]
                    taksit.odenen_faiz = odenmis["faiz"]
                    taksit.odenen_masraf = sum(
                        (odenmis[a] for a in GIDER_ALANLARI if a != "faiz"), SIFIR
                    )
                    taksit.odenen_tutar = BankaKrediService.taksit_toplam(odenmis)
                    if taksit.odenen_tutar == 0:
                        taksit.odeme_tarihi = None
                        taksit.odeme_belge_no = None
                        taksit.gider_belge_no = None
                        taksit.transaction_id = None
                        taksit.banka_dekont_no = None
                        taksit.odeme_hesap_id = None
                    taksit.durum = BankaKrediService.taksit_durum_hesapla(
                        taksit.vade_tarihi, taksit.odenen_tutar, plan["toplam"], None
                    )

            # Erken kapama iptali — kapatılan taksitler yeniden açılır
            if (odeme.odeme_turu or "") == "ERKEN_KAPAMA":
                for t in session.scalars(
                    select(BankaKrediTaksit).where(
                        BankaKrediTaksit.kredi_id == kredi.id,
                        BankaKrediTaksit.durum == "ERKEN_KAPATILDI",
                        BankaKrediTaksit.odeme_belge_no == odeme.odeme_belge_no,
                    )
                ).all():
                    odenmis = BankaKrediService._odenen_bilesenler(session, t.id)
                    plan = BankaKrediService.taksit_bilesenleri(t)
                    t.odeme_tarihi = None
                    t.odeme_belge_no = None
                    t.gider_belge_no = None
                    t.durum = BankaKrediService.taksit_durum_hesapla(
                        t.vade_tarihi,
                        BankaKrediService.taksit_toplam(odenmis),
                        plan["toplam"],
                        None,
                    )

            if (kredi.durum or "") == "KAPALI":
                kredi.durum = "KULLANDIRILDI"

            BankaKrediService._gunluk(
                session,
                "ODEME_IPTAL",
                kredi_id=kredi.id,
                taksit_id=odeme.taksit_id,
                odeme_id=odeme.id,
                transaction_id=odeme.transaction_id,
                detay=f"{odeme.odeme_belge_no} → {iptal_belge} | {gerekce}",
            )
            session.flush()
            return {
                "odeme_id": odeme.id,
                "odeme_belge_no": odeme.odeme_belge_no,
                "iptal_belge_no": iptal_belge,
                "iptal_odeme_id": ters.id,
                "toplam": toplam,
                "anapara": anapara,
                "neden": gerekce,
                "durum": "IPTAL",
            }

    # --- Raporlar ---

    @staticmethod
    def raporlar(rapor="bakiye_ozeti", banka_karti_id=None, kredi_id=None, gun=30) -> list[dict]:
        ad = (rapor or "bakiye_ozeti").strip().lower()
        if ad == "bakiye_ozeti":
            return BankaKrediService.rapor_bakiye_ozeti(banka_karti_id)
        if ad == "bankalara_gore":
            return BankaKrediService.rapor_bankalara_gore()
        if ad == "bekleyen":
            return BankaKrediService.bekleyen_taksitler(banka_karti_id, kredi_id)
        if ad == "odenen":
            return BankaKrediService.odenen_taksitler(banka_karti_id, kredi_id)
        if ad == "vadesi_gecen":
            return BankaKrediService.rapor_vadesi_gecen(banka_karti_id)
        if ad == "takvim":
            return BankaKrediService.rapor_takvim(gun, banka_karti_id)
        if ad == "maliyet_dagilimi":
            return BankaKrediService.rapor_maliyet_dagilimi(kredi_id)
        raise ValueError(f"Geçersiz rapor adı: {rapor}")

    @staticmethod
    def rapor_bakiye_ozeti(banka_karti_id=None) -> list[dict]:
        sonuc = []
        for k in BankaKrediService.kredi_listesi(banka_karti_id, aktif_only=False):
            sonuc.append({
                "kredi_id": k["id"],
                "belge_no": k["belge_no"],
                "banka_adi": k["banka_adi"],
                "kredi_adi": k["kredi_adi"],
                "ana_para": k["ana_para"],
                "kalan_anapara": k["kalan_anapara"],
                "kalan_borc": k["kalan_borc"],
                "aylik_taksit_tutari": k.get("aylik_taksit_tutari") or SIFIR,
                "gelecek_taksit_tarihi": k.get("gelecek_taksit_tarihi"),
                "odeme_gunu": (
                    k["gelecek_taksit_tarihi"].day
                    if k.get("gelecek_taksit_tarihi")
                    else None
                ),
                "kalan_taksit_sayisi": k["kalan_taksit_sayisi"],
                "durum": k["durum"],
            })
        return sonuc

    @staticmethod
    def rapor_bankalara_gore() -> list[dict]:
        toplamlar: dict[str, dict] = {}
        for k in BankaKrediService.kredi_listesi(aktif_only=False):
            ad = k["banka_adi"] or "—"
            kayit = toplamlar.setdefault(
                ad,
                {
                    "banka_adi": ad,
                    "kredi_sayisi": 0,
                    "ana_para": SIFIR,
                    "kalan_anapara": SIFIR,
                    "kalan_borc": SIFIR,
                },
            )
            kayit["kredi_sayisi"] += 1
            kayit["ana_para"] += k["ana_para"]
            kayit["kalan_anapara"] += k["kalan_anapara"]
            kayit["kalan_borc"] += k["kalan_borc"]
        return sorted(toplamlar.values(), key=lambda s: s["banka_adi"])

    @staticmethod
    def rapor_vadesi_gecen(banka_karti_id=None) -> list[dict]:
        bugun = date.today()
        return [
            s
            for s in BankaKrediService.bekleyen_taksitler(banka_karti_id)
            if s["vade_tarihi"] < bugun
        ]

    @staticmethod
    def rapor_takvim(gun=30, banka_karti_id=None) -> list[dict]:
        try:
            g = int(gun or 30)
        except (TypeError, ValueError) as hata:
            raise ValueError("Takvim gün sayısı 7, 30 veya 90 olmalıdır.") from hata
        if g not in (7, 30, 90):
            raise ValueError("Takvim gün sayısı 7, 30 veya 90 olmalıdır.")
        satirlar = BankaKrediService.bekleyen_taksitler(banka_karti_id, gun=g)
        takvim: dict[date, dict] = {}
        for s in satirlar:
            kayit = takvim.setdefault(
                s["vade_tarihi"],
                {
                    "vade_tarihi": s["vade_tarihi"],
                    "taksit_sayisi": 0,
                    "anapara": SIFIR,
                    "faiz": SIFIR,
                    "masraf": SIFIR,
                    "toplam": SIFIR,
                },
            )
            kayit["taksit_sayisi"] += 1
            kayit["anapara"] += s["anapara"]
            kayit["faiz"] += s["faiz"]
            kayit["masraf"] += sum(
                (s[a] for a in GIDER_ALANLARI if a != "faiz"), SIFIR
            )
            kayit["toplam"] += s["kalan"]
        return sorted(takvim.values(), key=lambda s: s["vade_tarihi"])

    @staticmethod
    def rapor_maliyet_dagilimi(kredi_id=None) -> list[dict]:
        with get_session() as session:
            q = select(BankaKrediOdeme).where(BankaKrediOdeme.durum == "AKTIF")
            if kredi_id:
                q = q.where(BankaKrediOdeme.kredi_id == int(kredi_id))
            odemeler = session.scalars(q).all()
            toplamlar = {alan: SIFIR for alan in BILESEN_ALANLARI}
            for o in odemeler:
                for alan in BILESEN_ALANLARI:
                    toplamlar[alan] += BankaKrediService._d(
                        getattr(o, alan, None) or 0, ALAN_ETIKET[alan]
                    )
            genel = sum(toplamlar.values(), SIFIR)
            sonuc = []
            for alan in BILESEN_ALANLARI:
                tutar = toplamlar[alan]
                oran = SIFIR
                if genel > 0:
                    oran = (tutar * Decimal("100") / genel).quantize(
                        KURUS, rounding=ROUND_HALF_UP
                    )
                sonuc.append({
                    "bilesen": alan,
                    "etiket": ALAN_ETIKET[alan],
                    "tutar": tutar,
                    "oran": oran,
                })
            return sonuc
