"""Fiyatlı Stok Ekstresi — okuma amaçlı görünüm (hareket kayıtlarını değiştirmez).

Tek doğruluk kaynağı: yürüyen kalan miktar + satır bazlı FIFO kalan değeri
(açık lot katmanlarının toplam maliyeti). Hareket kayıtlarını değiştirmez.
"""

from __future__ import annotations

import logging
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from sqlalchemy import select

from database.access import maliyet_izinli
from database.database import get_session
from database.models.stok import Depo, StokHareketi, StokKarti, StokLotu
from database.session_manager import oturum
from database.stok_service import CIKIS_HAREKETLERI, GIRIS_HAREKETLERI, StokService

_LOG = logging.getLogger(__name__)

_EKSTRE_GIRIS = set(GIRIS_HAREKETLERI) | {"İRSALİYE İADE GİRİŞ"}
_EKSTRE_CIKIS = set(CIKIS_HAREKETLERI) | {"İRSALİYE ÇIKIŞ"}


class FifoKatmanMotoru:
    """Hareket sırasına göre FIFO lot katmanlarını simüle eder (DB lotlarını değiştirmez)."""

    def __init__(self) -> None:
        self._lots: list[list[Decimal]] = []  # [kalan_miktar, birim_maliyet]
        self.negatif_acik: Decimal = Decimal("0")
        self.son_uyumsuzluk: str = ""

    def kopyala(self) -> "FifoKatmanMotoru":
        diger = FifoKatmanMotoru()
        diger._lots = [[m, c] for m, c in self._lots]
        diger.negatif_acik = self.negatif_acik
        return diger

    def lot_miktar(self) -> Decimal:
        return sum((m for m, _ in self._lots), Decimal("0"))

    def fifo_deger(self) -> Decimal:
        return sum((m * c for m, c in self._lots), Decimal("0"))

    def giris(self, miktar: Decimal, maliyet: Decimal) -> None:
        m = _d(miktar)
        c = _d(maliyet)
        if m <= 0:
            return
        if self.negatif_acik > 0:
            kapat = min(m, self.negatif_acik)
            self.negatif_acik -= kapat
            m -= kapat
            # Negatif stok kapatılan kısım için maliyet uydurulmaz.
        if m > 0:
            self._lots.append([m, c])

    def cikis(self, miktar: Decimal) -> tuple[Decimal, Decimal]:
        """(tüketilen_fifo_maliyet, karşılanamayan_miktar)."""
        kalan = _d(miktar)
        maliyet = Decimal("0")
        while kalan > 0 and self._lots:
            lot_m, lot_c = self._lots[0]
            al = min(kalan, lot_m)
            maliyet += al * lot_c
            lot_m -= al
            kalan -= al
            if lot_m <= 0:
                self._lots.pop(0)
            else:
                self._lots[0][0] = lot_m
        if kalan > 0:
            self.negatif_acik += kalan
        return maliyet, kalan

    def mutabakat(self, yuruyen_kalan: Decimal) -> str:
        """Yürüyen miktar ile açık lot toplamını karşılaştırır."""
        lot_m = self.lot_miktar()
        y = _d(yuruyen_kalan)
        if y < 0:
            if lot_m == 0:
                self.son_uyumsuzluk = ""
                return ""
            msg = (
                f"FIFO lot miktarı ({lot_m}) negatif yürüyen kalan ({y}) ile uyuşmuyor."
            )
            self.son_uyumsuzluk = msg
            _LOG.warning("FIFO mutabakat: %s", msg)
            return msg
        if lot_m != y:
            msg = f"FIFO lot miktarı ({lot_m}) yürüyen kalan ({y}) ile uyuşmuyor."
            self.son_uyumsuzluk = msg
            _LOG.warning("FIFO mutabakat: %s", msg)
            return msg
        self.son_uyumsuzluk = ""
        return ""

_CARI_YOK_TUR = {
    "GİRİŞ",
    "ÇIKIŞ",
    "TRANSFER GİRİŞ",
    "TRANSFER ÇIKIŞ",
    "SAYIM GİRİŞ",
    "SAYIM ÇIKIŞ",
    "PAKET GİRİŞ",
    "PAKET ÇIKIŞ",
}

_ETIKET = {
    "İADE GİRİŞ": "Satış İadesi",
    "SAYIM GİRİŞ": "Sayım",
    "SAYIM ÇIKIŞ": "Sayım",
    "TRANSFER GİRİŞ": "Transfer",
    "TRANSFER ÇIKIŞ": "Transfer",
    "PAKET GİRİŞ": "Paket",
    "PAKET ÇIKIŞ": "Paket",
    "İRSALİYE ÇIKIŞ": "İrsaliye",
    "İRSALİYE İADE GİRİŞ": "İrsaliye İade",
}


def _d(v, default="0") -> Decimal:
    try:
        return Decimal(str(v if v is not None else default))
    except Exception:
        return Decimal(default)


def _net_birim(satir, *, kdv_dahil: bool = False) -> Decimal | None:
    """Belge satırından KDV hariç, iskonto sonrası net birim fiyat (TL).

    Kaynak yoksa None (0,00 uydurulmaz). Gerçek sıfır fiyat Decimal('0') döner.
    tl_birim_fiyat varsa kur çevrimli TL birim fiyat; üzerine satır iskontoları uygulanır.
    """
    if satir is None:
        return None
    bf_tl = getattr(satir, "tl_birim_fiyat", None)
    bf = getattr(satir, "birim_fiyat", None)
    if bf_tl is None and bf is None:
        return None
    # tl_birim_fiyat > 0 ise döviz/TL çevrimli birim; yoksa belge birim fiyatı
    if bf_tl is not None and _d(bf_tl) != 0:
        birim = _d(bf_tl)
    elif bf is not None:
        birim = _d(bf)
    else:
        return None

    for alan in ("iskonto_orani", "iskonto_orani_2", "iskonto_orani_3"):
        oran = _d(getattr(satir, alan, 0) or 0)
        if oran:
            birim = birim * (Decimal("1") - oran / Decimal("100"))

    # KDV dahil saklanmışsa matraha indir (belge motoru varsayılanı KDV hariç)
    if kdv_dahil:
        kdv = _d(getattr(satir, "kdv_orani", 0) or 0)
        if kdv > 0:
            birim = birim / (Decimal("1") + kdv / Decimal("100"))

    return birim.quantize(Decimal("0.0001")) if birim != 0 else Decimal("0")


def _aktif_donem_tarihleri() -> tuple[date, date]:
    bugun = date.today()
    bas, bit = date(bugun.year, 1, 1), bugun
    if not oturum.period_id:
        return bas, bit
    try:
        from database.models.donem import Donem

        with get_session() as session:
            d = session.get(Donem, int(oturum.period_id))
            if d and d.baslangic_tarihi and d.bitis_tarihi:
                return d.baslangic_tarihi, min(d.bitis_tarihi, bugun)
    except Exception as exc:
        _LOG.warning("Dönem tarihleri okunamadı: %s", exc)
    return bas, bit


class FiyatliStokEkstreService:
    @staticmethod
    def varsayilan_tarih_araligi() -> tuple[date, date]:
        return _aktif_donem_tarihleri()

    @staticmethod
    def maliyet_gorunsun_mu() -> bool:
        try:
            return maliyet_izinli() or oturum.role_kod == "YONETICI"
        except Exception:
            return True

    @staticmethod
    def _belge_cari_haritasi(session, hareketler: list) -> dict[tuple[str, str], dict]:
        from database.models.alis_faturasi import AlisFaturasi, AlisFaturasiSatiri
        from database.models.alis_iade_faturasi import AlisIadeFaturasi, AlisIadeFaturasiSatiri
        from database.models.cari import Cari
        from database.models.satis_faturasi import SatisFaturasi, SatisFaturasiSatiri
        from database.models.satis_iade_faturasi import SatisIadeFaturasi, SatisIadeFaturasiSatiri

        nos_sf, nos_af, nos_si, nos_ai = set(), set(), set(), set()
        for h in hareketler:
            bn = (h.belge_no or "").strip()
            if not bn:
                continue
            tur = h.hareket_turu or ""
            if tur == "FATURA ÇIKIŞ":
                nos_sf.add(bn)
                nos_ai.add(bn)
            elif tur == "FATURA GİRİŞ":
                nos_af.add(bn)
            elif tur == "İADE GİRİŞ":
                nos_si.add(bn)

        sonuc: dict[tuple[str, str], dict] = {}

        def _cari_bilgi(cari_id):
            if not cari_id:
                return None, None, None
            c = session.get(Cari, int(cari_id))
            if not c:
                return None, None, None
            return int(c.id), (c.cari_kodu or "").strip(), (c.unvan or "").strip()

        if nos_sf:
            for f in session.scalars(
                select(SatisFaturasi).where(SatisFaturasi.fatura_no.in_(nos_sf))
            ).all():
                if getattr(f, "is_deleted", False):
                    continue
                if (f.durum or "").upper() in ("TASLAK", "İPTAL", "IPTAL"):
                    continue
                cid, kod, ad = _cari_bilgi(f.cari_id)
                sonuc[("FATURA ÇIKIŞ", f.fatura_no)] = {
                    "belge_turu": "Satış Faturası",
                    "belge_id": int(f.id),
                    "belge_modul": "satis_fatura",
                    "cari_id": cid,
                    "cari_kodu": kod or "",
                    "cari_adi": ad or "",
                    "para_birimi": getattr(f, "para_birimi", None) or "TRY",
                    "kur": _d(getattr(f, "kur", 1) or 1),
                    "durum": f.durum,
                    "_fatura": f,
                }

        if nos_af:
            for f in session.scalars(
                select(AlisFaturasi).where(AlisFaturasi.fatura_no.in_(nos_af))
            ).all():
                if (f.durum or "").upper() in ("TASLAK", "İPTAL", "IPTAL"):
                    continue
                cid, kod, ad = _cari_bilgi(f.cari_id)
                sonuc[("FATURA GİRİŞ", f.fatura_no)] = {
                    "belge_turu": "Alış Faturası",
                    "belge_id": int(f.id),
                    "belge_modul": "alis_fatura",
                    "cari_id": cid,
                    "cari_kodu": kod or "",
                    "cari_adi": ad or "",
                    "para_birimi": getattr(f, "para_birimi", None) or "TRY",
                    "kur": _d(getattr(f, "kur", 1) or 1),
                    "durum": f.durum,
                    "_fatura": f,
                }

        if nos_si:
            for f in session.scalars(
                select(SatisIadeFaturasi).where(SatisIadeFaturasi.iade_no.in_(nos_si))
            ).all():
                cid, kod, ad = _cari_bilgi(f.cari_id)
                sonuc[("İADE GİRİŞ", f.iade_no)] = {
                    "belge_turu": "Satış İadesi",
                    "belge_id": int(f.id),
                    "belge_modul": "satis_iade",
                    "cari_id": cid,
                    "cari_kodu": kod or "",
                    "cari_adi": ad or "",
                    "para_birimi": getattr(f, "para_birimi", None) or "TRY",
                    "kur": _d(getattr(f, "kur", 1) or 1),
                    "durum": getattr(f, "durum", None),
                    "_fatura": f,
                }

        if nos_ai:
            for f in session.scalars(
                select(AlisIadeFaturasi).where(AlisIadeFaturasi.iade_no.in_(nos_ai))
            ).all():
                cid, kod, ad = _cari_bilgi(f.cari_id)
                key = ("FATURA ÇIKIŞ", f.iade_no)
                if key not in sonuc:
                    sonuc[key] = {
                        "belge_turu": "Alış İadesi",
                        "belge_id": int(f.id),
                        "belge_modul": "alis_iade",
                        "cari_id": cid,
                        "cari_kodu": kod or "",
                        "cari_adi": ad or "",
                        "para_birimi": getattr(f, "para_birimi", None) or "TRY",
                        "kur": _d(getattr(f, "kur", 1) or 1),
                        "durum": getattr(f, "durum", None),
                        "_fatura": f,
                    }

        for meta in sonuc.values():
            fat = meta.pop("_fatura", None)
            meta["_satirlar"] = list(getattr(fat, "satirlar", None) or []) if fat else []

        for _key, meta in list(sonuc.items()):
            if meta.get("_satirlar"):
                continue
            modul = meta.get("belge_modul")
            bid = meta.get("belge_id")
            if not bid:
                continue
            if modul == "satis_fatura":
                meta["_satirlar"] = list(
                    session.scalars(
                        select(SatisFaturasiSatiri).where(SatisFaturasiSatiri.fatura_id == bid)
                    ).all()
                )
            elif modul == "alis_fatura":
                meta["_satirlar"] = list(
                    session.scalars(
                        select(AlisFaturasiSatiri).where(AlisFaturasiSatiri.fatura_id == bid)
                    ).all()
                )
            elif modul == "satis_iade":
                meta["_satirlar"] = list(
                    session.scalars(
                        select(SatisIadeFaturasiSatiri).where(
                            SatisIadeFaturasiSatiri.iade_id == bid
                        )
                    ).all()
                )
            elif modul == "alis_iade":
                meta["_satirlar"] = list(
                    session.scalars(
                        select(AlisIadeFaturasiSatiri).where(AlisIadeFaturasiSatiri.iade_id == bid)
                    ).all()
                )

        return sonuc

    @staticmethod
    def _yon(tur: str) -> str:
        if tur in _EKSTRE_GIRIS:
            return "giris"
        if tur in _EKSTRE_CIKIS:
            return "cikis"
        return "notr"

    @staticmethod
    def ekstre(
        stok_id: int,
        *,
        baslangic: date | None = None,
        bitis: date | None = None,
        depo_adi: str | None = None,
        hareket_turu: str | None = None,
        cari_id: int | None = None,
        belge_turu: str | None = None,
        belge_no: str | None = None,
        para_birimi: str | None = None,
        gosterim_birimi: str | None = None,
        kdv_dahil: bool = False,
        devir_satiri: bool = True,
    ) -> dict[str, Any]:
        if not stok_id:
            raise ValueError("Stok seçilmedi.")
        if baslangic is None or bitis is None:
            vb, ve = _aktif_donem_tarihleri()
            baslangic = baslangic or vb
            bitis = bitis or ve
        if baslangic > bitis:
            raise ValueError("Başlangıç tarihi bitişten büyük olamaz.")

        maliyet_ok = FiyatliStokEkstreService.maliyet_gorunsun_mu()
        uyarilar: list[str] = []

        with get_session() as session:
            stok = session.get(StokKarti, int(stok_id))
            if stok is None or getattr(stok, "is_deleted", False):
                raise ValueError("Stok kartı bulunamadı veya silinmiş.")

            ana_birim = (stok.birim or "Adet").strip() or "Adet"
            gosterim = (gosterim_birimi or ana_birim).strip() or ana_birim
            birimler = StokService.birimleri_dict_listesi(stok)

            def miktar_goster(m: Decimal) -> Decimal:
                if gosterim.casefold() == ana_birim.casefold():
                    return _d(m)
                try:
                    o = StokService.birim_donusum_onizleme(
                        m, ana_birim, gosterim, birimler, ana_birim
                    )
                    return _d(o["hedef_miktar"])
                except Exception as exc:
                    uyarilar.append(f"Birim dönüşümü uyarısı: {exc}")
                    return _d(m)

            q = (
                select(StokHareketi, Depo.ad, StokLotu.lot_no)
                .join(Depo, Depo.id == StokHareketi.depo_id)
                .outerjoin(StokLotu, StokLotu.id == StokHareketi.lot_id)
                .where(StokHareketi.stok_id == int(stok_id))
            )
            if depo_adi and depo_adi not in ("", "(Tümü)", "Tüm Depolar"):
                q = q.where(Depo.ad == depo_adi)
            q = q.order_by(
                StokHareketi.tarih.asc(),
                StokHareketi.olusturma_tarihi.asc(),
                StokHareketi.id.asc(),
            )
            ham = list(session.execute(q).all())
            hareketler = [h for h, _, _ in ham]
            depo_map = {h.id: (d or "") for h, d, _ in ham}
            lot_map = {h.id: (l or "") for h, _, l in ham}

            belge_map = FiyatliStokEkstreService._belge_cari_haritasi(session, hareketler)
            stok_kodu = (stok.stok_kodu or "").strip()

            # --- FIFO katman + yürüyen kalan (önce devir dönemi) ---
            fifo = FifoKatmanMotoru()
            kalan_m = Decimal("0")
            for h in hareketler:
                if h.tarih >= baslangic:
                    break
                yon = FiyatliStokEkstreService._yon(h.hareket_turu)
                m = _d(h.miktar)
                c = _d(h.birim_maliyet)
                if yon == "giris":
                    kalan_m += m
                    fifo.giris(m, c)
                elif yon == "cikis":
                    kalan_m -= m
                    fifo.cikis(m)

            mut_devir = fifo.mutabakat(kalan_m)
            if mut_devir:
                uyarilar.append(f"Devir mutabakat: {mut_devir}")

            devir_miktar = kalan_m
            devir_deger = fifo.fifo_deger()
            # Negatif bakiyede FIFO değeri 0; uydurma maliyet yok
            if devir_miktar < 0:
                uyarilar.append(
                    "Negatif açılış bakiyesi: FIFO değeri yalnızca açık lotlar üzerinden "
                    "(karşılanamayan kısım için maliyet üretilmedi)."
                )
            devir_birim = (
                (devir_deger / fifo.lot_miktar()) if fifo.lot_miktar() != 0 else Decimal("0")
            )

            satirlar: list[dict] = []
            if devir_satiri:
                satirlar.append(
                    {
                        "satir_turu": "devri",
                        "hareket_id": None,
                        "tarih": baslangic,
                        "saat": "",
                        "hareket_turu": "DEVRİ",
                        "etiket": "Devir",
                        "belge_turu": "Devir",
                        "belge_no": "",
                        "belge_modul": None,
                        "belge_id": None,
                        "cari_id": None,
                        "cari_kodu": "",
                        "cari_adi": "",
                        "depo": depo_adi
                        if depo_adi and depo_adi not in ("(Tümü)", "Tüm Depolar")
                        else "Tüm Depolar",
                        "lot_no": "",
                        "birim": ana_birim,
                        "aciklama": f"Açılış bakiyesi ({baslangic.strftime('%d.%m.%Y')} öncesi)",
                        "giren": Decimal("0"),
                        "cikan": Decimal("0"),
                        "kalan": miktar_goster(devir_miktar),
                        "giris_birim_fiyat": None,
                        "net_giris_birim_fiyat": None,
                        "cikis_satis_fiyat": None,
                        "net_cikis_birim_fiyat": None,
                        "cikis_maliyet_fiyat": (devir_birim if maliyet_ok else None),
                        "giris_tutari": None,
                        "cikis_tutari": None,
                        "cikis_maliyeti": None,
                        "kalan_deger": (devir_deger if maliyet_ok else None),
                        "fifo_hesaplanamadi": bool(devir_miktar < 0 and fifo.negatif_acik > 0),
                        "fiyat_kaynak": "",
                        "fiyat_yok": False,
                        "para_birimi": "TRY",
                        "kur": Decimal("1"),
                        "doviz_birim_fiyat": None,
                        "yon": "devri",
                        "uyari": mut_devir
                        or (
                            "⚠ Negatif stok bakiyesi"
                            if devir_miktar < 0
                            else ""
                        ),
                    }
                )

            kalan_d = devir_deger
            toplam_giren = Decimal("0")
            toplam_cikan = Decimal("0")
            toplam_giris_tutar = Decimal("0")
            toplam_cikis_maliyet = Decimal("0")
            toplam_cikis_satis = Decimal("0")

            belge_filtre = (belge_no or "").strip().casefold()
            pb_filtre = (para_birimi or "").strip().upper()
            if pb_filtre in ("", "TÜMÜ", "TUMU", "(TÜMÜ)"):
                pb_filtre = ""

            # Dönem içi: FIFO her zaman tüm hareketlerle ilerler (filtre sadece gösterimi keser).
            # Böylece yürüyen kalan / FIFO değeri filtreden bağımsız doğru kalır.
            gosterilecek_idler: set[int] | None = None
            if (
                (hareket_turu and hareket_turu not in ("", "(Tümü)"))
                or belge_filtre
                or (belge_turu and belge_turu not in ("", "(Tümü)"))
                or cari_id
                or pb_filtre
            ):
                gosterilecek_idler = set()
                for h in hareketler:
                    if h.tarih < baslangic or h.tarih > bitis:
                        continue
                    tur = h.hareket_turu or ""
                    bn = (h.belge_no or "").strip()
                    meta = belge_map.get((tur, bn), {})
                    if hareket_turu and hareket_turu not in ("", "(Tümü)") and tur != hareket_turu:
                        continue
                    if belge_filtre and belge_filtre not in bn.casefold():
                        continue
                    if belge_turu and belge_turu not in ("", "(Tümü)"):
                        bt = meta.get("belge_turu") or ""
                        if belge_turu == "Transfer" and "TRANSFER" not in tur:
                            continue
                        if belge_turu == "Sayım" and "SAYIM" not in tur:
                            continue
                        if belge_turu not in ("Transfer", "Sayım") and bt != belge_turu:
                            continue
                    if cari_id and meta.get("cari_id") != int(cari_id):
                        continue
                    if pb_filtre and (meta.get("para_birimi") or "TRY").upper() != pb_filtre:
                        continue
                    gosterilecek_idler.add(int(h.id))

            for h in hareketler:
                if h.tarih < baslangic or h.tarih > bitis:
                    continue
                tur = h.hareket_turu or ""
                bn = (h.belge_no or "").strip()
                meta = belge_map.get((tur, bn), {})

                yon = FiyatliStokEkstreService._yon(tur)
                m = _d(h.miktar)
                maliyet = _d(h.birim_maliyet)
                uyari = ""

                satir_orm = None
                for s in meta.get("_satirlar") or []:
                    if (getattr(s, "urun_kodu", None) or "").strip() == stok_kodu:
                        satir_orm = s
                        break
                # Ticari net birim: kaynak belge satırı (kart fiyatı değil)
                ticari = _net_birim(satir_orm, kdv_dahil=False) if satir_orm else None
                if kdv_dahil and ticari is not None and satir_orm is not None:
                    kdv = _d(getattr(satir_orm, "kdv_orani", 0) or 0)
                    ticari = ticari * (Decimal("1") + kdv / Decimal("100"))

                giren = m if yon == "giris" else Decimal("0")
                cikan = m if yon == "cikis" else Decimal("0")

                giris_bf = None
                cikis_sf = None
                cikis_mf = None
                giris_t = None
                cikis_t = None
                cikis_mal_t = None
                fifo_hesaplanamadi = False
                fiyat_kaynak = ""
                fiyat_yok = False

                if yon == "giris":
                    kalan_m += giren
                    fifo.giris(giren, maliyet)
                    if ticari is not None:
                        giris_bf = ticari
                        fiyat_kaynak = f"Belge satırı net (KDV hariç) · {bn}"
                    elif tur == "TRANSFER GİRİŞ" and maliyet != 0:
                        giris_bf = maliyet
                        fiyat_kaynak = f"Taşınan maliyet · {bn or tur}"
                    elif tur in ("SAYIM GİRİŞ", "GİRİŞ", "PAKET GİRİŞ") and maliyet != 0:
                        giris_bf = maliyet
                        fiyat_kaynak = f"Kayıtlı maliyet · {bn or tur}"
                    elif bn and satir_orm is None:
                        fiyat_yok = True
                        uyarilar.append(
                            f"{bn}: Net giriş fiyatı için kaynak belge satırı bulunamadı."
                        )
                    elif tur in (
                        "TRANSFER GİRİŞ",
                        "SAYIM GİRİŞ",
                        "GİRİŞ",
                        "PAKET GİRİŞ",
                        "İADE GİRİŞ",
                        "FATURA GİRİŞ",
                        "İRSALİYE İADE GİRİŞ",
                    ):
                        fiyat_yok = True
                    if maliyet_ok:
                        giris_t = giren * maliyet
                        toplam_giris_tutar += giren * maliyet
                elif yon == "cikis":
                    kalan_m -= cikan
                    fifo_cost, uncovered = fifo.cikis(cikan)
                    if ticari is not None:
                        cikis_sf = ticari
                        fiyat_kaynak = f"Belge satırı net (KDV hariç) · {bn}"
                    elif bn and satir_orm is None and tur in (
                        "FATURA ÇIKIŞ",
                        "İRSALİYE ÇIKIŞ",
                    ):
                        fiyat_yok = True
                        uyarilar.append(
                            f"{bn}: Net çıkış fiyatı için kaynak belge satırı bulunamadı."
                        )
                    elif tur in ("SAYIM ÇIKIŞ", "ÇIKIŞ", "TRANSFER ÇIKIŞ", "PAKET ÇIKIŞ"):
                        # Satış fiyatı uydurma yok; FIFO maliyeti ayrı kolonda
                        fiyat_yok = True
                    if uncovered > 0:
                        fifo_hesaplanamadi = True
                        uyari = (
                            f"⚠ Negatif stok: {uncovered} birim için FIFO maliyeti "
                            "hesaplanamadı (uydurma maliyet yok)."
                        )
                        uyarilar.append(f"{bn}: {uyari}")
                    if maliyet_ok:
                        cikis_mal_t = fifo_cost
                        toplam_cikis_maliyet += fifo_cost
                        if cikan > 0 and fifo_cost > 0:
                            cikis_mf = (fifo_cost / cikan).quantize(Decimal("0.0001"))
                        elif maliyet > 0 and uncovered == 0:
                            cikis_mf = maliyet
                        elif maliyet == 0 and tur in ("FATURA ÇIKIŞ", "İRSALİYE ÇIKIŞ") and not uyari:
                            uyari = (
                                "FIFO/maliyet verisi eksik; satış fiyatı maliyet olarak kullanılmadı."
                            )
                            uyarilar.append(f"{bn}: {uyari}")
                    if cikis_sf is not None:
                        cikis_t = cikan * cikis_sf
                        toplam_cikis_satis += cikis_t
                else:
                    uyari = "Yön belirsiz hareket; miktar kalan hesabına dahil edilmedi."
                    uyarilar.append(f"{bn}/{tur}: {uyari}")

                mut = fifo.mutabakat(kalan_m)
                if mut and not uyari:
                    uyari = f"⚠ {mut}"
                    uyarilar.append(uyari)

                kalan_d = fifo.fifo_deger()
                toplam_giren += giren
                toplam_cikan += cikan

                if gosterilecek_idler is not None and int(h.id) not in gosterilecek_idler:
                    continue

                cari_kodu = meta.get("cari_kodu") or ""
                cari_adi = meta.get("cari_adi") or ""
                if tur in _CARI_YOK_TUR and not cari_adi:
                    if "TRANSFER" in tur:
                        cari_adi = "Depo transferi"
                    elif "SAYIM" in tur:
                        cari_adi = "Sayım düzeltmesi"
                    elif tur in ("GİRİŞ", "ÇIKIŞ"):
                        cari_adi = "Manuel stok fişi"
                    elif "PAKET" in tur:
                        cari_adi = "Paket üretimi"

                saat = ""
                if h.olusturma_tarihi:
                    try:
                        saat = h.olusturma_tarihi.strftime("%H:%M:%S")
                    except Exception:
                        saat = ""

                belge_turu_goster = meta.get("belge_turu") or (
                    "Transfer"
                    if "TRANSFER" in tur
                    else "Sayım"
                    if "SAYIM" in tur
                    else "Stok Fişi"
                    if tur in ("GİRİŞ", "ÇIKIŞ")
                    else _ETIKET.get(tur, "")
                )

                aciklama = ""
                if lot_map.get(h.id):
                    aciklama = f"Lot: {lot_map.get(h.id)}"
                if uyari:
                    aciklama = f"{aciklama} | {uyari}".strip(" |")
                if fiyat_yok and not fiyat_kaynak:
                    aciklama = f"{aciklama} | Fiyat yok".strip(" |")

                satirlar.append(
                    {
                        "satir_turu": "hareket",
                        "hareket_id": int(h.id),
                        "tarih": h.tarih,
                        "saat": saat,
                        "hareket_turu": tur,
                        "etiket": _ETIKET.get(tur, ""),
                        "belge_turu": belge_turu_goster,
                        "belge_no": bn,
                        "belge_modul": meta.get("belge_modul"),
                        "belge_id": meta.get("belge_id"),
                        "cari_id": meta.get("cari_id"),
                        "cari_kodu": cari_kodu,
                        "cari_adi": cari_adi,
                        "depo": depo_map.get(h.id, ""),
                        "lot_no": lot_map.get(h.id, ""),
                        "birim": ana_birim,
                        "aciklama": aciklama,
                        "giren": miktar_goster(giren) if giren else Decimal("0"),
                        "cikan": miktar_goster(cikan) if cikan else Decimal("0"),
                        "kalan": miktar_goster(kalan_m),
                        # Net ticari fiyatlar (FIFO maliyetinden ayrı)
                        "giris_birim_fiyat": giris_bf,
                        "net_giris_birim_fiyat": giris_bf,
                        "cikis_satis_fiyat": cikis_sf,
                        "net_cikis_birim_fiyat": cikis_sf,
                        "cikis_maliyet_fiyat": cikis_mf if maliyet_ok else None,
                        "giris_tutari": giris_t if maliyet_ok else None,
                        "cikis_tutari": cikis_t,
                        "cikis_maliyeti": cikis_mal_t if maliyet_ok else None,
                        "kalan_deger": (kalan_d if maliyet_ok else None),
                        "fifo_hesaplanamadi": fifo_hesaplanamadi,
                        "fiyat_kaynak": fiyat_kaynak,
                        "fiyat_yok": fiyat_yok,
                        "para_birimi": meta.get("para_birimi") or "TRY",
                        "kur": meta.get("kur") or Decimal("1"),
                        "doviz_birim_fiyat": (
                            _d(getattr(satir_orm, "birim_fiyat_doviz", 0))
                            if satir_orm and getattr(satir_orm, "birim_fiyat_doviz", None)
                            else None
                        ),
                        "yon": yon,
                        "uyari": uyari,
                    }
                )

            lot_q = select(StokLotu).where(StokLotu.stok_id == stok.id)
            if depo_adi and depo_adi not in ("", "(Tümü)", "Tüm Depolar"):
                lot_q = lot_q.join(Depo, Depo.id == StokLotu.depo_id).where(Depo.ad == depo_adi)
            lotlar = list(session.scalars(lot_q).all())
            fifo_kalan = sum((_d(l.kalan_miktar) for l in lotlar), Decimal("0"))
            fifo_deger = sum(
                (_d(l.kalan_miktar) * _d(l.birim_maliyet) for l in lotlar), Decimal("0")
            )

            return {
                "stok_id": int(stok.id),
                "stok_kodu": stok.stok_kodu,
                "stok_adi": stok.stok_adi,
                "ana_birim": ana_birim,
                "gosterim_birimi": gosterim,
                "baslangic": baslangic,
                "bitis": bitis,
                "depo": depo_adi or "Tüm Depolar",
                "maliyet_yontemi": "FIFO",
                "maliyet_gorunur": maliyet_ok,
                "kdv_dahil": bool(kdv_dahil),
                "ozet": {
                    "devir_miktar": miktar_goster(devir_miktar),
                    "giren_miktar": miktar_goster(toplam_giren),
                    "cikan_miktar": miktar_goster(toplam_cikan),
                    "kalan_miktar": miktar_goster(kalan_m),
                    "devir_degeri": (devir_deger if maliyet_ok else None),
                    "giris_tutari": (toplam_giris_tutar if maliyet_ok else None),
                    "cikis_maliyeti": (toplam_cikis_maliyet if maliyet_ok else None),
                    "cikis_satis_tutari": toplam_cikis_satis,
                    "kalan_stok_degeri": (kalan_d if maliyet_ok else None),
                    "fifo_kalan_degeri": (kalan_d if maliyet_ok else None),
                    "fifo_lot_miktar": miktar_goster(fifo_kalan) if maliyet_ok else None,
                    "fifo_lot_deger": (fifo_deger if maliyet_ok else None),
                    "fifo_sim_lot_miktar": miktar_goster(fifo.lot_miktar()) if maliyet_ok else None,
                    "negatif_stok": bool(kalan_m < 0 or fifo.negatif_acik > 0),
                },
                "satirlar": satirlar,
                "uyarilar": list(dict.fromkeys(uyarilar)),
                "firma_kodu": getattr(oturum, "firma_kodu", None),
                "uretim_zamani": datetime.now(),
            }

    @staticmethod
    def hareket_ekstresi(
        stok_id: int,
        *,
        baslangic: date | None = None,
        bitis: date | None = None,
        depo_adi: str | None = None,
        belge_arama: str | None = None,
        hareket_turu: str | None = None,
        gosterim_birimi: str | None = None,
        tum_gecmis: bool = False,
    ) -> dict[str, Any]:
        """Stok Kartı > Hareketler sekmesi için FIFO ekstre (tek doğruluk kaynağı)."""
        if tum_gecmis or (baslangic is None and bitis is None):
            bas = baslangic or date(1970, 1, 1)
            bit = bitis or date.today()
            return FiyatliStokEkstreService.ekstre(
                stok_id,
                baslangic=bas,
                bitis=bit,
                depo_adi=depo_adi,
                belge_no=belge_arama,
                hareket_turu=hareket_turu,
                gosterim_birimi=gosterim_birimi,
                devir_satiri=False,
            )
        return FiyatliStokEkstreService.ekstre(
            stok_id,
            baslangic=baslangic,
            bitis=bitis,
            depo_adi=depo_adi,
            belge_no=belge_arama,
            hareket_turu=hareket_turu,
            gosterim_birimi=gosterim_birimi,
            devir_satiri=True,
        )

    @staticmethod
    def excel_aktar(sonuc: dict, yol: str) -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Fiyatlı Stok Ekstresi"
        ozet = sonuc.get("ozet") or {}
        ws.append(["Fiyatlı Stok Ekstresi"])
        ws.append([f"Stok: {sonuc.get('stok_kodu')} — {sonuc.get('stok_adi')}"])
        ws.append(
            [
                f"Tarih: {sonuc.get('baslangic')} – {sonuc.get('bitis')} | Depo: {sonuc.get('depo')} | "
                f"Maliyet: {sonuc.get('maliyet_yontemi')}"
            ]
        )
        ws.append(
            [
                "Devir",
                float(ozet.get("devir_miktar") or 0),
                "Giren",
                float(ozet.get("giren_miktar") or 0),
                "Çıkan",
                float(ozet.get("cikan_miktar") or 0),
                "Kalan",
                float(ozet.get("kalan_miktar") or 0),
            ]
        )
        if sonuc.get("maliyet_gorunur"):
            ws.append(
                [
                    "Devir Değeri",
                    float(ozet.get("devir_degeri") or 0),
                    "Giriş Tutarı",
                    float(ozet.get("giris_tutari") or 0),
                    "Çıkış Maliyeti",
                    float(ozet.get("cikis_maliyeti") or 0),
                    "Kalan Değer",
                    float(ozet.get("kalan_stok_degeri") or 0),
                ]
            )
        ws.append([])
        basliklar = [
            "Tarih",
            "Saat",
            "Hareket",
            "Belge Türü",
            "Belge No",
            "Cari Kodu",
            "Cari Kart Adı",
            "Depo",
            "Giriş Miktarı",
            "Net Giriş Birim Fiyatı",
            "Çıkış Miktarı",
            "Net Çıkış Birim Fiyatı",
            "Kalan Miktar",
            "FIFO Kalan Değeri",
            "Çıkış Maliyet Fiyatı",
            "Giriş Tutarı",
            "Çıkış Tutarı",
            "Çıkış Maliyeti",
            "Para Birimi",
            "Kur",
            "Fiyat Kaynağı",
            "Uyarı",
        ]
        ws.append(basliklar)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for s in sonuc.get("satirlar") or []:
            t = s.get("tarih")
            net_g = s.get("net_giris_birim_fiyat", s.get("giris_birim_fiyat"))
            net_c = s.get("net_cikis_birim_fiyat", s.get("cikis_satis_fiyat"))
            ws.append(
                [
                    t.strftime("%d.%m.%Y") if hasattr(t, "strftime") else str(t or ""),
                    s.get("saat") or "",
                    s.get("hareket_turu") or "",
                    s.get("belge_turu") or "",
                    s.get("belge_no") or "",
                    s.get("cari_kodu") or "",
                    s.get("cari_adi") or "",
                    s.get("depo") or "",
                    float(s.get("giren") or 0) if s.get("giren") else None,
                    float(net_g) if net_g is not None else None,
                    float(s.get("cikan") or 0) if s.get("cikan") else None,
                    float(net_c) if net_c is not None else None,
                    float(s.get("kalan") or 0),
                    float(s["kalan_deger"]) if s.get("kalan_deger") is not None else None,
                    float(s["cikis_maliyet_fiyat"])
                    if s.get("cikis_maliyet_fiyat") is not None
                    else None,
                    float(s["giris_tutari"]) if s.get("giris_tutari") is not None else None,
                    float(s["cikis_tutari"]) if s.get("cikis_tutari") is not None else None,
                    float(s["cikis_maliyeti"]) if s.get("cikis_maliyeti") is not None else None,
                    s.get("para_birimi") or "",
                    float(s.get("kur") or 1),
                    s.get("fiyat_kaynak") or "",
                    s.get("uyari") or "",
                ]
            )
        wb.save(yol)
        return yol

    @staticmethod
    def pdf_aktar(sonuc: dict, yol: str) -> str:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError:
            with open(yol, "w", encoding="utf-8") as f:
                f.write("Fiyatlı Stok Ekstresi\n")
                f.write(f"{sonuc.get('stok_kodu')} {sonuc.get('stok_adi')}\n")
                for s in sonuc.get("satirlar") or []:
                    f.write(
                        f"{s.get('tarih')}\t{s.get('hareket_turu')}\t{s.get('belge_no')}\t"
                        f"{s.get('cari_adi')}\t{s.get('giren')}\t{s.get('cikan')}\t{s.get('kalan')}\n"
                    )
            return yol

        doc = SimpleDocTemplate(
            yol,
            pagesize=landscape(A4),
            leftMargin=18,
            rightMargin=18,
            topMargin=24,
            bottomMargin=24,
        )
        styles = getSampleStyleSheet()
        story = []
        story.append(Paragraph("Fiyatlı Stok Ekstresi", styles["Title"]))
        story.append(
            Paragraph(
                f"{sonuc.get('stok_kodu')} — {sonuc.get('stok_adi')} | "
                f"{sonuc.get('baslangic')} – {sonuc.get('bitis')} | Depo: {sonuc.get('depo')} | "
                f"Maliyet: {sonuc.get('maliyet_yontemi')}",
                styles["Normal"],
            )
        )
        ozet = sonuc.get("ozet") or {}
        story.append(
            Paragraph(
                f"Devir: {ozet.get('devir_miktar')} | Giren: {ozet.get('giren_miktar')} | "
                f"Çıkan: {ozet.get('cikan_miktar')} | Kalan: {ozet.get('kalan_miktar')}",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 8))
        data = [
            [
                "Tarih",
                "Hareket",
                "Belge",
                "Cari",
                "Giriş",
                "Net Giriş F.",
                "Çıkış",
                "Net Çıkış F.",
                "Kalan",
                "FIFO Kalan",
            ]
        ]
        for s in sonuc.get("satirlar") or []:
            t = s.get("tarih")
            net_g = s.get("net_giris_birim_fiyat", s.get("giris_birim_fiyat"))
            net_c = s.get("net_cikis_birim_fiyat", s.get("cikis_satis_fiyat"))
            data.append(
                [
                    t.strftime("%d.%m.%Y") if hasattr(t, "strftime") else "",
                    (s.get("hareket_turu") or "")[:14],
                    (s.get("belge_no") or "")[:12],
                    (s.get("cari_adi") or "")[:16],
                    f"{s.get('giren') or ''}" if s.get("giren") else "",
                    f"{net_g}" if net_g is not None else ("Fiyat yok" if s.get("fiyat_yok") else ""),
                    f"{s.get('cikan') or ''}" if s.get("cikan") else "",
                    f"{net_c}" if net_c is not None else ("Fiyat yok" if s.get("fiyat_yok") and s.get("yon") == "cikis" else ""),
                    f"{s.get('kalan') or 0}",
                    f"{s.get('kalan_deger') or ''}",
                ]
            )
        tab = Table(data, repeatRows=1)
        tab.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
                ]
            )
        )
        story.append(tab)
        doc.build(story)
        return yol
