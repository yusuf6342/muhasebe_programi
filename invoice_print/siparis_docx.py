"""SATIŞ SİPARİŞ FORMU Word (.docx) çıktısı — düzenlenebilir metin ve tablolar (görüntü yok)."""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from database.irsaliye_customer_view import birim_toplamlari, firma_vergi_metni, miktar_metni, musteri_vergi_metni
from database.siparis_form_view import (
    SiparisFormViewModel,
    assert_siparis_form_safe,
    belge_bilgi_satirlari,
    kolonlar,
    satir_degerleri,
    teslimat_satirlari,
    tutar_satirlari,
)
from invoice_print.irsaliye_docx import (
    BEYAZ,
    KULLANILABILIR_MM,
    LACIVERT,
    METIN,
    SARI,
    SOLUK,
    _alan,
    _alt_cizgi,
    _golge,
    _kenarlik,
    _kutu_satiri,
    _paragraf,
    _run,
    _sabit_yukseklik,
    _satir_bolunmesin,
    _tablo_sabit,
    _ust_cizgi,
)

HIZA = {"c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT}


def kolon_genislikleri(vm: SiparisFormViewModel) -> list[float]:
    kol = kolonlar(vm)
    sabit = sum(g for _b, _h, g in kol if g)
    return [g if g else KULLANILABILIR_MM - sabit for _b, _h, g in kol]


def render_siparis_docx(vm: SiparisFormViewModel, hedef: Path) -> Path:
    assert_siparis_form_safe(vm)
    hedef = Path(hedef)
    hedef.parent.mkdir(parents=True, exist_ok=True)
    f = vm.firma or {}
    m = vm.musteri or {}

    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name = "Segoe UI"
    stil.font.size = Pt(9)
    bolum = doc.sections[0]
    bolum.page_width, bolum.page_height = Mm(210), Mm(297)
    bolum.left_margin = bolum.right_margin = Mm(12)
    bolum.top_margin = Mm(12)
    bolum.bottom_margin = Mm(18)
    bolum.footer_distance = Mm(7)

    alt = bolum.footer.paragraphs[0]
    sekmeler = alt.paragraph_format.tab_stops
    for varsayilan in (Inches(3.25), Inches(6.5)):
        sekmeler.add_tab_stop(varsayilan, WD_TAB_ALIGNMENT.CLEAR)
    sekmeler.add_tab_stop(Mm(KULLANILABILIR_MM), WD_TAB_ALIGNMENT.RIGHT)
    _ust_cizgi(alt, "F4C542", 8)
    _run(alt, f"Sipariş No: {vm.siparis_no}", boyut=7.5, renk=LACIVERT)
    _run(alt, "\tSayfa ", boyut=7.5, renk=LACIVERT)
    _alan(alt, "PAGE")
    _run(alt, " / ", boyut=7.5, renk=LACIVERT)
    _alan(alt, "NUMPAGES")

    # Antet
    antet = doc.add_table(rows=1, cols=2)
    _tablo_sabit(antet, (100, 86))
    sol, sag = antet.rows[0].cells
    p = _paragraf(sol, ilk=True)
    if vm.logo_data_uri and "," in vm.logo_data_uri:
        try:
            veri = base64.b64decode(vm.logo_data_uri.split(",", 1)[1])
            p.add_run().add_picture(io.BytesIO(veri), height=Mm(16))
            p = _paragraf(sol, once=2)
        except Exception:
            pass
    _run(p, f.get("unvan") or "", boyut=12, kalin=True, renk=LACIVERT)
    for satir in (
        f.get("adres"),
        "  ·  ".join(x for x in (f"Tel: {f['telefon']}" if f.get("telefon") else "", f.get("email") or "") if x),
        firma_vergi_metni(f),
    ):
        if satir:
            _run(_paragraf(sol), satir, boyut=8.5, renk=SOLUK)

    bp = _paragraf(sag, ilk=True, sonra=3)
    baslik_tablo = sag.add_table(rows=1, cols=1)
    bh = baslik_tablo.rows[0].cells[0]
    _golge(bh, "102A43")
    _kenarlik(bh, left={"sz": 24, "renk": "F4C542"})
    _run(_paragraf(bh, ilk=True, once=2, sonra=2), vm.belge_baslik, boyut=14, kalin=True, renk=BEYAZ)
    bp._p.getparent().remove(bp._p)
    bilgi = sag.add_table(rows=0, cols=2)
    for etiket, deger in belge_bilgi_satirlari(vm):
        c0, c1 = bilgi.add_row().cells
        _run(_paragraf(c0, ilk=True, once=1), etiket, boyut=8.5, renk=SOLUK)
        _run(_paragraf(c1, ilk=True, once=1), deger, boyut=9, kalin=True, renk=LACIVERT)
    _tablo_sabit(bilgi, (27, 57))
    _paragraf(sag).paragraph_format.line_spacing = Pt(2)

    serit = _paragraf(doc, once=2, sonra=6)
    serit.paragraph_format.line_spacing = Pt(4)
    _alt_cizgi(serit, "F4C542", 24)

    # Müşteri / teslimat kutuları
    kutular = doc.add_table(rows=1, cols=2)
    kutular.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tablo_sabit(kutular, (93, 93))
    _satir_bolunmesin(kutular.rows[0])
    k_m, k_t = kutular.rows[0].cells
    for hucre, baslik in ((k_m, "MÜŞTERİ"), (k_t, "TESLİMAT VE KOŞULLAR")):
        _golge(hucre, "F3F6F9")
        _kenarlik(hucre, top={"sz": 18, "renk": "102A43"}, left={}, bottom={}, right={})
        _run(_paragraf(hucre, ilk=True, once=2), baslik, boyut=9, kalin=True, renk=LACIVERT)
    _run(_paragraf(k_m, once=1), m.get("unvan") or "", boyut=10.5, kalin=True, renk=LACIVERT)
    _kutu_satiri(k_m, "Müşteri Kodu", m.get("kod") or "")
    vergi = musteri_vergi_metni(m)
    if vergi:
        _kutu_satiri(k_m, "Vergi Bilgisi", vergi)
    if m.get("adres"):
        _kutu_satiri(k_m, "Adres", m["adres"])
    if m.get("telefon"):
        _kutu_satiri(k_m, "Telefon", m["telefon"])
    for etiket, deger in teslimat_satirlari(vm):
        _kutu_satiri(k_t, etiket, deger)
    for hucre in (k_m, k_t):
        _paragraf(hucre, sonra=2)

    _paragraf(doc, sonra=4)

    # Ürün tablosu — başlık her sayfada tekrarlanır, satır bölünmez
    kol = kolonlar(vm)
    miktar_index = next(i for i, (b, _h, _g) in enumerate(kol) if b == "Miktar")
    tablo = doc.add_table(rows=1, cols=len(kol))
    for hucre, (metin, hz, _g) in zip(tablo.rows[0].cells, kol):
        _golge(hucre, "102A43")
        _kenarlik(hucre, bottom={"sz": 12, "renk": "F4C542"})
        _run(_paragraf(hucre, ilk=True, hiza=HIZA[hz], once=2, sonra=2), metin, boyut=8, kalin=True, renk=BEYAZ)
    _satir_bolunmesin(tablo.rows[0], baslik=True)
    for i, s in enumerate(vm.satirlar):
        satir = tablo.add_row()
        _satir_bolunmesin(satir)
        for j, (hucre, deger, (_b, hz, _g)) in enumerate(zip(satir.cells, satir_degerleri(vm, s), kol)):
            if i % 2:
                _golge(hucre, "F8FAFC")
            _kenarlik(hucre, bottom={"sz": 4, "renk": "D9E2EC"})
            p = _paragraf(hucre, ilk=True, hiza=HIZA[hz], once=1.5, sonra=1.5)
            _run(p, deger or "", boyut=8, kalin=(j == miktar_index))
            if j == 2 and s.aciklama:
                for parca in s.aciklama.splitlines():
                    _run(_paragraf(hucre, sonra=1.5), parca, boyut=7, renk=SOLUK)
    _tablo_sabit(tablo, kolon_genislikleri(vm))

    # Toplamlar: solda birim bazında miktar, sağda tutarlar (tek bölünmeyen satır)
    toplamlar = birim_toplamlari(vm.satirlar)
    _paragraf(doc, sonra=4)
    ozet = doc.add_table(rows=1, cols=2)
    _tablo_sabit(ozet, (110, 76))
    _satir_bolunmesin(ozet.rows[0])
    sol, sag = ozet.rows[0].cells
    _run(_paragraf(sol, ilk=True, once=1), f"Toplam Miktar: {vm.toplam_miktar_goster}", boyut=9.5, kalin=True, renk=LACIVERT)
    rozet_p = _paragraf(sol, once=2)
    for birim, miktar in toplamlar:
        r = _run(rozet_p, f" {miktar_metni(miktar)} {birim} ", boyut=9, kalin=True, renk=RGBColor(0x08, 0x1B, 0x2C))
        rpr = r._element.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "FFE89A")
        rpr.append(shd)
        _run(rozet_p, "  ", boyut=9)
    _run(
        _paragraf(sol, once=2),
        f"Kalem: {len(vm.satirlar)}" + (" · Farklı birimler ayrı toplanır" if len(toplamlar) > 1 else ""),
        boyut=8, renk=SOLUK,
    )
    tutar = sag.add_table(rows=0, cols=2)
    for etiket, deger, genel in tutar_satirlari(vm):
        c0, c1 = tutar.add_row().cells
        if genel:
            _golge(c0, "102A43")
            _golge(c1, "102A43")
        renk = SARI if genel else METIN
        boyut = 10 if genel else 9
        _run(_paragraf(c0, ilk=True, once=1, sonra=1), etiket, boyut=boyut, kalin=genel, renk=renk)
        _run(_paragraf(c1, ilk=True, hiza=WD_ALIGN_PARAGRAPH.RIGHT, once=1, sonra=1), deger, boyut=boyut, kalin=True, renk=renk)
    _tablo_sabit(tutar, (30, 44))
    sag.paragraphs[0]._p.getparent().remove(sag.paragraphs[0]._p)
    _paragraf(sag).paragraph_format.line_spacing = Pt(2)

    if vm.aciklama:
        np_ = _paragraf(doc, once=8, sonra=0)
        np_.paragraph_format.keep_with_next = True
        _run(np_, "Açıklama", boyut=9, kalin=True, renk=LACIVERT)
        for parca in vm.aciklama.splitlines() or [""]:
            _run(_paragraf(doc, sonra=0), parca, boyut=8.5)

    ara = _paragraf(doc, once=6, sonra=0)
    ara.paragraph_format.keep_with_next = True

    imza = doc.add_table(rows=1, cols=3)
    _tablo_sabit(imza, (90, 6, 90))
    _satir_bolunmesin(imza.rows[0])
    for hucre, baslik in ((imza.rows[0].cells[0], "SİPARİŞİ HAZIRLAYAN"), (imza.rows[0].cells[2], "MÜŞTERİ ONAYI")):
        kenar = {"sz": 10, "renk": "102A43"}
        _kenarlik(hucre, top=kenar, left=kenar, bottom=kenar, right=kenar)
        bas = _paragraf(hucre, ilk=True, once=0, sonra=4)
        pPr = bas._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "102A43")
        pPr.append(shd)
        _run(bas, f" {baslik}", boyut=9.5, kalin=True, renk=SARI)
        for etiket, bosluk in (("Adı Soyadı", 10), ("Tarih", 10), ("İmza", 26)):
            p = _paragraf(hucre, once=bosluk, sonra=2)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(86), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
            _run(p, f" {etiket}", boyut=8.5, renk=SOLUK)
            _run(p, "\t", boyut=8.5, renk=SOLUK)
    _kenarlik(imza.rows[0].cells[1])
    _sabit_yukseklik(imza.rows[0], 38)

    son = _paragraf(doc, once=2)
    _run(
        son,
        "Bu belge satış sipariş formudur; fatura veya irsaliye yerine geçmez."
        + (f" Oluşturulma: {vm.olusturma}" if vm.olusturma else ""),
        boyut=7.5, renk=RGBColor(0x94, 0xA3, 0xB8),
    )

    metin = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for satir in t.rows:
            for hucre in satir.cells:
                metin += "\n" + hucre.text
    assert_siparis_form_safe(vm, metin)
    gecici = hedef.with_name(f".{hedef.stem}_gecici.docx")
    doc.save(str(gecici))
    gecici.replace(hedef)
    return hedef
