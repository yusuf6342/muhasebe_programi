"""SATIN ALMA TALEP FORMU Word (.docx) çıktısı — düzenlenebilir metin ve tablolar."""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from database.irsaliye_customer_view import firma_vergi_metni
from database.talep_form_view import (
    TalepFormViewModel,
    belge_bilgi_satirlari,
    kolonlar,
    onay_metni,
    satir_degerleri,
    talep_bilgi_satirlari,
)
from invoice_print.irsaliye_docx import (
    BEYAZ,
    KULLANILABILIR_MM,
    LACIVERT,
    SARI,
    SOLUK,
    _alan,
    _alt_cizgi,
    _golge,
    _kenarlik,
    _kutu_satiri,
    _paragraf,
    _run,
    _sabit_yukseklik,
    _satir_bolunmesin,
    _tablo_sabit,
    _ust_cizgi,
)

HIZA = {"c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT}


def kolon_genislikleri(vm: TalepFormViewModel) -> list[float]:
    kol = kolonlar(vm)
    sabit = sum(g for _b, _h, g in kol if g)
    return [g if g else KULLANILABILIR_MM - sabit for _b, _h, g in kol]


def render_talep_docx(vm: TalepFormViewModel, hedef: Path) -> Path:
    hedef = Path(hedef)
    hedef.parent.mkdir(parents=True, exist_ok=True)
    f = vm.firma or {}

    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name = "Segoe UI"
    stil.font.size = Pt(9)
    bolum = doc.sections[0]
    bolum.page_width, bolum.page_height = Mm(210), Mm(297)
    bolum.left_margin = bolum.right_margin = Mm(12)
    bolum.top_margin = Mm(12)
    bolum.bottom_margin = Mm(18)
    bolum.footer_distance = Mm(7)

    if vm.filigran:
        ust = bolum.header.paragraphs[0]
        ust.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(ust, vm.filigran, boyut=28, kalin=True, renk=RGBColor(0xE8, 0xA0, 0xA0))

    alt = bolum.footer.paragraphs[0]
    sekmeler = alt.paragraph_format.tab_stops
    for varsayilan in (Inches(3.25), Inches(6.5)):
        sekmeler.add_tab_stop(varsayilan, WD_TAB_ALIGNMENT.CLEAR)
    sekmeler.add_tab_stop(Mm(KULLANILABILIR_MM), WD_TAB_ALIGNMENT.RIGHT)
    _ust_cizgi(alt, "F4C542", 8)
    _run(alt, f"Talep No: {vm.talep_no}", boyut=7.5, renk=LACIVERT)
    _run(alt, "\tSayfa ", boyut=7.5, renk=LACIVERT)
    _alan(alt, "PAGE")
    _run(alt, " / ", boyut=7.5, renk=LACIVERT)
    _alan(alt, "NUMPAGES")

    antet = doc.add_table(rows=1, cols=2)
    _tablo_sabit(antet, (100, 86))
    sol, sag = antet.rows[0].cells
    p = _paragraf(sol, ilk=True)
    if vm.logo_data_uri and "," in vm.logo_data_uri:
        try:
            veri = base64.b64decode(vm.logo_data_uri.split(",", 1)[1])
            p.add_run().add_picture(io.BytesIO(veri), height=Mm(16))
            p = _paragraf(sol, once=2)
        except Exception:
            pass
    _run(p, f.get("unvan") or "", boyut=12, kalin=True, renk=LACIVERT)
    for satir in (
        f.get("adres"),
        "  ·  ".join(x for x in (f"Tel: {f['telefon']}" if f.get("telefon") else "", f.get("email") or "") if x),
        firma_vergi_metni(f),
    ):
        if satir:
            _run(_paragraf(sol), satir, boyut=8.5, renk=SOLUK)

    bp = _paragraf(sag, ilk=True, sonra=3)
    baslik_tablo = sag.add_table(rows=1, cols=1)
    bh = baslik_tablo.rows[0].cells[0]
    _golge(bh, "102A43")
    _kenarlik(bh, left={"sz": 24, "renk": "F4C542"})
    _run(_paragraf(bh, ilk=True, once=2, sonra=2), vm.belge_baslik, boyut=13, kalin=True, renk=BEYAZ)
    bp._p.getparent().remove(bp._p)
    bilgi = sag.add_table(rows=0, cols=2)
    for etiket, deger in belge_bilgi_satirlari(vm):
        c0, c1 = bilgi.add_row().cells
        _run(_paragraf(c0, ilk=True, once=1), etiket, boyut=8.5, renk=SOLUK)
        _run(_paragraf(c1, ilk=True, once=1), deger, boyut=9, kalin=True, renk=LACIVERT)
    _tablo_sabit(bilgi, (27, 57))
    _paragraf(sag).paragraph_format.line_spacing = Pt(2)

    serit = _paragraf(doc, once=2, sonra=6)
    serit.paragraph_format.line_spacing = Pt(4)
    _alt_cizgi(serit, "F4C542", 24)

    kutu_t = doc.add_table(rows=1, cols=1)
    _tablo_sabit(kutu_t, (KULLANILABILIR_MM,))
    kutu = kutu_t.rows[0].cells[0]
    _golge(kutu, "F3F6F9")
    _kenarlik(kutu, top={"sz": 18, "renk": "102A43"}, left={}, bottom={}, right={})
    _run(_paragraf(kutu, ilk=True, once=2), "TALEP BİLGİLERİ", boyut=9, kalin=True, renk=LACIVERT)
    for etiket, deger in talep_bilgi_satirlari(vm):
        _kutu_satiri(kutu, etiket, deger)
    _paragraf(kutu, sonra=2)
    _paragraf(doc, sonra=4)

    kol = kolonlar(vm)
    tablo = doc.add_table(rows=1, cols=len(kol))
    for hucre, (metin, hz, _g) in zip(tablo.rows[0].cells, kol):
        _golge(hucre, "102A43")
        _kenarlik(hucre, bottom={"sz": 12, "renk": "F4C542"})
        _run(_paragraf(hucre, ilk=True, hiza=HIZA[hz], once=2, sonra=2), metin, boyut=8, kalin=True, renk=BEYAZ)
    _satir_bolunmesin(tablo.rows[0], baslik=True)
    for i, s in enumerate(vm.satirlar):
        satir = tablo.add_row()
        _satir_bolunmesin(satir)
        for j, (hucre, deger, (b, hz, _g)) in enumerate(zip(satir.cells, satir_degerleri(vm, s), kol)):
            if i % 2:
                _golge(hucre, "F8FAFC")
            _kenarlik(hucre, bottom={"sz": 4, "renk": "D9E2EC"})
            _run(_paragraf(hucre, ilk=True, hiza=HIZA[hz], once=1.5, sonra=1.5), deger or "", boyut=8,
                 kalin=(b == "Miktar"))
            if j == 2 and s.aciklama:
                for parca in s.aciklama.splitlines():
                    _run(_paragraf(hucre, sonra=1.5), parca, boyut=7, renk=SOLUK)
    _tablo_sabit(tablo, kolon_genislikleri(vm))

    ozet = _paragraf(doc, once=4)
    _run(ozet, f"Toplam Miktar: {vm.toplam_miktar_goster}   ·   Kalem: {len(vm.satirlar)}", boyut=9.5, kalin=True,
         renk=LACIVERT)
    if vm.fiyatli:
        _run(_paragraf(doc, once=2, hiza=WD_ALIGN_PARAGRAPH.RIGHT), f"Tahmini Toplam: {vm.tahmini_toplam_goster} TL",
             boyut=10, kalin=True, renk=LACIVERT)

    if vm.aciklama:
        np_ = _paragraf(doc, once=8, sonra=0)
        np_.paragraph_format.keep_with_next = True
        _run(np_, "Açıklama", boyut=9, kalin=True, renk=LACIVERT)
        for parca in vm.aciklama.splitlines() or [""]:
            _run(_paragraf(doc, sonra=0), parca, boyut=8.5)

    ara = _paragraf(doc, once=6, sonra=0)
    ara.paragraph_format.keep_with_next = True
    imza = doc.add_table(rows=1, cols=5)
    _tablo_sabit(imza, (60, 3, 60, 3, 60))
    _satir_bolunmesin(imza.rows[0])
    kutular = (
        (imza.rows[0].cells[0], "TALEP EDEN", vm.isteyen, ""),
        (imza.rows[0].cells[2], "KONTROL EDEN", "", ""),
        (imza.rows[0].cells[4], "ONAYLAYAN", vm.onaylayan, onay_metni(vm)),
    )
    for hucre, baslik, ad, dijital in kutular:
        kenar = {"sz": 10, "renk": "102A43"}
        _kenarlik(hucre, top=kenar, left=kenar, bottom=kenar, right=kenar)
        bas = _paragraf(hucre, ilk=True, once=0, sonra=4)
        pPr = bas._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "102A43")
        pPr.append(shd)
        _run(bas, f" {baslik}", boyut=9, kalin=True, renk=SARI)
        for etiket, deger, bosluk in (("Adı Soyadı", ad, 8), ("Tarih", "", 8), ("İmza", "", 20)):
            p = _paragraf(hucre, once=bosluk, sonra=2)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(57), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
            _run(p, f" {etiket}: ", boyut=8, renk=SOLUK)
            _run(p, deger or "\t", boyut=8, kalin=bool(deger), renk=LACIVERT)
        if dijital:
            _run(_paragraf(hucre, once=2, sonra=2), dijital, boyut=7, kalin=True, renk=RGBColor(0x0F, 0x7B, 0x3F))
    for i in (1, 3):
        _kenarlik(imza.rows[0].cells[i])
    _sabit_yukseklik(imza.rows[0], 40)

    _run(
        _paragraf(doc, once=2),
        "Bu belge şirket içi satın alma talep formudur; sipariş yerine geçmez."
        + (f" Oluşturulma: {vm.olusturma}" if vm.olusturma else ""),
        boyut=7.5, renk=RGBColor(0x94, 0xA3, 0xB8),
    )

    gecici = hedef.with_name(f".{hedef.stem}_gecici.docx")
    doc.save(str(gecici))
    gecici.replace(hedef)
    return hedef
