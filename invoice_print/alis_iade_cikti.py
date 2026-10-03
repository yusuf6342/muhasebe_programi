"""ALIŞ İADE FATURASI A4 çıktıları — PDF (Chromium) ve Word (.docx) aynı görünüm modelinden.

Tedarikçiye giden belgedir: iç stok maliyeti / FIFO bilgisi modele hiç alınmaz.
"""

from __future__ import annotations

import html
import logging
import os
import re
import tempfile
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from invoice_print.irsaliye_cikti import _sayfa_bilgisi_bas

_LOG = logging.getLogger("invoice_print.alis_iade")

YASAK_ALANLAR = ("maliyet", "fifo", "stok_maliyet", "kar")


@dataclass
class AlisIadeCiktiModeli:
    iade_no: str
    tarih: str
    durum: str
    durum_etiketi: str | None
    firma: dict[str, str]
    tedarikci: dict[str, str]
    bilgiler: list[tuple[str, str]]
    satirlar: list[dict[str, str]]
    toplamlar: list[tuple[str, str]]
    para_birimi: str = "TRY"
    notlar: list[str] = field(default_factory=list)


def _para(deger, pb: str = "TRY") -> str:
    d = Decimal(str(deger or 0)).quantize(Decimal("0.01"))
    metin = f"{d:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{metin} {'₺' if pb == 'TRY' else pb}"


def _miktar(deger) -> str:
    d = Decimal(str(deger or 0)).normalize()
    metin = f"{d:f}"
    return metin.replace(".", ",")


def cikti_modeli(iade_id: int) -> AlisIadeCiktiModeli:
    from database.alis_iade_faturasi_service import AlisIadeFaturasiService as S
    from database.database import get_session
    from database.models.alis_faturasi import AlisFaturasiSatiri
    from database.models.firma import Firma
    from sqlalchemy import select

    iade = S.getir(int(iade_id))
    if iade is None:
        raise ValueError("İade faturası bulunamadı.")
    pb = (iade.para_birimi or "TRY").upper()
    with get_session() as session:
        firma = session.scalar(select(Firma).order_by(Firma.id))
        kaynak_nolari = {}
        for s in iade.satirlar:
            if s.kaynak_fatura_satiri_id:
                ks = session.get(AlisFaturasiSatiri, int(s.kaynak_fatura_satiri_id))
                if ks is not None and ks.fatura is not None:
                    kaynak_nolari[s.id] = getattr(ks.fatura, "tedarikci_fatura_no", None) or ks.fatura.fatura_no
        firma_bilgi = {
            "unvan": getattr(firma, "unvan", "") or "",
            "vergi": " / ".join(x for x in (getattr(firma, "vergi_dairesi", None), getattr(firma, "vergi_no", None)) if x),
            "adres": getattr(firma, "adres", "") or "",
        }
    cari = iade.cari
    ted = {
        "unvan": getattr(cari, "unvan", "") or "",
        "kod": getattr(cari, "cari_kodu", "") or "",
        "vergi": " / ".join(x for x in (getattr(cari, "vergi_dairesi", None), getattr(cari, "vergi_numarasi", None)) if x),
        "adres": getattr(cari, "adres", "") or "",
        "telefon": getattr(cari, "telefon", "") or "",
    }
    durum = iade.durum or "TASLAK"
    etiket = {"TASLAK": "TASLAK — ONAYLANMAMIŞ BELGE", "İPTAL": "İPTAL EDİLDİ"}.get(durum)
    bilgiler = [("İade No", iade.iade_no), ("İade Tarihi", f"{iade.iade_tarihi:%d.%m.%Y}"),
                ("Depo", iade.depo or "")]
    if iade.kaynak_fatura is not None:
        bilgiler.append(("Kaynak Alış", iade.kaynak_fatura.fatura_no))
    if pb != "TRY":
        bilgiler.append(("Para Birimi / Kur", f"{pb} / {Decimal(str(iade.kur or 1)).normalize():f}".replace(".", ",")))
    if iade.iade_nedeni:
        bilgiler.append(("İade Nedeni", iade.iade_nedeni))
    satirlar = []
    for sira, s in enumerate(iade.satirlar, start=1):
        t = S.satir_tutarlari(s)
        iskontolar = [Decimal(str(getattr(s, a, 0) or 0)) for a in ("iskonto_orani", "iskonto_orani_2", "iskonto_orani_3")]
        isk = "+".join(f"%{_miktar(x)}" for x in iskontolar if x > 0) or "—"
        satirlar.append({
            "sira": str(sira), "kod": s.urun_kodu or "", "barkod": s.barkod or "", "ad": s.urun_adi or "",
            "miktar": _miktar(s.miktar), "birim": s.birim or "Adet",
            "fiyat": _para(s.birim_fiyat, "TRY"), "iskonto": isk,
            "net_fiyat": _para(t["net_birim"], "TRY"), "kdv": f"%{_miktar(s.kdv_orani)}",
            "net_tutar": _para(t["net"], "TRY"), "kaynak": kaynak_nolari.get(s.id, ""),
        })
    t = S.belge_toplami(iade)
    toplamlar = [("Brüt Tutar", _para(t["ara_toplam"])), ("İskonto", _para(t["iskonto"])),
                 ("Net Tutar", _para(t["ara_toplam"] - t["iskonto"])), ("KDV", _para(t["kdv"])),
                 ("GENEL TOPLAM", _para(t["genel_toplam"]))]
    if pb != "TRY" and Decimal(str(iade.doviz_ara_toplam or 0)) > 0:
        toplamlar.append((f"Döviz Net ({pb})", _para(iade.doviz_ara_toplam, pb)))
    notlar = [x for x in (iade.aciklama,) if x]
    if durum == "İPTAL" and iade.iptal_nedeni:
        notlar.append(f"İptal nedeni: {iade.iptal_nedeni}")
    vm = AlisIadeCiktiModeli(
        iade_no=iade.iade_no, tarih=f"{iade.iade_tarihi:%d.%m.%Y}", durum=durum, durum_etiketi=etiket,
        firma=firma_bilgi, tedarikci=ted, bilgiler=bilgiler, satirlar=satirlar, toplamlar=toplamlar,
        para_birimi=pb, notlar=notlar)
    model_guvenli_mi(vm)
    return vm


def model_guvenli_mi(vm: AlisIadeCiktiModeli) -> None:
    """Tedarikçi çıktısına iç maliyet anahtarı sızmasın."""
    for s in vm.satirlar:
        for anahtar in s:
            if any(y in anahtar for y in YASAK_ALANLAR):
                raise AssertionError(f"Tedarikçi çıktısında iç alan: {anahtar}")


def varsayilan_dosya_adi(vm: AlisIadeCiktiModeli, uzanti: str) -> str:
    no = re.sub(r"[^\w\-]+", "_", vm.iade_no or "Yeni").strip("_") or "Yeni"
    try:
        tarih = datetime.strptime(vm.tarih, "%d.%m.%Y").strftime("%Y-%m-%d")
    except (TypeError, ValueError):
        tarih = datetime.now().strftime("%Y-%m-%d")
    return f"Alis_Iade_Faturasi_{no}_{tarih}.{uzanti.lstrip('.')}"


def alt_bilgi_metni(vm: AlisIadeCiktiModeli) -> str:
    return f"Alış İade Faturası {vm.iade_no}" + (f" · {vm.durum_etiketi}" if vm.durum_etiketi else "")


# ------------------------------------------------------------------ HTML / PDF
KOLONLAR = (("sira", "#", "4%", "c"), ("kod", "Kod", "10%", "l"), ("ad", "Ürün", "24%", "l"),
            ("miktar", "Miktar", "7%", "r"), ("birim", "Birim", "6%", "l"), ("fiyat", "Birim Fiyat", "10%", "r"),
            ("iskonto", "İsk.", "7%", "r"), ("net_fiyat", "Net Fiyat", "10%", "r"), ("kdv", "KDV", "5%", "r"),
            ("net_tutar", "Net Tutar", "11%", "r"), ("kaynak", "Kaynak Fatura", "10%", "l"))


def render_html(vm: AlisIadeCiktiModeli) -> str:
    e = html.escape
    filigran = (f'<div class="filigran">{e(vm.durum_etiketi)}</div>' if vm.durum_etiketi else "")
    bilgi = "".join(f"<tr><td class='et'>{e(k)}</td><td>{e(v)}</td></tr>" for k, v in vm.bilgiler)
    basliklar = "".join(f'<th style="width:{w}" class="{h}">{e(b)}</th>' for _a, b, w, h in KOLONLAR)
    satirlar = "".join(
        "<tr>" + "".join(f'<td class="{h}">{e(s.get(a, ""))}</td>' for a, _b, _w, h in KOLONLAR) + "</tr>"
        for s in vm.satirlar)
    toplam = "".join(
        f"<tr class='{'genel' if i == len(vm.toplamlar) - 1 or k == 'GENEL TOPLAM' else ''}'>"
        f"<td>{e(k)}</td><td class='r'>{e(v)}</td></tr>" for i, (k, v) in enumerate(vm.toplamlar))
    notlar = "".join(f"<p>{e(n)}</p>" for n in vm.notlar)
    durum_bandi = f'<div class="bant">{e(vm.durum_etiketi)}</div>' if vm.durum_etiketi else ""
    return f"""<!DOCTYPE html><html lang="tr"><head><meta charset="utf-8"><title>Alış İade Faturası {e(vm.iade_no)}</title>
<style>
@page {{ size: A4; margin: 12mm 10mm 16mm 10mm; }}
body {{ font-family: 'Segoe UI', Arial, sans-serif; color: #172B4D; font-size: 9pt; margin: 0; }}
.ust {{ display: flex; justify-content: space-between; border-bottom: 3px solid #F4C542; padding-bottom: 6px; }}
.ust h1 {{ color: #102A43; font-size: 17pt; margin: 0; }}
.ust .alt {{ color: #627D98; font-size: 9pt; }}
.firma {{ text-align: right; font-size: 8.5pt; }}
.kutular {{ display: flex; gap: 8px; margin: 8px 0; }}
.kutu {{ flex: 1; border: 1px solid #D9E2EC; padding: 6px 8px; }}
.kutu h3 {{ margin: 0 0 4px 0; font-size: 8.5pt; color: #102A43; text-transform: uppercase; }}
.kutu table td {{ padding: 1px 4px 1px 0; }} .et {{ color: #627D98; white-space: nowrap; }}
table.satirlar {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
table.satirlar thead {{ display: table-header-group; }}
table.satirlar th {{ background: #102A43; color: #fff; font-size: 8pt; padding: 4px 3px; }}
table.satirlar td {{ border-bottom: 1px solid #E4E9F0; padding: 3px; font-size: 8pt; word-wrap: break-word; }}
table.satirlar tr {{ page-break-inside: avoid; }}
.r {{ text-align: right; }} .c {{ text-align: center; }} .l {{ text-align: left; }}
.toplam {{ width: 45%; margin-left: auto; margin-top: 8px; border-collapse: collapse; page-break-inside: avoid; }}
.toplam td {{ padding: 3px 6px; border-bottom: 1px solid #E4E9F0; }}
.toplam tr.genel td {{ background: #F4C542; color: #102A43; font-weight: 700; font-size: 10pt; }}
.bant {{ background: #DC2626; color: #fff; text-align: center; font-weight: 700; padding: 3px; margin-top: 6px; }}
.filigran {{ position: fixed; top: 40%; left: 5%; width: 90%; text-align: center; font-size: 46pt;
  color: rgba(220, 38, 38, 0.13); transform: rotate(-25deg); font-weight: 800; z-index: -1; }}
.notlar {{ margin-top: 8px; font-size: 8.5pt; }}
.imza {{ display: flex; justify-content: space-between; margin-top: 28px; page-break-inside: avoid; }}
.imza div {{ width: 40%; border-top: 1px solid #627D98; text-align: center; padding-top: 4px; color: #627D98; }}
</style></head><body>{filigran}
<div class="ust"><div><h1>ALIŞ İADE FATURASI</h1><div class="alt">Tedarikçiye İade · {e(vm.iade_no)} · {e(vm.tarih)}</div></div>
<div class="firma"><b>{e(vm.firma.get('unvan', ''))}</b><br>{e(vm.firma.get('vergi', ''))}<br>{e(vm.firma.get('adres', ''))}</div></div>
{durum_bandi}
<div class="kutular"><div class="kutu"><h3>Tedarikçi</h3><b>{e(vm.tedarikci.get('unvan', ''))}</b><br>
{e(vm.tedarikci.get('vergi', ''))}<br>{e(vm.tedarikci.get('adres', ''))}<br>{e(vm.tedarikci.get('telefon', ''))}</div>
<div class="kutu"><h3>Belge Bilgileri</h3><table>{bilgi}</table></div></div>
<table class="satirlar"><thead><tr>{basliklar}</tr></thead><tbody>{satirlar}</tbody></table>
<table class="toplam">{toplam}</table>
<div class="notlar">{notlar}</div>
<div class="imza"><div>Teslim Eden</div><div>Teslim Alan (Tedarikçi)</div></div>
</body></html>"""


def pdf_olustur(vm: AlisIadeCiktiModeli, hedef: Path) -> Path:
    from invoice_print.pdf_service import html_metnini_pdfe_cevir

    hedef = Path(hedef)
    hedef.parent.mkdir(parents=True, exist_ok=True)
    gecici = hedef.with_name(f".{hedef.stem}_{uuid.uuid4().hex[:6]}.pdf")
    try:
        html_metnini_pdfe_cevir(render_html(vm), gecici, belge_adi="Alış iade faturası")
        try:
            _sayfa_bilgisi_bas(gecici, alt_bilgi_metni(vm))
        except ImportError:
            _LOG.warning("pymupdf yok; sayfa numarası basılamadı: %s", hedef)
        os.replace(gecici, hedef)
    except PermissionError as exc:
        raise ValueError(f"PDF kaydedilemedi; dosya başka bir programda açık olabilir:\n{hedef}") from exc
    finally:
        if gecici.exists():
            try:
                gecici.unlink()
            except OSError:
                pass
    return hedef


# ------------------------------------------------------------------ Word

def docx_olustur(vm: AlisIadeCiktiModeli, hedef: Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, RGBColor

    from invoice_print.irsaliye_docx import (
        BEYAZ,
        LACIVERT,
        SOLUK,
        _alan,
        _alt_cizgi,
        _golge,
        _paragraf,
        _run,
        _satir_bolunmesin,
        _tablo_sabit,
    )

    kolon_mm = (8, 18, 44, 12, 11, 19, 11, 19, 9, 21, 14)
    doc = Document()
    bolum = doc.sections[0]
    bolum.page_height, bolum.page_width = Mm(297), Mm(210)
    bolum.left_margin = bolum.right_margin = Mm(10)
    bolum.top_margin, bolum.bottom_margin = Mm(14), Mm(14)
    # Sayfa üst bilgisi her sayfada tekrar eder: belge adı, no, durum
    ust = bolum.header.paragraphs[0]
    _run(ust, "ALIŞ İADE FATURASI  ", boyut=11, kalin=True, renk=LACIVERT)
    _run(ust, f"{vm.iade_no} · {vm.tarih}", boyut=9, renk=SOLUK)
    if vm.durum_etiketi:
        _run(ust, f"  [{vm.durum_etiketi}]", boyut=9, kalin=True, renk=RGBColor(0xDC, 0x26, 0x26))
    _alt_cizgi(ust, "F4C542", 12)
    alt = bolum.footer.paragraphs[0]
    alt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(alt, f"{vm.firma.get('unvan', '')} · Sayfa ", boyut=7.5, renk=SOLUK)
    _alan(alt, "PAGE")
    _run(alt, " / ", boyut=7.5, renk=SOLUK)
    _alan(alt, "NUMPAGES")

    p = _paragraf(doc, sonra=2)
    _run(p, vm.firma.get("unvan", ""), boyut=10, kalin=True)
    if vm.firma.get("vergi"):
        _run(p, f"  {vm.firma['vergi']}", boyut=8, renk=SOLUK)
    kutular = doc.add_table(rows=1, cols=2)
    _tablo_sabit(kutular, (93, 93))
    sol, sag = kutular.rows[0].cells
    p = _paragraf(sol, ilk=True)
    _run(p, "TEDARİKÇİ\n", boyut=8, kalin=True, renk=SOLUK)
    _run(p, vm.tedarikci.get("unvan", ""), boyut=9.5, kalin=True)
    for anahtar in ("vergi", "adres", "telefon"):
        if vm.tedarikci.get(anahtar):
            _run(_paragraf(sol), vm.tedarikci[anahtar], boyut=8)
    p = _paragraf(sag, ilk=True)
    _run(p, "BELGE BİLGİLERİ", boyut=8, kalin=True, renk=SOLUK)
    for k, v in vm.bilgiler:
        p = _paragraf(sag)
        _run(p, f"{k}: ", boyut=8.5, renk=SOLUK)
        _run(p, v, boyut=8.5, kalin=True)

    tablo = doc.add_table(rows=1, cols=len(KOLONLAR))
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
    tablo.style = "Table Grid"
    baslik = tablo.rows[0]
    _satir_bolunmesin(baslik, baslik=True)
    for hucre, (_a, b, _w, _h) in zip(baslik.cells, KOLONLAR):
        _golge(hucre, "102A43")
        _run(_paragraf(hucre, ilk=True), b, boyut=7.5, kalin=True, renk=BEYAZ)
    hizalar = {"r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": None}
    for s in vm.satirlar:
        satir = tablo.add_row()
        _satir_bolunmesin(satir)
        for hucre, (a, _b, _w, h) in zip(satir.cells, KOLONLAR):
            _run(_paragraf(hucre, ilk=True, hiza=hizalar[h]), s.get(a, ""), boyut=7.5)
    _tablo_sabit(tablo, kolon_mm)

    toplam = doc.add_table(rows=0, cols=2)
    toplam.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for k, v in vm.toplamlar:
        r = toplam.add_row()
        _satir_bolunmesin(r)
        genel = k == "GENEL TOPLAM"
        if genel:
            for h in r.cells:
                _golge(h, "F4C542")
        _run(_paragraf(r.cells[0], ilk=True), k, boyut=8.5, kalin=genel)
        _run(_paragraf(r.cells[1], ilk=True, hiza=WD_ALIGN_PARAGRAPH.RIGHT), v, boyut=8.5, kalin=True)
    _tablo_sabit(toplam, (45, 40))
    for n in vm.notlar:
        _run(_paragraf(doc, once=4), n, boyut=8.5)
    hedef = Path(hedef)
    hedef.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(hedef))
    except PermissionError as exc:
        raise ValueError(f"Word dosyası kaydedilemedi; dosya başka bir programda açık olabilir:\n{hedef}") from exc
    return hedef


def cikti_klasoru() -> Path:
    klasor = Path(tempfile.gettempdir()) / "muhasebe_alis_iade"
    klasor.mkdir(parents=True, exist_ok=True)
    return klasor


def cikti_uret(iade_id: int, tur: str, hedef: Path | None = None) -> Path:
    """tur: 'pdf' | 'docx'. Hedef verilmezse geçici klasöre benzersiz adla yazar."""
    vm = cikti_modeli(iade_id)
    tur = tur.lower().lstrip(".")
    if hedef is None:
        ad = varsayilan_dosya_adi(vm, tur).replace(f".{tur}", f"_{datetime.now():%H%M%S}_{uuid.uuid4().hex[:4]}.{tur}")
        hedef = cikti_klasoru() / ad
    if tur == "pdf":
        return pdf_olustur(vm, hedef)
    if tur == "docx":
        return docx_olustur(vm, hedef)
    raise ValueError(f"Bilinmeyen çıktı türü: {tur}")
