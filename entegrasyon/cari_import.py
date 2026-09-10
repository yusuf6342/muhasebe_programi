"""EvoBulut cari satırlarını yerel Cari modeline aktarma (API JSON / Excel / CSV)."""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from database.cari_service import CariService

# CLI yolu main.py model bootstrap'ını atlar; SatisFaturasi ilişkileri için gerekli
from database.models.satis_siparisi import SatisSiparisi  # noqa: F401
from database.models.satis_irsaliyesi import SatisIrsaliyesi  # noqa: F401
from database.models.satis_faturasi import SatisFaturasi  # noqa: F401

# Excel/CSV başlık → Cari alanı (küçük harf, aksansız eşleşme)
_BASLIK_HARITASI: dict[str, str] = {
    "cari_kodu": "cari_kodu",
    "carikodu": "cari_kodu",
    "musteri_kodu": "cari_kodu",
    "musterikodu": "cari_kodu",
    "hesap_kodu": "cari_kodu",
    "hesapkodu": "cari_kodu",
    "kod": "cari_kodu",
    "a_kod": "cari_kodu",
    "r.a_kod": "cari_kodu",
    "unvan": "unvan",
    "unvani": "unvan",
    "firma_adi": "unvan",
    "firmaadi": "unvan",
    "musteri_adi": "unvan",
    "musteriadi": "unvan",
    "adi": "unvan",
    "ad": "unvan",
    "a_ad": "unvan",
    "r.a_ad": "unvan",
    "a_resmi_ad": "unvan",
    "r.a_resmi_ad": "unvan",
    "musteri": "unvan",
    "cari_turu": "cari_turu",
    "caritur": "cari_turu",
    "tur": "cari_turu",
    "tip": "cari_turu",
    "musteri_tipi": "cari_turu",
    "mustip.a_adi": "cari_turu",
    "vergi_dairesi": "vergi_dairesi",
    "vergidairesi": "vergi_dairesi",
    "a_vergidairesi": "vergi_dairesi",
    "r.a_vergidairesi": "vergi_dairesi",
    "vergi_numarasi": "vergi_numarasi",
    "verginosu": "vergi_numarasi",
    "vergi_no": "vergi_numarasi",
    "vkn": "vergi_numarasi",
    "a_verginosu": "vergi_numarasi",
    "r.a_verginosu": "vergi_numarasi",
    "tc_kimlik": "tc_kimlik",
    "tckn": "tc_kimlik",
    "tc": "tc_kimlik",
    "a_tcno": "tc_kimlik",
    "r.a_tcno": "tc_kimlik",
    "telefon": "telefon",
    "tel": "telefon",
    "gsm": "telefon",
    "telefon2": "telefon2",
    "telefon3": "telefon3",
    "email": "email",
    "e_posta": "email",
    "eposta": "email",
    "a_e_posta": "email",
    "r.a_e_posta": "email",
    "musteri_grubu": "musteri_grubu",
    "grup": "musteri_grubu",
    "g.a_adi": "musteri_grubu",
    "adres": "adres",
    "il": "il",
    "sehir": "il",
    "s2.a_adi": "il",
    "ilce": "ilce",
    "ilce.a_adi": "ilce",
    "ozel_notlar": "ozel_notlar",
    "not": "ozel_notlar",
    "a_not": "ozel_notlar",
    "r.a_not": "ozel_notlar",
}

_CARI_ALANLARI = {
    "cari_kodu",
    "unvan",
    "cari_turu",
    "vergi_dairesi",
    "vergi_numarasi",
    "tc_kimlik",
    "telefon",
    "telefon2",
    "telefon3",
    "email",
    "musteri_grubu",
    "adres",
    "il",
    "ilce",
    "ozel_notlar",
    "aktif",
}


@dataclass
class ImportSonuc:
    eklenen: int = 0
    guncellenen: int = 0
    atlanan: int = 0
    cekilen: int = 0
    hatalar: list[str] = field(default_factory=list)


def _normalize_baslik(baslik: str) -> str:
    metin = unicodedata.normalize("NFKD", str(baslik or "").strip())
    metin = "".join(c for c in metin if not unicodedata.combining(c))
    metin = metin.casefold().replace("ı", "i")
    metin = re.sub(r"[^a-z0-9.]+", "_", metin)
    return metin.strip("_")


def _temiz(deger: Any) -> str:
    if deger is None:
        return ""
    return str(deger).strip()


def _telefonlari_ayir(ham: str) -> tuple[str, str, str]:
    parcalar = [p.strip() for p in re.split(r"[|,;/]+", ham) if p.strip()]
    while len(parcalar) < 3:
        parcalar.append("")
    return parcalar[0][:30], parcalar[1][:30], parcalar[2][:30]


def _cari_turu_normalize(ham: str, varsayilan: str = "Müşteri") -> str:
    t = _normalize_baslik(ham)
    if not t:
        return varsayilan
    if t.startswith("tedarik") or t in {"t", "2", "tedarikci"}:
        return "Tedarikçi"
    if t.startswith("muster") or t in {"m", "1", "0", "musteri"}:
        return "Müşteri"
    if "tedarik" in t:
        return "Tedarikçi"
    return varsayilan


def map_evobulut_api_satir(satir: dict[str, Any], varsayilan_tur: str = "Müşteri") -> dict[str, Any]:
    """Cari Liste API satırını CariService alanlarına çevirir."""

    def al(*anahtarlar: str) -> str:
        for a in anahtarlar:
            if a in satir and _temiz(satir[a]):
                return _temiz(satir[a])
        return ""

    kod = al("R.a_kod", "a_kod", "cari_kodu")
    resmi = al("R.a_resmi_ad", "a_resmi_ad")
    ad = al("R.a_ad", "a_ad", "musteri")
    soy = al("R.a_soy", "a_soy")
    unvan = resmi or ad
    if soy and soy not in unvan:
        unvan = f"{unvan} {soy}".strip() if unvan else soy
    if not unvan:
        unvan = al("musteri", "text")
        if "----" in unvan:
            unvan = unvan.split("----", 1)[0].strip()

    tel1, tel2, tel3 = _telefonlari_ayir(al("Telefon", "telefon"))
    vkn = re.sub(r"\D", "", al("R.a_verginosu", "a_verginosu", "vergi_numarasi"))
    tc = re.sub(r"\D", "", al("R.a_tcno", "a_tcno", "tc_kimlik"))
    if len(vkn) == 11 and not tc:
        tc, vkn = vkn, ""
    if len(vkn) != 10:
        vkn = ""
    if len(tc) != 11:
        tc = ""

    tur = _cari_turu_normalize(
        al("MUSTIP.a_adi", "cari_turu", "a_mus_tip"),
        varsayilan=varsayilan_tur,
    )

    veriler: dict[str, Any] = {
        "cari_kodu": kod[:20] if kod else "",
        "unvan": unvan[:200],
        "cari_turu": tur,
        "vergi_dairesi": al("R.a_vergidairesi", "a_vergidairesi")[:100] or None,
        "vergi_numarasi": vkn or None,
        "tc_kimlik": tc or None,
        "telefon": tel1 or None,
        "telefon2": tel2 or None,
        "telefon3": tel3 or None,
        "email": al("R.a_e_posta", "a_e_posta", "email")[:150] or None,
        "musteri_grubu": al("G.a_adi", "musteri_grubu")[:100] or None,
        "il": al("S2.a_adi", "il")[:50] or None,
        "ilce": al("ILCE.a_adi", "ilce")[:50] or None,
        "ozel_notlar": al("R.a_not", "a_not")[:1000] or None,
        "aktif": True,
    }
    return {k: v for k, v in veriler.items() if v is not None and v != ""}


def map_excel_satir(
    satir: dict[str, Any],
    varsayilan_tur: str = "Müşteri",
) -> dict[str, Any]:
    ham: dict[str, Any] = {}
    for baslik, deger in satir.items():
        alan = _BASLIK_HARITASI.get(_normalize_baslik(str(baslik)))
        if not alan or alan not in _CARI_ALANLARI:
            continue
        if _temiz(deger):
            ham[alan] = _temiz(deger)

    if "telefon" in ham:
        t1, t2, t3 = _telefonlari_ayir(str(ham["telefon"]))
        ham["telefon"] = t1
        ham.setdefault("telefon2", t2)
        ham.setdefault("telefon3", t3)

    if "cari_turu" in ham:
        ham["cari_turu"] = _cari_turu_normalize(str(ham["cari_turu"]), varsayilan_tur)
    else:
        ham["cari_turu"] = varsayilan_tur

    if "vergi_numarasi" in ham:
        vkn = re.sub(r"\D", "", str(ham["vergi_numarasi"]))
        ham["vergi_numarasi"] = vkn if len(vkn) == 10 else None
    if "tc_kimlik" in ham:
        tc = re.sub(r"\D", "", str(ham["tc_kimlik"]))
        ham["tc_kimlik"] = tc if len(tc) == 11 else None

    ham["aktif"] = True
    if "unvan" in ham:
        ham["unvan"] = str(ham["unvan"])[:200]
    if "cari_kodu" in ham:
        ham["cari_kodu"] = str(ham["cari_kodu"])[:20]
    return {k: v for k, v in ham.items() if v is not None and v != ""}


def _risk_limit_temizle(veriler: dict[str, Any]) -> dict[str, Any]:
    """Cari modeline uymayan anahtarları at."""
    return {k: v for k, v in veriler.items() if k in _CARI_ALANLARI or k == "aktif"}


def kaydet_veya_guncelle(veriler: dict[str, Any]) -> str:
    """Returns: 'eklendi' | 'guncellendi' | 'atlandi'."""
    veriler = _risk_limit_temizle(dict(veriler))
    unvan = _temiz(veriler.get("unvan"))
    if not unvan:
        return "atlandi"
    veriler["unvan"] = unvan
    veriler.setdefault("cari_turu", "Müşteri")
    kod = _temiz(veriler.get("cari_kodu"))
    if kod:
        mevcut = CariService.kod_ile_getir(kod)
        if mevcut:
            guncelle = {k: v for k, v in veriler.items() if k != "cari_kodu"}
            CariService.guncelle(mevcut.id, guncelle)
            return "guncellendi"
        veriler["cari_kodu"] = kod
    else:
        veriler.pop("cari_kodu", None)
    CariService.ekle(veriler)
    return "eklendi"


def aktar_satirlari(
    satirlar: Iterable[dict[str, Any]],
    *,
    kaynak: str = "api",
    varsayilan_tur: str = "Müşteri",
) -> ImportSonuc:
    sonuc = ImportSonuc()
    for i, satir in enumerate(satirlar, start=1):
        try:
            if kaynak == "api":
                veriler = map_evobulut_api_satir(satir, varsayilan_tur)
            else:
                veriler = map_excel_satir(satir, varsayilan_tur)
            durum = kaydet_veya_guncelle(veriler)
            if durum == "eklendi":
                sonuc.eklenen += 1
            elif durum == "guncellendi":
                sonuc.guncellenen += 1
            else:
                sonuc.atlanan += 1
        except Exception as exc:  # noqa: BLE001 — satır bazlı hata topla
            sonuc.hatalar.append(f"Satır {i}: {exc}")
            sonuc.atlanan += 1
    return sonuc


def oku_excel_veya_csv(yol: str | Path) -> list[dict[str, Any]]:
    path = Path(yol)
    if not path.is_file():
        raise FileNotFoundError(f"Dosya bulunamadı: {path}")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            basliklar = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            return []
        sonuc: list[dict[str, Any]] = []
        for row in rows:
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            sonuc.append(
                {
                    basliklar[i]: row[i]
                    for i in range(min(len(basliklar), len(row)))
                    if basliklar[i]
                }
            )
        return sonuc
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            okuyucu = csv.DictReader(f)
            return [dict(r) for r in okuyucu]
    raise ValueError("Desteklenen formatlar: .xlsx, .xlsm, .csv")


def aktar_dosyadan(yol: str | Path, varsayilan_tur: str = "Müşteri") -> ImportSonuc:
    return aktar_satirlari(
        oku_excel_veya_csv(yol),
        kaynak="excel",
        varsayilan_tur=varsayilan_tur,
    )


def aktar_api_den(varsayilan_tur: str = "Müşteri") -> ImportSonuc:
    from entegrasyon.evobulut_client import EvobulutClient, load_credentials

    client = EvobulutClient(load_credentials())
    client.login()
    satirlar = client.tum_carileri_cek()
    sonuc = aktar_satirlari(satirlar, kaynak="api", varsayilan_tur=varsayilan_tur)
    sonuc.cekilen = len(satirlar)
    return sonuc
