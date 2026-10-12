"""Kaynağı verilmeyen devir kalemleri için doldurulacak Excel şablonları.

Kart/hesap listeleri plandan önceden yazılır; tutar sütunları boş bırakılır (uydurulmaz).

    python -m entegrasyon.evo_devir.sablonlar --anlik A.json --bakiye B.json --cikti C:\\CinMuhasebeEvoDevir\\sablonlar
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASLIK = Font(bold=True, color="FFFFFF")
DOLGU = PatternFill("solid", fgColor="1F4E78")


def _sayfa(wb, ad, aciklama, sutunlar, satirlar=()):
    ws = wb.create_sheet(ad[:31])
    ws.append([aciklama])
    ws["A1"].font = Font(italic=True)
    ws.append(sutunlar)
    for i, _ in enumerate(sutunlar, start=1):
        h = ws.cell(row=2, column=i)
        h.font, h.fill = BASLIK, DOLGU
        ws.column_dimensions[h.column_letter].width = max(14, len(str(sutunlar[i - 1])) + 4)
    for s in satirlar:
        ws.append(list(s))
    ws.freeze_panes = "A3"


KREDI_SUTUNLARI = [
    "Evo id", "Kredi hesabı", "Aktarılan kalan borç", "Kredi adı / sözleşme no",
    "Kredi türü (İşletme/Spot/Taksitli/Taşıt/Diğer; boş = Diğer)", "Kullandırım tarihi",
    "Taksit no", "Vade tarihi", "Anapara", "Faiz", "BSMV", "KKDF",
]
CEK_SUTUNLARI = [
    "Tür (Alınan çek/Verilen çek/Alınan senet/Verilen senet)", "Belge no", "Cari Evo id", "Cari unvan (bilgi)",
    "Düzenleme tarihi", "Vade tarihi", "Tutar", "Para birimi (boş = TL)", "Kur (döviz ise)",
    "Durum (Portföyde/Bankada tahsilde/Tedarikçiye verildi; boş = Portföyde veya Tedarikçiye verildi)",
    "Banka", "Keşideci / borçlu", "Açıklama",
]

SAYFA_ADI = {
    "Kredi ödeme": "Kredi_Odeme_Plani", "KMH": "KMH_Limit", "Kredi kartı": "Kredi_Karti",
    "POS": "POS_Komisyon", "Çek": "Cek_Senet", "avans": "Avans_Depozito", "Vergi": "Vergi_SGK_Diger",
    "Sabit": "Sabit_Kiymet", "Sermaye": "GM_Acilis", "Depo": "Depo_Stok", "FIFO": "-",
    "grup": "Stok_Ek_Bilgi", "risk": "Cari_Ek_Bilgi", "Pasif": "Pasif_Kartlar", "sipariş": "-",
}


def sablon_sayfasi(kalem: str) -> str:
    return next((v for anahtar, v in SAYFA_ADI.items() if anahtar.casefold() in kalem.casefold()), "-")


def olustur(plan, cikti: Path) -> Path:
    from entegrasyon.evo_devir.eslestirme import AKTARILANLAR, TAMAMLANDI

    wb = Workbook()
    ozet = wb.active
    ozet.title = "Durum"
    ozet.append(["Kalem", "Durum", "Neden", "Şablon sayfası", "Doldurulacak alanlar",
                 "Cin'de aktarılacağı yer", "Çift mali etki kuralı"])
    for c in ozet[1]:
        c.font, c.fill = BASLIK, DOLGU
    ozet_satirlari = []
    for k in plan.kaynak_verilmedi:
        ozet_satirlari.append([k["kalem"], k["durum"], k["neden"], sablon_sayfasi(k["kalem"]), None,
                               k["cin_hedef"], k["cift_etki"]])
    for kalem, yer in AKTARILANLAR:
        ozet.append([kalem, TAMAMLANDI, "Evo son durumu 06.10.2026", "-", "-", yer, "Tek açılış hareketi"])
    for s in ozet_satirlari:
        ozet.append(s)
    for harf, gen in zip("ABCDEFG", (55, 20, 60, 20, 60, 55, 80)):
        ozet.column_dimensions[harf].width = gen

    banka = [f for f in plan.finans if f["evo_tur"] == "banka"]
    _sayfa(wb, "Kredi_Odeme_Plani",
           "Devirden sonra ödenecek her taksit bir satır; Evo id her satırda yazılır. Faiz/BSMV/KKDF yoksa 0 yazın "
           "(boş hücre eksik sayılır). Bir hesabın anapara toplamı 'Aktarılan kalan borç' ile aynı olmalı. "
           "Bakiye zaten aktarıldı: plan yeni borç yazmaz.",
           KREDI_SUTUNLARI,
           [(f["evo_id"], f["ad_maskeli"], str(f["hedefler"].get("KREDILER", ""))) for f in banka
            if f["sinif"] == "KREDILER" and f["hedefler"].get("KREDILER")])
    _sayfa(wb, "KMH_Limit", "Mevduat/POS bağlı hesapların KMH limiti.",
           ["Evo id", "Hesap", "KMH limiti", "Faiz oranı %", "Not"],
           [(f["evo_id"], f["ad_maskeli"]) for f in banka if f["sinif"] in ("MEVDUAT", "POS_BAGLI")])
    _sayfa(wb, "Kredi_Karti", "Kart limiti, kesim/son ödeme günü ve devirden sonra ödenecek taksitler.",
           ["Evo id", "Kart", "Son 4 hane", "Aktarılan borç", "Limit", "Hesap kesim günü", "Son ödeme günü",
            "Ekstre dönem borcu", "Gelecek taksit toplamı"],
           [(f["evo_id"], f["ad_maskeli"], f.get("kart_son4", ""), str(f["hedefler"].get("KREDI_KARTI", "")))
            for f in banka if f["sinif"] == "KREDI_KARTI"])
    _sayfa(wb, "POS_Komisyon", "POS taksit komisyon oranları (1-12).",
           ["Evo id", "POS hesabı", "Valör günü"] + [f"{n} taksit %" for n in range(1, 13)],
           [(f["evo_id"], f["ad_maskeli"], f.get("pos_valor_gun")) for f in banka if f["sinif"] == "POS_BAGLI"])
    _sayfa(wb, "Cek_Senet",
           "Devir tarihinde portföydeki / ödenecek çek ve senetler. Tutar Evo cari bakiyesinde zaten var: "
           "cari hareketi yazılmaz. Cari Evo id için Cari_Listesi sayfasına bakın.",
           CEK_SUTUNLARI)
    _sayfa(wb, "Cari_Listesi", "Bilgi sayfası (doldurulmaz): Cek_Senet sayfasındaki 'Cari Evo id' için.",
           ["Cari Evo id", "Unvan", "Tür"],
           sorted(((c["evo_id"], c["veriler"].get("unvan"), c["tur"]) for c in plan.cariler),
                  key=lambda r: str(r[1] or "").casefold()))
    _sayfa(wb, "Avans_Depozito", "Cari bakiyesinin avans/depozito kısmı (cari açılışından ayrılacak tutar).",
           ["Cari Evo id", "Unvan", "Tür (Alınan avans/Verilen avans/Depozito)", "Tutar", "Not"])
    _sayfa(wb, "Vergi_SGK_Diger", "Devir tarihindeki vergi, SGK, personel ve diğer borç/alacaklar.",
           ["Kalem", "Hesap kodu (öneri)", "Borç", "Alacak", "Vade", "Not"])
    _sayfa(wb, "Sabit_Kiymet", "Demirbaş/taşıt vb. ve birikmiş amortisman.",
           ["Kıymet adı", "Edinim tarihi", "Maliyet", "Birikmiş amortisman", "Faydalı ömür (yıl)", "Not"])
    _sayfa(wb, "GM_Acilis", "Genel muhasebe açılış fişi (sermaye, geçmiş dönem sonuçları vb.).",
           ["Hesap kodu", "Hesap adı", "Borç", "Alacak", "Açıklama"])
    _sayfa(wb, "Depo_Stok", "Depo bazında miktar; boşsa tümü ANA DEPO'da kalır.",
           ["Stok kodu", "Stok adı", "Depo", "Miktar"])
    _sayfa(wb, "Stok_Ek_Bilgi", "Ana/tali/alt grup, ek birim ve katsayı, ek barkodlar.",
           ["Stok kodu", "Ana grup", "Tali grup", "Alt grup", "Ek birim", "Katsayı", "Ek barkod"])
    _sayfa(wb, "Cari_Ek_Bilgi", "Risk limiti, adres, pasif durumu.",
           ["Cari Evo id", "Unvan", "Risk limiti", "Adres", "Pasif mi (E/H)"])
    _sayfa(wb, "Pasif_Kartlar", "Evo listesi pasif kartları döndürmedi; pasif kartların Evo dışa aktarımı.",
           ["Tür (Cari/Stok)", "Evo id", "Kod", "Ad", "Bakiye/Miktar"])
    for satir in ozet.iter_rows(min_row=2):
        sayfa = satir[3].value
        if sayfa and sayfa != "-" and sayfa in wb.sheetnames:
            satir[4].value = ", ".join(str(c.value) for c in wb[sayfa][2] if c.value)
    cikti.mkdir(parents=True, exist_ok=True)
    yol = cikti / f"Evo_Devir_Eksik_Kaynak_Sablonlari_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(yol)
    return yol


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser()
    p.add_argument("--anlik", required=True)
    p.add_argument("--bakiye", required=True)
    p.add_argument("--devir", default="06.10.2026")
    p.add_argument("--veri", default=r"C:\CinMuhasebeEvoDevir\data")
    p.add_argument("--cikti", default=r"C:\CinMuhasebeEvoDevir\sablonlar")
    a = p.parse_args()
    from entegrasyon.evo_devir.oturum import veri_klasoru_sec

    veri_klasoru_sec(a.veri)
    from entegrasyon.evo_devir.eslestirme import plan_olustur

    plan = plan_olustur(a.anlik, datetime.strptime(a.devir, "%d.%m.%Y").date(), a.bakiye)
    print(olustur(plan, Path(a.cikti)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
