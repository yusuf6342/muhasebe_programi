"""Evo devir ayrıntıları: kredi ödeme planı, çek/senet, KMH / kredi kartı / POS tanımları.

Bakiyeler ``aktarim`` ile açılış hareketi olarak yazıldı. Bu modül aynı tutar için yeni kasa/banka
hareketi, cari hareket veya muhasebe fişi yazmaz; ayrıntıyı ilgili açılış hesabına bağlar:

* Kredi: ``BankaKredisi`` + bekleyen ``BankaKrediTaksit`` satırları, Evo kredi banka kartına. Hesabın
  kalan anapara toplamı aktarılan KREDİLER bakiyesine eşit değilse o hesabın planı yazılmaz. Sonraki
  taksit ödemesi KREDİLER'i anapara kadar bir kez azaltır.
* Çek/senet: ``CekSenetEvrak`` + KAYIT hareketi; cari hareketi ve fiş yok (tutar cari açılışında var).
* KMH limiti, kredi kartı limiti / kesim / son ödeme günü, POS valör ve komisyon: yalnız tanım alanı.

Boş zorunlu hücre "eksik" sayılır, o kayıt yazılmaz; tahmini tutar/tarih veya sıfır konmaz. Aynı
ayrıntı (tür + anahtar) etkin bir aktarımda varsa yeniden yazılmaz.

    python -m entegrasyon.evo_devir.detay_aktarim onizle  --sablon S.xlsx
    python -m entegrasyon.evo_devir.detay_aktarim aktar   --sablon S.xlsx
    python -m entegrasyon.evo_devir.detay_aktarim kontrol
    python -m entegrasyon.evo_devir.detay_aktarim geri_al --batch EVODETAY-...
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from entegrasyon.evo_devir import oturum as evo_oturum

VARSAYILAN_VERI = os.path.expandvars(r"%LOCALAPPDATA%\CinMuhasebeEvoDevir\data")
PROGRAM_KLASORU = os.path.expandvars(r"%LOCALAPPDATA%\Programs\CinMuhasebeEvoDevir")
MODUL = "evo_devir_detay"
DEVIR_ISARETI = "[EVO DEVİR]"
KULLANICI = "Evo devir aktarımı"
SIFIR = Decimal("0")
KURUS = Decimal("0.01")

# Anahtarlar turkce_normalize biçiminde (İ/ı/i → i, ş → s, ç → c ...)
KREDI_TURU = {"isletme": "ISLETME", "spot": "SPOT", "taksitli": "TAKSITLI", "tasit": "TASIT", "diger": "DIGER"}
CEK_TURU = {"alinan cek": ("CEK", "ALINAN"), "verilen cek": ("CEK", "VERILEN"),
            "alinan senet": ("SENET", "ALINAN"), "verilen senet": ("SENET", "VERILEN")}


class Eksik(ValueError):
    pass


# ---------------------------------------------------------------- hücre okuma

def _anahtar(baslik) -> str:
    return str(baslik).split(" (")[0].strip().casefold()


def _bos(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _secim(v) -> str:
    from database.turkce_normalize import turkce_normalize

    return " ".join(turkce_normalize(_metin(v)).split())


def _metin(v) -> str:
    if _bos(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _sayi(v, alan: str, ondalik: int = 2) -> Decimal:
    if _bos(v):
        raise Eksik(alan)
    try:
        if isinstance(v, (int, float, Decimal)):
            d = Decimal(str(v))
        else:
            metin = str(v).strip().replace(" ", "")
            if "," in metin:
                metin = metin.replace(".", "").replace(",", ".")
            d = Decimal(metin)
    except InvalidOperation:
        raise ValueError(f"{alan} sayı değil: {v!r}") from None
    adim = Decimal(1).scaleb(-ondalik)
    if d != d.quantize(adim):
        raise ValueError(f"{alan} en fazla {ondalik} ondalık olmalı: {v!r}")
    if d < 0:
        raise ValueError(f"{alan} negatif olamaz: {v!r}")
    return d.quantize(adim)


def _tam(v, alan: str, en_az: int, en_cok: int) -> int:
    d = _sayi(v, alan, 0)
    if not en_az <= d <= en_cok:
        raise ValueError(f"{alan} {en_az}-{en_cok} arasında olmalı: {v!r}")
    return int(d)


def _tarih(v, alan: str) -> date:
    if _bos(v):
        raise Eksik(alan)
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for bicim in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(v).strip(), bicim).date()
        except ValueError:
            continue
    raise ValueError(f"{alan} GG.AA.YYYY olmalı: {v!r}")


def sayfa_oku(wb, ad: str) -> list[tuple[int, dict]]:
    """Başlıklar 2. satırda; 3. satırdan itibaren boş olmayan satırlar."""
    if ad not in wb.sheetnames:
        return []
    ws = wb[ad]
    basliklar = [_anahtar(c.value) if not _bos(c.value) else None for c in ws[2]]
    satirlar = []
    for satir in ws.iter_rows(min_row=3):
        degerler = {b: c.value for b, c in zip(basliklar, satir) if b}
        if any(not _bos(v) for v in degerler.values()):
            satirlar.append((satir[0].row, degerler))
    return satirlar


def _sablon_sha(yol: Path) -> str:
    return hashlib.sha256(yol.read_bytes()).hexdigest()


def _kisa(metin: str) -> str:
    return hashlib.sha1(metin.encode("utf-8")).hexdigest()[:8]


def _json(deger) -> str:
    return json.dumps(deger, ensure_ascii=False, default=str)


# ---------------------------------------------------------------- haritalar

def _detay_harita() -> dict[tuple[str, str], dict]:
    from sqlalchemy import select

    from database.database import get_session
    from excel_aktarim.models import ImportBatch, ImportRow

    harita = {}
    with get_session() as s:
        q = (select(ImportRow, ImportBatch.batch_kodu)
             .join(ImportBatch, ImportBatch.id == ImportRow.batch_id)
             .where(ImportBatch.modul == MODUL, ImportBatch.durum.in_(("basarili", "kismi")),
                    ImportRow.durum == "uygulandi"))
        for row, kod in s.execute(q).all():
            ham = json.loads(row.ham_json or "{}")
            harita[(ham.get("tur"), ham.get("anahtar"))] = {"hedef_id": row.hedef_id, "batch": kod,
                                                           "tablo": row.hedef_tablo}
    return harita


def _alt_hesap(s, kart_id: int, alt: str):
    from sqlalchemy import select

    from database.models.finans import FinansHesabi

    return s.scalar(select(FinansHesabi).where(FinansHesabi.banka_karti_id == kart_id,
                                               FinansHesabi.alt_hesap_turu == alt))


# ---------------------------------------------------------------- hazırlık (yazmadan)

class Hazirlik:
    """Şablonu okur, doğrular, yazılacak işleri ve satır sonuçlarını üretir. Veritabanına yazmaz."""

    def __init__(self, sablon: Path, devir: date):
        from openpyxl import load_workbook

        self.sablon = sablon
        self.devir = devir
        self.wb = load_workbook(sablon, data_only=True, read_only=False)
        from entegrasyon.evo_devir.aktarim import _etkin_harita

        self.evo = _etkin_harita()
        self.detay = _detay_harita()
        self.isler: list[dict] = []
        self.rapor: list[dict] = []

    def not_et(self, sayfa, satir, tur, anahtar, durum, mesaj):
        self.rapor.append({"sayfa": sayfa, "satir": satir, "tur": tur, "anahtar": anahtar,
                           "durum": durum, "mesaj": mesaj})

    def is_ekle(self, is_: dict):
        onceki = self.detay.get((is_["tur"], is_["anahtar"]))
        if onceki:
            is_["onceki"] = onceki
        self.isler.append(is_)

    def banka_karti(self, evo_id: str) -> tuple[int, dict]:
        es = self.evo.get(("banka", evo_id))
        if not es or not es.get("hedef_id"):
            raise ValueError(f"Evo banka id {evo_id} devir aktarımında yok")
        return int(es["hedef_id"]), es.get("ek") or {}

    def hepsi(self):
        from database.database import get_session

        with get_session() as s:
            self.krediler(s)
            self.cekler(s)
            self.kmh()
            self.kredi_kartlari(s)
            self.pos()
        self.kapsam_disi()
        return self

    # -- kredi
    def krediler(self, s):
        from sqlalchemy import select

        from database.banka_kredi_service import BankaKrediService
        from database.finans_service import FinansService
        from database.models.finans import BankaKredisi

        sayfa = "Kredi_Odeme_Plani"
        zorunlu = ("kredi adı / sözleşme no", "kullandırım tarihi", "taksit no", "vade tarihi", "anapara",
                   "faiz", "bsmv", "kkdf")
        hesaplar: dict[str, dict] = {}
        for no, r in sayfa_oku(self.wb, sayfa):
            evo_id = _metin(r.get("evo id"))
            if all(_bos(r.get(k)) for k in zorunlu):
                continue  # hesap ve aktarılan borç önceden yazılı bilgi satırı
            h = hesaplar.setdefault(evo_id, {"gruplar": {}, "sorun": [], "satirlar": []})
            h["satirlar"].append(no)
            eksik = [k for k in zorunlu if _bos(r.get(k))] + ([] if evo_id else ["evo id"])
            if eksik:
                h["sorun"].append(f"satır {no}: boş {', '.join(eksik)}")
                self.not_et(sayfa, no, "kredi", evo_id, "eksik", "Boş: " + ", ".join(eksik))
                continue
            try:
                ad = _metin(r["kredi adı / sözleşme no"])
                tur_metin = _secim(r.get("kredi türü"))
                tur = KREDI_TURU.get(tur_metin, "DIGER" if not tur_metin else None)
                if tur is None:
                    raise ValueError(f"Kredi türü tanınmadı: {r.get('kredi türü')!r}")
                kullandirim = _tarih(r["kullandırım tarihi"], "Kullandırım tarihi")
                if kullandirim > self.devir:
                    raise ValueError("Kullandırım devirden sonra; bu kredi normal ekrandan girilir")
                taksit = {"taksit_no": _tam(r["taksit no"], "Taksit no", 1, 600),
                          "vade_tarihi": _tarih(r["vade tarihi"], "Vade tarihi"),
                          "anapara": _sayi(r["anapara"], "Anapara"), "faiz": _sayi(r["faiz"], "Faiz"),
                          "bsmv": _sayi(r["bsmv"], "BSMV"), "kkdf": _sayi(r["kkdf"], "KKDF"), "satir": no}
            except ValueError as hata:
                h["sorun"].append(f"satır {no}: {hata}")
                self.not_et(sayfa, no, "kredi", evo_id, "hata", str(hata))
                continue
            g = h["gruplar"].setdefault(ad.casefold(), {"ad": ad, "kredi_turu": tur, "kullandirim": kullandirim,
                                                        "taksitler": []})
            if (g["kredi_turu"], g["kullandirim"]) != (tur, kullandirim):
                h["sorun"].append(f"satır {no}: aynı kredide farklı tür/kullandırım tarihi")
                self.not_et(sayfa, no, "kredi", evo_id, "hata", "Aynı kredide farklı tür/kullandırım tarihi")
                continue
            if any(t["taksit_no"] == taksit["taksit_no"] for t in g["taksitler"]):
                h["sorun"].append(f"satır {no}: taksit no tekrar ediyor")
                self.not_et(sayfa, no, "kredi", evo_id, "hata", "Taksit no tekrar ediyor")
                continue
            g["taksitler"].append(taksit)

        for evo_id, h in hesaplar.items():
            if not evo_id:
                continue
            if h["sorun"]:
                self.not_et(sayfa, h["satirlar"], "kredi_hesap", evo_id, "eksik",
                            "Hesabın planı yazılmadı (bütün satırlar tamamlanmalı): " + "; ".join(h["sorun"])[:900])
                continue
            try:
                kart_id, _ = self.banka_karti(evo_id)
                kredi_hesap = _alt_hesap(s, kart_id, "KREDILER")
                if kredi_hesap is None:
                    raise ValueError("Bu Evo hesabında KREDİLER alt hesabı yok")
                bakiye = Decimal(str(FinansService.bakiye(kredi_hesap))).quantize(KURUS)
                mevcut = s.scalars(select(BankaKredisi).where(BankaKredisi.banka_karti_id == kart_id,
                                                               BankaKredisi.durum == "KULLANDIRILDI")).all()
                mevcut_kalan = sum((BankaKrediService._kalan_anapara(s, k) for k in mevcut), SIFIR)
            except ValueError as hata:
                self.not_et(sayfa, h["satirlar"], "kredi_hesap", evo_id, "hata", str(hata))
                continue
            yeni, toplam_yeni = [], SIFIR
            for anahtar_ad, g in h["gruplar"].items():
                g["taksitler"].sort(key=lambda t: t["taksit_no"])
                ana = sum((t["anapara"] for t in g["taksitler"]), SIFIR)
                is_ = {"tur": "kredi", "anahtar": f"{evo_id}|{anahtar_ad}", "sayfa": sayfa,
                       "satirlar": [t["satir"] for t in g["taksitler"]], "evo_id": evo_id, "kart_id": kart_id,
                       "kredi_hesap_id": kredi_hesap.id, "ana_para": ana, **g}
                if ana <= 0:
                    self.not_et(sayfa, is_["satirlar"], "kredi", is_["anahtar"], "hata", "Anapara toplamı sıfır")
                    continue
                if (is_["tur"], is_["anahtar"]) in self.detay:
                    self.is_ekle(is_)
                    continue
                yeni.append(is_)
                toplam_yeni += ana
            if not yeni:
                continue
            if mevcut_kalan + toplam_yeni != bakiye:
                self.not_et(sayfa, h["satirlar"], "kredi_hesap", evo_id, "hata",
                            f"Kalan anapara toplamı {mevcut_kalan + toplam_yeni} ≠ aktarılan KREDİLER bakiyesi "
                            f"{bakiye} (mevcut kredi kalanı {mevcut_kalan}); plan yazılmadı")
                continue
            for is_ in yeni:
                self.is_ekle(is_)

    # -- çek / senet
    def cekler(self, s):
        from sqlalchemy import select

        from database.cek_senet_service import CekSenetService
        from database.models.cek_senet import (
            DURUM_BANKAYA_TAHSILE, DURUM_PORTFOYDE, DURUM_TEDARIKCIYE_VERILDI, ISLEM_YONU_ALINAN, CekSenetEvrak,
        )

        sayfa = "Cek_Senet"
        zorunlu = ("tür", "belge no", "cari evo id", "düzenleme tarihi", "vade tarihi", "tutar")
        gorulen = set()
        for no, r in sayfa_oku(self.wb, sayfa):
            eksik = [k for k in zorunlu if _bos(r.get(k))]
            belge = _metin(r.get("belge no"))
            if eksik:
                self.not_et(sayfa, no, "cek", belge, "eksik", "Boş: " + ", ".join(eksik))
                continue
            try:
                tur = CEK_TURU.get(_secim(r["tür"]))
                if tur is None:
                    raise ValueError(f"Tür tanınmadı: {r['tür']!r}")
                basit, yon = tur
                evo_cari = _metin(r["cari evo id"])
                es = self.evo.get(("cari", evo_cari))
                if not es or not es.get("hedef_id"):
                    raise ValueError(f"Cari Evo id {evo_cari} devir aktarımında yok")
                cari_id = int(es["hedef_id"])
                duzenleme = _tarih(r["düzenleme tarihi"], "Düzenleme tarihi")
                if duzenleme > self.devir:
                    raise ValueError("Düzenleme devirden sonra; bu evrak normal ekrandan girilir")
                pb = _metin(r.get("para birimi")).upper() or "TL"
                pb = "TL" if pb in ("TRY", "TL") else pb
                kur = _sayi(r.get("kur"), "Kur", 6) if pb != "TL" else Decimal("1")
                durum_metin = _secim(r.get("durum"))
                if yon == ISLEM_YONU_ALINAN:
                    durum = {"": DURUM_PORTFOYDE, "portfoyde": DURUM_PORTFOYDE,
                             "bankada tahsilde": DURUM_BANKAYA_TAHSILE}.get(durum_metin)
                else:
                    durum = {"": DURUM_TEDARIKCIYE_VERILDI,
                             "tedarikciye verildi": DURUM_TEDARIKCIYE_VERILDI}.get(durum_metin)
                if durum is None:
                    raise ValueError(f"Durum bu tür için geçersiz: {r.get('durum')!r}")
                veri = CekSenetService._verileri_dogrula({
                    "islem_yonu": yon, "basit_tur": basit, "evrak_no": belge, "duzenleme_tarihi": duzenleme,
                    "vade_tarihi": _tarih(r["vade tarihi"], "Vade tarihi"),
                    "tutar": _sayi(r["tutar"], "Tutar"), "doviz_turu": pb, "kur": kur, "cari_id": cari_id,
                    "banka_adi": _metin(r.get("banka")),
                    ("kesideci_adi" if yon == ISLEM_YONU_ALINAN else "lehtar"): _metin(r.get("keşideci / borçlu")),
                    "aciklama": _metin(r.get("açıklama")), "created_by": KULLANICI,
                })
            except Eksik as hata:
                self.not_et(sayfa, no, "cek", belge, "eksik", f"Boş: {hata}")
                continue
            except ValueError as hata:
                self.not_et(sayfa, no, "cek", belge, "hata", str(hata))
                continue
            anahtar = f"{yon}|{basit}|{belge.casefold()}|{cari_id}"
            if anahtar in gorulen:
                self.not_et(sayfa, no, "cek", anahtar, "hata", "Şablonda aynı evrak iki kez var")
                continue
            gorulen.add(anahtar)
            if ("cek", anahtar) not in self.detay:
                var = s.scalar(select(CekSenetEvrak.id).where(
                    CekSenetEvrak.evrak_no == veri["evrak_no"], CekSenetEvrak.islem_yonu == yon,
                    CekSenetEvrak.cari_id == cari_id, CekSenetEvrak.aktif.is_(True)))
                if var:
                    self.not_et(sayfa, no, "cek", anahtar, "hata",
                                f"Bu evrak zaten kayıtlı (evrak id {var}); devir kaydı yazılmadı")
                    continue
            self.is_ekle({"tur": "cek", "anahtar": anahtar, "sayfa": sayfa, "satirlar": [no],
                          "veri": veri, "durum": durum})

    # -- KMH limiti
    def kmh(self):
        sayfa = "KMH_Limit"
        for no, r in sayfa_oku(self.wb, sayfa):
            evo_id = _metin(r.get("evo id"))
            if _bos(r.get("kmh limiti")) and _bos(r.get("faiz oranı %")) and _bos(r.get("not")):
                continue
            if _bos(r.get("kmh limiti")):
                self.not_et(sayfa, no, "kmh", evo_id, "eksik", "Boş: kmh limiti")
                continue
            try:
                kart_id, _ = self.banka_karti(evo_id)
                limit = _sayi(r["kmh limiti"], "KMH limiti")
            except ValueError as hata:
                self.not_et(sayfa, no, "kmh", evo_id, "hata", str(hata))
                continue
            if not _bos(r.get("faiz oranı %")):
                self.not_et(sayfa, no, "kmh", evo_id, "bilgi", "Cin'de KMH faiz oranı alanı yok; yazılmadı")
            self.is_ekle({"tur": "kmh", "anahtar": f"{evo_id}|{limit}", "sayfa": sayfa, "satirlar": [no],
                          "tablo": "banka_kartlari", "hedef_id": kart_id, "alanlar": {"kmh_limiti": limit}})

    # -- kredi kartı tanımı
    def kredi_kartlari(self, s):
        from database.finans_service import FinansService

        sayfa = "Kredi_Karti"
        for no, r in sayfa_oku(self.wb, sayfa):
            evo_id = _metin(r.get("evo id"))
            alanlar_ham = {"kart_limiti": r.get("limit"), "hesap_kesim_gunu": r.get("hesap kesim günü"),
                           "son_odeme_gunu": r.get("son ödeme günü")}
            ekstre, gelecek = r.get("ekstre dönem borcu"), r.get("gelecek taksit toplamı")
            if all(_bos(v) for v in (*alanlar_ham.values(), ekstre, gelecek)):
                continue
            try:
                kart_id, ek = self.banka_karti(evo_id)
                kk_id = ek.get("kredi_karti_id")
                if not kk_id:
                    raise ValueError("Bu Evo hesabında kredi kartı tanımı yok")
                alanlar = {}
                if not _bos(alanlar_ham["kart_limiti"]):
                    alanlar["kart_limiti"] = _sayi(alanlar_ham["kart_limiti"], "Limit")
                for alan, etiket in (("hesap_kesim_gunu", "Hesap kesim günü"), ("son_odeme_gunu", "Son ödeme günü")):
                    if not _bos(alanlar_ham[alan]):
                        alanlar[alan] = _tam(alanlar_ham[alan], etiket, 1, 31)
                if not _bos(ekstre) or not _bos(gelecek):
                    if _bos(ekstre) or _bos(gelecek):
                        self.not_et(sayfa, no, "kk_borc", evo_id, "eksik",
                                    "Ekstre dönem borcu ve gelecek taksit toplamı birlikte girilmeli (kontrol için)")
                    else:
                        hesap = _alt_hesap(s, kart_id, "KREDI_KARTI")
                        borc = Decimal(str(FinansService.bakiye(hesap))).quantize(KURUS) if hesap else SIFIR
                        toplam = _sayi(ekstre, "Ekstre dönem borcu") + _sayi(gelecek, "Gelecek taksit toplamı")
                        self.not_et(sayfa, no, "kk_borc", evo_id, "bilgi" if toplam == borc else "uyari",
                                    f"Ekstre + gelecek taksit {toplam} / aktarılan KK borcu {borc}"
                                    f"{'' if toplam == borc else ' — FARK'}; kayıt yazılmaz")
            except ValueError as hata:
                self.not_et(sayfa, no, "kredi_karti", evo_id, "hata", str(hata))
                continue
            if alanlar:
                self.is_ekle({"tur": "kredi_karti", "anahtar": f"{evo_id}|{_json(alanlar)}", "sayfa": sayfa,
                              "satirlar": [no], "tablo": "kredi_karti_tanimlari", "hedef_id": int(kk_id),
                              "alanlar": alanlar})

    # -- POS
    def pos(self):
        sayfa = "POS_Komisyon"
        for no, r in sayfa_oku(self.wb, sayfa):
            evo_id = _metin(r.get("evo id"))
            oranlar_ham = {n: r.get(f"{n} taksit %") for n in range(1, 13)}
            valor_dolu = not _bos(r.get("valör günü"))
            if not any(not _bos(v) for v in oranlar_ham.values()):
                continue  # valör günü devirde yazıldı; oran yoksa yapılacak iş yok
            try:
                kart_id, _ = self.banka_karti(evo_id)
                oranlar = {n: _sayi(v, f"{n} taksit %", 4) for n, v in oranlar_ham.items() if not _bos(v)}
                if any(o > 100 for o in oranlar.values()):
                    raise ValueError("Komisyon oranı %100'den büyük olamaz")
                alanlar = {"pos_valor_gun": _tam(r["valör günü"], "Valör günü", 0, 60)} if valor_dolu else {}
            except ValueError as hata:
                self.not_et(sayfa, no, "pos", evo_id, "hata", str(hata))
                continue
            self.is_ekle({"tur": "pos", "anahtar": f"{evo_id}|{_json(alanlar)}|{_json(oranlar)}", "sayfa": sayfa,
                          "satirlar": [no], "tablo": "banka_kartlari", "hedef_id": kart_id, "alanlar": alanlar,
                          "oranlar": oranlar})

    def kapsam_disi(self):
        for ad in ("Avans_Depozito", "Vergi_SGK_Diger", "Sabit_Kiymet", "GM_Acilis", "Depo_Stok",
                   "Stok_Ek_Bilgi", "Cari_Ek_Bilgi", "Pasif_Kartlar"):
            dolu = sayfa_oku(self.wb, ad)
            if dolu:
                self.not_et(ad, [n for n, _ in dolu], "kapsam_disi", ad, "bekliyor",
                            f"{len(dolu)} satır dolu; bu araç bu sayfayı aktarmaz (ayrı adım)")


# ---------------------------------------------------------------- yazıcılar (her iş kendi işleminde)

def _kredi_yaz(s, is_: dict, devir: date) -> tuple[str, int, dict]:
    from database.banka_kredi_service import BankaKrediService
    from database.models.finans import BankaKredisi, BankaKrediTaksit

    taksitler = is_["taksitler"]
    kredi = BankaKredisi(
        belge_no=f"EVB-KR-{is_['evo_id']}-{_kisa(is_['anahtar'])}",
        banka_karti_id=is_["kart_id"], kredi_adi=is_["ad"][:120], kredi_turu=is_["kredi_turu"],
        ana_para=is_["ana_para"], toplam_faiz=sum((t["faiz"] for t in taksitler), SIFIR),
        toplam_masraf=sum((t["bsmv"] + t["kkdf"] for t in taksitler), SIFIR), taksit_sayisi=len(taksitler),
        kullandirim_tarihi=is_["kullandirim"], ilk_taksit_tarihi=min(t["vade_tarihi"] for t in taksitler),
        bitis_tarihi=max(t["vade_tarihi"] for t in taksitler), odeme_hesap_turu="MEVDUAT",
        aciklama=(f"{DEVIR_ISARETI} {devir:%d.%m.%Y} itibarıyla kalan anapara. Borç KREDİLER açılışında "
                  f"(EVB-FA-B{is_['evo_id']}-KREDILER); kullandırım hareketi ve fişi yazılmadı.")[:500],
        durum="KULLANDIRILDI", plan_versiyon=1, aktif=True, para_birimi="TRY",
    )
    s.add(kredi)
    s.flush()
    for t in taksitler:
        tk = BankaKrediTaksit(kredi_id=kredi.id, taksit_no=t["taksit_no"], vade_tarihi=t["vade_tarihi"],
                              durum="BEKLIYOR", odenen_anapara=SIFIR, odenen_faiz=SIFIR, odenen_masraf=SIFIR,
                              odenen_tutar=SIFIR)
        BankaKrediService._taksit_bilesen_yaz(tk, {k: t[k] for k in ("anapara", "faiz", "bsmv", "kkdf")})
        tk.durum = BankaKrediService.taksit_durum_hesapla(
            tk.vade_tarihi, SIFIR, BankaKrediService.taksit_bilesenleri(tk)["toplam"])
        s.add(tk)
    BankaKrediService._gunluk(s, "EVO_DEVIR", kredi_id=kredi.id,
                              detay=f"{kredi.belge_no} / kalan anapara {kredi.ana_para} / {len(taksitler)} taksit")
    s.flush()
    return "banka_kredileri", kredi.id, {"belge_no": kredi.belge_no, "ana_para": str(kredi.ana_para),
                                         "taksit": len(taksitler)}


def _cek_yaz(s, is_: dict, devir: date) -> tuple[str, int, dict]:
    from database.cek_senet_service import CekSenetService
    from database.models.cek_senet import DURUM_BANKAYA_TAHSILE, DURUM_PORTFOYDE, CekSenetEvrak

    v, durum = is_["veri"], is_["durum"]
    yer = {DURUM_PORTFOYDE: "Portföy", DURUM_BANKAYA_TAHSILE: "Bankada tahsilde"}.get(durum, "Tedarikçi / Verilen")
    portfoy = CekSenetService.portfoy_no_uret(v["evrak_turu"], v["islem_yonu"], yil=v["duzenleme_tarihi"].year,
                                              session=s)
    aciklama = f"{DEVIR_ISARETI} Tutar cari açılış bakiyesinde; cari hareketi ve fiş yazılmadı."
    if v["aciklama"]:
        aciklama = f"{aciklama} {v['aciklama']}"
    alanlar = {k: v[k] for k in (
        "cari_id", "evrak_turu", "islem_yonu", "evrak_no", "seri_no", "duzenleme_tarihi", "vade_tarihi",
        "doviz_turu", "doviz_tutari", "kur", "tl_tutari", "banka_adi", "banka_subesi", "sube_kodu", "hesap_no",
        "iban", "cek_no", "kesideci_adi", "kesideci_vergi_tc", "keside_yeri", "hesap_sahibi", "lehtar",
        "teslim_eden", "senet_no", "borclu_adi", "borclu_vergi_tc", "kefil", "duzenleme_yeri", "odeme_yeri",
        "ozel_not", "created_by")}
    evrak = CekSenetEvrak(**alanlar, portfoy_no=portfoy, tahsil_edilen_tutar=SIFIR, kalan_tutar=v["tl_tutari"],
                          durum=durum, bulundugu_yer=yer, aciklama=aciklama[:500], aktif=True,
                          created_at=datetime.now(), son_islem_tarihi=devir)
    s.add(evrak)
    s.flush()
    CekSenetService._hareket_ekle(s, evrak, onceki_durum=None, yeni_durum=durum, islem_turu="KAYIT",
                                  tutar=v["tl_tutari"], kullanici=KULLANICI,
                                  aciklama=f"Evo devir {devir:%d.%m.%Y}: cari hareketi ve fiş yazılmadı")
    s.flush()
    return "cek_senet_evraklari", evrak.id, {"portfoy_no": portfoy, "tl_tutari": str(v["tl_tutari"])}


def _alan_yaz(s, is_: dict, degisim: list) -> tuple[str, int, dict]:
    from sqlalchemy import select

    from database.models.finans import BankaKarti, KrediKartiTanimi, PosTaksitKomisyon

    model = {"banka_kartlari": BankaKarti, "kredi_karti_tanimlari": KrediKartiTanimi}[is_["tablo"]]
    nesne = s.get(model, is_["hedef_id"])
    if nesne is None:
        raise ValueError(f"{is_['tablo']} id {is_['hedef_id']} bulunamadı")
    onceki = {k: getattr(nesne, k) for k in is_["alanlar"]}
    for k, deger in is_["alanlar"].items():
        setattr(nesne, k, deger)
    if is_["alanlar"]:
        degisim.append((is_["tablo"], nesne.id, onceki, is_["alanlar"]))
    for n, oran in (is_.get("oranlar") or {}).items():
        mevcut = s.scalar(select(PosTaksitKomisyon).where(PosTaksitKomisyon.banka_karti_id == nesne.id,
                                                          PosTaksitKomisyon.taksit_sayisi == n))
        if mevcut is None:
            s.add(PosTaksitKomisyon(banka_karti_id=nesne.id, taksit_sayisi=n, komisyon_orani=oran))
        else:
            degisim.append(("pos_taksit_komisyonlari", mevcut.id, {"komisyon_orani": mevcut.komisyon_orani},
                            {"komisyon_orani": oran}))
            mevcut.komisyon_orani = oran
    s.flush()
    return is_["tablo"], nesne.id, {"onceki": onceki, "yeni": is_["alanlar"], "oranlar": is_.get("oranlar")}


# ---------------------------------------------------------------- komutlar

def _program_acik_mi() -> list[str]:
    try:
        cikti = subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             "Get-Process CinMuhasebe -ErrorAction SilentlyContinue | ForEach-Object { $_.Path }"],
            capture_output=True, text=True, timeout=30).stdout
    except (OSError, subprocess.SubprocessError):
        return []
    kok = os.path.normcase(PROGRAM_KLASORU)
    return [y.strip() for y in cikti.splitlines() if y.strip() and os.path.normcase(y.strip()).startswith(kok)]


def onizle(a) -> dict:
    h = Hazirlik(Path(a.sablon), _gun(a.devir)).hepsi()
    return _ozet(h, yazildi=False)


def _ozet(h: Hazirlik, yazildi: bool, uygulanan: list | None = None) -> dict:
    durumlar: dict[str, int] = {}
    for r in h.rapor:
        durumlar[r["durum"]] = durumlar.get(r["durum"], 0) + 1
    isler = [{"tur": i["tur"], "anahtar": i["anahtar"], "sayfa": i["sayfa"], "satirlar": i["satirlar"],
              "durum": ("onceki_aktarimda_var" if i.get("onceki") else ("yazilacak" if not yazildi else "?"))}
             for i in h.isler]
    return {"sablon": str(h.sablon), "devir": h.devir.isoformat(), "satir_durumlari": durumlar,
            "isler": uygulanan if uygulanan is not None else isler, "rapor": h.rapor}


def aktar(a) -> dict:
    from database.banka_kredi_service import BankaKrediService
    from database.cek_senet_service import CekSenetService
    from database.database import get_session
    from entegrasyon.evo_devir.aktarim import _en_buyuk_idler, _yeni_satirlar
    from excel_aktarim.models import ImportBatch, ImportChange, ImportRow

    acik = _program_acik_mi()
    if acik:
        raise SystemExit(f"Önce Evo Devir Deneme programını kapatın: {acik}")
    BankaKrediService.schema_hazirla()
    CekSenetService.schema_hazirla()
    sablon = Path(a.sablon)
    devir = _gun(a.devir)
    h = Hazirlik(sablon, devir).hepsi()
    kod = f"EVODETAY-{datetime.now():%Y%m%d%H%M%S}"
    with get_session() as s:
        batch = ImportBatch(batch_kodu=kod, modul=MODUL, import_tipi="evo_devir_ayrinti", dosya_adi=sablon.name,
                            dosya_yolu=str(sablon)[:500], durum="taslak", kullanici_adi=KULLANICI,
                            basladi_at=datetime.now(),
                            notlar=_json({"sablon_sha256": _sablon_sha(sablon), "devir_tarihi": devir.isoformat()}))
        s.add(batch)
        s.flush()
        batch_id = batch.id
        once = _en_buyuk_idler(s)

    sonuclar, degisim = [], []
    for is_ in h.isler:
        if is_.get("onceki"):
            sonuclar.append({**_is_ozeti(is_), "durum": "atla",
                             "mesaj": f"Önceki aktarımda yazıldı ({is_['onceki']['batch']})",
                             "tablo": is_["onceki"]["tablo"], "hedef_id": is_["onceki"]["hedef_id"], "ek": {}})
            continue
        try:
            with get_session() as s:
                if is_["tur"] == "kredi":
                    tablo, hid, ek = _kredi_yaz(s, is_, devir)
                elif is_["tur"] == "cek":
                    tablo, hid, ek = _cek_yaz(s, is_, devir)
                else:
                    yerel: list = []
                    tablo, hid, ek = _alan_yaz(s, is_, yerel)
            if is_["tur"] not in ("kredi", "cek"):
                degisim.extend(yerel)
            sonuclar.append({**_is_ozeti(is_), "durum": "uygulandi", "mesaj": "", "tablo": tablo, "hedef_id": hid,
                             "ek": ek})
        except Exception as hata:  # iş hatası: diğer işler sürer
            sonuclar.append({**_is_ozeti(is_), "durum": "hata", "mesaj": f"{type(hata).__name__}: {hata}",
                             "tablo": None, "hedef_id": None, "ek": {}})

    with get_session() as s:
        for tablo, idler in _yeni_satirlar(s, once).items():
            s.add_all(ImportChange(batch_id=batch_id, hedef_tablo=tablo, hedef_id=i, islem="insert") for i in idler)
        for tablo, hid, onceki, sonraki in degisim:
            s.add(ImportChange(batch_id=batch_id, hedef_tablo=tablo, hedef_id=hid, islem="update",
                               onceki_json=_json(onceki), sonraki_json=_json(sonraki)))
        for no, r in enumerate(sonuclar, start=1):
            s.add(ImportRow(batch_id=batch_id, satir_no=no,
                            ham_json=_json({"tur": r["tur"], "anahtar": r["anahtar"], "sayfa": r["sayfa"],
                                            "satirlar": r["satirlar"]}),
                            eslenen_json=_json(r["ek"]), durum=r["durum"], mesaj=(r["mesaj"] or "")[:2000] or None,
                            hedef_tablo=r["tablo"], hedef_id=r["hedef_id"],
                            islem="insert" if r["durum"] == "uygulandi" else "skip"))
        b = s.get(ImportBatch, batch_id)
        b.toplam_satir = len(sonuclar)
        b.basarili_satir = b.eklenen_satir = sum(r["durum"] == "uygulandi" for r in sonuclar)
        b.hata_satir = sum(r["durum"] == "hata" for r in sonuclar)
        b.atlanan_satir = sum(r["durum"] == "atla" for r in sonuclar)
        b.durum = "kismi" if b.hata_satir else "basarili"
        b.bitti_at = datetime.now()
        b.hata_ozeti = _json([r for r in h.rapor if r["durum"] in ("eksik", "hata")])[:4000] or None
    sonuc = _ozet(h, yazildi=True, uygulanan=sonuclar)
    sonuc["batch"] = kod
    sonuc["kontrol"] = kontrol(a)
    return sonuc


def _is_ozeti(is_: dict) -> dict:
    return {k: is_[k] for k in ("tur", "anahtar", "sayfa", "satirlar")}


def geri_al(a) -> dict:
    """Aktarımın eklediği satırları siler, güncellediği alanları eski değerine döndürür.

    Kredide ödeme, çek/senette KAYIT dışında hareket varsa geri alınmaz (gerçek işlem silinmesin).
    """
    from sqlalchemy import select

    from database.database import Base, get_session
    from database.models.cek_senet import CekSenetHareket
    from database.models.finans import BankaKrediOdeme
    from entegrasyon.evo_devir.aktarim import _satirlari_sil
    from excel_aktarim.models import ImportBatch, ImportChange, ImportRow

    with get_session() as s:
        b = s.scalar(select(ImportBatch).where(ImportBatch.batch_kodu == a.batch))
        if b is None or b.modul != MODUL:
            raise SystemExit(f"Ayrıntı aktarımı bulunamadı: {a.batch}")
        if b.durum not in ("basarili", "kismi"):
            raise SystemExit(f"Aktarım durumu {b.durum}; geri alınamaz")
        degisim = s.scalars(select(ImportChange).where(ImportChange.batch_id == b.id,
                                                       ImportChange.geri_alindi.is_(False))).all()
        ekle = {}
        for ch in degisim:
            if ch.islem == "insert":
                ekle.setdefault(ch.hedef_tablo, []).append(int(ch.hedef_id))
        kredi_idler = ekle.get("banka_kredileri", [])
        cek_idler = ekle.get("cek_senet_evraklari", [])
        if kredi_idler and s.scalar(select(BankaKrediOdeme.id).where(BankaKrediOdeme.kredi_id.in_(kredi_idler))):
            raise SystemExit("Devir kredilerinde ödeme kaydı var; önce ödemeleri iptal edin")
        if cek_idler and s.scalar(select(CekSenetHareket.id).where(CekSenetHareket.evrak_id.in_(cek_idler),
                                                                   CekSenetHareket.islem_turu != "KAYIT")):
            raise SystemExit("Devir çek/senetlerinde sonraki işlem var; geri alınamaz")
        silinen = _satirlari_sil(s, ekle)
        tablolar = {t.name: t for t in Base.metadata.sorted_tables}
        donen = 0
        for ch in degisim:
            if ch.islem == "update" and ch.hedef_tablo in tablolar:
                onceki = json.loads(ch.onceki_json or "{}")
                t = tablolar[ch.hedef_tablo]
                s.execute(t.update().where(t.c.id == ch.hedef_id).values(**onceki))
                donen += 1
            ch.geri_alindi = True
        for r in s.scalars(select(ImportRow).where(ImportRow.batch_id == b.id, ImportRow.durum == "uygulandi")):
            r.durum = "geri_alindi"
        b.durum = "geri_alindi"
        b.notlar = (b.notlar or "") + f"\nGeri alındı: {datetime.now().isoformat(timespec='seconds')}"
    return {"batch": a.batch, "silinen": silinen, "eski_degerine_donen_alan": donen}


def kontrol(a) -> dict:
    """Ayrıntıların bakiyelere ikinci kez etki etmediğini ve açılış hesaplarıyla tutarlılığını denetler."""
    from sqlalchemy import func, select

    from database.banka_kredi_service import BankaKrediService
    from database.database import get_session
    from database.finans_service import FinansService
    from database.models.cek_senet import CekSenetEvrak, CekSenetHareket
    from database.models.finans import BankaKredisi, FinansHareketi
    from database.models.genel_muhasebe import MuhasebeFisi
    from entegrasyon.evo_devir.aktarim import _etkin_harita

    devir = _gun(a.devir)
    evo = _etkin_harita()
    detay = _detay_harita()
    devir_kredi = sorted({v["hedef_id"] for (t, _), v in detay.items() if t == "kredi"})
    devir_cek = sorted({v["hedef_id"] for (t, _), v in detay.items() if t == "cek"})
    sonuc: dict = {"kredi_hesaplari": [], "sorunlar": []}
    with get_session() as s:
        for (tur, evo_id), es in sorted(evo.items(), key=lambda x: str(x[0])):
            if tur != "banka" or not es.get("hedef_id"):
                continue
            hesap = _alt_hesap(s, int(es["hedef_id"]), "KREDILER")
            if hesap is None:
                continue
            bakiye = Decimal(str(FinansService.bakiye(hesap))).quantize(KURUS)
            krediler = s.scalars(select(BankaKredisi).where(BankaKredisi.banka_karti_id == int(es["hedef_id"]),
                                                             BankaKredisi.durum == "KULLANDIRILDI")).all()
            if bakiye == 0 and not krediler:
                continue
            kalan = sum((BankaKrediService._kalan_anapara(s, k) for k in krediler), SIFIR)
            durum = "tutarli" if kalan == bakiye else ("plan_bekleniyor" if not krediler else "FARK")
            sonuc["kredi_hesaplari"].append({"evo_id": evo_id, "kredi_bakiyesi": str(bakiye),
                                             "plan_kalan_anapara": str(kalan), "kredi": len(krediler),
                                             "durum": durum})
            if durum == "FARK":
                sonuc["sorunlar"].append(f"Evo {evo_id}: KREDİLER {bakiye} ≠ plan kalanı {kalan}")

        belge = [b for (b,) in s.execute(select(BankaKredisi.belge_no).where(BankaKredisi.id.in_(devir_kredi)))]
        kull = s.scalar(select(func.count()).select_from(FinansHareketi).where(FinansHareketi.belge_no.in_(belge)))
        kredi_fis = s.scalar(select(func.count()).select_from(MuhasebeFisi).where(
            MuhasebeFisi.kaynak_turu.like("%kullandirim%"), MuhasebeFisi.kaynak_id.in_(devir_kredi)))
        kayitlar = s.scalars(select(CekSenetHareket).where(CekSenetHareket.evrak_id.in_(devir_cek),
                                                           CekSenetHareket.islem_turu == "KAYIT")).all()
        cek_cari = sum(1 for h in kayitlar if h.cari_hareket_id or h.muhasebe_fisi_id)
        cek_fis = s.scalar(select(func.count()).select_from(MuhasebeFisi).where(
            MuhasebeFisi.kaynak_turu.like("cek_senet%"), MuhasebeFisi.kaynak_id.in_([h.id for h in kayitlar])))
        mukerrer = s.execute(select(CekSenetEvrak.evrak_no, CekSenetEvrak.islem_yonu, CekSenetEvrak.cari_id,
                                    func.count()).where(CekSenetEvrak.aktif.is_(True))
                             .group_by(CekSenetEvrak.evrak_no, CekSenetEvrak.islem_yonu, CekSenetEvrak.cari_id)
                             .having(func.count() > 1)).all()
        normal_cek = s.execute(select(CekSenetEvrak.id, CekSenetEvrak.evrak_no).join(
            CekSenetHareket, CekSenetHareket.evrak_id == CekSenetEvrak.id).where(
            CekSenetHareket.islem_turu == "KAYIT", CekSenetHareket.cari_hareket_id.is_not(None),
            CekSenetEvrak.duzenleme_tarihi <= devir, CekSenetEvrak.id.not_in(devir_cek or [0]))).all()
        normal_kredi = s.execute(select(BankaKredisi.id, BankaKredisi.belge_no).where(
            BankaKredisi.kullandirim_tarihi <= devir, BankaKredisi.id.not_in(devir_kredi or [0]),
            BankaKredisi.durum != "IPTAL")).all()
    sonuc.update({
        "devir_kredi": len(devir_kredi), "devir_cek_senet": len(devir_cek),
        "devir_kredisine_kullandirim_hareketi": kull, "devir_kredisine_muhasebe_fisi": kredi_fis,
        "devir_cekine_cari_hareket_veya_fis": cek_cari + (cek_fis or 0),
        "mukerrer_cek_senet": [list(map(str, r)) for r in mukerrer],
        "devir_oncesi_tarihli_normal_cek_kaydi": [list(map(str, r)) for r in normal_cek],
        "devir_oncesi_tarihli_normal_kredi_kaydi": [list(map(str, r)) for r in normal_kredi],
    })
    for anahtar, mesaj in (("devir_kredisine_kullandirim_hareketi", "Devir kredisine kullandırım hareketi var"),
                           ("devir_kredisine_muhasebe_fisi", "Devir kredisine kullandırım fişi var"),
                           ("devir_cekine_cari_hareket_veya_fis", "Devir çek/senedine cari hareket/fiş var")):
        if sonuc[anahtar]:
            sonuc["sorunlar"].append(mesaj)
    if mukerrer:
        sonuc["sorunlar"].append("Mükerrer çek/senet evrakı")
    if normal_cek:
        sonuc["sorunlar"].append("Devir öncesi tarihli çek/senet normal ekrandan girilmiş: cariye ikinci etki riski")
    if normal_kredi:
        sonuc["sorunlar"].append("Devir öncesi kullandırılmış kredi normal ekrandan girilmiş: KREDİLER'e ikinci etki")
    sonuc["temiz"] = not sonuc["sorunlar"]
    return sonuc


def _gun(metin: str) -> date:
    return datetime.strptime(metin, "%d.%m.%Y").date()


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser()
    p.add_argument("komut", choices=("onizle", "aktar", "kontrol", "geri_al"))
    p.add_argument("--sablon")
    p.add_argument("--devir", default="06.10.2026")
    p.add_argument("--batch")
    p.add_argument("--veri", default=VARSAYILAN_VERI)
    p.add_argument("--unvan", default=evo_oturum.VARSAYILAN_UNVAN)
    p.add_argument("--rapor", help="Sonucu bu JSON dosyasına da yaz")
    a = p.parse_args()
    if a.komut in ("onizle", "aktar") and not a.sablon:
        p.error("--sablon gerekli")
    if a.komut == "geri_al" and not a.batch:
        p.error("--batch gerekli")

    evo_oturum.veri_klasoru_sec(a.veri)
    bilgi = evo_oturum.firma_oturumu_ac(a.unvan)
    print(f"Firma: {bilgi['unvan']} ({bilgi['kod']}) {bilgi['yol']}", flush=True)
    sonuc = {"onizle": onizle, "aktar": aktar, "kontrol": kontrol, "geri_al": geri_al}[a.komut](a)
    metin = json.dumps(sonuc, ensure_ascii=False, indent=1, default=str)
    if a.rapor:
        Path(a.rapor).parent.mkdir(parents=True, exist_ok=True)
        Path(a.rapor).write_text(metin, encoding="utf-8")
    print(metin)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
