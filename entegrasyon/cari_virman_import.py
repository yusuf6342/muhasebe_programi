"""Cari virman aktarımı.

EvoBulut OpenAPI'de Cari Virman için REST uç noktası YOKTUR (modül id 1097 menüde
geçer; /CariVirman/base vb. 404/ERR). Bu yüzden aktarım Excel/CSV ile yapılır.

Beklenen sütunlar (başlık esnek):
  kaynak_kodu / kaynak / from
  hedef_kodu / hedef / to
  tarih
  tutar
  aciklama (opsiyonel)
  belge_no (opsiyonel; yoksa EVB-VRM-{satır} veya otomatik VRM-)

Yerel: CariService.virman_yap (belge_no destekli).
"""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from database.cari_service import CariService

# CLI bootstrap
from database.models.satis_siparisi import SatisSiparisi  # noqa: F401
from database.models.satis_irsaliyesi import SatisIrsaliyesi  # noqa: F401
from database.models.satis_faturasi import SatisFaturasi  # noqa: F401


@dataclass
class VirmanImportSonuc:
    cekilen: int = 0
    olusturulan: int = 0
    atlanan: int = 0
    hatalar: list[str] = field(default_factory=list)


def _temiz(deger: Any) -> str:
    if deger is None:
        return ""
    return str(deger).strip()


def _baslik_norm(ad: str) -> str:
    s = unicodedata.normalize("NFKD", ad or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.casefold().replace("ı", "i")
    return re.sub(r"[^a-z0-9]+", "", s)


_ALAN = {
    "kaynakkodu": "kaynak_kodu",
    "kaynak": "kaynak_kodu",
    "from": "kaynak_kodu",
    "fromkod": "kaynak_kodu",
    "carikaynak": "kaynak_kodu",
    "hedefkodu": "hedef_kodu",
    "hedef": "hedef_kodu",
    "to": "hedef_kodu",
    "tokod": "hedef_kodu",
    "carihedef": "hedef_kodu",
    "tarih": "tarih",
    "islemtarihi": "tarih",
    "tutar": "tutar",
    "miktar": "tutar",
    "aciklama": "aciklama",
    "ack": "aciklama",
    "not": "aciklama",
    "belgeno": "belge_no",
    "belge": "belge_no",
    "fisno": "belge_no",
}


def _tarih(deger: Any) -> date | None:
    metin = _temiz(deger)
    if not metin:
        return None
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(
                metin[:19] if " " in metin and "%H" in fmt else metin[:10],
                fmt,
            ).date()
        except ValueError:
            continue
    return None


def _tutar(deger: Any) -> Decimal | None:
    metin = _temiz(deger).replace(" ", "").replace(",", ".")
    if not metin:
        return None
    try:
        return Decimal(metin)
    except (InvalidOperation, ValueError):
        return None


def _satirlari_oku(yol: Path) -> list[dict[str, Any]]:
    suf = yol.suffix.lower()
    if suf in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel için openpyxl gerekli: pip install openpyxl") from exc
        wb = load_workbook(yol, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        basliklar = [_baslik_norm(str(h or "")) for h in rows[0]]
        sonuc = []
        for row in rows[1:]:
            ham: dict[str, Any] = {}
            for i, h in enumerate(basliklar):
                if not h or i >= len(row):
                    continue
                alan = _ALAN.get(h)
                if alan:
                    ham[alan] = row[i]
            if any(_temiz(v) for v in ham.values()):
                sonuc.append(ham)
        return sonuc

    # CSV
    with yol.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        sonuc = []
        for row in reader:
            ham = {}
            for k, v in row.items():
                alan = _ALAN.get(_baslik_norm(k or ""))
                if alan:
                    ham[alan] = v
            if any(_temiz(v) for v in ham.values()):
                sonuc.append(ham)
        return sonuc


def aktar_dosyadan(yol: str | Path) -> VirmanImportSonuc:
    path = Path(yol)
    if not path.is_file():
        raise FileNotFoundError(str(path))
    satirlar = _satirlari_oku(path)
    sonuc = VirmanImportSonuc(cekilen=len(satirlar))
    for i, satir in enumerate(satirlar, start=1):
        try:
            kaynak_kod = _temiz(satir.get("kaynak_kodu"))
            hedef_kod = _temiz(satir.get("hedef_kodu"))
            if not kaynak_kod or not hedef_kod:
                raise ValueError("kaynak_kodu ve hedef_kodu zorunlu")
            kaynak = CariService.kod_ile_getir(kaynak_kod)
            hedef = CariService.kod_ile_getir(hedef_kod)
            if kaynak is None:
                raise ValueError(f"Kaynak cari yok: {kaynak_kod}")
            if hedef is None:
                raise ValueError(f"Hedef cari yok: {hedef_kod}")
            tarih = _tarih(satir.get("tarih")) or date.today()
            tutar = _tutar(satir.get("tutar"))
            if tutar is None or tutar <= 0:
                raise ValueError("Geçersiz tutar")
            belge = _temiz(satir.get("belge_no")) or f"EVB-VRM-{i}"
            if CariService.islem_belge_var_mi(belge):
                sonuc.atlanan += 1
                continue
            CariService.virman_yap(
                kaynak.id,
                hedef.id,
                tarih,
                tutar,
                aciklama=_temiz(satir.get("aciklama")) or None,
                belge_no=belge[:50],
            )
            sonuc.olusturulan += 1
        except Exception as exc:  # noqa: BLE001
            sonuc.hatalar.append(f"Satır {i}: {exc}")
            sonuc.atlanan += 1
    return sonuc


def api_destekleniyor_mu() -> bool:
    """EvoBulut Cari Virman REST yok — her zaman False."""
    return False
