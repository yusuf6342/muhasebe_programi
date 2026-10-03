"""Geçmişte kartta TANIMSIZ birimle kesilmiş belge satırları — ürün bazında inceleme listesi.

SALT OKUNUR (``mode=ro``). Geçmiş satırlar değiştirilmez; yeni belgelerde bu birimler artık
kaydı engeller. Liste, kullanıcının hangi ürünlerde birim katsayısı tanımlaması gerektiğini gösterir.

Çalıştırma:
    .venv\\Scripts\\python.exe tools\\stok_tanimsiz_birim_listesi.py <db_kopyası> [--cikti dosya.xlsx]
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from tools.stok_birim_inceleme_listesi import satirlari_topla  # noqa: E402

VARSAYILAN_CIKTI = ROOT / "inceleme" / "stoklar" / "Tanimsiz_Birim_Satirlari_Inceleme.xlsx"
TANIMLAMA_YOLU = "Stok Kartı → «Barkod & Birimler» sekmesi → Alternatif Birimler"


def tanimsiz_satirlar(db: Path) -> list[dict]:
    return [s for s in satirlari_topla(db) if s["kart_katsayisi"] is None]


def sorun_aciklamasi(s: dict) -> str:
    islenen = "stoğa 1:1 işlenmiş" if s["stok_etkisi"] is not None else "stoğa işlenmemiş"
    return (f"«{s['belge_birimi'] or '(boş)'}» birimi stok kartında tanımlı değil (temel birim "
            f"{s['temel_birim'] or '?'}); {islenen}. {TANIMLAMA_YOLU} bölümünde katsayı tanımlanmalı.")


def excel_yaz(satirlar: list[dict], cikti: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    def sayi(v):
        return float(v) if isinstance(v, Decimal) else v

    wb = Workbook()
    ozet = wb.active
    ozet.title = "Ürün Özeti"
    ozet.append(["Stok kodu", "Stok adı", "Temel birim", "Tanımsız birim(ler)", "Satır", "Belge sayısı",
                 "Toplam belge miktarı", "Stoğa işlenmiş satır", "İlk tarih", "Son tarih", "Yapılacak"])
    grup: dict[str, list[dict]] = defaultdict(list)
    for s in satirlar:
        grup[s["stok_kodu"]].append(s)
    for kod, liste in sorted(grup.items(), key=lambda x: (-len(x[1]), x[0])):
        birimler = sorted({s["belge_birimi"] or "" for s in liste})
        tarihler = sorted(str(s["tarih"]) for s in liste if s["tarih"])
        ozet.append([
            kod, liste[0]["stok_adi"], liste[0]["temel_birim"], ", ".join(birimler), len(liste),
            len({(s["belge_turu"], s["belge_no"]) for s in liste}),
            sayi(sum((s["miktar"] for s in liste), Decimal("0"))),
            sum(1 for s in liste if s["stok_etkisi"] is not None),
            tarihler[0] if tarihler else "", tarihler[-1] if tarihler else "",
            f"{TANIMLAMA_YOLU}: " + ", ".join(f"1 {b} = ? {liste[0]['temel_birim']}" for b in birimler),
        ])

    detay = wb.create_sheet("Satırlar")
    basliklar = [
        ("belge_no", "Belge no"), ("tarih", "Tarih"), ("stok_kodu", "Ürün kodu"), ("stok_adi", "Ürün adı"),
        ("belge_birimi", "Birim"), ("miktar", "Miktar"), ("sorun", "Sorun açıklaması"),
        ("belge_turu", "Belge türü"), ("durum", "Durum"), ("cari", "Cari"), ("temel_birim", "Temel birim"),
        ("kullanilan_katsayi", "Stoğa işlenen katsayı"), ("kullanilan_aciklama", "Katsayı kaynağı"),
        ("stok_etkisi", "Stok etkisi (belge+ürün)"), ("belge_id", "Belge ID"), ("satir_id", "Satır ID"),
    ]
    detay.append([b for _, b in basliklar])
    for s in satirlar:
        detay.append([sayi(sorun_aciklamasi(s) if k == "sorun" else s[k]) for k, _ in basliklar])

    bilgi = wb.create_sheet("Açıklama")
    for satir in (
        ["Salt okunur tarama; hiçbir kayıt değiştirilmedi. Gerçek veritabanının kopyasından üretildi."],
        ["Bu satırlarda belge birimi stok kartında tanımlı değil; geçmişte stoğa 1:1 işlenmişti."],
        ["Yeni alış/satış faturası, irsaliye, sipariş, teklif, satın alma talebi, iade ve hızlı satışta "
         "tanımsız birim artık kaydı engeller."],
        [f"Düzeltme yolu: {TANIMLAMA_YOLU} bölümünde birim ve katsayıyı tanımlayın."],
        ["Geçmiş stok farkı (1:1 işlenmiş satırlar) için ayrı kullanıcı kararı gerekir; otomatik düzeltme yapılmadı."],
    ):
        bilgi.append(satir)

    for ws in (ozet, detay):
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
    ozet.column_dimensions["B"].width = 40
    ozet.column_dimensions["K"].width = 70
    bilgi.column_dimensions["A"].width = 110
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("db")
    ap.add_argument("--cikti", default=str(VARSAYILAN_CIKTI))
    a = ap.parse_args()
    satirlar = tanimsiz_satirlar(Path(a.db))
    excel_yaz(satirlar, Path(a.cikti))
    tur: dict[str, int] = defaultdict(int)
    for s in satirlar:
        tur[s["belge_turu"]] += 1
    print(f"Tanımsız birimli satır: {len(satirlar)}  ürün: {len({s['stok_kodu'] for s in satirlar})}")
    print("Belge türü:", dict(tur))
    print("Birimler:", dict(sorted(((b, sum(1 for s in satirlar if s['belge_birimi'] == b))
                                    for b in {s['belge_birimi'] for s in satirlar}), key=lambda x: -x[1])))
    print("Çıktı:", a.cikti)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
