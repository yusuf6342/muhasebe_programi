"""Tanımsız «Adet» birimli belge satırlarının ürünleri için Adet katsayısı önerisi ve (yalnız kesin) tanımı.

Kesin sayılma koşulu (hepsi): ürün adında tek ve açık adet katsayısı ("5 Lİ", "PK:5", "x12", "500 Adet"),
aynı belge türünde (satış/satış, alış/alış) ana birim ile Adet fiyat oranının bu katsayıyla ±%10 uyumu ve
hiçbir belge/kart fiyat oranının çelişmemesi. Tek kanıt veya çelişki → belirsiz. Toplu 1 verilmez.

Varsayılan kuru çalışmadır: Excel üretir, veri yazmaz. ``--uygula`` yalnız kesin olanlara karta alternatif
birim ekler (``StokService._birimleri_sessiona_yaz``, tek transaction); geçmiş belge satırları değişmez.
Gerçek DB'de CinMuhasebe.exe kontrolü + doğrulanmış yedek zorunludur.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from statistics import median

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("MUHASEBE_DB_DIR", tempfile.mkdtemp(prefix="cin_adet_"))

from tools.stok_birim_inceleme_listesi import satirlari_topla, tr_kucuk  # noqa: E402

VARSAYILAN_DB = Path(os.environ["LOCALAPPDATA"]) / "MuhasebeProgrami" / "data" / "muhasebe.db"
CIKTI = ROOT / "inceleme" / "stoklar" / "Adet_Katsayi_Onerileri.xlsx"
TOLERANS = Decimal("0.10")
SATIR_TABLOLARI = ("satis_faturasi_satirlari", "alis_faturasi_satirlari", "satis_iade_faturasi_satirlari")
AD_DESENLERI = (
    re.compile(r"(?<![\d.,])(\d+)\s*['’]?\s*(?:L[İI]|LI|LU|LÜ|li|lı|lu|lü)\b", re.IGNORECASE),
    re.compile(r"\bPK\s*[:.]?\s*(\d+)", re.IGNORECASE),
    re.compile(r"(?<![\d.,])[xX×]\s?(\d+)\b"),
    re.compile(r"(?<![\d.,])(\d+)\s*AD(?:ET)?\b", re.IGNORECASE),
)


def adet_anahtari(b) -> str:
    return tr_kucuk(b).replace("ı", "i")


def ad_katsayilari(ad: str) -> list[tuple[int, str]]:
    bulunan = []
    for desen in AD_DESENLERI:
        for m in desen.finditer(ad or ""):
            n = int(m.group(1))
            if n > 1:
                bulunan.append((n, m.group(0).strip()))
    return bulunan


def _oran(a, b):
    return (Decimal(str(a)) / Decimal(str(b))) if a and b else None


def urun_kanitlari(db: Path) -> tuple[list[dict], dict]:
    tum = [s for s in satirlari_topla(db) if s["kart_katsayisi"] is None]
    ham = Counter(s["belge_birimi"] for s in tum if adet_anahtari(s["belge_birimi"]) == "adet")
    adet = [s for s in tum if adet_anahtari(s["belge_birimi"]) == "adet"]
    sayim = {"tanimsiz_satir": len(tum), "tanimsiz_tur": dict(Counter(s["belge_turu"] for s in tum)),
             "adet_satir": len(adet), "adet_ham_yazimlar": dict(ham),
             "adet_urun": len({s["stok_kodu"] for s in adet})}
    c = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    grup = defaultdict(list)
    for s in adet:
        grup[s["stok_kodu"]].append(s)
    sonuc = []
    for kod, ss in sorted(grup.items()):
        k = c.execute("SELECT id, stok_adi, birim, aciklama FROM stok_kartlari WHERE stok_kodu=?", (kod,)).fetchone()
        alt = [f"{r['birim_adi']}={Decimal(str(r['carpan'])).normalize()}{'' if r['aktif'] else ' (pasif)'}"
               for r in c.execute("SELECT birim_adi, carpan, aktif FROM stok_birimleri WHERE stok_id=?", (k["id"],))]
        kart_fiyat = {r["fiyat_adi"]: Decimal(str(r["tutar"])) for r in
                      c.execute("SELECT fiyat_adi, tutar FROM stok_fiyatlari WHERE stok_id=?", (k["id"],))}
        fiyat = defaultdict(list)
        for tur, tab in (("satış", "satis_faturasi_satirlari"), ("alış", "alis_faturasi_satirlari")):
            for r in c.execute(f"SELECT birim, birim_fiyat FROM {tab} WHERE urun_kodu=? AND birim_fiyat>0", (kod,)):
                fiyat[(tur, adet_anahtari(r["birim"]))].append(Decimal(str(r["birim_fiyat"])))
        ana = adet_anahtari(k["birim"])
        belge_oran, kart_oran = {}, {}
        for tur in ("satış", "alış"):
            a, b = fiyat.get((tur, "adet")), fiyat.get((tur, ana))
            if a and b:
                belge_oran[tur] = _oran(median(b), median(a))
        if fiyat.get(("satış", "adet")) and kart_fiyat.get("SATIŞ FİYATI 1"):
            kart_oran["satış"] = _oran(kart_fiyat["SATIŞ FİYATI 1"], median(fiyat[("satış", "adet")]))
        if fiyat.get(("alış", "adet")) and kart_fiyat.get("ALIŞ FİYATI"):
            kart_oran["alış"] = _oran(kart_fiyat["ALIŞ FİYATI"], median(fiyat[("alış", "adet")]))

        adk = ad_katsayilari(k["stok_adi"])
        notk = ad_katsayilari(k["aciklama"] or "")
        n_kume = {n for n, _ in adk + notk}
        kanit, guven, oneri, grup_adi = [], "belirsiz", None, "b"
        if ana == "adet":
            grup_adi = "a"
        if adk:
            kanit.append("Ürün adında: " + ", ".join(m for _, m in adk))
        if notk:
            kanit.append("Kart notunda: " + ", ".join(m for _, m in notk))
        if alt:
            kanit.append("Mevcut alternatifler: " + ", ".join(alt))
        oran_metin = [f"belge/{t} {o.quantize(Decimal('0.01'))}" for t, o in belge_oran.items()]
        oran_metin += [f"kart/{t} {o.quantize(Decimal('0.01'))}" for t, o in kart_oran.items()]
        if oran_metin:
            kanit.append("Fiyat oranı (ana/Adet): " + "; ".join(oran_metin))
        if ana in ("ad",):
            kanit.append("Ana birim «Ad» Adet kısaltması olabilir; kesin eşitlik için fiyat kanıtı gerekir")
        if len(n_kume) == 1:
            n = Decimal(next(iter(n_kume)))
            uyumlu = [o for o in belge_oran.values() if abs(o / n - 1) <= TOLERANS]
            celiskili = [o for o in list(belge_oran.values()) + list(kart_oran.values()) if abs(o / n - 1) > TOLERANS]
            if uyumlu and not celiskili:
                guven, oneri = "kesin", n
                kanit.append(f"Ad katsayısı {n} ile belge fiyat oranı uyumlu (±%10), çelişki yok")
            elif celiskili:
                kanit.append(f"Ad katsayısı {n} ile fiyat oranı ÇELİŞİYOR")
            else:
                kanit.append("Tek kanıt (yalnız ad); belge fiyat ilişkisi yok")
        elif len(n_kume) > 1:
            kanit.append("Adda birden çok farklı sayı — belirsiz")
        else:
            kanit.append("Adda/kartta açık adet katsayısı yok" + (" — yalnız fiyat oranı (tek kanıt)" if belge_oran or kart_oran else ""))
        if grup_adi == "a":
            # Ana birim Adet: Türkçe/harf varyantı kod düzeyinde eşlenir, katsayı 1
            guven, oneri = "kesin", None
            kanit.append("Ana birim Adet (yazım varyantı) — kod normalizasyonu ile katsayı 1, veri eklenmez")
        sonuc.append({
            "stok_id": k["id"], "stok_kodu": kod, "stok_adi": k["stok_adi"], "ana_birim": k["birim"],
            "alternatifler": ", ".join(alt) or "-", "adet_satir": len(ss),
            "belge_turleri": dict(Counter(s["belge_turu"] for s in ss)), "grup": grup_adi,
            "onerilen_ana_basina_adet": oneri,
            "onerilen_adet_carpani": (Decimal("1") / oneri).quantize(Decimal("0.000001")) if oneri else None,
            "kanit": " | ".join(kanit), "guven": guven,
        })
    c.close()
    return sonuc, sayim


def gecmis_etki(db: Path, kesinler: list[dict]) -> list[dict]:
    """Kesin ürünlerin geçmiş Adet satırları: kayıtlı katsayı yoksa raporlar yeni katsayıyı kullanır."""
    c = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    out = []
    for u in kesinler:
        for tab, no_tab, fk in (("satis_faturasi_satirlari", "satis_faturalari", "fatura_id"),
                                ("alis_faturasi_satirlari", "alis_faturalari", "fatura_id")):
            kol = {r[1] for r in c.execute(f"PRAGMA table_info({tab})")}
            fifo = "s.fifo_birim_maliyeti" if "fifo_birim_maliyeti" in kol else "NULL"
            bc = "s.birim_carpani" if "birim_carpani" in kol else "NULL"
            for r in c.execute(f"SELECT b.fatura_no, b.fatura_tarihi, s.miktar, s.birim, s.birim_fiyat, {fifo} fifo, "
                               f"{bc} bc FROM {tab} s JOIN {no_tab} b ON b.id=s.{fk} WHERE s.urun_kodu=?",
                               (u["stok_kodu"],)):
                if adet_anahtari(r["birim"]) != "adet":
                    continue
                m = Decimal(str(r["miktar"] or 0))
                f = Decimal(str(r["fifo"] or 0))
                out.append({
                    "stok_kodu": u["stok_kodu"], "belge": r["fatura_no"], "tablo": tab, "tarih": r["fatura_tarihi"],
                    "miktar": m, "birim_fiyat": r["birim_fiyat"], "kayitli_katsayi": r["bc"],
                    "rapor_maliyeti_simdi": (m * f) if r["bc"] is None else None,
                    "rapor_maliyeti_tanim_sonrasi": (m * f * u["onerilen_adet_carpani"]) if r["bc"] is None else None,
                    "stok_hareketi_1e1": m,
                    "fiziksel_dogru_ana_birim": m * u["onerilen_adet_carpani"],
                })
    c.close()
    return out


def excel_yaz(urunler, sayim, etki, cikti: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    def v(x):
        return float(x) if isinstance(x, Decimal) else (json.dumps(x, ensure_ascii=False) if isinstance(x, dict) else x)

    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    kesin = [u for u in urunler if u["guven"] == "kesin" and u["onerilen_adet_carpani"]]
    for satir in (
        ["Tanımsız birimli satır", sayim["tanimsiz_satir"]],
        ["Türlere göre", json.dumps(sayim["tanimsiz_tur"], ensure_ascii=False)],
        ["«Adet» satırı (büyük/küçük harf + Türkçe I normalize)", sayim["adet_satir"]],
        ["Ham yazımlar", json.dumps(sayim["adet_ham_yazimlar"], ensure_ascii=False)],
        ["«Adet» satırlı ürün", sayim["adet_urun"]],
        ["(a) Ana birimi Adet (varyant)", sum(u["grup"] == "a" for u in urunler)],
        ["(b) Ana birimi başka", sum(u["grup"] == "b" for u in urunler)],
        ["Kesin (Adet katsayısı tanımlanabilir)", len(kesin)],
        ["Belirsiz (inceleme listesinde kalır)", sum(u["guven"] == "belirsiz" for u in urunler)],
        [],
        ["Not", "Katsayı = 1 Adet kaç ana birim (stok kartı alternatif birim «carpan» alanı)."],
        ["Not", "Geçmiş belge satırları değişmez. Kayıtlı katsayısı olmayan eski satırları kâr analizi ve "
                "muhasebeleştirme kartın güncel katsayısıyla yorumlar — «Geçmiş satır etkisi» sayfasına bakın."],
    ):
        oz.append(satir)
    oz.column_dimensions["A"].width = 55
    oz.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Ürünler")
    basl = [("stok_kodu", "Stok kodu"), ("stok_adi", "Stok adı"), ("ana_birim", "Ana birim"),
            ("alternatifler", "Mevcut alternatifler"), ("adet_satir", "Tanımsız Adet satırı"),
            ("belge_turleri", "Belge türleri"), ("grup", "Grup (a/b)"),
            ("onerilen_ana_basina_adet", "1 ana birim = ? Adet"),
            ("onerilen_adet_carpani", "Önerilen Adet katsayısı (1 Adet = ? ana)"), ("guven", "Güven"),
            ("kanit", "Kanıt")]
    ws.append([b for _, b in basl])
    sari = PatternFill("solid", fgColor="FFF2CC")
    yesil = PatternFill("solid", fgColor="E2EFDA")
    for u in sorted(urunler, key=lambda x: (x["guven"] != "kesin", x["stok_kodu"])):
        ws.append([v(u[a]) for a, _ in basl])
        for hucre in ws[ws.max_row]:
            hucre.fill = yesil if u["guven"] == "kesin" else sari
    for i, w in enumerate((16, 48, 10, 22, 10, 30, 8, 12, 16, 10, 120), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    es = wb.create_sheet("Geçmiş satır etkisi")
    eb = [("stok_kodu", "Stok kodu"), ("belge", "Belge"), ("tablo", "Tablo"), ("tarih", "Tarih"),
          ("miktar", "Miktar (Adet)"), ("birim_fiyat", "Birim fiyat"), ("kayitli_katsayi", "Kayıtlı katsayı"),
          ("stok_hareketi_1e1", "Stoğa işlenen (ana birim, 1:1)"),
          ("fiziksel_dogru_ana_birim", "Katsayıyla ana birim karşılığı"),
          ("rapor_maliyeti_simdi", "Kâr analizi maliyeti (şimdi)"),
          ("rapor_maliyeti_tanim_sonrasi", "Kâr analizi maliyeti (tanım sonrası)")]
    es.append([b for _, b in eb])
    for e in etki:
        es.append([v(e[a]) for a, _ in eb])
    for ws_ in (ws, es):
        for h in ws_[1]:
            h.font = Font(bold=True)
        ws_.freeze_panes = "A2"
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def uygulama_calisiyor() -> bool:
    try:
        cikti = subprocess.run(["tasklist", "/FI", "IMAGENAME eq CinMuhasebe.exe"],
                               capture_output=True, text=True, timeout=20).stdout
    except Exception:
        return False
    return "CinMuhasebe.exe" in cikti


def satir_ozeti(db: Path) -> dict:
    """Geçmiş belge satırlarının değişmediğini kanıtlamak için tablo özetleri (sayı + SHA256)."""
    c = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    out = {}
    for t in SATIR_TABLOLARI + ("stok_hareketleri", "stok_lotlari"):
        h = hashlib.sha256()
        n = 0
        for r in c.execute(f'SELECT * FROM "{t}" ORDER BY id'):
            h.update(repr(r).encode())
            n += 1
        out[t] = {"satir": n, "sha256": h.hexdigest()}
    out["stok_birimleri"] = c.execute("SELECT COUNT(*) FROM stok_birimleri").fetchone()[0]
    c.close()
    return out


def tanimla(db: Path, kesinler: list[dict]) -> list[str]:
    from sqlalchemy import create_engine, select
    from sqlalchemy.orm import Session, selectinload

    from database.models.stok import StokKarti
    from database.stok_service import StokService, _birim_orm_to_dict, birim_anahtari

    motor = create_engine(f"sqlite:///{db.resolve().as_posix()}", connect_args={"timeout": 30})
    eklenen = []
    try:
        with Session(motor) as s, s.begin():
            for u in kesinler:
                stok = s.scalar(select(StokKarti).where(StokKarti.id == u["stok_id"])
                                .options(selectinload(StokKarti.birimler), selectinload(StokKarti.barkodlar)))
                if any(birim_anahtari(b.birim_adi) == birim_anahtari("Adet") for b in stok.birimler):
                    continue
                mevcut = [_birim_orm_to_dict(b) for b in stok.birimler]
                c = str(u["onerilen_adet_carpani"])
                StokService._birimleri_sessiona_yaz(s, stok, mevcut + [{
                    "birim_adi": "Adet", "carpan": c, "referans_birim": stok.birim, "referans_carpan": c,
                    "fiyat_modu": "otomatik", "ondalik": 0,
                }], ana_birim=stok.birim)
                eklenen.append(f"{stok.stok_kodu}: 1 Adet = {c} {stok.birim}")
            s.flush()
    finally:
        motor.dispose()
    return eklenen


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default=str(VARSAYILAN_DB))
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--dry-run", action="store_true", default=True)
    g.add_argument("--uygula", action="store_true")
    ap.add_argument("--cikti", default=str(CIKTI))
    ap.add_argument("--sonuc", default=None)
    a = ap.parse_args()
    db = Path(a.db)
    gercek = db.resolve() == VARSAYILAN_DB.resolve()

    urunler, sayim = urun_kanitlari(db)
    kesinler = [u for u in urunler if u["guven"] == "kesin" and u["onerilen_adet_carpani"]]
    etki = gecmis_etki(db, kesinler)
    excel_yaz(urunler, sayim, etki, Path(a.cikti))
    rapor = {"zaman": datetime.now().isoformat(timespec="seconds"), "db": str(db), "gercek_db": gercek,
             "uygula": a.uygula, "sayim": sayim,
             "kesin": [(u["stok_kodu"], u["stok_adi"], str(u["onerilen_adet_carpani"])) for u in kesinler],
             "belirsiz_urun": sum(u["guven"] == "belirsiz" for u in urunler),
             "gecmis_etki_satir": len(etki),
             "gecmis_etki_kayitli_katsayisiz": sum(e["kayitli_katsayi"] is None for e in etki)}
    kod = 0
    if a.uygula and kesinler:
        if gercek:
            if uygulama_calisiyor():
                rapor["durum"] = "program açık, uygulanmadı"
                kod = 3
            else:
                from database.gecis_guvenligi import GecisYedekHatasi, dogrulanmis_yedek
                try:
                    y = dogrulanmis_yedek(db, "adet_katsayi")
                    rapor["yedek"] = str(y.yol)
                except GecisYedekHatasi as hata:
                    rapor["durum"] = f"yedek doğrulanamadı, uygulanmadı: {hata}"
                    kod = 4
        if kod == 0:
            once = satir_ozeti(db)
            try:
                rapor["eklenen"] = tanimla(db, kesinler)
            except Exception as hata:
                rapor["durum"] = f"geri alındı: {hata}"
                kod = 5
            else:
                sonra = satir_ozeti(db)
                degisen = [t for t in once if t != "stok_birimleri" and once[t] != sonra[t]]
                rapor["stok_birimleri"] = (once["stok_birimleri"], sonra["stok_birimleri"])
                rapor["degismeyen_tablolar"] = [t for t in once if t != "stok_birimleri" and t not in degisen]
                c = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
                rapor["integrity_sonrasi"] = c.execute("PRAGMA integrity_check").fetchone()[0]
                c.close()
                sorun = degisen + ([] if sonra["stok_birimleri"] - once["stok_birimleri"] == len(rapor["eklenen"])
                                   else ["stok_birimleri sayısı beklenmedik"])
                if rapor["integrity_sonrasi"] != "ok":
                    sorun.append("integrity_check")
                rapor["durum"] = "uygulandı, doğrulama tamam" if not sorun else f"uygulandı, SORUN: {sorun}"
                kod = 0 if not sorun else 6
    else:
        rapor["durum"] = "kuru çalışma: Excel üretildi, veri yazılmadı"
    if a.sonuc:
        Path(a.sonuc).write_text(json.dumps(rapor, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps(rapor, ensure_ascii=False, indent=1, default=str))
    return kod


if __name__ == "__main__":
    raise SystemExit(main())
