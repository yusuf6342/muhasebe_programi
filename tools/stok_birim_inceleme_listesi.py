"""Geçmişte kart temel biriminden farklı birimle kesilmiş belge satırları — ürün bazlı inceleme listesi.

SALT OKUNUR: veritabanı ``mode=ro`` ile açılır, hiçbir kayıt değiştirilmez.
Eski satırlar 1:1 (belge miktarı = temel birim miktarı) stoğa işlenmişti; bu liste yalnızca
kullanıcının hangi ürünlerde birim tanımı/düzeltme gerektiğine karar vermesi içindir.

Çalıştırma:
    .venv\\Scripts\\python.exe tools\\stok_birim_inceleme_listesi.py [db_yolu] [--cikti dosya.xlsx]
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
VARSAYILAN_DB = Path.home() / "AppData/Local/MuhasebeProgrami/data/muhasebe.db"
VARSAYILAN_CIKTI = ROOT / "inceleme" / "stoklar" / "Farkli_Birim_Satirlari_Inceleme.xlsx"

BELGELER = (
    # (tür, satır tablosu, belge tablosu, fk, no kolonu, tarih kolonu, stok hareket türü, yön)
    ("Alış faturası", "alis_faturasi_satirlari", "alis_faturalari", "fatura_id", "fatura_no",
     "fatura_tarihi", "FATURA GİRİŞ", "+"),
    ("Satış faturası", "satis_faturasi_satirlari", "satis_faturalari", "fatura_id", "fatura_no",
     "fatura_tarihi", "FATURA ÇIKIŞ", "-"),
    ("Alış iade faturası", "alis_iade_faturasi_satirlari", "alis_iade_faturalari", "iade_id", "iade_no",
     "iade_tarihi", "FATURA ÇIKIŞ", "-"),
    ("Satış iade faturası", "satis_iade_faturasi_satirlari", "satis_iade_faturalari", "iade_id", "iade_no",
     "iade_tarihi", "İADE GİRİŞ", "+"),
)


def tr_kucuk(metin: str | None) -> str:
    return (metin or "").strip().replace("İ", "i").replace("I", "ı").lower()


def kolonlar(cur, tablo: str) -> set[str]:
    return {r[1] for r in cur.execute(f'PRAGMA table_info("{tablo}")')}


def satirlari_topla(db: Path) -> list[dict]:
    con = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    cur = con.cursor()
    kartlar = {
        r["stok_kodu"]: dict(r)
        for r in cur.execute("SELECT id, stok_kodu, stok_adi, birim FROM stok_kartlari")
    }
    birimler: dict[int, dict[str, Decimal]] = defaultdict(dict)
    for r in cur.execute("SELECT stok_id, birim_adi, carpan, aktif FROM stok_birimleri"):
        if r["aktif"] in (0, False):
            continue
        c = Decimal(str(r["carpan"] or 0))
        if c > 0:
            birimler[r["stok_id"]][tr_kucuk(r["birim_adi"])] = c
    cariler = {r[0]: r[1] for r in cur.execute("SELECT id, unvan FROM cari_kartlar")}

    sonuc: list[dict] = []
    for tur, s_tab, b_tab, fk, no_kol, tarih_kol, h_tur, yon in BELGELER:
        if not kolonlar(cur, s_tab):
            continue
        hareket: dict[tuple[str, int], Decimal] = defaultdict(Decimal)
        for r in cur.execute(
            "SELECT belge_no, stok_id, SUM(miktar) FROM stok_hareketleri WHERE hareket_turu=? GROUP BY 1, 2",
            (h_tur,),
        ):
            hareket[(r[0], r[1])] = Decimal(str(r[2] or 0))
        satir_sayisi: dict[tuple[int, str], int] = defaultdict(int)
        for r in cur.execute(f"SELECT {fk}, urun_kodu FROM {s_tab}"):
            satir_sayisi[(r[0], r[1])] += 1
        bc = kolonlar(cur, s_tab)
        carpan_kol = "s.birim_carpani" if "birim_carpani" in bc else "NULL"
        for r in cur.execute(
            f"""SELECT s.id satir_id, s.{fk} belge_id, s.urun_kodu, s.urun_adi, s.miktar, s.birim,
                       {carpan_kol} kayitli_carpan, b.{no_kol} belge_no, b.{tarih_kol} tarih,
                       b.durum, b.depo, b.cari_id
                FROM {s_tab} s JOIN {b_tab} b ON b.id = s.{fk}"""
        ):
            kart = kartlar.get(r["urun_kodu"])
            if kart is None:
                continue
            temel = kart["birim"] or "Adet"
            if tr_kucuk(r["birim"]) == tr_kucuk(temel):
                continue
            tanimli = birimler.get(kart["id"], {}).get(tr_kucuk(r["birim"]))
            miktar = Decimal(str(r["miktar"] or 0))
            etki = hareket.get((r["belge_no"], kart["id"]))
            ayni_urun_satir = satir_sayisi[(r["belge_id"], r["urun_kodu"])]
            if r["kayitli_carpan"] is not None:
                kullanilan = Decimal(str(r["kayitli_carpan"]))
                kullanilan_ack = "Belgede kayıtlı katsayı"
            elif etki is None:
                kullanilan = None
                kullanilan_ack = "Stok hareketi yok (taslak/iptal/irsaliyeli)"
            elif ayni_urun_satir == 1 and miktar and etki == miktar:
                kullanilan = Decimal("1")
                kullanilan_ack = "1:1 işlenmiş (hareket = belge miktarı)"
            elif ayni_urun_satir == 1 and miktar:
                kullanilan = (etki / miktar).quantize(Decimal("0.0001"))
                kullanilan_ack = "Hareket/belge oranı"
            else:
                kullanilan = None
                kullanilan_ack = "Aynı üründe birden çok satır — belge toplamına bakın"

            if etki is None:
                sinif = "STOK ETKİSİZ"
                aciklama = "Bu satır stoğa hareket yazmamış; stok miktarını etkilemez."
            elif tanimli is not None and kullanilan is not None and kullanilan != tanimli:
                sinif = "KESİN"
                fark = miktar * (tanimli - kullanilan)
                aciklama = (
                    f"Kartta «{r['birim']}» = {tanimli.normalize()} {temel} tanımlı; stoğa "
                    f"{kullanilan.normalize()} katsayısıyla işlenmiş. Olası fark: "
                    f"{yon}{fark.normalize()} {temel} (kullanıcı kararı gerekir)."
                )
            elif tanimli is not None and kullanilan == tanimli:
                sinif = "TUTARLI"
                aciklama = "Kart katsayısıyla işlenmiş; işlem gerekmez."
            elif tanimli is not None:
                sinif = "BELİRSİZ"
                aciklama = (
                    f"Kartta katsayı {tanimli.normalize()} tanımlı ama satırın kullandığı katsayı "
                    f"belge toplamından ayrıştırılamadı."
                )
            else:
                sinif = "BELİRSİZ"
                aciklama = (
                    f"«{r['birim']}» birimi kartta tanımlı değil; 1 {r['birim']} kaç {temel} "
                    f"bilinmiyor. Kart → Birimler'de tanımlanmalı (yeni belgelerde zorunlu)."
                )
            sonuc.append({
                "stok_kodu": kart["stok_kodu"],
                "stok_adi": kart["stok_adi"],
                "belge_turu": tur,
                "belge_id": r["belge_id"],
                "belge_no": r["belge_no"],
                "satir_id": r["satir_id"],
                "tarih": r["tarih"],
                "durum": r["durum"],
                "depo": r["depo"],
                "cari": cariler.get(r["cari_id"], ""),
                "miktar": miktar,
                "belge_birimi": r["birim"],
                "temel_birim": temel,
                "kart_katsayisi": tanimli,
                "kullanilan_katsayi": kullanilan,
                "kullanilan_aciklama": kullanilan_ack,
                "stok_etkisi": (Decimal(yon + "1") * etki) if etki is not None else None,
                "ayni_urun_satir": ayni_urun_satir,
                "sinif": sinif,
                "aciklama": aciklama,
            })
    con.close()
    sonuc.sort(key=lambda s: (s["stok_kodu"], str(s["tarih"]), s["belge_turu"], s["satir_id"]))
    return sonuc


def excel_yaz(satirlar: list[dict], cikti: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    def sayi(v):
        return float(v) if isinstance(v, Decimal) else v

    wb = Workbook()
    ozet = wb.active
    ozet.title = "Ürün Özeti"
    ozet_baslik = ["Stok kodu", "Stok adı", "Temel birim", "Belge birimleri", "Satır", "KESİN", "BELİRSİZ",
                   "TUTARLI", "STOK ETKİSİZ", "Önerilen işlem"]
    ozet.append(ozet_baslik)
    grup: dict[str, list[dict]] = defaultdict(list)
    for s in satirlar:
        grup[s["stok_kodu"]].append(s)
    for kod, liste in sorted(grup.items()):
        say = defaultdict(int)
        for s in liste:
            say[s["sinif"]] += 1
        oneri = (
            "Birim katsayısı ile geçmiş stok farkını kontrol edin" if say["KESİN"]
            else "Kart → Birimler'de katsayı tanımlayın" if say["BELİRSİZ"]
            else "İşlem gerekmez"
        )
        birimler = ", ".join(sorted({s["belge_birimi"] or "" for s in liste}))
        ozet.append([kod, liste[0]["stok_adi"], liste[0]["temel_birim"], birimler, len(liste), say["KESİN"],
                     say["BELİRSİZ"], say["TUTARLI"], say["STOK ETKİSİZ"], oneri])

    detay = wb.create_sheet("Satırlar")
    basliklar = [
        ("stok_kodu", "Stok kodu"), ("stok_adi", "Stok adı"), ("belge_turu", "Belge türü"),
        ("belge_id", "Belge ID"), ("belge_no", "Belge no"), ("satir_id", "Satır ID"), ("tarih", "Tarih"),
        ("durum", "Durum"), ("depo", "Depo"), ("cari", "Cari"), ("miktar", "Belge miktarı"),
        ("belge_birimi", "Belge birimi"), ("temel_birim", "Temel birim"),
        ("kart_katsayisi", "Kartta tanımlı katsayı"), ("kullanilan_katsayi", "Kullanılan katsayı"),
        ("kullanilan_aciklama", "Katsayı kaynağı"), ("stok_etkisi", "Stok etkisi (belge+ürün toplamı)"),
        ("ayni_urun_satir", "Belgede aynı ürün satırı"), ("sinif", "Sınıf"), ("aciklama", "Açıklama"),
    ]
    detay.append([b for _, b in basliklar])
    renk = {"KESİN": "FCE4D6", "BELİRSİZ": "FFF2CC", "TUTARLI": "E2EFDA", "STOK ETKİSİZ": "EDEDED"}
    for s in satirlar:
        detay.append([sayi(s[k]) for k, _ in basliklar])
        dolgu = PatternFill("solid", fgColor=renk.get(s["sinif"], "FFFFFF"))
        detay.cell(row=detay.max_row, column=19).fill = dolgu

    bilgi = wb.create_sheet("Açıklama")
    for satir in (
        ["Bu liste salt okunur bir taramadır; hiçbir kayıt değiştirilmemiştir."],
        ["Eski belgeler kart biriminden farklı birimle kesildiğinde stoğa 1:1 işlenmişti."],
        ["Yeni belgelerde katsayı kayıt anında belgeye yazılır; tanımsız birimle stok hareketi oluşturulmaz."],
        ["KESİN: kartta katsayı var ve eski hareket farklı katsayıyla işlenmiş (olası stok farkı hesaplandı)."],
        ["BELİRSİZ: birim kartta tanımsız veya satır katsayısı ayrıştırılamadı — kullanıcı kararı gerekir."],
        ["TUTARLI: kart katsayısıyla aynı işlenmiş."],
        ["STOK ETKİSİZ: belge stok hareketi yazmamış (taslak/iptal/irsaliyeye bağlı)."],
    ):
        bilgi.append(satir)

    for ws in (ozet, detay):
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
    bilgi.column_dimensions["A"].width = 110
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("db", nargs="?", default=str(VARSAYILAN_DB))
    ap.add_argument("--cikti", default=str(VARSAYILAN_CIKTI))
    a = ap.parse_args()
    satirlar = satirlari_topla(Path(a.db))
    excel_yaz(satirlar, Path(a.cikti))
    say = defaultdict(int)
    for s in satirlar:
        say[s["sinif"]] += 1
    tur = defaultdict(int)
    for s in satirlar:
        tur[s["belge_turu"]] += 1
    print(f"Toplam satır: {len(satirlar)}  ürün: {len({s['stok_kodu'] for s in satirlar})}")
    print("Belge türü:", dict(tur))
    print("Sınıf:", dict(say))
    print("Çıktı:", a.cikti)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
